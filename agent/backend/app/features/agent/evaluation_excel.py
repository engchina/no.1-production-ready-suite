"""品質評価の評価セットの Excel（#776）。NL2SQL の SQL生成評価と同じく、評価ケースを Excel で
取り込み・書き出しする（1 行 1 ケース。列は ケースID・質問・期待する回答の要点・期待するツール・
ツールの順（「順どおり」なら期待するツールを書いた順に呼ぶ。#1845））。
"""

from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]
from openpyxl.styles import Alignment, Font, PatternFill  # type: ignore[import-untyped]
from pydantic import ValidationError

from app.features.agent.evaluation import (
    EVALUATION_MAX_CASES,
    EvaluationCase,
    EvaluationSet,
    number_cases,
)

EXCEL_MAX_BYTES = 2 * 1024 * 1024
SHEET_TITLE = "評価ケース"
HEADERS = ("ケースID", "質問", "期待する回答の要点", "期待するツール", "ツールの順")
# 「ツールの順」の列の値（#1845）。書き出しは IN_ORDER_LABEL、読み込みは IN_ORDER_VALUES のどれか。
IN_ORDER_LABEL = "順どおり"
IN_ORDER_VALUES = frozenset({IN_ORDER_LABEL, "順番どおり", "はい", "yes", "true", "1", "in_order"})
# 見出しの別名（英語の列名でも読めるようにする）。
_ALIASES: dict[str, tuple[str, ...]] = {
    "id": ("ケースid", "ケース id", "id", "case_id"),
    "question": ("質問", "question"),
    "expected": ("期待する回答の要点", "期待する回答", "expected"),
    "expected_tools": ("期待するツール", "expected_tools", "tools"),
    "expected_tools_in_order": ("ツールの順", "期待するツールの順", "expected_tools_in_order"),
}
_SAMPLE = EvaluationCase(
    id="expense-deadline",
    question="経費精算の締め日はいつですか？",
    expected="毎月 25 日が締め日で、過ぎた分は翌月の精算になること",
    expected_tools=["rag_retrieve_evidence", "rag_validate_answer"],
    expected_tools_in_order=True,
)


class EvaluationExcelError(ValueError):
    """Excel を読めない（利用者向けの日本語）。"""


def _workbook(cases: list[EvaluationCase]) -> bytes:
    book = Workbook()
    sheet = book.active or book.create_sheet()
    sheet.title = SHEET_TITLE
    sheet.append(list(HEADERS))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
    for case in cases:
        sheet.append(
            [
                case.id,
                case.question,
                case.expected,
                ", ".join(case.expected_tools),
                IN_ORDER_LABEL if case.expected_tools_in_order and case.expected_tools else "",
            ]
        )
    for column, width in zip("ABCDE", (18, 48, 56, 28, 12), strict=True):
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def template_xlsx() -> bytes:
    return _workbook([_SAMPLE])


def export_set_xlsx(evaluation_set: EvaluationSet) -> bytes:
    return _workbook(list(evaluation_set.cases))


def _header_map(row: tuple[Any, ...]) -> dict[str, int]:
    found: dict[str, int] = {}
    for index, value in enumerate(row):
        label = str(value or "").strip().lower()
        for key, aliases in _ALIASES.items():
            if label in aliases and key not in found:
                found[key] = index
    return found


def _text(row: tuple[Any, ...], index: int | None) -> str:
    if index is None or index >= len(row) or row[index] is None:
        return ""
    return str(row[index]).strip()


def parse_cases_xlsx(data: bytes) -> list[EvaluationCase]:
    """Excel の評価ケースを読む（最初のシート。空の行は飛ばす）。"""
    if len(data) > EXCEL_MAX_BYTES:
        raise EvaluationExcelError("Excel のファイルは 2MB までです。")
    try:
        book = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - 形式の誤りは利用者へ返す
        raise EvaluationExcelError("Excel（.xlsx）として読めません。") from exc
    sheet = book.worksheets[0] if book.worksheets else None
    if sheet is None:
        raise EvaluationExcelError("シートがありません。")
    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None)
    columns = _header_map(tuple(header or ()))
    if "question" not in columns or "expected" not in columns:
        raise EvaluationExcelError(
            "1 行目に「質問」と「期待する回答の要点」の列が必要です"
            "（テンプレートを使ってください）。"
        )
    cases: list[EvaluationCase] = []
    for number, row in enumerate(rows, start=2):
        values = tuple(row)
        question = _text(values, columns.get("question"))
        expected = _text(values, columns.get("expected"))
        if not question and not expected:
            continue
        if len(cases) >= EVALUATION_MAX_CASES:
            raise EvaluationExcelError(f"評価ケースは {EVALUATION_MAX_CASES} 件までです。")
        tools = _text(values, columns.get("expected_tools"))
        in_order = _text(values, columns.get("expected_tools_in_order")).lower() in IN_ORDER_VALUES
        try:
            cases.append(
                EvaluationCase(
                    id=_text(values, columns.get("id")),
                    question=question,
                    expected=expected,
                    expected_tools=[item for item in tools.replace("、", ",").split(",")],
                    expected_tools_in_order=in_order,
                )
            )
        except ValidationError as exc:
            raise EvaluationExcelError(
                f"{number} 行目を読めません（質問と期待する回答の要点は必須です）。"
            ) from exc
    if not cases:
        raise EvaluationExcelError("評価ケースが 1 件もありません。")
    try:
        return number_cases(cases)
    except ValueError as exc:
        raise EvaluationExcelError(str(exc)) from exc
