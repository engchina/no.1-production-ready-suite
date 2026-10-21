"""品質評価の Excel は入力した文を数式に変えず、往復で保持する。"""

from io import BytesIO

from openpyxl import load_workbook  # type: ignore[import-untyped]

from app.features.agent.evaluation import EvaluationCase, EvaluationSet
from app.features.agent.evaluation_excel import export_set_xlsx, parse_cases_xlsx


def test_formula_like_text_round_trips_as_literal_cells() -> None:
    case = EvaluationCase(
        id="=synthetic-id",
        question="=1+1",
        expected="=synthetic-expected",
        expected_tools=["=synthetic-tool"],
    )
    evaluation_set = EvaluationSet(agent_id="synthetic", name="合成の評価セット", cases=[case])
    data = export_set_xlsx(evaluation_set)
    book = load_workbook(BytesIO(data), data_only=False)
    cells = list(book.worksheets[0].iter_rows(min_row=2))[0]
    assert all(cell.data_type == "s" for cell in cells[:4])
    [parsed] = parse_cases_xlsx(data)
    assert parsed.model_dump() == case.model_dump()
