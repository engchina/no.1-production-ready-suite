"""AI オントロジー構築(ontology_build)のテスト。

LLM は fake、store は InMemoryOntologyStore。job → Markdown Draft → publish と、
既存 proposal endpoint の accept → publish 経路、スコープ外候補の warnings 落ちを検証する。
"""

from __future__ import annotations

import hashlib
import io
import json
import time
from typing import Any

import pytest

import app.features.nl2sql.ontology_build as ontology_build_module
from app.features.nl2sql.models import (
    Nl2SqlProfile,
    SchemaCatalog,
    SchemaColumn,
    SchemaConstraintDetail,
    SchemaTable,
    SchemaViewDependency,
)
from app.features.nl2sql.ontology_build import (
    OntologyBuildService,
    build_schema_context_from_catalog,
    parse_qa_workbook,
    render_ontology_build_markdown,
)
from app.features.nl2sql.ontology_catalog import SchemaOntology, catalog_schema_fingerprint
from app.features.nl2sql.ontology_models import (
    BusinessRuleDefinition,
    BusinessRuleExecutionMode,
    BusinessRuleKind,
    BusinessRuleSeverity,
    EnumValueDefinition,
    JoinCondition,
    OntologyBuildJob,
    OntologyBuildStatus,
    OntologyBuildStep,
    OntologyBuildStepName,
    OntologyBuildStepStatus,
    OntologyEdge,
    OntologyEdgeKind,
    OntologyNode,
    OntologyNodeKind,
    OntologyProposalKind,
    OntologyProvenance,
    OntologyReviewStatus,
    OntologyRevision,
    OntologyRevisionStatus,
    OntologySourceDocument,
    OntologySourceKind,
    OntologySourceStatus,
    PhysicalColumnRef,
    PhysicalMapping,
    PhysicalObjectRef,
    ProfileOntologyView,
    QaPair,
    RelationshipCardinality,
)
from app.features.nl2sql.ontology_router import (
    OntologyApiRuntime,
    OntologyDraftRequest,
    OntologyMarkdownDraftPatch,
    OntologyPublishRequest,
)
from app.features.nl2sql.ontology_service import OntologyNotFoundError
from app.features.nl2sql.ontology_store import InMemoryOntologyStore, OntologyVersionConflict
from app.settings import get_settings


class _FakeLegacyNl2SqlService:
    def __init__(self) -> None:
        self._enterprise_ai_client: Any = None
        self.profile = Nl2SqlProfile(
            id="sales",
            name="販売分析",
            allowed_tables=["APP.ORDERS", "APP.CUSTOMERS"],
            default_row_limit=100,
        )
        self.catalog = SchemaCatalog(
            refreshed_at="2026-07-11T00:00:00Z",
            tables=[
                SchemaTable(
                    table_name="ORDERS",
                    logical_name="受注",
                    owner="APP",
                    columns=[
                        SchemaColumn(column_name="ID", logical_name="受注 ID", data_type="NUMBER"),
                        SchemaColumn(
                            column_name="CUSTOMER_ID", logical_name="顧客 ID", data_type="NUMBER"
                        ),
                        SchemaColumn(
                            column_name="AMOUNT", logical_name="受注金額", data_type="NUMBER"
                        ),
                    ],
                    constraint_details=[
                        SchemaConstraintDetail(
                            constraint_name="FK_ORDERS_CUSTOMER",
                            constraint_type="R",
                            owner="APP",
                            table_name="ORDERS",
                            columns=["CUSTOMER_ID"],
                            referenced_owner="APP",
                            referenced_table="CUSTOMERS",
                            referenced_columns=["ID"],
                        )
                    ],
                ),
                SchemaTable(
                    table_name="CUSTOMERS",
                    logical_name="顧客",
                    owner="APP",
                    columns=[
                        SchemaColumn(column_name="ID", logical_name="顧客 ID", data_type="NUMBER"),
                        SchemaColumn(
                            column_name="NAME", logical_name="顧客名", data_type="VARCHAR2"
                        ),
                    ],
                    constraint_details=[
                        SchemaConstraintDetail(
                            constraint_name="PK_CUSTOMERS",
                            constraint_type="P",
                            owner="APP",
                            table_name="CUSTOMERS",
                            columns=["ID"],
                        )
                    ],
                ),
            ],
        )

    def get_catalog(self) -> SchemaCatalog:
        return self.catalog

    def get_profile(self, profile_id: str) -> Nl2SqlProfile:
        if profile_id != self.profile.id:
            raise ValueError("profile not found")
        return self.profile


class _FakeEnterpriseAiClient:
    def __init__(self, payload: str, *, configured: bool = True) -> None:
        self.payload = payload
        self.configured = configured
        self.calls: list[str] = []
        self.contexts: list[str] = []

    def is_configured(self) -> bool:
        return self.configured

    def model_id(self) -> str:
        return "fake-enterprise-ai"

    def generate(self, *, prompt: str, context: str, system_prompt: str) -> str:
        self.calls.append(prompt)
        self.contexts.append(context)
        return self.payload


class _FakeOntologySourceStorage:
    def __init__(self, contents: dict[str, bytes]) -> None:
        self.contents = contents

    def load(self, document: OntologySourceDocument) -> bytes:
        return self.contents[document.id]


_EXTRACTION = {
    "entities": [
        {
            "object_name": "APP.ORDERS",
            "business_name_ja": "受注",
            "description_ja": "受注トランザクション",
            "aliases": ["注文"],
            "confidence": 0.9,
        },
        {
            "object_name": "APP.SECRET",
            "business_name_ja": "秘密",
            "description_ja": "profile 範囲外",
            "aliases": [],
            "confidence": 0.9,
        },
    ],
    "relationships": [
        {
            "source_object": "APP.ORDERS",
            "target_object": "APP.CUSTOMERS",
            "relationship_name_ja": "顧客を参照",
            "cardinality": "many_to_one",
            "join_conditions": [
                {"left": "APP.ORDERS.CUSTOMER_ID", "right": "APP.CUSTOMERS.ID", "operator": "="}
            ],
            "evidence_ja": "Q/A の JOIN 句",
            "confidence": 0.8,
        }
    ],
    "metrics": [
        {
            "metric_name_ja": "受注金額合計",
            "expression_sql": "SUM(AMOUNT)",
            "aggregation": "sum",
            "base_columns": ["APP.ORDERS.AMOUNT"],
            "unit": "円",
            "description_ja": "受注金額の合計",
            "evidence_ja": "",
            "confidence": 0.7,
        }
    ],
    "synonyms": [{"target": "APP.ORDERS", "aliases": ["オーダー"], "evidence_ja": ""}],
    "warnings_ja": [],
}
_FENCED_PAYLOAD = "以下が抽出結果です。\n" + json.dumps(_EXTRACTION, ensure_ascii=False) + "\n以上"

_QA_SQL = (
    "SELECT C.NAME, SUM(O.AMOUNT) FROM APP.ORDERS O "
    "JOIN APP.CUSTOMERS C ON O.CUSTOMER_ID = C.ID GROUP BY C.NAME"
)


def _xlsx_bytes(rows: list[list[str]]) -> bytes:
    import openpyxl  # type: ignore[import-untyped]

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _source_document(
    source_id: str,
    filename: str,
    content: bytes,
    *,
    media_type: str = "text/markdown",
) -> OntologySourceDocument:
    return OntologySourceDocument(
        id=source_id,
        profile_id="sales",
        filename=filename,
        media_type=media_type,
        size_bytes=len(content),
        sha256=hashlib.sha256(content).hexdigest(),
        storage_uri=f"/tmp/{filename}",
    )


def _wait_for_job(service: OntologyBuildService, job_id: str) -> Any:
    for _ in range(500):
        job = service.get(job_id)
        if job is not None and job.status in {
            OntologyBuildStatus.SUCCEEDED,
            OntologyBuildStatus.FAILED,
        }:
            return job
        time.sleep(0.01)
    raise AssertionError("ontology build job did not finish")


def _seed_build_proposals(runtime: OntologyApiRuntime) -> list[Any]:
    from app.features.nl2sql.ontology_build import convert_extraction_to_proposals, parse_extraction

    view, ontology = runtime.profile_view("sales")
    drafts, warnings = convert_extraction_to_proposals(
        parse_extraction(_FENCED_PAYLOAD),
        ontology=ontology,
        view=view,
        job_id="seed-job",
        inferred_by="test",
        qa_sql_texts=[_QA_SQL],
    )
    assert warnings
    return [
        runtime.create_build_proposal(
            profile_id="sales",
            job_id="seed-job",
            title_ja=draft.title_ja,
            description_ja=draft.description_ja,
            kind=draft.kind,
            proposal_payload=draft.payload,
            base_revision_id=ontology.revision.id,
        )
        for draft in drafts
    ]


def test_profile_delete_cancels_queued_build_and_worker_cannot_restart_it(
    harness: tuple[OntologyApiRuntime, InMemoryOntologyStore, _FakeLegacyNl2SqlService],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    runtime, store, legacy = harness
    legacy._enterprise_ai_client = _FakeEnterpriseAiClient(_FENCED_PAYLOAD)
    monkeypatch.setattr(get_settings(), "nl2sql_ontology_worker_mode", "external")
    service = OntologyBuildService(runtime)
    started = service.start("sales", business_text="販売業務")

    assert service.cancel_profile_jobs("sales") == 1
    cancelled = service.run_persisted(started.id)

    assert cancelled.status == OntologyBuildStatus.CANCELLED
    assert all(step.status == OntologyBuildStepStatus.SKIPPED for step in cancelled.steps)
    assert store.list_documents("proposals", {"profile_id": "sales"}) == []


@pytest.fixture
def harness() -> tuple[OntologyApiRuntime, InMemoryOntologyStore, _FakeLegacyNl2SqlService]:
    store = InMemoryOntologyStore()
    legacy = _FakeLegacyNl2SqlService()
    return OntologyApiRuntime(legacy_service=legacy, store=store), store, legacy


# --- Q/A workbook ----------------------------------------------------------------------------


def test_parse_qa_workbook_xlsx_and_csv() -> None:
    xlsx = _xlsx_bytes(
        [
            ["質問", "SQL", "備考"],
            ["顧客別の売上は?", _QA_SQL, "月次で利用"],
            ["削除して", "DELETE FROM APP.ORDERS", ""],
            ["", "", ""],
        ]
    )
    pairs, warnings = parse_qa_workbook("qa.xlsx", xlsx)
    assert [pair.question for pair in pairs] == ["顧客別の売上は?"]
    assert pairs[0].sql == _QA_SQL
    assert pairs[0].note_ja == "月次で利用"
    assert any("SELECT/WITH 以外" in warning for warning in warnings)

    csv_content = "QUESTION,SQL\n受注件数,SELECT COUNT(*) FROM APP.ORDERS\n".encode()
    csv_pairs, csv_warnings = parse_qa_workbook("qa.csv", csv_content)
    assert len(csv_pairs) == 1
    assert csv_warnings == []


def test_parse_qa_workbook_keeps_more_than_two_hundred_valid_rows() -> None:
    rows = [
        "QUESTION,SQL",
        *[
            f"質問 {index},SELECT {index} AS VALUE FROM APP.ORDERS"
            for index in range(205)
        ],
    ]
    pairs, warnings = parse_qa_workbook("qa.csv", "\n".join(rows).encode())

    assert warnings == []
    assert len(pairs) == 205
    assert pairs[-1].question == "質問 204"


def test_parse_qa_workbook_rejects_missing_headers_and_unknown_suffix() -> None:
    pairs, warnings = parse_qa_workbook("qa.csv", b"A,B\n1,2\n")
    assert pairs == []
    assert any("QUESTION" in warning for warning in warnings)

    pairs, warnings = parse_qa_workbook("qa.pdf", b"binary")
    assert pairs == []
    assert any("未対応の形式" in warning for warning in warnings)

    pairs, warnings = parse_qa_workbook(
        "qa.tsv", b"QUESTION\tSQL\nq\tSELECT COUNT(*) FROM APP.ORDERS\n"
    )
    assert pairs == []
    assert any("未対応の形式" in warning for warning in warnings)


# --- DB catalog schema context ---------------------------------------------------------------


def test_build_schema_context_from_profile_and_db_catalog_includes_fk_and_views() -> None:
    catalog = SchemaCatalog(
        refreshed_at="2026-07-12T00:00:00Z",
        tables=[
            SchemaTable(
                owner="APP",
                table_name="CUSTOMERS",
                logical_name="顧客",
                columns=[
                    SchemaColumn(column_name="ID", logical_name="顧客 ID", data_type="NUMBER")
                ],
                constraint_details=[
                    SchemaConstraintDetail(
                        constraint_name="PK_CUSTOMERS",
                        constraint_type="P",
                        owner="APP",
                        table_name="CUSTOMERS",
                        columns=["ID"],
                    )
                ],
            ),
            SchemaTable(
                owner="APP",
                table_name="ORDERS",
                logical_name="受注",
                columns=[
                    SchemaColumn(column_name="ID", logical_name="受注 ID", data_type="NUMBER"),
                    SchemaColumn(
                        column_name="CUSTOMER_ID", logical_name="顧客 ID", data_type="NUMBER"
                    ),
                ],
                constraint_details=[
                    SchemaConstraintDetail(
                        constraint_name="FK_ORDERS_CUSTOMER",
                        constraint_type="R",
                        owner="APP",
                        table_name="ORDERS",
                        columns=["CUSTOMER_ID"],
                        referenced_owner="APP",
                        referenced_table="CUSTOMERS",
                        referenced_columns=["ID"],
                    )
                ],
            ),
            SchemaTable(
                owner="APP",
                table_name="V_ORDER_CUSTOMER",
                table_type="view",
                logical_name="受注顧客ビュー",
                columns=[
                    SchemaColumn(column_name="ORDER_ID", logical_name="受注 ID", data_type="NUMBER")
                ],
            ),
        ],
        view_dependencies=[
            SchemaViewDependency(
                owner="APP",
                view_name="V_ORDER_CUSTOMER",
                referenced_owner="APP",
                referenced_name="ORDERS",
                referenced_type="TABLE",
            )
        ],
    )
    profile = Nl2SqlProfile(
        id="sales",
        name="販売分析",
        allowed_tables=["APP.ORDERS", "APP.CUSTOMERS"],
        allowed_views=["APP.V_ORDER_CUSTOMER"],
    )

    prepared = build_schema_context_from_catalog(profile, catalog)
    context = json.loads(prepared.schema_context)

    assert prepared.errors == []
    assert {item["object"] for item in context["objects"]} == {
        "APP.CUSTOMERS",
        "APP.ORDERS",
        "APP.V_ORDER_CUSTOMER",
    }
    assert any(
        relationship["kind"] == "foreign_key"
        and relationship["join_conditions"][0]["expression"]
        == "APP.ORDERS.CUSTOMER_ID = APP.CUSTOMERS.ID"
        for relationship in context["relationships"]
    )
    assert any(
        relationship["kind"] == "view_dependency"
        and relationship["source_object"] == "APP.V_ORDER_CUSTOMER"
        and relationship["target_object"] == "APP.ORDERS"
        for relationship in context["relationships"]
    )


def test_build_schema_context_fails_ambiguous_unqualified_profile_object() -> None:
    catalog = SchemaCatalog(
        refreshed_at="2026-07-12T00:00:00Z",
        tables=[
            SchemaTable(owner="APP", table_name="ORDERS", logical_name="受注"),
            SchemaTable(owner="CRM", table_name="ORDERS", logical_name="CRM 受注"),
        ],
    )
    profile = Nl2SqlProfile(id="sales", name="販売分析", allowed_tables=["ORDERS"])

    prepared = build_schema_context_from_catalog(profile, catalog)

    assert prepared.object_count == 0
    assert any("複数 object" in error for error in prepared.errors)


# --- job → Markdown Draft ---------------------------------------------------------------------


def test_build_job_creates_markdown_draft_and_drops_outside_candidates(
    harness: tuple[OntologyApiRuntime, InMemoryOntologyStore, _FakeLegacyNl2SqlService],
) -> None:
    runtime, _store, legacy = harness
    client = _FakeEnterpriseAiClient(_FENCED_PAYLOAD)
    legacy._enterprise_ai_client = client
    service = OntologyBuildService(runtime)

    qa_pairs, _ = parse_qa_workbook("qa.csv", f"QUESTION,SQL\n顧客別売上,{_QA_SQL}\n".encode())
    job = service.start(
        "sales",
        business_text="受注は顧客に紐づく。売上は受注金額の合計。",
        qa_pairs=qa_pairs,
    )
    finished = _wait_for_job(service, job.id)

    assert finished.status == OntologyBuildStatus.SUCCEEDED
    assert [step.name for step in finished.steps] == [
        OntologyBuildStepName.SCHEMA_CONTEXT,
        OntologyBuildStepName.SCHEMA_NAMING,
        OntologyBuildStepName.QA_EXTRACTION,
        OntologyBuildStepName.TEXT_EXTRACTION,
        OntologyBuildStepName.PROPOSAL_REGISTRATION,
    ]
    assert all(step.status == OntologyBuildStepStatus.SUCCEEDED for step in finished.steps)
    # 各ステップに開始・終了時刻が入り、アクティビティタイムラインが時系列で積まれる
    assert all(
        step.started_at is not None and step.finished_at is not None for step in finished.steps
    )
    assert finished.started_at is not None
    assert len(finished.events) >= 5
    event_times = [event.at for event in finished.events]
    assert event_times == sorted(event_times)
    assert any("スキーマ情報を準備しました" in event.message_ja for event in finished.events)
    assert any("Markdown Draft" in event.message_ja for event in finished.events)
    # スコープ外(APP.SECRET)は draft graph 化されず warnings に落ちる
    assert any("APP.SECRET" in warning for warning in finished.warnings_ja)
    assert finished.proposal_ids == []
    assert finished.draft_revision_id
    assert finished.draft_etag
    schema_context = json.loads(client.contexts[0])
    assert {item["object"] for item in schema_context["objects"]} == {
        "APP.CUSTOMERS",
        "APP.ORDERS",
    }
    orders = next(item for item in schema_context["objects"] if item["object"] == "APP.ORDERS")
    assert any(
        column["column"] == "CUSTOMER_ID" and column["data_type"] == "NUMBER"
        for column in orders["columns"]
    )
    assert any(
        relationship["source_object"] == "APP.ORDERS"
        and relationship["target_object"] == "APP.CUSTOMERS"
        and relationship["join_conditions"][0]["expression"]
        == "APP.ORDERS.CUSTOMER_ID = APP.CUSTOMERS.ID"
        for relationship in schema_context["relationships"]
    )
    assert finished.markdown_output.startswith("# Ontology Draft")
    assert "## Physical Objects" in finished.markdown_output
    assert "`APP.ORDERS` (table)" in finished.markdown_output
    assert "business_name: 受注" in finished.markdown_output
    assert "## Entities" in finished.markdown_output
    assert "受注 (`APP.ORDERS`)" in finished.markdown_output
    assert "## Relationships / Join" in finished.markdown_output
    assert "顧客を参照" in finished.markdown_output
    assert "allowed_path: true" in finished.markdown_output
    assert "`APP.ORDERS.CUSTOMER_ID = APP.CUSTOMERS.ID`" in finished.markdown_output
    assert "## Metrics" in finished.markdown_output
    assert "受注金額合計" in finished.markdown_output
    assert "APP.ORDERS.AMOUNT" in finished.markdown_output
    assert "## Business Rules / Enum Values" in finished.markdown_output
    assert "## Synonyms" in finished.markdown_output
    assert "オーダー" in finished.markdown_output
    assert "## Evidence / Warnings" in finished.markdown_output
    assert "APP.SECRET" in finished.markdown_output

    assert runtime.list_profile_proposals("sales") == []
    state = runtime.ontology_markdown_state("sales")
    assert state.draft_revision is not None
    assert state.draft_revision.id == finished.draft_revision_id
    assert state.draft_etag == finished.draft_etag
    assert state.draft_markdown == finished.markdown_output
    saved_state = runtime.save_ontology_markdown_draft(
        "sales",
        OntologyMarkdownDraftPatch(
            markdown=f"{state.draft_markdown}\n## Manual Notes\n- 確認済み\n",
            base_etag=state.draft_etag,
        ),
    )
    assert saved_state.draft_etag != state.draft_etag
    assert "Manual Notes" in saved_state.draft_markdown
    with pytest.raises(OntologyVersionConflict):
        runtime.save_ontology_markdown_draft(
            "sales",
            OntologyMarkdownDraftPatch(markdown="stale", base_etag=state.draft_etag),
        )
    draft = runtime.ontology_revision(finished.draft_revision_id)
    kinds = {node.kind for node in draft.nodes}
    assert OntologyNodeKind.BUSINESS_ENTITY in kinds
    assert OntologyNodeKind.METRIC in kinds
    relationship_edges = [
        edge for edge in draft.edges if edge.kind == OntologyEdgeKind.BUSINESS_RELATIONSHIP
    ]
    assert len(relationship_edges) == 1
    assert relationship_edges[0].review_status == OntologyReviewStatus.APPROVED
    orders_entity = next(
        node
        for node in draft.nodes
        if node.kind == OntologyNodeKind.BUSINESS_ENTITY and node.technical_name == "APP.ORDERS"
    )
    assert set(orders_entity.aliases) >= {"注文", "オーダー"}


def test_build_job_batches_all_source_chunks_without_omission(
    harness: tuple[OntologyApiRuntime, InMemoryOntologyStore, _FakeLegacyNl2SqlService],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    runtime, _store, legacy = harness
    client = _FakeEnterpriseAiClient(_FENCED_PAYLOAD)
    legacy._enterprise_ai_client = client
    monkeypatch.setattr(ontology_build_module, "_ONTOLOGY_BUILD_LLM_CONTEXT_MAX_CHARS", 6_500)

    contents: dict[str, bytes] = {}
    source_documents: list[OntologySourceDocument] = []
    markers: list[str] = []
    for source_index in range(6):
        lines: list[str] = []
        for chunk_index in range(3):
            marker = f"FULL_SOURCE_MARKER_{source_index}_{chunk_index}"
            markers.append(marker)
            lines.append(f"{marker} " + ("受注は顧客に紐づく。" * 80))
        content = "\n".join(lines).encode()
        source_id = f"ontology_source_{source_index}"
        contents[source_id] = content
        source_documents.append(_source_document(source_id, f"source-{source_index}.md", content))

    service = OntologyBuildService(
        runtime,
        source_storage=_FakeOntologySourceStorage(contents),  # type: ignore[arg-type]
    )
    queued = service.start(
        "sales",
        run_schema_naming=False,
        run_qa_extraction=False,
        source_documents=source_documents,
    )
    finished = _wait_for_job(service, queued.id)

    assert finished.status == OntologyBuildStatus.SUCCEEDED
    text_contexts = [
        context for context in client.contexts if "business_text_chunks" in context
    ]
    assert len(text_contexts) > 1
    assert all(
        ontology_build_module._llm_call_chars(client.calls[index], context)
        <= ontology_build_module._ONTOLOGY_BUILD_LLM_CONTEXT_MAX_CHARS
        for index, context in enumerate(client.contexts)
    )
    combined_context = "\n".join(text_contexts)
    for marker in markers:
        assert marker in combined_context
    processed_chunks = sum(
        len(json.loads(context)["business_text_chunks"]) for context in text_contexts
    )
    assert processed_chunks == len(markers)
    assert any("chunk batch" in event.message_ja for event in finished.events)


def test_build_job_batches_more_than_two_hundred_qa_pairs(
    harness: tuple[OntologyApiRuntime, InMemoryOntologyStore, _FakeLegacyNl2SqlService],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    runtime, _store, legacy = harness
    client = _FakeEnterpriseAiClient(_FENCED_PAYLOAD)
    legacy._enterprise_ai_client = client
    monkeypatch.setattr(ontology_build_module, "_ONTOLOGY_BUILD_LLM_CONTEXT_MAX_CHARS", 8_500)
    qa_pairs = [
        QaPair(
            question=f"顧客別売上 {index}",
            sql=(
                "SELECT C.NAME, SUM(O.AMOUNT) FROM APP.ORDERS O "
                "JOIN APP.CUSTOMERS C ON O.CUSTOMER_ID = C.ID "
                f"WHERE O.ID >= {index} GROUP BY C.NAME"
            ),
        )
        for index in range(205)
    ]
    service = OntologyBuildService(runtime)

    queued = service.start(
        "sales",
        qa_pairs=qa_pairs,
        run_schema_naming=False,
        run_text_extraction=False,
    )
    finished = _wait_for_job(service, queued.id)

    assert finished.status == OntologyBuildStatus.SUCCEEDED
    qa_contexts = [context for context in client.contexts if '"qa_pairs"' in context]
    assert len(qa_contexts) > 1
    sent_pairs = [
        pair
        for context in qa_contexts
        for pair in json.loads(context)["qa_pairs"]
    ]
    assert len(sent_pairs) == 205
    assert sent_pairs[-1]["question"] == "顧客別売上 204"
    assert any("Q/A batch" in event.message_ja for event in finished.events)


def test_build_job_fails_when_pdf_page_requires_ocr_without_image_runner(
    harness: tuple[OntologyApiRuntime, InMemoryOntologyStore, _FakeLegacyNl2SqlService],
) -> None:
    from pypdf import PdfWriter

    runtime, _store, legacy = harness
    legacy._enterprise_ai_client = _FakeEnterpriseAiClient(_FENCED_PAYLOAD)
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    buffer = io.BytesIO()
    writer.write(buffer)
    content = buffer.getvalue()
    source_id = "ontology_source_blank_pdf"
    source = _source_document(
        source_id,
        "scan.pdf",
        content,
        media_type="application/pdf",
    )
    service = OntologyBuildService(
        runtime,
        source_storage=_FakeOntologySourceStorage({source_id: content}),  # type: ignore[arg-type]
    )

    queued = service.start(
        "sales",
        run_schema_naming=False,
        run_qa_extraction=False,
        source_documents=[source],
    )
    finished = _wait_for_job(service, queued.id)

    assert finished.status == OntologyBuildStatus.FAILED
    assert "page:1" in finished.error_message_ja
    assert "OCR 設定" in finished.error_message_ja
    source_progress = finished.sources[0]
    assert source_progress.status == OntologySourceStatus.FAILED
    assert "page:1" in source_progress.error_message_ja


def test_markdown_draft_merges_profile_view_overrides_rules_and_enums() -> None:
    provenance = OntologyProvenance(source_kind=OntologySourceKind.MANUAL, source_id="test")
    revision = OntologyRevision(
        id="ontology_revision_test",
        version=4,
        status=OntologyRevisionStatus.PUBLISHED,
        schema_fingerprint="fp-test",
        etag="rev-etag-test",
    )
    orders_ref = PhysicalObjectRef(
        node_id="table-orders",
        owner="APP",
        object_name="ORDERS",
        object_type="table",
    )
    customers_ref = PhysicalObjectRef(
        node_id="table-customers",
        owner="APP",
        object_name="CUSTOMERS",
        object_type="table",
    )
    status_column = PhysicalColumnRef(
        node_id="column-order-status",
        owner="APP",
        object_name="ORDERS",
        column_name="STATUS",
        ordinal=3,
    )
    customer_id_column = PhysicalColumnRef(
        node_id="column-order-customer",
        owner="APP",
        object_name="ORDERS",
        column_name="CUSTOMER_ID",
        ordinal=2,
    )
    customer_pk_column = PhysicalColumnRef(
        node_id="column-customer-id",
        owner="APP",
        object_name="CUSTOMERS",
        column_name="ID",
        ordinal=1,
    )
    orders = OntologyNode(
        id="table-orders",
        revision_id=revision.id,
        kind=OntologyNodeKind.TABLE,
        technical_name="APP.ORDERS",
        business_name_ja="受注",
        description_ja="受注トランザクション",
        physical_mappings=[PhysicalMapping(object_ref=orders_ref)],
        provenance=provenance,
        review_status=OntologyReviewStatus.APPROVED,
    )
    customers = OntologyNode(
        id="table-customers",
        revision_id=revision.id,
        kind=OntologyNodeKind.TABLE,
        technical_name="APP.CUSTOMERS",
        business_name_ja="顧客",
        description_ja="顧客マスタ",
        physical_mappings=[PhysicalMapping(object_ref=customers_ref)],
        provenance=provenance,
        review_status=OntologyReviewStatus.APPROVED,
    )
    status_property = OntologyNode(
        id="property-order-status",
        revision_id=revision.id,
        kind=OntologyNodeKind.PROPERTY,
        technical_name="APP.ORDERS.STATUS",
        business_name_ja="受注状態",
        description_ja="受注の状態",
        physical_mappings=[PhysicalMapping(object_ref=orders_ref, column_refs=[status_column])],
        provenance=provenance,
        review_status=OntologyReviewStatus.APPROVED,
    )
    rule = OntologyNode(
        id="business-rule-order-status",
        revision_id=revision.id,
        kind=OntologyNodeKind.BUSINESS_RULE,
        technical_name="order_status_required",
        business_name_ja="受注状態必須",
        description_ja="受注状態を空にしない。",
        provenance=provenance,
        review_status=OntologyReviewStatus.APPROVED,
        business_rule_definition=BusinessRuleDefinition(
            rule_kind=BusinessRuleKind.VALIDATION,
            statement_ja="受注状態を必須にします。",
            applies_to_node_ids=["table-orders"],
            severity=BusinessRuleSeverity.WARNING,
            execution_mode=BusinessRuleExecutionMode.DOCUMENTATION,
        ),
    )
    enum_value = OntologyNode(
        id="enum-order-confirmed",
        revision_id=revision.id,
        kind=OntologyNodeKind.ENUM_VALUE,
        technical_name="CONFIRMED",
        business_name_ja="確定済み",
        description_ja="確定した受注。",
        provenance=provenance,
        review_status=OntologyReviewStatus.APPROVED,
        enum_value_definition=EnumValueDefinition(
            code="CONFIRMED",
            label_ja="確定済み",
            aliases=["確定"],
            physical_literal="C",
            property_node_id="property-order-status",
        ),
    )
    relationship = OntologyEdge(
        id="edge-order-customer",
        revision_id=revision.id,
        kind=OntologyEdgeKind.FOREIGN_KEY,
        source_node_id="table-orders",
        target_node_id="table-customers",
        relationship_name_ja="受注の顧客",
        description_ja="Oracle FK_ORDERS_CUSTOMER",
        cardinality=RelationshipCardinality.MANY_TO_ONE,
        join_conditions=[
            JoinCondition(left=customer_id_column, right=customer_pk_column, operator="=")
        ],
        provenance=provenance,
        review_status=OntologyReviewStatus.APPROVED,
    )
    ontology = SchemaOntology(
        revision=revision,
        nodes=[orders, customers, status_property, rule, enum_value],
        edges=[relationship],
    )
    view = ProfileOntologyView(
        id="profile-view-test",
        profile_id="sales",
        ontology_revision_id=revision.id,
        etag="view-etag-test",
        node_ids=[node.id for node in ontology.nodes],
        edge_ids=[relationship.id],
        physical_objects=[orders_ref, customers_ref],
        table_usages_ja={"table-orders": "受注分析の主表"},
        allowed_path_ids=[relationship.id],
        draft_node_overrides=[{"node_id": "table-orders", "business_name_ja": "受注明細"}],
        draft_edge_overrides=[
            {
                "edge_id": relationship.id,
                "cardinality": "one_to_many",
                "allowed_path": True,
            }
        ],
    )

    markdown = render_ontology_build_markdown(
        profile_id="sales",
        schema_context=json.dumps({"objects": [], "relationships": []}),
        drafts=[],
        warnings=["APP.MISSING を公開 Ontology に解決できません。"],
        source_count=0,
        qa_pair_count=0,
        business_text_present=False,
        ontology=ontology,
        profile_view=view,
    )

    assert "## Physical Objects" in markdown
    assert "`APP.ORDERS` (table)" in markdown
    assert "business_name: 受注明細" in markdown
    assert "usage: 受注分析の主表" in markdown
    assert "受注明細 (`APP.ORDERS`)" in markdown
    assert "## Relationships / Join" in markdown
    assert "cardinality: one_to_many" in markdown
    assert "allowed_path: true" in markdown
    assert "`APP.ORDERS.CUSTOMER_ID = APP.CUSTOMERS.ID`" in markdown
    assert "## Business Rules / Enum Values" in markdown
    assert "statement: 受注状態を必須にします。" in markdown
    assert "applies_to: 受注明細" in markdown
    assert "code: CONFIRMED" in markdown
    assert "literal: `C`" in markdown
    assert "property: 受注状態" in markdown
    assert "APP.MISSING" in markdown


def test_build_job_uses_ontology_schema_fingerprint_for_registration_conflict(
    harness: tuple[OntologyApiRuntime, InMemoryOntologyStore, _FakeLegacyNl2SqlService],
) -> None:
    """catalog head の fingerprint 口径差だけでは schema drift 扱いにしない。"""

    runtime, _store, legacy = harness
    legacy._enterprise_ai_client = _FakeEnterpriseAiClient(_FENCED_PAYLOAD)
    legacy.catalog = legacy.catalog.model_copy(
        update={"schema_fingerprint": "oracle-style-different-hash"},
        deep=True,
    )
    assert legacy.catalog.schema_fingerprint != catalog_schema_fingerprint(legacy.catalog)
    service = OntologyBuildService(runtime)

    job = service.start("sales", business_text="受注は顧客に紐づく。")
    finished = _wait_for_job(service, job.id)

    assert finished.status == OntologyBuildStatus.SUCCEEDED
    assert finished.proposal_ids == []
    assert finished.draft_revision_id
    assert not (
        finished.error_message_ja
        and "AI 構築中に DB schema catalog が更新されました" in finished.error_message_ja
    )
    assert runtime.ontology_markdown_state("sales").draft_revision is not None
    assert runtime.list_profile_proposals("sales") == []


def test_convert_extraction_warns_on_unknown_cardinality(
    harness: tuple[OntologyApiRuntime, InMemoryOntologyStore, _FakeLegacyNl2SqlService],
) -> None:
    """cardinality 未確定(unknown)の関係候補は提案化しつつ warning を付ける。"""

    from app.features.nl2sql.ontology_build import convert_extraction_to_proposals
    from app.features.nl2sql.ontology_models import OntologyBuildExtraction

    runtime, _store, _legacy = harness
    view, ontology = runtime.profile_view("sales")
    extraction = OntologyBuildExtraction.model_validate(
        {
            "relationships": [
                {
                    "source_object": "APP.ORDERS",
                    "target_object": "APP.CUSTOMERS",
                    "relationship_name_ja": "顧客を参照",
                    "cardinality": "unknown",
                    "join_conditions": [
                        {"left": "APP.ORDERS.CUSTOMER_ID", "right": "APP.CUSTOMERS.ID"}
                    ],
                }
            ]
        }
    )
    drafts, warnings = convert_extraction_to_proposals(
        extraction,
        ontology=ontology,
        view=view,
        job_id="job-1",
        inferred_by="test",
    )
    assert any(draft.kind == OntologyProposalKind.RELATIONSHIP for draft in drafts)
    assert any("cardinality が未確定" in warning for warning in warnings)


def test_extraction_prompt_contains_playground_rules() -> None:
    """Playground 由来の抽出ルール(名詞/動詞・cardinality 必須・主識別子)を明文化する。"""

    from app.features.nl2sql.ontology_build import _EXTRACTION_SYSTEM_PROMPT

    assert "名詞をエンティティ候補" in _EXTRACTION_SYSTEM_PROMPT
    assert "one_to_one / one_to_many / many_to_one / many_to_many" in _EXTRACTION_SYSTEM_PROMPT
    assert "主識別子" in _EXTRACTION_SYSTEM_PROMPT


def test_build_job_fails_gracefully_without_enterprise_ai(
    harness: tuple[OntologyApiRuntime, InMemoryOntologyStore, _FakeLegacyNl2SqlService],
) -> None:
    runtime, _store, _legacy = harness
    service = OntologyBuildService(runtime)
    job = service.start("sales", business_text="テスト")
    finished = _wait_for_job(service, job.id)

    assert finished.status == OntologyBuildStatus.FAILED
    assert "Enterprise AI" in finished.error_message_ja
    assert finished.proposal_ids == []
    assert all(step.status == OntologyBuildStepStatus.SKIPPED for step in finished.steps)


def test_build_job_survives_invalid_llm_json(
    harness: tuple[OntologyApiRuntime, InMemoryOntologyStore, _FakeLegacyNl2SqlService],
) -> None:
    runtime, _store, legacy = harness
    legacy._enterprise_ai_client = _FakeEnterpriseAiClient("これは JSON ではありません")
    service = OntologyBuildService(runtime)
    job = service.start("sales", business_text="テスト", run_schema_naming=False)
    finished = _wait_for_job(service, job.id)

    text_step = next(
        step for step in finished.steps if step.name == OntologyBuildStepName.TEXT_EXTRACTION
    )
    assert text_step.status == OntologyBuildStepStatus.FAILED
    assert any("抽出に失敗" in warning for warning in finished.warnings_ja)
    assert finished.proposal_ids == []


def test_build_job_fails_gracefully_when_markdown_draft_save_times_out(
    harness: tuple[OntologyApiRuntime, InMemoryOntologyStore, _FakeLegacyNl2SqlService],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    runtime, _store, legacy = harness
    legacy._enterprise_ai_client = _FakeEnterpriseAiClient(_FENCED_PAYLOAD)

    def fail_create_build_markdown_draft(**kwargs: Any) -> Any:
        on_progress = kwargs.get("on_progress")
        if callable(on_progress):
            on_progress("Draft revision を保存しています…")
        raise TimeoutError("Oracle round-trip timeout")

    monkeypatch.setattr(
        runtime,
        "create_build_markdown_draft",
        fail_create_build_markdown_draft,
    )
    service = OntologyBuildService(runtime)
    job = service.start("sales", business_text="受注は顧客に紐づく。", run_schema_naming=False)
    finished = _wait_for_job(service, job.id)

    assert finished.status == OntologyBuildStatus.FAILED
    registration = next(
        step
        for step in finished.steps
        if step.name == OntologyBuildStepName.PROPOSAL_REGISTRATION
    )
    assert registration.status == OntologyBuildStepStatus.FAILED
    assert registration.detail_ja == "Markdown Draft の保存に失敗しました。"
    assert "Oracle round-trip timeout" in finished.error_message_ja
    assert any("Draft revision を保存しています" in event.message_ja for event in finished.events)


def test_build_job_fails_gracefully_when_final_status_save_times_out(
    harness: tuple[OntologyApiRuntime, InMemoryOntologyStore, _FakeLegacyNl2SqlService],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    runtime, store, legacy = harness
    legacy._enterprise_ai_client = _FakeEnterpriseAiClient(_FENCED_PAYLOAD)
    original_save_document = store.save_document

    def fail_final_status_save(
        collection: str,
        document: Any,
        *,
        expected_etag: str | None = None,
    ) -> dict[str, Any]:
        if collection == "jobs" and document.get("status") == OntologyBuildStatus.SUCCEEDED.value:
            raise TimeoutError("Oracle final job timeout")
        return original_save_document(collection, document, expected_etag=expected_etag)

    monkeypatch.setattr(store, "save_document", fail_final_status_save)
    service = OntologyBuildService(runtime)
    job = service.start("sales", business_text="受注は顧客に紐づく。", run_schema_naming=False)
    finished = _wait_for_job(service, job.id)

    assert finished.status == OntologyBuildStatus.FAILED
    registration = next(
        step
        for step in finished.steps
        if step.name == OntologyBuildStepName.PROPOSAL_REGISTRATION
    )
    assert registration.status == OntologyBuildStepStatus.FAILED
    assert registration.detail_ja == "構築 job の完了状態の保存に失敗しました。"
    assert "Oracle final job timeout" in finished.error_message_ja
    assert runtime.ontology_markdown_state("sales").draft_revision is not None


def test_get_build_job_normalizes_succeeded_markdown_job_with_running_final_step(
    harness: tuple[OntologyApiRuntime, InMemoryOntologyStore, _FakeLegacyNl2SqlService],
) -> None:
    runtime, store, _legacy = harness
    stuck = OntologyBuildJob(
        id="ontology_build_stuck_final_step",
        profile_id="sales",
        status=OntologyBuildStatus.SUCCEEDED,
        steps=[
            OntologyBuildStep(
                name=OntologyBuildStepName.SCHEMA_CONTEXT,
                status=OntologyBuildStepStatus.SUCCEEDED,
            ),
            OntologyBuildStep(
                name=OntologyBuildStepName.PROPOSAL_REGISTRATION,
                status=OntologyBuildStepStatus.RUNNING,
                detail_ja="構築 job の完了状態を保存しています…",
            ),
        ],
        draft_revision_id="ontology_revision_draft_stuck",
        draft_etag="markdown-etag-stuck",
        markdown_output="# Ontology Draft\n",
    )
    stuck.started_at = stuck.created_at
    stuck.finished_at = stuck.created_at
    for step in stuck.steps:
        step.started_at = stuck.created_at
    store.save_document(
        "jobs",
        {
            "job_id": stuck.id,
            "job_type": "build",
            "profile_id": stuck.profile_id,
            "status": stuck.status.value,
            "payload": stuck.model_dump(mode="json"),
            "input_payload": {},
        },
    )
    service = OntologyBuildService(runtime)

    normalized = service.get(stuck.id)

    assert normalized is not None
    assert normalized.status == OntologyBuildStatus.SUCCEEDED
    registration = next(
        step
        for step in normalized.steps
        if step.name == OntologyBuildStepName.PROPOSAL_REGISTRATION
    )
    assert registration.status == OntologyBuildStepStatus.SUCCEEDED
    assert registration.detail_ja == "Markdown Draft を生成しました。"
    assert registration.finished_at == normalized.finished_at
    raw_document = store.get_document("jobs", {"job_id": stuck.id})
    assert raw_document is not None
    raw_job = OntologyBuildJob.model_validate(raw_document["payload"])
    raw_registration = next(
        step
        for step in raw_job.steps
        if step.name == OntologyBuildStepName.PROPOSAL_REGISTRATION
    )
    assert raw_registration.status == OntologyBuildStepStatus.RUNNING


def test_list_profile_jobs_returns_newest_first_with_limit(
    harness: tuple[OntologyApiRuntime, InMemoryOntologyStore, _FakeLegacyNl2SqlService],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    runtime, _store, _legacy = harness
    monkeypatch.setattr(get_settings(), "nl2sql_ontology_worker_mode", "external")
    service = OntologyBuildService(runtime)
    first = service.start("sales", business_text="一件目")
    second = service.start("sales", business_text="二件目")
    third = service.start("sales", business_text="三件目")

    jobs = service.list_profile_jobs("sales", limit=2)
    assert [job.id for job in jobs] == [third.id, second.id]
    all_jobs = service.list_profile_jobs("sales", limit=10)
    assert [job.id for job in all_jobs] == [third.id, second.id, first.id]


def test_cancel_single_job_is_idempotent_and_blocks_worker(
    harness: tuple[OntologyApiRuntime, InMemoryOntologyStore, _FakeLegacyNl2SqlService],
    monkeypatch: pytest.MonkeyPatch,
) -> None:

    runtime, store, legacy = harness
    legacy._enterprise_ai_client = _FakeEnterpriseAiClient(_FENCED_PAYLOAD)
    monkeypatch.setattr(get_settings(), "nl2sql_ontology_worker_mode", "external")
    service = OntologyBuildService(runtime)
    job = service.start("sales", business_text="販売業務")

    cancelled = service.cancel(job.id)
    assert cancelled.status == OntologyBuildStatus.CANCELLED
    assert any("利用者の操作" in event.message_ja for event in cancelled.events)
    # CANCELLED への再キャンセルは no-op 成功
    assert service.cancel(job.id).status == OntologyBuildStatus.CANCELLED
    # worker が後から実行しても cancelled のまま(proposal 生成なし)
    result = service.run_persisted(job.id)
    assert result.status == OntologyBuildStatus.CANCELLED
    assert store.list_documents("proposals", {"profile_id": "sales"}) == []


def test_cancel_finished_or_missing_job_raises(
    harness: tuple[OntologyApiRuntime, InMemoryOntologyStore, _FakeLegacyNl2SqlService],
) -> None:
    from app.features.nl2sql.ontology_service import (
        OntologyNotFoundError,
        OntologyStateConflictError,
    )

    runtime, _store, legacy = harness
    legacy._enterprise_ai_client = _FakeEnterpriseAiClient(_FENCED_PAYLOAD)
    service = OntologyBuildService(runtime)
    job = service.start("sales", business_text="販売業務")
    _wait_for_job(service, job.id)

    with pytest.raises(OntologyStateConflictError):
        service.cancel(job.id)
    with pytest.raises(OntologyNotFoundError):
        service.cancel("ontology_build_missing")


def test_retry_reuses_persisted_inputs_and_toggles(
    harness: tuple[OntologyApiRuntime, InMemoryOntologyStore, _FakeLegacyNl2SqlService],
) -> None:
    runtime, _store, legacy = harness
    # 1 回目は Enterprise AI 未設定で失敗させる
    service = OntologyBuildService(runtime)
    failed = service.start(
        "sales",
        business_text="受注は顧客に紐づく。",
        run_schema_naming=False,
    )
    finished = _wait_for_job(service, failed.id)
    assert finished.status == OntologyBuildStatus.FAILED

    # 設定を直して retry → 保存済み入力(業務テキスト・トグル)で新規 job が成功する
    legacy._enterprise_ai_client = _FakeEnterpriseAiClient(_FENCED_PAYLOAD)
    retried = service.retry(failed.id)
    assert retried.id != failed.id
    result = _wait_for_job(service, retried.id)
    assert result.status == OntologyBuildStatus.SUCCEEDED
    step_names = {step.name for step in result.steps}
    assert OntologyBuildStepName.SCHEMA_NAMING not in step_names
    assert OntologyBuildStepName.TEXT_EXTRACTION in step_names
    # 二度押しは idempotency で同じ再実行 job に合流する
    assert service.retry(failed.id).id == retried.id


def test_retry_rejects_non_terminal_and_missing_jobs(
    harness: tuple[OntologyApiRuntime, InMemoryOntologyStore, _FakeLegacyNl2SqlService],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from app.features.nl2sql.ontology_service import (
        OntologyNotFoundError,
        OntologyStateConflictError,
    )

    runtime, _store, legacy = harness
    legacy._enterprise_ai_client = _FakeEnterpriseAiClient(_FENCED_PAYLOAD)
    monkeypatch.setattr(get_settings(), "nl2sql_ontology_worker_mode", "external")
    service = OntologyBuildService(runtime)
    queued = service.start("sales", business_text="販売業務")

    with pytest.raises(OntologyStateConflictError):
        service.retry(queued.id)
    with pytest.raises(OntologyNotFoundError):
        service.retry("ontology_build_missing")


def test_get_prefers_store_in_external_worker_mode(
    harness: tuple[OntologyApiRuntime, InMemoryOntologyStore, _FakeLegacyNl2SqlService],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """API プロセスの古い in-memory コピーではなく worker が書いた store 進捗を返す。"""

    runtime, store, _legacy = harness
    monkeypatch.setattr(get_settings(), "nl2sql_ontology_worker_mode", "external")
    service = OntologyBuildService(runtime)
    job = service.start("sales", business_text="販売業務")

    # worker 側の進捗書込を模擬(store の status を running へ)
    document = store.get_document("jobs", {"job_id": job.id})
    assert document is not None
    payload = dict(document["payload"])
    payload["status"] = OntologyBuildStatus.RUNNING.value
    store.save_document(
        "jobs",
        {
            **{k: v for k, v in document.items() if k not in {"etag", "created_at", "updated_at"}},
            "status": OntologyBuildStatus.RUNNING.value,
            "payload": payload,
        },
        expected_etag=str(document["etag"]),
    )

    refreshed = service.get(job.id)
    assert refreshed is not None
    assert refreshed.status == OntologyBuildStatus.RUNNING


def test_start_rejects_unknown_profile(
    harness: tuple[OntologyApiRuntime, InMemoryOntologyStore, _FakeLegacyNl2SqlService],
) -> None:
    runtime, _store, _legacy = harness
    service = OntologyBuildService(runtime)
    with pytest.raises(OntologyNotFoundError):
        service.start("unknown-profile")


# --- accept → draft → publish → 再起動復元 -----------------------------------------------------


def test_accept_applies_upserts_accumulates_and_publishes(
    harness: tuple[OntologyApiRuntime, InMemoryOntologyStore, _FakeLegacyNl2SqlService],
) -> None:
    runtime, store, legacy = harness
    proposals = _seed_build_proposals(runtime)
    relationship = next(
        proposal for proposal in proposals if proposal.kind == OntologyProposalKind.RELATIONSHIP
    )
    mapping = next(
        proposal for proposal in proposals if proposal.kind == OntologyProposalKind.MAPPING
    )
    metric = next(
        proposal
        for proposal in proposals
        if proposal.kind == OntologyProposalKind.METRIC_DEFINITION
    )

    # 関係提案の accept → 業務ノード + 承認済み関係が draft に入る
    review = runtime.accept_proposal(relationship.id)
    assert review.draft is not None
    draft = review.draft
    business_edges = [
        edge for edge in draft.edges if edge.kind == OntologyEdgeKind.BUSINESS_RELATIONSHIP
    ]
    assert len(business_edges) == 1
    assert business_edges[0].review_status == OntologyReviewStatus.APPROVED
    assert business_edges[0].join_conditions[0].left.column_name == "CUSTOMER_ID"

    # 続けて命名提案を accept → 直前の draft に積み上がり、合成 endpoint が上書きされる
    review2 = runtime.accept_proposal(mapping.id)
    assert review2.draft is not None
    draft2 = review2.draft
    entity_nodes = [node for node in draft2.nodes if node.kind == OntologyNodeKind.BUSINESS_ENTITY]
    orders_entity = next(node for node in entity_nodes if node.technical_name == "APP.ORDERS")
    assert orders_entity.business_name_ja == "受注"
    assert "オーダー" in orders_entity.aliases
    assert not orders_entity.metadata.get("synthetic_endpoint")
    # 関係提案の内容も残っている(最新 revision へ積み上げ)
    assert any(edge.kind == OntologyEdgeKind.BUSINESS_RELATIONSHIP for edge in draft2.edges)

    # 指標提案も accept → metric_definition が node metadata に入る
    review3 = runtime.accept_proposal(metric.id)
    assert review3.draft is not None
    metric_nodes = [node for node in review3.draft.nodes if node.kind == OntologyNodeKind.METRIC]
    assert len(metric_nodes) == 1
    assert metric_nodes[0].metadata["metric_definition"]["expression_sql"] == "SUM(AMOUNT)"

    # publish は全業務要素 APPROVED のため成功する
    published = runtime.publish_ontology_revision(
        review3.draft.revision.id,
        OntologyPublishRequest(etag=review3.draft.revision.etag),
    )
    assert published.revision.status.value == "published"

    # 再起動(同じ store)でも ontology_build 由来 proposal の復元が落ちない
    restarted = OntologyApiRuntime(legacy_service=legacy, store=store)
    restored = restarted.list_profile_proposals("sales")
    assert {proposal.id for proposal in restored} >= {relationship.id, mapping.id, metric.id}


def test_build_job_fails_fast_when_db_profile_scope_is_empty(
    harness: tuple[OntologyApiRuntime, InMemoryOntologyStore, _FakeLegacyNl2SqlService],
) -> None:
    """DB catalog で profile scope が空のときは LLM を呼ばずに明確なエラーで失敗する。"""

    runtime, _store, legacy = harness
    client = _FakeEnterpriseAiClient(_FENCED_PAYLOAD)
    legacy._enterprise_ai_client = client
    # 空カタログ → profile の対象オブジェクトが DB schema catalog に解決できない状態
    legacy.catalog = SchemaCatalog(refreshed_at="2026-07-12T00:00:00Z", tables=[])

    service = OntologyBuildService(runtime)
    job = service.start("sales", business_text="受注は顧客に紐づく。")
    finished = _wait_for_job(service, job.id)

    assert finished.status == OntologyBuildStatus.FAILED
    assert "DB schema catalog" in finished.error_message_ja
    assert "DB 構造を再取得" in finished.error_message_ja
    assert client.calls == []
    schema_step = next(
        step for step in finished.steps if step.name == OntologyBuildStepName.SCHEMA_CONTEXT
    )
    assert schema_step.status == OntologyBuildStepStatus.FAILED
    assert schema_step.detail_ja == "profile 範囲に DB 表・ビューがありません。"
    assert all(
        step.status == OntologyBuildStepStatus.SKIPPED
        for step in finished.steps
        if step.name != OntologyBuildStepName.SCHEMA_CONTEXT
    )


def test_build_job_uses_latest_schema_when_published_profile_scope_is_stale(
    harness: tuple[OntologyApiRuntime, InMemoryOntologyStore, _FakeLegacyNl2SqlService],
) -> None:
    """published Ontology が古くても AI input は DB catalog から作る。"""

    runtime, _store, legacy = harness
    _view, published_base = runtime.profile_view("sales")
    table_node = next(
        node
        for node in published_base.nodes
        if node.kind == OntologyNodeKind.TABLE and node.technical_name == "APP.ORDERS"
    )
    business_node = OntologyNode(
        id="business_orders_pinned_for_build",
        revision_id=published_base.revision.id,
        kind=OntologyNodeKind.BUSINESS_ENTITY,
        technical_name="APP.ORDERS",
        business_name_ja="受注",
        physical_mappings=[PhysicalMapping(object_ref=table_node.physical_mappings[0].object_ref)],
        provenance=OntologyProvenance(source_kind=OntologySourceKind.MANUAL),
        review_status=OntologyReviewStatus.APPROVED,
    )
    draft = runtime.create_ontology_draft(
        published_base.revision.id,
        OntologyDraftRequest(
            base_etag=published_base.revision.etag,
            node_upserts=[business_node],
        ),
    )
    published_with_business = runtime.publish_ontology_revision(
        draft.revision.id,
        OntologyPublishRequest(etag=draft.revision.etag),
    )

    invoice_table = SchemaTable(
        table_name="INVOICES",
        logical_name="請求",
        owner="APP",
        columns=[
            SchemaColumn(column_name="ID", logical_name="請求 ID", data_type="NUMBER"),
            SchemaColumn(column_name="AMOUNT", logical_name="請求金額", data_type="NUMBER"),
        ],
    )
    legacy.catalog = legacy.catalog.model_copy(
        update={"tables": [*legacy.catalog.tables, invoice_table]}
    )
    legacy.profile = legacy.profile.model_copy(update={"allowed_tables": ["APP.INVOICES"]})

    pinned_view, pinned_ontology = runtime.profile_view("sales")
    assert pinned_ontology.revision.id == published_with_business.revision.id
    pinned_node_ids = set(pinned_view.node_ids)
    assert not any(
        node.kind in {OntologyNodeKind.TABLE, OntologyNodeKind.VIEW} and node.id in pinned_node_ids
        for node in pinned_ontology.nodes
    )

    invoice_extraction = {
        "entities": [
            {
                "object_name": "APP.INVOICES",
                "business_name_ja": "請求",
                "description_ja": "顧客への請求を表す。",
                "aliases": ["インボイス"],
                "confidence": 0.9,
            }
        ],
        "relationships": [],
        "metrics": [
            {
                "metric_name_ja": "請求金額合計",
                "expression_sql": "SUM(AMOUNT)",
                "aggregation": "sum",
                "base_columns": ["APP.INVOICES.AMOUNT"],
                "unit": "円",
                "description_ja": "請求金額の合計",
                "evidence_ja": "業務説明",
                "confidence": 0.8,
            }
        ],
        "synonyms": [],
        "warnings_ja": [],
    }
    client = _FakeEnterpriseAiClient(
        "抽出結果:\n" + json.dumps(invoice_extraction, ensure_ascii=False)
    )
    legacy._enterprise_ai_client = client
    service = OntologyBuildService(runtime)
    job = service.start("sales", business_text="請求は顧客への請求金額を管理する。")
    finished = _wait_for_job(service, job.id)

    assert finished.status == OntologyBuildStatus.SUCCEEDED
    assert not any("公開 Ontology" in warning for warning in finished.warnings_ja)
    assert not any("公開 Ontology" in event.message_ja for event in finished.events)
    assert any("DB から profile 範囲" in event.message_ja for event in finished.events)
    schema_context = json.loads(client.contexts[0])
    assert {item["object"] for item in schema_context["objects"]} == {"APP.INVOICES"}
    assert runtime.list_profile_proposals("sales") == []
    state = runtime.ontology_markdown_state("sales")
    assert state.draft_revision is not None
    assert state.draft_revision.id == finished.draft_revision_id
    assert state.draft_revision.schema_fingerprint != (
        published_with_business.revision.schema_fingerprint
    )

    query_view_after, query_ontology_after = runtime.profile_view("sales")
    assert query_ontology_after.revision.id == published_with_business.revision.id
    assert query_view_after.ontology_revision_id == published_with_business.revision.id


def test_build_job_does_not_read_profile_view_for_ai_schema_input(
    harness: tuple[OntologyApiRuntime, InMemoryOntologyStore, _FakeLegacyNl2SqlService],
) -> None:
    """AI input の schema は DB catalog 直読みにし、profile view API へ依存しない。"""

    runtime, _store, legacy = harness
    legacy._enterprise_ai_client = _FakeEnterpriseAiClient(_FENCED_PAYLOAD)

    class _NoProfileViewRuntime:
        def __getattr__(self, name: str) -> Any:
            return getattr(runtime, name)

        def profile_view(self, profile_id: str) -> Any:
            raise AssertionError(f"profile_view must not be used for AI build: {profile_id}")

        def profile_view_for_build(self, profile_id: str) -> Any:
            raise AssertionError(
                f"profile_view_for_build must not be used for AI build: {profile_id}"
            )

    service = OntologyBuildService(_NoProfileViewRuntime())
    job = service.start("sales", business_text="受注は顧客に紐づく。")
    finished = _wait_for_job(service, job.id)

    assert finished.status == OntologyBuildStatus.SUCCEEDED


def test_start_returns_immediately_even_if_build_schema_context_is_slow(
    harness: tuple[OntologyApiRuntime, InMemoryOntologyStore, _FakeLegacyNl2SqlService],
) -> None:
    """start() は重い DB schema context 準備を待たずに job を返す。"""

    import threading as _threading

    runtime, _store, legacy = harness
    legacy._enterprise_ai_client = _FakeEnterpriseAiClient(_FENCED_PAYLOAD)
    release = _threading.Event()
    original_prepare_build_schema_context = runtime.prepare_build_schema_context

    class _SlowRuntime:
        def __getattr__(self, name: str) -> Any:
            return getattr(runtime, name)

        def prepare_build_schema_context(self, profile_id: str) -> Any:
            release.wait(timeout=5)
            return original_prepare_build_schema_context(profile_id)

    service = OntologyBuildService(_SlowRuntime())
    started = time.monotonic()
    job = service.start("sales", business_text="受注は顧客に紐づく。")
    elapsed = time.monotonic() - started

    try:
        # prepare_build_schema_context がブロックしていても POST(start)は即時に返る
        assert elapsed < 1.0
        assert job.status in {OntologyBuildStatus.QUEUED, OntologyBuildStatus.RUNNING}
        assert job.steps[0].name == OntologyBuildStepName.SCHEMA_CONTEXT
        # ポーリングで「スキーマ情報の準備」が実行中と観測できる
        for _ in range(100):
            snapshot = service.get(job.id)
            assert snapshot is not None
            schema_step = snapshot.steps[0]
            if schema_step.status == OntologyBuildStepStatus.RUNNING:
                break
            time.sleep(0.01)
        assert schema_step.status == OntologyBuildStepStatus.RUNNING
    finally:
        release.set()
    finished = _wait_for_job(service, job.id)
    assert finished.status == OntologyBuildStatus.SUCCEEDED


def test_accept_ignores_stale_drafts_from_previous_schema_generations(
    harness: tuple[OntologyApiRuntime, InMemoryOntologyStore, _FakeLegacyNl2SqlService],
) -> None:
    """過去スキーマ世代の draft が store に残っていても accept が 409 にならない。"""

    runtime, _store, legacy = harness
    legacy._enterprise_ai_client = _FakeEnterpriseAiClient(_FENCED_PAYLOAD)
    # 旧スキーマ世代で提案を承認して draft を作る(store に残留する)
    old_proposal = _seed_build_proposals(runtime)[0]
    runtime.accept_proposal(old_proposal.id)

    # スキーマ drift(列追加で fingerprint が変わる)→ 公開世代が変わる
    legacy.catalog = legacy.catalog.model_copy(
        update={
            "tables": [
                legacy.catalog.tables[0].model_copy(
                    update={
                        "columns": [
                            *legacy.catalog.tables[0].columns,
                            SchemaColumn(
                                column_name="CREATED_AT",
                                logical_name="作成日時",
                                data_type="TIMESTAMP",
                            ),
                        ]
                    }
                ),
                *legacy.catalog.tables[1:],
            ]
        },
        deep=True,
    )
    # 新世代で提案を作る
    new_proposals = _seed_build_proposals(runtime)
    assert new_proposals

    # 旧世代 draft(business 定義持ち)が store に残っていても、新提案の accept は成功する
    review = runtime.accept_proposal(new_proposals[0].id)
    assert review.draft is not None
    _view, published_graph = runtime.profile_view("sales")
    assert published_graph is not None
    published_fp = published_graph.revision.schema_fingerprint
    assert review.draft.revision.schema_fingerprint == published_fp


def test_batch_accept_creates_single_draft_for_all_proposals(
    harness: tuple[OntologyApiRuntime, InMemoryOntologyStore, _FakeLegacyNl2SqlService],
) -> None:
    runtime, _store, _legacy = harness
    proposals = _seed_build_proposals(runtime)
    assert len(proposals) >= 3

    accepted, draft = runtime.accept_proposals([proposal.id for proposal in proposals])

    # すべて同じ draft revision に反映され、全提案が accepted になる
    assert {proposal.status.value for proposal in accepted} == {"accepted"}
    assert {proposal.proposal_payload.values["draft_revision_id"] for proposal in accepted} == {
        draft.revision.id
    }
    # 命名提案の業務名が合成 endpoint に上書きされない
    entity_nodes = [node for node in draft.nodes if node.kind == OntologyNodeKind.BUSINESS_ENTITY]
    orders_entity = next(node for node in entity_nodes if node.technical_name == "APP.ORDERS")
    assert orders_entity.business_name_ja == "受注"
    assert any(edge.kind == OntologyEdgeKind.BUSINESS_RELATIONSHIP for edge in draft.edges)
    # publish まで通る
    published = runtime.publish_ontology_revision(
        draft.revision.id, OntologyPublishRequest(etag=draft.revision.etag)
    )
    assert published.revision.status.value == "published"


def test_rerun_replaces_markdown_draft_without_creating_proposals(
    harness: tuple[OntologyApiRuntime, InMemoryOntologyStore, _FakeLegacyNl2SqlService],
) -> None:
    """AI 構築を再実行すると proposal は作らず、最新 Markdown Draft が差し替わる。"""

    runtime, store, legacy = harness
    legacy._enterprise_ai_client = _FakeEnterpriseAiClient(_FENCED_PAYLOAD)
    service = OntologyBuildService(runtime)

    first = _wait_for_job(service, service.start("sales", business_text="受注は顧客に紐づく。").id)
    assert first.proposal_ids == []
    first_state = runtime.ontology_markdown_state("sales")
    assert first_state.draft_revision is not None
    assert first_state.draft_revision.id == first.draft_revision_id

    second = _wait_for_job(service, service.start("sales", business_text="受注は顧客に紐づく。").id)
    assert second.status == OntologyBuildStatus.SUCCEEDED
    assert second.proposal_ids == []
    second_state = runtime.ontology_markdown_state("sales")
    assert second_state.draft_revision is not None
    assert second_state.draft_revision.id == second.draft_revision_id
    assert second_state.draft_revision.id != first_state.draft_revision.id
    assert second_state.draft_markdown == second.markdown_output
    assert runtime.list_profile_proposals("sales") == []

    # 再起動(同じ store)でも proposal は復活せず、最新 Draft を読める。
    restarted = OntologyApiRuntime(legacy_service=legacy, store=store)
    assert restarted.list_profile_proposals("sales") == []
    assert restarted.ontology_markdown_state("sales").draft_revision is not None


def test_start_prunes_oldest_finished_jobs(
    harness: tuple[OntologyApiRuntime, InMemoryOntologyStore, _FakeLegacyNl2SqlService],
) -> None:
    from datetime import timedelta

    from app.features.nl2sql.ontology_build import _MAX_FINISHED_JOBS
    from app.features.nl2sql.ontology_models import OntologyBuildJob, utc_now

    runtime, _store, legacy = harness
    legacy._enterprise_ai_client = _FakeEnterpriseAiClient(_FENCED_PAYLOAD)
    service = OntologyBuildService(runtime)

    # 完了 job を上限 +2 件、実行中 job を 1 件直接注入する(実 job を回すと遅いため)。
    base = utc_now()
    for index in range(_MAX_FINISHED_JOBS + 2):
        job = OntologyBuildJob(
            id=f"ontology_build_old_{index:03d}",
            profile_id="sales",
            status=OntologyBuildStatus.SUCCEEDED,
            finished_at=base + timedelta(seconds=index),
        )
        service._jobs[job.id] = job
    running = OntologyBuildJob(
        id="ontology_build_running",
        profile_id="sales",
        status=OntologyBuildStatus.RUNNING,
    )
    service._jobs[running.id] = running

    started = service.start("sales", business_text="受注は顧客に紐づく。")

    # 最古の完了 2 件だけが prune され、実行中・新規 job は保護される。
    assert service.get("ontology_build_old_000") is None
    assert service.get("ontology_build_old_001") is None
    assert service.get(f"ontology_build_old_{_MAX_FINISHED_JOBS + 1:03d}") is not None
    assert service.get(running.id) is not None
    assert service.get(started.id) is not None
    # worker thread の終了を待って teardown を安定させる。
    _wait_for_job(service, started.id)


def test_external_worker_rehydrates_persisted_build_input(
    harness: tuple[OntologyApiRuntime, InMemoryOntologyStore, _FakeLegacyNl2SqlService],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    runtime, store, legacy = harness
    legacy._enterprise_ai_client = _FakeEnterpriseAiClient(_FENCED_PAYLOAD)
    monkeypatch.setattr(get_settings(), "nl2sql_ontology_worker_mode", "external")
    api_service = OntologyBuildService(runtime)
    queued = api_service.start(
        "sales",
        business_text="受注は顧客に紐づく。",
        qa_pairs=[QaPair(question="顧客別売上", sql=_QA_SQL)],
        idempotency_key="external-build-1",
    )
    assert queued.status == OntologyBuildStatus.QUEUED
    persisted = store.get_job(queued.id)
    assert persisted is not None
    assert persisted["input_payload"]["business_text"] == "受注は顧客に紐づく。"

    worker_service = OntologyBuildService(runtime)
    finished = worker_service.run_persisted(queued.id)
    assert finished.status == OntologyBuildStatus.SUCCEEDED
    assert finished.proposal_ids == []
    assert finished.draft_revision_id
    assert runtime.ontology_markdown_state("sales").draft_revision is not None


def test_start_persists_business_text_and_source_document_refs(
    harness: tuple[OntologyApiRuntime, InMemoryOntologyStore, _FakeLegacyNl2SqlService],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    runtime, store, _legacy = harness
    monkeypatch.setattr(get_settings(), "nl2sql_ontology_worker_mode", "external")
    service = OntologyBuildService(runtime)
    source_documents = [
        OntologySourceDocument(
            id="ontology_source_rules",
            profile_id="sales",
            filename="rules.md",
            media_type="text/markdown",
            size_bytes=12,
            sha256="0" * 64,
            storage_uri="/tmp/rules.md",
        ),
        OntologySourceDocument(
            id="ontology_source_qa",
            profile_id="sales",
            filename="qa_cases.csv",
            media_type="text/csv",
            size_bytes=48,
            sha256="1" * 64,
            storage_uri="/tmp/qa_cases.csv",
        ),
    ]

    queued = service.start(
        "sales",
        business_text="受注は顧客に紐づく。",
        run_qa_extraction=True,
        run_text_extraction=True,
        source_documents=source_documents,
        idempotency_key="persist-build-inputs-1",
    )

    assert queued.source_document_ids == ["ontology_source_rules", "ontology_source_qa"]
    assert [source.filename for source in queued.sources] == ["rules.md", "qa_cases.csv"]
    persisted = store.get_job(queued.id)
    assert persisted is not None
    assert persisted["input_payload"]["business_text"] == "受注は顧客に紐づく。"
    assert persisted["payload"]["source_document_ids"] == [
        "ontology_source_rules",
        "ontology_source_qa",
    ]
    stored_sources = store.list_documents("source_documents", {"profile_id": "sales"})
    assert {document["payload"]["filename"] for document in stored_sources} == {
        "rules.md",
        "qa_cases.csv",
    }
