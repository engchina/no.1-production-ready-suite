"""Excel（.xlsx / .xlsm / .xls）の読み方（#1221・#2516・#2526。#2525 で RAG の前処理から移した）。

`read_excel` は原本の bytes と選択肢（`ExcelOptions`）だけで、シートごとに「表頭・元の行番号・
セル範囲・表示の値」を持つ記録（`SheetRecordsDocument`）を返す。製品のサービス・HTTP・DB に
依存しない（RAG の前処理 `excel_to_jsonl` も、チャットの添付も同じ読み方にする）。openpyxl /
xlrd は関数の中で読み、使う側が依存に入れる。読み方は engchina/no.1-rag の前処理
（`utils/office_preprocess_util.py`）に合わせる。

- シートの読み方（`mode`）: `table` は 1 行 1 block。`procedure` は手順書（番号の列と題名の列）
  を、番号と題名のある行から次の手順までの複数行で 1 block にし、詳細の無い題名の行は章にする。
  `form` は帳票（方眼紙）を「ラベル: 値」の項目ごとに 1 block にする（`layout`。#2516）。
  `auto` はまず構造の特徴の点数で帳票かを決め（`form_detection_threshold` 以上なら form）、
  帳票でなければ、番号と題名の列があれば procedure、無ければデータの行の埋まり方（中央値 60%
  以上）で table / procedure を決める。判定の根拠と信頼度はシートの `layout_detection` に残す。
- 1 シートの複数の表（#2516）: 空の行・空の列で分かれた表を領域に分け（`split_tables`）、領域ごとに
  版式を判定して表頭を推定し、別々の表（`table_index`・`region`・`title`）として出す。表頭の行を
  指定したとき（`header_row`）は分けない。
- 変換の本体（`read_excel`）は前処理のサービスに依存せず、原本の bytes と選択肢だけで呼べる。
- 図形・テキストボックスの文字（#2526）: `.xlsx` は drawing の XML を直接読み（`shapes`）、
  アンカーの範囲を 1 つの単位（結合セル）にしてセルの格子に置く。表・手順書・帳票の読み方は
  そのまま使い、記録の `shape_cells` に置いた範囲を残す。`.xls` は読まず、図形の文字があれば
  診断を残す。
- 非表示の行・列（#2526）: 既定は読まない（`include_hidden_cells`）。オートフィルターで隠れた行は
  利用者が隠した行ではないので読む。読まなかった行・列は診断に残す。

- `.xlsx` は openpyxl で読む（結合セルと表示の書式を読むため read_only にしない）。数式は
  キャッシュ値を読み、キャッシュの無い数式のセルは空にせず数式の文字列にして診断を残す。マクロは
  実行しない。
- `.xls` は xlrd で読む（書式・結合を読むため formatting_info を使う）。日付は日付、真偽は
  TRUE/FALSE。
- 表頭は選択肢の `header_row` を優先し、無ければ先頭の行から採点して選ぶ（信頼度と理由を残し、
  低いときは警告）。表頭より上の行（説明など）は preamble として残す。
- 表の下の注記（#3566）: 表の後ろの、セルが 1 つ（横の結合を含む）で注記の印（「※1」など。
  選択肢 `note_markers`）で始まるか長い文の行は、表の行・表頭にせず `notes` に残す。注記の印を
  セルに含む行は、その注記の文を記録（`SheetBlock.notes`）にも持つ。手順書は今までどおり手順の
  続きの行として読む。
- 空行を消す前に元の行番号を持つ。結合セルは、表頭の中では結合の範囲を埋め、データでは同じ行の
  横の結合だけを埋める（縦の結合の値を後の行へ漏らさない。手順の境目を増やさない）。普通の空のセルは
  埋めない。既定は表示されているシートだけを読み、選択肢でシートの指定・除外と列の除外ができる。
- 列の役割（#1281）: 設定値の表の列を、表頭の語（既定値・記入例・現在値・推奨値・設定範囲・
  説明 など。英語の語も）で決め、`SheetColumn.role` に残す。製品は値の列の本文に役割の表示を
  付ける（`column_display_name`）。語が完全に一致すれば信頼度 1.0、表頭の末尾が語なら 0.8。
  0.75 未満・複数の役割に当たる・
  意味が資料ごとに違う語（「設定値」「値」）の列には付けず、診断を残す（推測しない）。表頭の推定の
  信頼度が低いシート・役割の列が 1 つだけのシートでは語の判定を使わない。選択肢の `column_roles`
  （列名か列の記号 → 役割）は判定より優先する。
- 形式判定は magic bytes(ZIP=xlsx / OLE2=xls)優先、失敗時は拡張子フォールバック。
- 未対応・依存欠如・解析失敗・空のときは passthrough(変換せず原本を使う)へ縮退する。
"""

from __future__ import annotations

import datetime as _dt
import io
import re
import unicodedata
from collections.abc import Mapping, Set
from dataclasses import dataclass
from typing import Any, cast

from .grid import Grid, cell_ref, column_letter, is_textual
from .layout import (
    NOTE_MIN_CHARS,
    detect_form_layout,
    form_blocks,
    is_numbering,
    note_marker,
    split_regions,
)
from .models import (
    COLUMN_ROLE_NONE,
    ColumnRole,
    ExcelOptions,
    ExcelRange,
    HeaderDetection,
    LayoutDetection,
    SheetBlock,
    SheetColumn,
    SheetDiagnostic,
    SheetNoteRow,
    SheetPreambleRow,
    SheetRecords,
    SheetRecordsDocument,
    SkippedSheet,
    column_display_name,
    parse_excel_range,
)
from .shapes import SheetShapes, xls_shape_text_counts, xlsx_sheet_shapes

# xlsx は ZIP(PK\x03\x04)、xls は OLE2 複合ドキュメント(D0CF11E0...)。
_XLSX_MAGIC = b"PK\x03\x04"
_XLS_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
# 表頭を探す範囲（先頭の空でない行の数）と、表頭の後の行の一致を見る数。
_HEADER_SCAN_ROWS = 30
_HEADER_LOOKAHEAD_ROWS = 5
# これより低い信頼度の表頭は警告にする（黙って推測しない）。
HEADER_LOW_CONFIDENCE = 0.6
# 手順書の表頭に多い語（表頭の行の採点で加点する。no.1-rag と同じ）。
PROCEDURE_HEADER_HINTS = ("作業項目", "作業内容", "コマンド", "確認ポイント", "確認条件", "備考")
# auto で table にする、データの行の埋まり方（中央値）の下限と、見る行の数。
TABLE_DENSITY_THRESHOLD = 0.6
_DENSITY_SAMPLE_ROWS = 50
# 列の役割（#1281）を決める表頭の語（正規化した形。`_normalize_header` を参照）。
# 語の選び方: 設定値の表で意味がほぼ 1 つに決まる語だけを置く。資料ごとに意味の違う語は
# `COLUMN_ROLE_UNCERTAIN_TERMS` に置き、役割を付けずに診断だけを残す。
COLUMN_ROLE_TERMS: dict[str, tuple[str, ...]] = {
    "default": (
        "既定値",
        "既定",
        "デフォルト",
        "デフォルト値",
        "初期値",
        "初期設定",
        "初期設定値",
        "出荷時設定",
        "出荷時の値",
        "出荷時設定値",
        "標準値",
        "default",
        "defaultvalue",
        "defaults",
        "initialvalue",
        "factorydefault",
    ),
    "example": (
        "例",
        "記入例",
        "入力例",
        "設定例",
        "記述例",
        "記載例",
        "値の例",
        "例示",
        "例示値",
        "サンプル",
        "サンプル値",
        "example",
        "examples",
        "examplevalue",
        "sample",
        "samplevalue",
        "eg",
    ),
    "current": (
        "現在値",
        "現在の値",
        "現在の設定値",
        "現在設定値",
        "現行値",
        "現行設定値",
        "実設定値",
        "本番値",
        "本番設定値",
        "実績値",
        "currentvalue",
        "currentsetting",
        "actualvalue",
        "configuredvalue",
    ),
    "recommended": (
        "推奨値",
        "推奨",
        "推奨設定",
        "推奨設定値",
        "recommended",
        "recommendedvalue",
    ),
    "allowed": (
        "設定範囲",
        "設定可能範囲",
        "設定可能値",
        "設定できる値",
        "値の範囲",
        "有効範囲",
        "許容値",
        "許容範囲",
        "取りうる値",
        "取り得る値",
        "有効値",
        "選択肢",
        "allowedvalues",
        "validvalues",
        "validrange",
    ),
    "definition": (
        "説明",
        "定義",
        "内容",
        "概要",
        "意味",
        "用途",
        "解説",
        "項目説明",
        "パラメータ説明",
        "description",
        "definition",
        "meaning",
    ),
}
# 表頭の一部に含まれるときに見る語の最短の長さ（「タイムアウトの既定値」「Windows 既定値」）。
# 短い語（「例」「既定」「内容」「推奨」）は「事例」「既定外」「作業内容」のように別の意味の語にも
# 含まれるので、表頭全体が一致したときだけ使う。
_COLUMN_ROLE_SUFFIX_MIN_CHARS = 3
# 資料ごとに意味の違う語（「設定値」は現場の値のことも、選べる値の一覧のこともある）。
# 役割を付けない。
COLUMN_ROLE_UNCERTAIN_TERMS = ("設定値", "値", "設定", "value", "values", "setting")
COLUMN_ROLE_EXACT_CONFIDENCE = 1.0
COLUMN_ROLE_SUFFIX_CONFIDENCE = 0.8
# これより低い信頼度の判定は役割にしない（黙って推測しない）。
COLUMN_ROLE_MIN_CONFIDENCE = 0.75
# 役割が決まった列（判定と指定の合計）がこれより少ないシートでは、判定を使わない。設定値の表は
# 説明の列と値の列、または複数の値の列を持つ。1 列だけ当たる表（検体の「サンプル」、作業の
# 「内容」）は設定値の表とは限らない。
COLUMN_ROLE_MIN_DETECTED_COLUMNS = 2
_HEADER_BRACKETS = re.compile(r"[（(［\[【〔<＜][^）)］\]】〕>＞]*[）)］\]】〕>＞]")
_HEADER_NOISE = re.compile(r"[\s_\-・:：.。、,/／※*＊]+")


def openpyxl_available() -> bool:
    try:
        import openpyxl  # noqa: F401
    except Exception:
        return False
    return True


def xlrd_available() -> bool:
    try:
        import xlrd  # noqa: F401
    except Exception:
        return False
    return True


@dataclass(frozen=True)
class ExcelReadResult:
    """Excel を行の記録にした結果（``document`` が None なら ``failure`` が縮退の理由）。"""

    document: SheetRecordsDocument | None
    warnings: tuple[str, ...] = ()
    failure: str | None = None


def _is_xlsx(source_bytes: bytes, extension: str | None) -> bool:
    if source_bytes[:4] == _XLSX_MAGIC:
        return True
    if source_bytes[:8] == _XLS_MAGIC:
        return False
    return (extension or "").strip().lower().lstrip(".") == "xlsx"


def read_excel(
    source_bytes: bytes,
    *,
    extension: str | None = None,
    options: ExcelOptions | Mapping[str, Any] | None = None,
) -> ExcelReadResult:
    """Excel（.xlsx / .xlsm / .xls）の原本を行の記録（`SheetRecordsDocument`）にする。

    前処理のサービスに依存しない（原本の bytes と選択肢だけで呼べる）。形式は magic bytes で決め、
    分からなければ ``extension`` を使う。選択肢が不正なら ``ValidationError`` を送出する。
    """
    if not isinstance(options, ExcelOptions):
        options = ExcelOptions.model_validate(dict(options or {}))
    if not source_bytes:
        return ExcelReadResult(None, failure="excel_empty")
    is_xlsx = _is_xlsx(source_bytes, extension)
    include_hidden = options.include_hidden_cells
    grids, failure = (
        _read_xlsx(source_bytes, include_hidden=include_hidden)
        if is_xlsx
        else _read_xls(source_bytes, include_hidden=include_hidden)
    )
    if grids is None:
        return ExcelReadResult(None, failure=failure or "excel_parse_failed")
    warnings: list[str] = []
    sheets: list[SheetRecords] = []
    skipped: list[SkippedSheet] = []
    requested = [name.strip() for name in options.sheets if name.strip()]
    available = {grid.name for grid in grids}
    for name in requested:
        if name not in available:
            warnings.append(f"excel_sheet_not_found:{name}")
    excluded_sheets = {name.strip().casefold() for name in options.exclude_sheets if name.strip()}
    ranges = options.parsed_ranges()
    for name in sorted({item.sheet for item in ranges if item.sheet is not None}):
        if name not in available:
            warnings.append(f"excel_range_sheet_not_found:{name}")
    for grid in grids:
        if requested and grid.name not in requested:
            skipped.append(SkippedSheet(name=grid.name, reason="not_selected"))
            continue
        if grid.name.strip().casefold() in excluded_sheets:
            skipped.append(SkippedSheet(name=grid.name, reason="excluded"))
            continue
        if grid.hidden and not options.include_hidden_sheets and grid.name not in requested:
            skipped.append(SkippedSheet(name=grid.name, reason="hidden"))
            continue
        sheet_ranges = _ranges_for(grid.name, ranges)
        tables = sheet_tables(_crop(grid, sheet_ranges) if sheet_ranges else grid, options)
        if not tables:
            skipped.append(SkippedSheet(name=grid.name, reason="empty"))
            continue
        diagnostics = [item for sheet in tables for item in sheet.diagnostics]
        if any(item.code == "header_low_confidence" for item in diagnostics):
            warnings.append(f"excel_header_low_confidence:{grid.name}")
        missing = sum(1 for item in diagnostics if item.code == "formula_without_cached_value")
        if missing:
            warnings.append(f"excel_formula_without_cached_value:{grid.name}:{missing}")
        sheets.extend(tables)
    for key in options.column_roles:
        # 指定した列がどのシートにも無ければ警告する（列名の書き違いに気付けるように）。
        if not any(
            key.casefold() in {column.name.casefold(), column.column.casefold()}
            for sheet in sheets
            for column in sheet.columns
        ):
            warnings.append(f"excel_column_role_target_not_found:{key}")
    if not any(sheet.blocks or sheet.preamble for sheet in sheets):
        return ExcelReadResult(None, failure="excel_no_rows")
    document = SheetRecordsDocument(
        source_format="xlsx" if is_xlsx else "xls", sheets=sheets, skipped_sheets=skipped
    )
    return ExcelReadResult(document, warnings=tuple(warnings))


# ---- 読み取り ------------------------------------------------------------------------


def _read_xlsx(
    source_bytes: bytes, *, include_hidden: bool = False
) -> tuple[list[Grid] | None, str | None]:
    try:
        import openpyxl
    except Exception:
        return None, "openpyxl_unavailable"
    try:
        values_book = openpyxl.load_workbook(io.BytesIO(source_bytes), data_only=True)
        formula_book = openpyxl.load_workbook(io.BytesIO(source_bytes), data_only=False)
    except Exception:
        return None, "excel_open_failed"
    grids: list[Grid] = []
    try:
        for worksheet in values_book.worksheets:
            formulas = formula_book[worksheet.title]
            grid = Grid(name=worksheet.title, hidden=worksheet.sheet_state != "visible")
            _read_xlsx_column_widths(worksheet, grid)
            _read_xlsx_row_heights(worksheet, grid)
            _read_xlsx_hidden(worksheet, grid)
            skip = set() if include_hidden else grid.hidden_rows - grid.filter_rows
            skip_columns = set() if include_hidden else grid.hidden_columns
            for row in worksheet.iter_rows():
                for cell in row:
                    if cell.row in skip or cell.column in skip_columns:
                        if cell.value is not None and str(cell.value).strip():
                            grid.hidden_cells.add((cell.row, cell.column))
                        continue
                    formula = formulas.cell(row=cell.row, column=cell.column).value
                    is_formula = isinstance(formula, str) and formula.startswith("=")
                    if cell.value is None:
                        if is_formula:
                            # キャッシュの無い数式は空にせず、数式の文字列にして診断を残す。
                            grid.missing_formula_cells.add((cell.row, cell.column))
                            grid.cells[(cell.row, cell.column)] = str(formula)
                        continue
                    text = _format_value(cell.value, cell.number_format, is_date=cell.is_date)
                    if text:
                        grid.cells[(cell.row, cell.column)] = text
                        if _xlsx_label_hint(cell):
                            grid.label_hints.add((cell.row, cell.column))
            grid.merges = [
                (rng.min_row, rng.min_col, rng.max_row, rng.max_col)
                for rng in worksheet.merged_cells.ranges
            ]
            grids.append(grid)
    except Exception:
        return None, "excel_read_failed"
    finally:
        values_book.close()
        formula_book.close()
    shapes = xlsx_sheet_shapes(source_bytes)
    for grid in grids:
        _place_shapes(grid, shapes.get(grid.name), include_hidden=include_hidden)
    return grids, None


def _read_xlsx_hidden(worksheet: Any, grid: Grid) -> None:
    """非表示の行・列と、オートフィルターの範囲のデータの行（隠れていても読む行）。"""
    grid.hidden_rows = {
        index for index, dimension in worksheet.row_dimensions.items() if dimension.hidden
    }
    for dimension in worksheet.column_dimensions.values():
        if dimension.hidden and dimension.min and dimension.max:
            # 使っていない列をまとめて隠す指定（G:XFD）は、値のある最後の列までにする。
            end = min(dimension.max, worksheet.max_column)
            grid.hidden_columns.update(range(dimension.min, end + 1))
    ref = getattr(worksheet.auto_filter, "ref", None)
    if ref and grid.hidden_rows:
        try:
            parsed = parse_excel_range(ref if ":" in ref else f"{ref}:{ref}")
        except ValueError:
            return
        # 表頭（範囲の先頭の行）の下で隠れた行は、フィルターの条件で隠れた行とみなす。
        grid.filter_rows = {
            row for row in grid.hidden_rows if parsed.min_row < row <= parsed.max_row
        }


def _place_shapes(grid: Grid, shapes: SheetShapes | None, *, include_hidden: bool) -> None:
    """図形の文字をセルの格子に置く（非表示の行・列だけにかかる図形は読まない）。"""
    if shapes is None:
        return
    hidden_rows = set() if include_hidden else grid.hidden_rows - grid.filter_rows
    hidden_columns = set() if include_hidden else grid.hidden_columns
    for shape in shapes.shapes:
        box = (shape.top, shape.left, shape.bottom, shape.right)
        if shape.size is not None:
            # 大きさで置く図形（oneCellAnchor / absoluteAnchor）は、列幅・行の高さで範囲にする。
            box = grid.anchor_box(shape.top, shape.left, *shape.size)
        rows = set(range(box[0], box[2] + 1))
        columns = set(range(box[1], box[3] + 1))
        if rows <= hidden_rows or columns <= hidden_columns:
            continue
        if grid.place_shape_text(box, shape.text) is None:
            # 範囲のすべてのセルに値がある: 置き場所が無い。文字は診断に残す。
            grid.notes.append(("shape_text_not_placed", cell_ref(*box), shape.text[:200]))


def _read_xlsx_row_heights(worksheet: Any, grid: Grid) -> None:
    """指定された行の高さ（ポイント）と、シートの既定の行の高さ（図形の大きさの換算に使う）。"""
    default = getattr(worksheet.sheet_format, "defaultRowHeight", None)
    if default:
        grid.default_row_height = float(default)
    for index, dimension in worksheet.row_dimensions.items():
        if dimension.ht:
            grid.row_heights[index] = float(dimension.ht)


def _read_xlsx_column_widths(worksheet: Any, grid: Grid) -> None:
    """指定された列幅（帳票の方眼紙の手がかり）と、シートの既定の列幅を読む。"""
    default = getattr(worksheet.sheet_format, "defaultColWidth", None)
    if default:
        grid.default_column_width = float(default)
    for dimension in worksheet.column_dimensions.values():
        if not dimension.customWidth or not dimension.width:
            continue
        start = dimension.min or 0
        end = dimension.max or start
        # 全列の指定（A:XFD）で時間をかけない。値のある列の外は判定に使わない。
        for col in range(start, min(end, 16384) + 1) if end - start < 1024 else ():
            grid.column_widths[col] = float(dimension.width)


def _xlsx_label_hint(cell: Any) -> bool:
    """塗りつぶしか太字のセル（帳票のラベルの手がかり）。"""
    try:
        fill = cell.fill
        if fill is not None and fill.fill_type not in (None, "none"):
            return True
        return bool(cell.font is not None and cell.font.b)
    except Exception:
        return False


def _read_xls(
    source_bytes: bytes, *, include_hidden: bool = False
) -> tuple[list[Grid] | None, str | None]:
    try:
        import xlrd
    except Exception:
        return None, "xlrd_unavailable"
    try:
        book = xlrd.open_workbook(file_contents=source_bytes, formatting_info=True)
    except Exception:
        try:
            book = xlrd.open_workbook(file_contents=source_bytes)
        except Exception:
            return None, "excel_open_failed"
    grids: list[Grid] = []
    shape_counts = xls_shape_text_counts(source_bytes) or []
    try:
        for index, sheet in enumerate(book.sheets()):
            grid = Grid(name=sheet.name, hidden=getattr(sheet, "visibility", 0) != 0)
            _read_xls_column_widths(sheet, grid)
            _read_xls_hidden(sheet, grid)
            for row in range(sheet.nrows):
                for col in range(sheet.ncols):
                    text = _xls_cell_text(book, sheet, row, col, xlrd)
                    if (
                        text
                        and not include_hidden
                        and (row + 1 in grid.hidden_rows or col + 1 in grid.hidden_columns)
                    ):
                        grid.hidden_cells.add((row + 1, col + 1))
                        continue
                    if text:
                        grid.cells[(row + 1, col + 1)] = text
                        if _xls_label_hint(book, sheet, row, col):
                            grid.label_hints.add((row + 1, col + 1))
            grid.merges = [
                (rlo + 1, clo + 1, rhi, chi) for rlo, rhi, clo, chi in sheet.merged_cells
            ]
            if index < len(shape_counts) and shape_counts[index]:
                # xlrd は図形を読まない。図形の文字があることだけを残す。
                grid.notes.append(("shape_text_not_supported_xls", None, str(shape_counts[index])))
            grids.append(grid)
    except Exception:
        return None, "excel_read_failed"
    return grids, None


def _read_xls_column_widths(sheet: Any, grid: Grid) -> None:
    """.xls の列幅（1/256 文字の単位）。formatting_info で開いたときだけ読める。"""
    try:
        if sheet.defcolwidth:
            grid.default_column_width = float(sheet.defcolwidth)
        elif sheet.standardwidth:
            grid.default_column_width = float(sheet.standardwidth) / 256
        for col, info in sheet.colinfo_map.items():
            if info.width:
                grid.column_widths[col + 1] = float(info.width) / 256
    except Exception:
        return


def _read_xls_hidden(sheet: Any, grid: Grid) -> None:
    """.xls の非表示の行・列（formatting_info で開いたときだけ読める）。"""
    try:
        grid.hidden_rows = {
            row + 1 for row, info in sheet.rowinfo_map.items() if getattr(info, "hidden", 0)
        }
        grid.hidden_columns = {
            col + 1 for col, info in sheet.colinfo_map.items() if getattr(info, "hidden", 0)
        }
    except Exception:
        return


def _xls_label_hint(book: Any, sheet: Any, row: int, col: int) -> bool:
    """.xls の塗りつぶしか太字のセル。"""
    try:
        xf = book.xf_list[sheet.cell_xf_index(row, col)]
        if xf.background.fill_pattern:
            return True
        return bool(book.font_list[xf.font_index].bold)
    except Exception:
        return False


def _xls_cell_text(book: Any, sheet: Any, row: int, col: int, xlrd: Any) -> str:
    cell = sheet.cell(row, col)
    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
        return ""
    if cell.ctype == xlrd.XL_CELL_TEXT:
        return str(cell.value)
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return "TRUE" if cell.value else "FALSE"
    if cell.ctype == xlrd.XL_CELL_ERROR:
        return str(xlrd.error_text_from_code.get(cell.value, "#ERROR"))
    if cell.ctype == xlrd.XL_CELL_DATE:
        try:
            value = xlrd.xldate_as_datetime(cell.value, book.datemode)
        except Exception:
            return _general_number(float(cell.value))
        return _format_datetime(value)
    return _format_number(float(cell.value), _xls_number_format(book, sheet, row, col))


def _xls_number_format(book: Any, sheet: Any, row: int, col: int) -> str:
    try:
        xf = book.xf_list[sheet.cell_xf_index(row, col)]
        return str(book.format_map[xf.format_key].format_str)
    except Exception:
        return "General"


# ---- 値の文字列化 ----------------------------------------------------------------------


def _format_value(value: object, number_format: str | None, *, is_date: bool) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, _dt.datetime | _dt.date | _dt.time):
        return _format_datetime(value)
    if isinstance(value, int | float):
        if is_date:
            return _general_number(float(value))
        return _format_number(value, number_format or "General")
    return str(value).strip()


def _format_datetime(value: _dt.datetime | _dt.date | _dt.time) -> str:
    if isinstance(value, _dt.datetime):
        if value.time() == _dt.time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    return value.isoformat()


def _general_number(value: int | float) -> str:
    if isinstance(value, int):
        return str(value)
    if value.is_integer():
        return str(int(value))
    # 浮動小数の誤差（0.1+0.2 など）を出さない。指数表記にもしない。
    text = f"{value:.15f}".rstrip("0").rstrip(".")
    return text or "0"


def _format_number(value: int | float, number_format: str) -> str:
    """表示の書式（ゼロ埋め・桁区切り・小数桁・百分率）で数値を文字列にする。"""
    pattern = (number_format or "General").split(";", 1)[0].strip()
    if pattern in {"", "General", "@"} or not re.fullmatch(r"[0#,.%]+", pattern):
        return _general_number(value)
    percent = pattern.endswith("%")
    core = pattern.rstrip("%")
    integer_part, _, decimal_part = core.partition(".")
    decimals = len(decimal_part.replace(",", ""))
    thousands = "," in integer_part
    min_digits = integer_part.replace(",", "").replace("#", "").count("0")
    number = float(value) * (100 if percent else 1)
    sign = "-" if number < 0 else ""
    if thousands:
        body = f"{abs(number):,.{decimals}f}"
    else:
        body = f"{abs(number):.{decimals}f}"
        whole, dot, fraction = body.partition(".")
        body = whole.zfill(min_digits) + dot + fraction
    return f"{sign}{body}{'%' if percent else ''}"


# ---- 表頭・行 --------------------------------------------------------------------------


def _ranges_for(sheet_name: str, ranges: list[ExcelRange]) -> list[ExcelRange]:
    """シートに当てる範囲（シート名を付けた範囲があればそれ、無ければシート名の無い範囲）。"""
    named = [item for item in ranges if item.sheet == sheet_name]
    return named or [item for item in ranges if item.sheet is None]


def _crop(grid: Grid, ranges: list[ExcelRange]) -> Grid:
    """範囲の外のセルを除いたシート（#1229）。

    結合セルは範囲と重なる部分に切り詰め、左上が範囲の外でも値は切り詰めた左上へ移す。
    """

    def inside(row: int, col: int) -> bool:
        return any(item.contains(row, col) for item in ranges)

    cropped = Grid(
        name=grid.name,
        hidden=grid.hidden,
        cells={key: value for key, value in grid.cells.items() if inside(*key)},
        missing_formula_cells={key for key in grid.missing_formula_cells if inside(*key)},
        column_widths=dict(grid.column_widths),
        default_column_width=grid.default_column_width,
        label_hints={key for key in grid.label_hints if inside(*key)},
        shape_boxes={key: box for key, box in grid.shape_boxes.items() if inside(*key)},
        hidden_cells={key for key in grid.hidden_cells if inside(*key)},
        hidden_rows=set(grid.hidden_rows),
        hidden_columns=set(grid.hidden_columns),
        filter_rows=set(grid.filter_rows),
        notes=[note for note in grid.notes if note[1] is None or _note_inside(note[1], inside)],
    )
    for min_row, min_col, max_row, max_col in grid.merges:
        value = grid.cells.get((min_row, min_col), "")
        for item in ranges:
            top, left = max(min_row, item.min_row), max(min_col, item.min_col)
            bottom, right = min(max_row, item.max_row), min(max_col, item.max_col)
            if top > bottom or left > right:
                continue
            cropped.merges.append((top, left, bottom, right))
            if value:
                cropped.cells.setdefault((top, left), value)
    return cropped


def _note_inside(ref: str, inside: Any) -> bool:
    """セルの範囲の左上が範囲の中にあるか。"""
    start = parse_excel_range(f"{ref}:{ref}" if ":" not in ref else ref)
    return bool(inside(start.min_row, start.min_col))


def _apply_merges(grid: Grid, header_rows: range) -> None:
    """結合セルの値を埋める。

    表頭の行の中では結合の範囲（横・縦）を左上の値で埋める。データでは同じ行の横の結合だけを埋め、
    縦の結合の値は後の行へ漏らさない（no.1-rag と同じ。手順の境目を増やさない）。データの行で埋めた
    セルは ``grid.merge_filled`` に残し、記録には入れない（#2531）。
    """
    for min_row, min_col, max_row, max_col in grid.merges:
        value = grid.cells.get((min_row, min_col), "")
        if not value:
            continue
        for row in range(min_row, max_row + 1):
            in_header = row in header_rows and min_row in header_rows
            if row != min_row and not in_header:
                continue
            for col in range(min_col, max_col + 1):
                if (row, col) in grid.cells:
                    continue
                grid.cells[(row, col)] = value
                if not in_header:
                    # データの行の横の結合で埋めたセル。版式・表頭の判定には使い、記録には同じ値を
                    # 繰り返さない（#2531）。
                    grid.merge_filled.add((row, col))
    grid.invalidate_rows()


def _infer_header_rows(grid: Grid, header_row: int) -> tuple[int, str | None]:
    """表頭の段数を、表頭の行の横の結合から推定する（#2532。推定しなければ 1 と None）。

    表頭の行に横の結合（上の段の見出し）があり、すぐ下の行が結合の範囲のすべての列に文字
    （数値・日付でない）を持ち、その下の行に数値・日付があるときだけ 2 段にする。
    """
    spans = [
        (left, right)
        for top, left, bottom, right in grid.merges
        if top == bottom == header_row and right > left and grid.cells.get((top, left), "").strip()
    ]
    if not spans:
        return 1, None
    lower = header_row + 1
    lower_values = [grid.cells[(lower, col)] for col in grid.row_columns(lower)]
    if not lower_values or not all(is_textual(value) for value in lower_values):
        return 1, None
    for left, right in spans:
        for col in range(left, right + 1):
            if not grid.cells.get((lower, col), "").strip():
                return 1, None
    data = [grid.cells[(lower + 1, col)] for col in grid.row_columns(lower + 1)]
    if not any(not is_textual(value) for value in data):
        return 1, None
    return 2, "上の段の横の結合から表頭を 2 行と推定"


def _detect_header(grid: Grid, nonempty_rows: list[int]) -> tuple[int, float, str]:
    widest = max((len(grid.row_columns(row)) for row in nonempty_rows), default=1)
    best: tuple[float, int, str] | None = None
    for position, row in enumerate(nonempty_rows[:_HEADER_SCAN_ROWS]):
        columns = grid.row_columns(row)
        values = [grid.cells[(row, col)] for col in columns]
        if not values:
            continue
        textish = sum(1 for value in values if is_textual(value)) / len(values)
        unique = len(set(values)) / len(values)
        coverage = len(columns) / widest
        following = nonempty_rows[position + 1 : position + 1 + _HEADER_LOOKAHEAD_ROWS]
        if following:
            consistency = sum(
                len(set(grid.row_columns(other)) & set(columns)) / len(columns)
                for other in following
            ) / len(following)
        else:
            consistency = 0.0
        # 1 つのセルだけの行（説明・表題）は、表がそれ 1 列でない限り表頭にしない。
        if len(columns) == 1 and widest > 1:
            coverage = 0.0
        hints = sum(1 for hint in PROCEDURE_HEADER_HINTS if hint in " ".join(values))
        score = 0.3 * textish + 0.15 * unique + 0.3 * consistency + 0.25 * coverage
        score += min(0.2, 0.1 * hints)
        # 数値・日付だけの行はデータの行らしい（表頭の無い表）。文字の割合で下げる。
        score *= 0.5 + 0.5 * textish
        reason = (
            f"{len(columns)} 列・文字 {textish:.0%}・重複なし {unique:.0%}・"
            f"後続の行の一致 {consistency:.0%}"
        )
        if best is None or score > best[0] + 1e-9:
            best = (score, row, reason)
    if best is None:
        return nonempty_rows[0], 0.0, "表頭の候補がありません"
    score, row, reason = best
    return row, round(min(max(score, 0.0), 1.0), 3), reason


def _column_names(raw: list[tuple[int, str, str]]) -> list[str]:
    """空の表頭は ``column_<列>``、重複は ``<名前>__<列>``（no.1-rag と同じ）。"""
    names: list[str] = []
    seen: set[str] = set()
    for _, letter, name in raw:
        candidate = name or f"column_{letter}"
        if candidate.casefold() in seen:
            candidate = f"{candidate}__{letter}"
        seen.add(candidate.casefold())
        names.append(candidate)
    return names


def _excluded(name: str, letter: str, excludes: set[str]) -> bool:
    return name.casefold() in excludes or letter.casefold() in excludes


def _procedure_columns(
    grid: Grid, header_row: int, data_rows: list[int], columns: list[int]
) -> tuple[int, int] | None:
    """手順書の番号の列と題名の列を探す（表頭のある最初の列が題名、その左で一緒に埋まる列が番号）。"""
    labeled = [col for col in columns if grid.cells.get((header_row, col), "").strip()]
    if not labeled:
        return None
    title_column = labeled[0]
    candidates: list[tuple[int, int]] = []
    for col in range(1, title_column):
        together = sum(
            1
            for row in data_rows
            if grid.cells.get((row, col), "").strip()
            and grid.cells.get((row, title_column), "").strip()
        )
        if together >= 2:
            candidates.append((together, col))
    if not candidates:
        return None
    return max(candidates)[1], title_column


def _density(grid: Grid, data_rows: list[int], columns: list[int]) -> float:
    sample = data_rows[:_DENSITY_SAMPLE_ROWS]
    if not sample or not columns:
        return 0.0
    ratios = sorted(
        sum(1 for col in columns if grid.cells.get((row, col), "").strip()) / len(columns)
        for row in sample
    )
    middle = len(ratios) // 2
    if len(ratios) % 2:
        return ratios[middle]
    return (ratios[middle - 1] + ratios[middle]) / 2


def _row_range(cols: list[int], row_start: int, row_end: int) -> str:
    return f"{column_letter(cols[0])}{row_start}:{column_letter(cols[-1])}{row_end}"


def _table_blocks(
    grid: Grid, data_rows: list[int], columns: list[int], names: dict[int, str]
) -> list[SheetBlock]:
    blocks: list[SheetBlock] = []
    for row in data_rows:
        values = {
            names[col]: value
            for col in columns
            if (row, col) not in grid.merge_filled
            and (value := grid.cells.get((row, col), "").strip())
        }
        if values:
            blocks.append(
                SheetBlock(
                    kind="row",
                    row_start=row,
                    row_end=row,
                    cell_range=_row_range(columns, row, row),
                    values=values,
                    shape_cells=grid.shape_refs(row, columns[0], row, columns[-1]),
                )
            )
    return blocks


def _procedure_blocks(
    grid: Grid,
    data_rows: list[int],
    columns: list[int],
    names: dict[int, str],
    index_column: int,
    title_column: int,
) -> list[SheetBlock]:
    """番号と題名のある行を手順の始まりにし、次の手順の前までの行を 1 block にする。"""
    all_columns = sorted({*columns, index_column, title_column})
    detail_columns = [col for col in all_columns if col not in {index_column, title_column}]
    boundaries = [
        row
        for row in data_rows
        if grid.cells.get((row, index_column), "").strip()
        and grid.cells.get((row, title_column), "").strip()
    ]
    blocks: list[SheetBlock] = []
    section_title: str | None = None
    for position, row_start in enumerate(boundaries):
        title = grid.cells[(row_start, title_column)].strip()
        has_detail = any(grid.cells.get((row_start, col), "").strip() for col in detail_columns)
        if not has_detail:
            # 詳細の無い題名の行は章の見出し（以降の手順の section_path に入れる）。
            section_title = title
            continue
        next_start = boundaries[position + 1] if position + 1 < len(boundaries) else None
        rows = [
            row
            for row in data_rows
            if row >= row_start and (next_start is None or row < next_start)
        ]
        lines = [f"セクション: {section_title}"] if section_title else []
        last_row = row_start
        for row in rows:
            parts = [
                f"{names.get(col, f'column_{column_letter(col)}')}: {value}"
                for col in all_columns
                if (row, col) not in grid.merge_filled
                and (value := grid.cells.get((row, col), "").strip())
            ]
            if parts:
                lines.append(" | ".join(parts))
                last_row = row
        blocks.append(
            SheetBlock(
                kind="procedure_step",
                row_start=row_start,
                row_end=last_row,
                cell_range=_row_range(all_columns, row_start, last_row),
                section_path=[value for value in (section_title, title) if value],
                lines=lines,
                shape_cells=grid.shape_refs(row_start, all_columns[0], last_row, all_columns[-1]),
            )
        )
    return blocks


def _normalize_header(text: str) -> str:
    """表頭の語の比較用の形（全角半角・大小をそろえ、括弧の単位・注記と区切りを除く）。"""
    folded = unicodedata.normalize("NFKC", text).casefold()
    folded = _HEADER_BRACKETS.sub("", folded)
    return _HEADER_NOISE.sub("", folded)


def detect_column_role(header: str) -> tuple[str | None, float, str | None, str | None]:
    """表頭の語から列の役割を決める（役割・信頼度・当たった語・役割を付けない理由）。

    複数行の表頭（「設定 / 既定値」）は下の行（具体的な語）から順に見て、最初に当たった行で決める。
    """
    for part in reversed([item for item in header.split(" / ") if item.strip()]):
        normalized = _normalize_header(part)
        if not normalized:
            continue
        exact = [role for role, terms in COLUMN_ROLE_TERMS.items() if normalized in terms]
        if len(exact) == 1:
            return exact[0], COLUMN_ROLE_EXACT_CONFIDENCE, part.strip(), None
        if normalized in COLUMN_ROLE_UNCERTAIN_TERMS:
            return None, 0.0, part.strip(), "uncertain"
        # 長い語（3 文字以上）が表頭のどこかに含まれる役割を集める。2 つ以上なら「既定値の説明」
        # 「既定値／推奨値」のように 1 つに決められない。1 つでも、表頭の末尾がその語でなければ
        # （「既定値の変更履歴」）値の列とは言えないので付けない。
        # 短い語（2 文字。「説明」「推奨」）は役割を決めるのには使わないが、末尾にあれば
        # 1 つに決められない理由にはする（「既定値の説明」）。
        contained: dict[str, bool] = {}
        short_suffix: set[str] = set()
        for role, terms in COLUMN_ROLE_TERMS.items():
            for term in terms:
                if len(term) >= _COLUMN_ROLE_SUFFIX_MIN_CHARS and term in normalized:
                    contained[role] = contained.get(role, False) or normalized.endswith(term)
                elif len(term) >= 2 and normalized.endswith(term):
                    short_suffix.add(role)
        if contained and len(set(contained) | short_suffix) > 1:
            return None, 0.0, part.strip(), "ambiguous"
        if len(contained) == 1:
            role, at_end = next(iter(contained.items()))
            if at_end:
                return role, COLUMN_ROLE_SUFFIX_CONFIDENCE, part.strip(), None
            return None, 0.0, part.strip(), "uncertain"
    return None, 0.0, None, None


def _column_roles(
    columns: list[tuple[int, str, str, str]],
    configured: Mapping[str, str],
    *,
    detect: bool,
) -> tuple[list[SheetColumn], list[SheetDiagnostic]]:
    """列（列番号・記号・列名・表頭の文字）の役割を、指定 → 表頭の語の順に決める。"""
    by_key = {key.casefold(): (key, role) for key, role in configured.items()}
    result: list[SheetColumn] = []
    diagnostics: list[SheetDiagnostic] = []
    detected: list[int] = []
    for _, letter, name, header in columns:
        match = by_key.get(name.casefold()) or by_key.get(letter.casefold())
        if match is not None:
            key, configured_role = match
            if configured_role == COLUMN_ROLE_NONE:
                result.append(SheetColumn(name=name, column=letter))
            else:
                result.append(
                    SheetColumn(
                        name=name,
                        column=letter,
                        role=cast(ColumnRole, configured_role),
                        role_method="configured",
                        role_term=key,
                    )
                )
            continue
        result.append(SheetColumn(name=name, column=letter))
        if not detect:
            continue
        role, confidence, term, reason = detect_column_role(header)
        if role is not None and confidence >= COLUMN_ROLE_MIN_CONFIDENCE:
            result[-1] = SheetColumn(
                name=name,
                column=letter,
                role=cast(ColumnRole, role),
                role_method="detected",
                role_term=term,
            )
            detected.append(len(result) - 1)
        elif reason is not None:
            diagnostics.append(
                SheetDiagnostic(code=f"column_role_{reason}", detail=f"{letter}:{term or name}")
            )
    configured_count = sum(1 for item in result if item.role_method == "configured")
    if len(detected) > 0 and len(detected) + configured_count < COLUMN_ROLE_MIN_DETECTED_COLUMNS:
        # 1 列だけ当たったシートは設定値の表と言い切れない。役割を外して診断だけを残す
        # （指定した役割の列があれば、設定値の表として判定を使う）。
        for index in detected:
            column = result[index]
            diagnostics.append(
                SheetDiagnostic(
                    code="column_role_not_corroborated",
                    detail=f"{column.column}:{column.role_term or column.name}",
                )
            )
            result[index] = SheetColumn(name=column.name, column=column.column)
    return result, diagnostics


def _sheet_records(grid: Grid, options: ExcelOptions) -> SheetRecords | None:
    nonempty_rows = sorted({row for (row, _), value in grid.cells.items() if value.strip()})
    if not nonempty_rows:
        return None
    diagnostics = _grid_diagnostics(grid)
    header_count = options.header_row_count
    if options.header_row is not None:
        header_row = options.header_row
        detection = HeaderDetection(
            method="configured", confidence=1.0, reason=f"{header_row} 行目を表頭に指定"
        )
        if not grid.row_columns(header_row):
            diagnostics.append(SheetDiagnostic(code="header_row_empty", detail=str(header_row)))
    else:
        header_row, confidence, reason = _detect_header(grid, nonempty_rows)
        detection = HeaderDetection(method="detected", confidence=confidence, reason=reason)
        if confidence < HEADER_LOW_CONFIDENCE:
            diagnostics.append(
                SheetDiagnostic(code="header_low_confidence", detail=f"{confidence:.2f}")
            )
        if "header_row_count" not in options.model_fields_set:
            header_count, note = _infer_header_rows(grid, header_row)
            if note:
                detection = detection.model_copy(update={"reason": f"{detection.reason}・{note}"})
    header_rows = range(header_row, header_row + header_count)
    data_start = header_row + header_count
    _apply_merges(grid, header_rows)

    detect_roles = options.column_role_detection == "auto"
    if (
        detect_roles
        and detection.method == "detected"
        and detection.confidence < HEADER_LOW_CONFIDENCE
    ):
        # 表頭が確かでないシートでは、表頭の語で列の役割を決めない（指定の役割だけを使う）。
        detect_roles = False
        diagnostics.append(SheetDiagnostic(code="column_roles_skipped_low_header_confidence"))
    layout = _sheet_columns(grid, header_row, header_rows, options, detect_roles=detect_roles)
    # 列の役割の診断の位置（表の下の注記を除いて列を決め直したときも同じ位置に置く）。
    role_diagnostics_at = len(diagnostics)

    preamble: list[SheetPreambleRow] = []
    for row in nonempty_rows:
        if row >= header_row:
            break
        cols = grid.row_columns(row)
        # 横の結合で埋めた同じ値は 1 つにする（結合した見出しの文を繰り返さない）。
        texts: list[str] = []
        for col in cols:
            value = grid.cells[(row, col)].strip()
            if not texts or texts[-1] != value:
                texts.append(value)
        text = " ".join(texts)
        preamble.append(
            SheetPreambleRow(
                row_number=row,
                cell_range=_row_range(cols, row, row),
                text=text,
                shape_cells=grid.shape_refs(row, cols[0], row, cols[-1]),
            )
        )

    columns = layout.columns
    data_rows = _data_rows(grid, nonempty_rows, data_start, columns)
    procedure = _procedure_columns(grid, header_row, data_rows, columns) if columns else None
    mode = options.mode
    if mode == "auto":
        if procedure is not None:
            mode = "procedure"
        else:
            mode = (
                "table"
                if _density(grid, data_rows, columns) >= TABLE_DENSITY_THRESHOLD
                else "procedure"
            )
    notes: list[SheetNoteRow] = []
    if mode == "procedure" and procedure is not None:
        # 手順書は、表の後ろの 1 つのセルの行も手順の続きの行として読む（注記に分けない）。
        display = {
            col: column_display_name(name, layout.roles.get(name))
            for col, name in layout.names.items()
        }
        blocks = _procedure_blocks(grid, data_rows, columns, display, *procedure)
    else:
        # 番号と題名の列が無い手順書は、no.1-rag と同じく表として 1 行ずつ読む。
        if mode == "procedure":
            diagnostics.append(SheetDiagnostic(code="procedure_columns_not_found"))
        mode = "table"
        notes = _table_notes(grid, data_rows, columns, options.note_marker_patterns())
        if notes:
            # 注記の行は表の行・列にしない（注記だけが使う列を列名にしない）。
            note_rows = {note.row_number for note in notes}
            layout = _sheet_columns(
                grid,
                header_row,
                header_rows,
                options,
                detect_roles=detect_roles,
                skip_rows=note_rows,
            )
            columns = layout.columns
            data_rows = [
                row
                for row in _data_rows(grid, nonempty_rows, data_start, columns)
                if row not in note_rows
            ]
        blocks = _table_blocks(grid, data_rows, columns, layout.names)
        _attach_notes(blocks, notes)
    diagnostics[role_diagnostics_at:role_diagnostics_at] = layout.role_diagnostics
    return SheetRecords(
        name=grid.name,
        header_row=header_row,
        header_row_count=header_count,
        header_detection=detection,
        mode=mode,
        columns=layout.sheet_columns,
        preamble=preamble,
        blocks=blocks,
        notes=notes,
        diagnostics=diagnostics,
    )


@dataclass(frozen=True)
class _SheetColumns:
    """表の列（列番号・列名・役割）。"""

    columns: list[int]
    names: dict[int, str]
    sheet_columns: list[SheetColumn]
    roles: dict[str, str]
    role_diagnostics: list[SheetDiagnostic]


def _sheet_columns(
    grid: Grid,
    header_row: int,
    header_rows: range,
    options: ExcelOptions,
    *,
    detect_roles: bool,
    skip_rows: Set[int] = frozenset(),
) -> _SheetColumns:
    """表頭の行から下の値のある列を表の列にする（``skip_rows`` の行の値は列にしない）。"""
    used_columns = sorted(
        {
            col
            for (row, col), value in grid.cells.items()
            if row >= header_row and row not in skip_rows and value.strip()
        }
    )
    excludes = {item.strip().casefold() for item in options.exclude_columns if item.strip()}
    raw_names: list[tuple[int, str, str]] = []
    header_paths: dict[str, list[str]] = {}
    for col in used_columns:
        letter = column_letter(col)
        parts: list[str] = []
        for row in header_rows:
            value = grid.cells.get((row, col), "").strip()
            if value and (not parts or parts[-1] != value):
                parts.append(value)
        name = " / ".join(parts)
        if not _excluded(name, letter, excludes):
            raw_names.append((col, letter, name))
            header_paths[letter] = parts
    names = dict(zip([col for col, _, _ in raw_names], _column_names(raw_names), strict=True))
    sheet_columns, role_diagnostics = _column_roles(
        [(col, column_letter(col), names[col], header) for col, _, header in raw_names],
        options.column_roles,
        detect=detect_roles,
    )
    # 表頭の段の文字（「 / 」でつないだ列名からは、セルの中の「 / 」と区別して戻せない）。
    sheet_columns = [
        item.model_copy(update={"header_path": header_paths.get(item.column, [])})
        for item in sheet_columns
    ]
    return _SheetColumns(
        columns=[col for col, _, _ in raw_names],
        names=names,
        sheet_columns=sheet_columns,
        roles={item.name: item.role for item in sheet_columns if item.role},
        role_diagnostics=role_diagnostics,
    )


def _data_rows(
    grid: Grid, nonempty_rows: list[int], data_start: int, columns: list[int]
) -> list[int]:
    return [
        row
        for row in nonempty_rows
        if row >= data_start and any(grid.cells.get((row, col), "").strip() for col in columns)
    ]


# ---- 表の下の注記（#3566）--------------------------------------------------------------


def _table_notes(
    grid: Grid, data_rows: list[int], columns: list[int], patterns: list[re.Pattern[str]]
) -> list[SheetNoteRow]:
    """表の後ろの注記の行（表の行・表頭にしない）。

    表の最後の行から上へ、セルが 1 つ（横の結合を含む）で、注記の印（``ExcelOptions.note_markers``）
    で始まるか長い文（``NOTE_MIN_CHARS`` 以上）の行を続けて集める。印の無い長い文の行は、横に結合
    しているか、表との間に空の行があるときだけ注記にする（データの行の 1 つの長いセルと区別する）。
    1 列の表と、表の行が残らないときは注記にしない。
    """
    if len(columns) < 2:
        return []
    candidates: list[tuple[int, str, str | None, list[int]]] = []
    for row in reversed(data_rows):
        cols = grid.row_columns(row)
        units = [col for col in cols if (row, col) not in grid.merge_filled]
        if len(units) != 1:
            break
        text = grid.cells[(row, units[0])].strip()
        marker = note_marker(text, patterns)
        if marker is None and len(text) < NOTE_MIN_CHARS:
            break
        candidates.insert(0, (row, text, marker, cols))
    while candidates:
        table_rows = [row for row in data_rows if row < candidates[0][0]]
        if not table_rows:
            return []
        separated = candidates[0][0] - table_rows[-1] > 1
        weak = [
            position
            for position, (_, _, marker, cols) in enumerate(candidates)
            if marker is None and len(cols) < 2 and not separated
        ]
        if not weak:
            break
        candidates = candidates[weak[-1] + 1 :]
    return [
        SheetNoteRow(
            row_number=row,
            cell_range=_row_range(cols, row, row),
            text=text,
            marker=marker,
            shape_cells=grid.shape_refs(row, cols[0], row, cols[-1]),
        )
        for row, text, marker, cols in candidates
    ]


def _note_reference(marker: str | None) -> re.Pattern[str] | None:
    """表のセルから注記を指す印の照合（数字の付いた記号・語の印だけ。例: ※1、注2）。

    番号だけの印（「(1)」「1)」）と数字の無い印（「※」）は、セルの中の番号・記号と区別できない
    ので行に付けない（表の注記の記録には残る）。
    """
    if not marker:
        return None
    reference = "".join(
        char for char in marker if not char.isspace() and char not in "()（）[]［］"
    ).rstrip(":：.．、")
    if not any(char.isdigit() for char in reference) or is_numbering(reference):
        return None
    # 語で始まる印（「注1」）は語の途中（「…注1」）に当てない。数字の続き（「※10」）にも当てない。
    before = r"(?<!\w)" if re.match(r"\w", reference) else ""
    return re.compile(before + re.escape(reference) + r"(?!\d)")


def _attach_notes(blocks: list[SheetBlock], notes: list[SheetNoteRow]) -> None:
    """注記の印をセルに含む表の行に、その注記の文を付ける（1 行 = 1 記録の本文に注記を含める）。"""
    references = [
        (pattern, note.text)
        for note in notes
        if (pattern := _note_reference(note.marker)) is not None
    ]
    if not references:
        return
    for block in blocks:
        cells = ["".join(value.split()) for value in block.values.values()]
        block.notes = [
            text for pattern, text in references if any(pattern.search(cell) for cell in cells)
        ]


# ---- 版式（帳票・1 シートの複数の表。#2516）------------------------------------------------


def _grid_diagnostics(grid: Grid) -> list[SheetDiagnostic]:
    """読み取りの診断（キャッシュの無い数式・読まなかった非表示の行と列・図形の文字）。"""
    diagnostics = [
        SheetDiagnostic(code="formula_without_cached_value", cell=f"{column_letter(col)}{row}")
        for row, col in sorted(grid.missing_formula_cells)
    ]
    hidden_rows = grid.hidden_rows - grid.filter_rows
    rows = sorted({row for row, _ in grid.hidden_cells if row in hidden_rows})
    columns = sorted({col for _, col in grid.hidden_cells if col in grid.hidden_columns})
    if rows:
        diagnostics.append(
            SheetDiagnostic(code="hidden_rows_skipped", detail=",".join(map(str, rows)))
        )
    if columns:
        diagnostics.append(
            SheetDiagnostic(
                code="hidden_columns_skipped", detail=",".join(map(column_letter, columns))
            )
        )
    filtered = sorted(grid.filter_rows & {row for row, _ in grid.cells})
    if filtered:
        diagnostics.append(
            SheetDiagnostic(code="hidden_filter_rows_read", detail=",".join(map(str, filtered)))
        )
    diagnostics.extend(
        SheetDiagnostic(code=code, cell=cell, detail=detail) for code, cell, detail in grid.notes
    )
    return diagnostics


def _form_records(grid: Grid, detection: LayoutDetection) -> SheetRecords | None:
    """帳票のシート（の領域）を「ラベル: 値」の項目の記録にする。"""
    blocks, diagnostics = form_blocks(grid)
    if not blocks:
        return None
    return SheetRecords(
        name=grid.name,
        header_row=None,
        mode="form",
        layout_detection=detection,
        blocks=blocks,
        diagnostics=[*_grid_diagnostics(grid), *_layout_diagnostics(detection), *diagnostics],
    )


def _layout_diagnostics(detection: LayoutDetection | None) -> list[SheetDiagnostic]:
    """帳票と判定したときの診断（表・手順書の判定の根拠は ``layout_detection`` だけに残す）。"""
    if detection is None or detection.method != "detected":
        return []
    return [
        SheetDiagnostic(
            code=f"layout_detected_{detection.layout}",
            detail=f"{detection.confidence:.2f}: {detection.reason}",
        )
    ]


def _detect(grid: Grid, options: ExcelOptions) -> tuple[bool, float, dict[str, float], str]:
    """auto の帳票の判定（帳票か・点数・特徴・理由）。"""
    score, features, reason = detect_form_layout(grid)
    return score >= options.form_detection_threshold, score, features, reason


def _table_records(
    grid: Grid, options: ExcelOptions, form_score: tuple[float, dict[str, float], str] | None
) -> SheetRecords | None:
    """表・手順書として読む（auto のときは帳票でないと判定した根拠を残す）。"""
    sheet = _sheet_records(grid, options)
    if sheet is None or form_score is None:
        return sheet
    score, features, reason = form_score
    detection = LayoutDetection(
        layout=sheet.mode,
        method="detected",
        confidence=round(1 - score, 3),
        reason=f"帳票の点数 {score:.2f} が閾値 {options.form_detection_threshold:.2f} 未満"
        f"（{reason}）",
        features=features,
    )
    return sheet.model_copy(update={"layout_detection": detection})


def _region_records(
    grid: Grid,
    options: ExcelOptions,
    detected: tuple[bool, float, dict[str, float], str] | None = None,
) -> SheetRecords | None:
    """1 つの領域（シート全体か、分けた表）を版式に合わせて読む（``detected`` は判定済み）。"""
    if options.mode == "form":
        detection = LayoutDetection(
            layout="form", method="configured", confidence=1.0, reason="読み方に帳票を指定"
        )
        return _form_records(grid, detection)
    if options.mode != "auto":
        return _sheet_records(grid, options)
    is_form, score, features, reason = detected or _detect(grid, options)
    if is_form:
        detection = LayoutDetection(
            layout="form", method="detected", confidence=score, reason=reason, features=features
        )
        return _form_records(grid, detection) or _table_records(
            grid, options, (score, features, reason)
        )
    sheet = _table_records(grid, options, (score, features, reason))
    if sheet is not None and not sheet.blocks and not sheet.preamble:
        # 表頭の行だけの領域（表の外の文・図形の文字）: 表頭にすると文字が記録に残らないので、
        # 帳票の項目として読む（#2526）。
        detection = LayoutDetection(
            layout="form",
            method="detected",
            confidence=score,
            reason=f"データの行が無い（表頭の行だけの）領域を項目として読む（{reason}）",
            features=features,
        )
        return _form_records(grid, detection) or sheet
    return sheet


def sheet_tables(grid: Grid, options: ExcelOptions) -> list[SheetRecords]:
    """1 シートを記録にする（複数の表に分けたときは表ごとに 1 つ。空なら空の一覧）。

    シート全体が帳票なら分けない（帳票の空の行・列は表の区切りではない）。分けるのは、表頭の
    行を指定していない（``header_row`` が無い）ときだけ。
    """
    if not any(value.strip() for value in grid.cells.values()):
        return []
    detected = _detect(grid, options) if options.mode == "auto" else None
    whole_form = options.mode == "form" or (detected is not None and detected[0])
    regions = (
        split_regions(
            grid,
            gap_rows=options.table_gap_rows,
            gap_columns=options.table_gap_columns,
            note_markers=options.note_marker_patterns(),
        )
        if options.split_tables and options.header_row is None and not whole_form
        else []
    )
    if len(regions) <= 1:
        sheet = _region_records(grid, options, detected)
        return [sheet] if sheet is not None else []
    tables: list[SheetRecords] = []
    for top, left, bottom, right in regions:
        region = ExcelRange(sheet=None, min_row=top, min_col=left, max_row=bottom, max_col=right)
        cropped = _crop(grid, [region])
        if tables:
            # セルの位置の無いシートの診断は、最初の表にだけ残す。
            cropped.notes = [note for note in cropped.notes if note[1] is not None]
        sheet = _region_records(cropped, options)
        if sheet is None:
            continue
        title = sheet.preamble[0].text if sheet.preamble else None
        tables.append(
            sheet.model_copy(
                update={
                    "table_index": len(tables) + 1,
                    "region": f"{column_letter(left)}{top}:{column_letter(right)}{bottom}",
                    "title": title,
                }
            )
        )
    return tables
