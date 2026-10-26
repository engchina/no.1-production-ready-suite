"""Excel の読み方（pr_backend_core.spreadsheet。#2525）の検証。

テストのデータはすべて合成（ラベル・値は匿名の「項目A」「値A」）。RAG の前処理の変換
（JSONL）は rag/services/preprocess/excel_to_jsonl のテストが確かめる。
"""

from __future__ import annotations

import io
import subprocess
import sys
from typing import Any

import openpyxl
import pytest
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import ValidationError

from pr_backend_core.spreadsheet import (
    ExcelOptions,
    SheetRecords,
    column_display_name,
    openpyxl_available,
    read_excel,
)

_FILL = PatternFill("solid", fgColor="DDDDDD")


def _bytes(workbook: openpyxl.Workbook) -> bytes:
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _put(worksheet: Worksheet, ref: str, value: Any, *, fill: bool = False) -> None:
    first = ref.split(":", 1)[0]
    worksheet[first] = value
    if ":" in ref:
        worksheet.merge_cells(ref)
    if fill:
        worksheet[first].fill = _FILL


def _sheets(source: bytes, **options: Any) -> list[SheetRecords]:
    result = read_excel(source, options=options)
    assert result.failure is None, result.failure
    assert result.document is not None
    return result.document.sheets


def test_importing_the_package_does_not_import_openpyxl() -> None:
    """openpyxl / xlrd は使う側が入れる（関数の中で読む）。"""
    code = (
        "import sys, pr_backend_core.spreadsheet as s; "
        "print('openpyxl' in sys.modules, 'xlrd' in sys.modules)"
    )
    output = subprocess.run(
        [sys.executable, "-c", code], check=True, capture_output=True, text=True
    ).stdout
    assert output.strip() == "False False"
    assert openpyxl_available() is True


def test_table_sheet_keeps_header_rows_and_cell_ranges() -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "表"
    worksheet.append(["列A", "列B", "列C"])
    for index in range(3):
        worksheet.append([f"a{index}", index, index * 10])
    (sheet,) = _sheets(_bytes(workbook))
    assert sheet.mode == "table"
    assert sheet.header_row == 1
    assert [column.name for column in sheet.columns] == ["列A", "列B", "列C"]
    assert [(block.cell_range, block.values) for block in sheet.blocks[:1]] == [
        ("A2:C2", {"列A": "a0", "列B": "0", "列C": "0"})
    ]
    assert sheet.layout_detection is not None
    assert sheet.layout_detection.layout == "table"


def test_procedure_sheet_groups_steps() -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "手順"
    worksheet.append(["", "作業項目", "作業内容"])
    worksheet.append([1, "準備", None])
    worksheet.append([1.1, "項目A", "内容A"])
    worksheet.append([None, None, "内容A の続き"])
    worksheet.append([1.2, "項目B", "内容B"])
    (sheet,) = _sheets(_bytes(workbook))
    assert sheet.mode == "procedure"
    assert [(block.section_path, block.row_start, block.row_end) for block in sheet.blocks] == [
        (["準備", "項目A"], 3, 4),
        (["準備", "項目B"], 5, 5),
    ]


def test_form_sheet_is_detected_and_read_as_label_value_fields() -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "様式"
    for col in range(1, 31):
        worksheet.column_dimensions[get_column_letter(col)].width = 2.5
    _put(worksheet, "B1:AC1", "様式A")
    _put(worksheet, "B3:E3", "項目A", fill=True)
    _put(worksheet, "F3:N3", "値A")
    _put(worksheet, "P3:S3", "項目B", fill=True)
    _put(worksheet, "T3:AC3", "値B")
    _put(worksheet, "B4:E4", "項目C", fill=True)
    _put(worksheet, "F4:AC4", "値C")
    (sheet,) = _sheets(_bytes(workbook))
    assert sheet.mode == "form"
    assert sheet.layout_detection is not None
    assert sheet.layout_detection.method == "detected"
    assert sheet.layout_detection.confidence >= 0.6
    assert [
        (block.label, block.value, block.label_cell, block.value_cell, block.section_path)
        for block in sheet.blocks
    ] == [
        ("項目A", "値A", "B3:E3", "F3:N3", ["様式A"]),
        ("項目B", "値B", "P3:S3", "T3:AC3", ["様式A"]),
        ("項目C", "値C", "B4:E4", "F4:AC4", ["様式A"]),
    ]


def test_sheet_with_two_tables_is_split() -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "表"
    worksheet.append(["表の題名1"])
    worksheet.append(["列A", "列B"])
    worksheet.append(["a1", 1])
    worksheet.append([])
    worksheet.append([])
    worksheet.append(["項目X", "項目Y", "項目Z"])
    worksheet.append(["x1", "y1", "z1"])
    sheets = _sheets(_bytes(workbook))
    assert [(sheet.table_index, sheet.region, sheet.title) for sheet in sheets] == [
        (1, "A1:B3", "表の題名1"),
        (2, "A6:C7", None),
    ]
    assert [column.name for column in sheets[1].columns] == ["項目X", "項目Y", "項目Z"]
    assert len(_sheets(_bytes(workbook), split_tables=False)) == 1


def test_options_are_validated_and_bad_input_degrades() -> None:
    with pytest.raises(ValidationError):
        read_excel(b"x", options={"mode": "unknown"})
    assert ExcelOptions().mode == "auto"
    assert read_excel(b"").failure == "excel_empty"
    assert read_excel(b"not an excel file", extension="xlsx").failure == "excel_open_failed"


def test_xls_is_read_with_widths_and_label_hints() -> None:
    import xlwt  # type: ignore[import-untyped]  # .xls の見本を作る（dev の依存）

    book = xlwt.Workbook()
    sheet = book.add_sheet("旧様式")
    for col in range(30):
        sheet.col(col).width = 256 * 2
    label = xlwt.easyxf("pattern: pattern solid, fore_colour ice_blue; font: bold on")
    sheet.write_merge(2, 2, 1, 5, "項目A", label)
    sheet.write_merge(2, 2, 6, 28, "値A", xlwt.easyxf(""))
    sheet.write_merge(3, 3, 1, 5, "項目B", label)
    sheet.write_merge(3, 3, 6, 10, 25, xlwt.easyxf(""))
    sheet.write_merge(3, 3, 11, 15, "項目C", label)
    sheet.write_merge(3, 3, 16, 28, "値C", xlwt.easyxf(""))
    buffer = io.BytesIO()
    book.save(buffer)
    (records,) = _sheets(buffer.getvalue())
    assert records.mode == "form"
    assert [(block.label, block.value) for block in records.blocks] == [
        ("項目A", "値A"),
        ("項目B", "25"),
        ("項目C", "値C"),
    ]


def test_column_display_name_marks_value_columns_only() -> None:
    assert column_display_name("既定値", "default") == "既定値［資料の既定値］"
    assert column_display_name("説明", "definition") == "説明"
    assert column_display_name("列A", None) == "列A"


# ---- 見本で見つけた読み方の不具合（#2527〜#2532）--------------------------------------------


class _CountingCells(dict[tuple[int, int], str]):
    """セルの全体の走査（items）の回数を数える。"""

    scans = 0

    def items(self) -> Any:
        type(self).scans += 1
        return super().items()


def test_header_detection_scans_cells_a_constant_number_of_times() -> None:
    """表頭の推定は行ごとに全セルを走査しない（行の数 × セルの数に比例しない。#2527）。"""
    from pr_backend_core.spreadsheet.grid import Grid
    from pr_backend_core.spreadsheet.reader import _detect_header

    def scans(rows: int) -> int:
        cells = _CountingCells()
        cells.update({(1, col): f"列{col}" for col in range(1, 6)})
        for row in range(2, rows + 2):
            cells.update({(row, col): str(row * col) for col in range(1, 6)})
        grid = Grid(name="表", hidden=False, cells=cells)
        _CountingCells.scans = 0
        _detect_header(grid, list(range(1, rows + 2)))
        return _CountingCells.scans

    assert scans(2000) == scans(50) <= 2


def test_data_row_after_blank_row_stays_in_the_table_above() -> None:
    """空の 1 行の後の 1 行のデータは上の表の続きで、下の表の題名にしない（#2528）。"""
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "集計"
    worksheet["A1"] = "表の題名1"
    worksheet.append(["列A", "列B", "列C"])
    worksheet.append(["a1", 1, 2])
    worksheet.append(["a2", 3, 4])
    worksheet.append([])
    worksheet.append(["a3", 5, 6])
    worksheet.append([])
    worksheet.append([])
    worksheet.append(["表の題名2"])
    worksheet.append(["項目X", "項目Y"])
    worksheet.append(["x1", "y1"])
    worksheet["E2"], worksheet["F2"] = "区分", "件数"
    worksheet["E3"], worksheet["F3"] = "k1", 1
    worksheet["E4"], worksheet["F4"] = "k2", 2
    sheets = _sheets(_bytes(workbook))
    assert [(sheet.region, sheet.title, sheet.header_row) for sheet in sheets] == [
        ("A1:C6", "表の題名1", 2),
        ("E2:F4", None, 2),
        ("A9:B11", "表の題名2", 10),
    ]
    assert [block.row_start for block in sheets[0].blocks] == [3, 4, 6]
    assert sheets[2].blocks[0].values == {"項目X": "x1", "項目Y": "y1"}


def _grid_form(worksheet: Worksheet, columns: int = 24) -> None:
    for col in range(1, columns + 1):
        worksheet.column_dimensions[get_column_letter(col)].width = 3


def test_inline_label_value_cells_and_stamp_boxes() -> None:
    """1 つのセルの「ラベル：値」は単独の項目、押印欄の枠の上のラベルは空欄（#2529）。"""
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "様式"
    _grid_form(worksheet)
    _put(worksheet, "B1:X1", "様式A")
    _put(worksheet, "B2:L2", "項目A：2026/04/15")
    _put(worksheet, "M2:X2", "項目B：値B")
    _put(worksheet, "B3:L3", "時刻 10:30")
    _put(worksheet, "B4:E4", "項目C", fill=True)
    _put(worksheet, "F4:X4", "値C")
    for ref, text in (("P6:R6", "枠A"), ("S6:U6", "枠B"), ("V6:X6", "枠C")):
        _put(worksheet, ref, text, fill=True)
    for ref in ("P7:R9", "S7:U9", "V7:X9"):
        worksheet.merge_cells(ref)
    (sheet,) = _sheets(_bytes(workbook))
    assert sheet.mode == "form"
    assert [(block.label, block.value, block.cell_range) for block in sheet.blocks] == [
        ("項目A", "2026/04/15", "B2:L2"),
        ("項目B", "値B", "M2:X2"),
        ("項目C", "値C", "B4:X4"),
    ]
    # 時刻のコロンでは分けない（行に 1 つだけの単位なので、以降の項目の見出しになる）。
    assert sheet.blocks[2].section_path == ["時刻 10:30"]
    empty = {item.code: item.detail for item in sheet.diagnostics}
    assert empty["form_empty_fields"] == "3"


def test_normal_width_forms_are_detected() -> None:
    """普通の列幅でも、結合した値がラベルのすぐ右に並ぶ帳票は form にする（#2530）。"""
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "届出"
    _put(worksheet, "A1:F1", "様式B")
    _put(worksheet, "A3", "項目A：")
    _put(worksheet, "B3:C3", "値A")
    _put(worksheet, "A4", "項目B：")
    _put(worksheet, "B4:F4", "値B")
    _put(worksheet, "A5", "項目C：")
    _put(worksheet, "B5:F5", "値C")
    _put(worksheet, "A6", "項目D：")
    _put(worksheet, "B6:F7", "値D")
    (sheet,) = _sheets(_bytes(workbook))
    assert sheet.mode == "form"
    assert sheet.layout_detection is not None
    assert sheet.layout_detection.features["narrow"] == 0.0
    # ラベルの末尾のコロンは除く。
    assert [(block.label, block.value) for block in sheet.blocks] == [
        ("項目A", "値A"),
        ("項目B", "値B"),
        ("項目C", "値C"),
        ("項目D", "値D"),
    ]


def test_table_with_shaded_header_stays_table() -> None:
    """塗りつぶしの表頭の下に値の並ぶ表は帳票にしない（#2530）。"""
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "表"
    for col, text in enumerate(("列A", "列B", "列C", "列D"), start=1):
        _put(worksheet, f"{get_column_letter(col)}1", text, fill=True)
    for row in range(2, 8):
        worksheet.append([f"a{row}", row, row * 2, f"b{row}"])
    (sheet,) = _sheets(_bytes(workbook))
    assert sheet.mode == "table"
    assert sheet.layout_detection is not None
    assert sheet.layout_detection.features["label_value"] == 0.0


def test_horizontally_merged_data_cell_is_not_repeated() -> None:
    """データの行の横の結合の値は、結合した各列に繰り返さない（#2531）。"""
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "表"
    worksheet.append(["列A", "列B", "列C", "列D"])
    worksheet.append(["a1", "b1", "c1", "d1"])
    worksheet.append(["a2", "長い説明", None, None])
    worksheet.merge_cells("B3:D3")
    (sheet,) = _sheets(_bytes(workbook), mode="table")
    assert sheet.blocks[1].values == {"列A": "a2", "列B": "長い説明"}
    assert sheet.blocks[1].cell_range == "A3:D3"


def test_two_row_header_with_merged_upper_row_is_inferred() -> None:
    """上の段を結合した 2 段の表頭を推定する（指定があれば指定に従う。#2532）。"""
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "実績"
    _put(worksheet, "A1:A2", "拠点")
    _put(worksheet, "B1:C1", "上期")
    _put(worksheet, "D1:E1", "下期")
    for col, text in zip("BCDE", ("件数", "金額", "件数", "金額"), strict=True):
        worksheet[f"{col}2"] = text
    for row in range(3, 8):
        worksheet.append([f"拠点{row}", row, row * 10, row + 1, row * 20])
    (sheet,) = _sheets(_bytes(workbook))
    assert sheet.header_row_count == 2
    assert [column.name for column in sheet.columns] == [
        "拠点",
        "上期 / 件数",
        "上期 / 金額",
        "下期 / 件数",
        "下期 / 金額",
    ]
    assert [block.row_start for block in sheet.blocks] == [3, 4, 5, 6, 7]
    assert sheet.header_detection is not None
    assert "2 行" in sheet.header_detection.reason
    (configured,) = _sheets(_bytes(workbook), header_row_count=1)
    assert configured.header_row_count == 1
    assert configured.blocks[0].row_start == 2


def test_columns_keep_the_rows_of_a_multi_row_header() -> None:
    """列は表頭の段の文字を上の段から持つ（「 / 」を含むセルも 1 段のまま。#3262）。"""
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    _put(worksheet, "A1:A2", "項目A")
    _put(worksheet, "B1:C1", "区分 / 甲")
    worksheet["B2"] = "値A"
    worksheet["C2"] = "値B"
    worksheet["D2"] = "値C"
    for row in range(3, 7):
        worksheet.append([f"行{row}", row, row * 10, row * 100])
    (sheet,) = _sheets(_bytes(workbook), header_row=1, header_row_count=2)
    assert [(column.column, column.header_path) for column in sheet.columns] == [
        ("A", ["項目A"]),
        ("B", ["区分 / 甲", "値A"]),
        ("C", ["区分 / 甲", "値B"]),
        ("D", ["値C"]),
    ]
    # 列名は段を「 / 」でつないだもので、セルの中の「 / 」と区別できない。
    assert sheet.columns[1].name == "区分 / 甲 / 値A"
    (single,) = _sheets(_bytes(workbook), header_row_count=1, header_row=2)
    assert [column.header_path for column in single.columns] == [[], ["値A"], ["値B"], ["値C"]]


# ---- 表の下の注記（#3566）。文はすべて合成 -------------------------------------------------

_NOTE_1 = "※1 値は合成の注記の一つ目の文で、表の行の値の条件を説明します。"
_NOTE_2 = "※2 合成の注記の二つ目の文です。"


def _table(worksheet: Worksheet) -> None:
    worksheet.append(["項目", "値", "区分"])
    worksheet.append(["項目A", "10※1", "甲"])
    worksheet.append(["項目B", "20", "乙"])
    worksheet.append(["項目C", "30※2", "丙"])


def test_note_rows_below_a_table_are_notes_not_a_new_table() -> None:
    """空の行を挟んだ注記の行は、新しい表・表頭にせず上の表の注記にする。"""
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    _table(worksheet)
    worksheet["A6"] = _NOTE_1
    worksheet["A7"] = _NOTE_2
    (sheet,) = _sheets(_bytes(workbook))
    assert sheet.header_row == 1
    assert [column.name for column in sheet.columns] == ["項目", "値", "区分"]
    assert [block.row_start for block in sheet.blocks] == [2, 3, 4]
    assert [(note.row_number, note.cell_range, note.marker) for note in sheet.notes] == [
        (6, "A6:A6", "※1"),
        (7, "A7:A7", "※2"),
    ]
    assert [note.text for note in sheet.notes] == [_NOTE_1, _NOTE_2]
    # 注記の印をセルに含む行は、その注記の文を持つ（1 行 = 1 記録の本文に入る）。
    assert [block.notes for block in sheet.blocks] == [[_NOTE_1], [], [_NOTE_2]]


def test_note_rows_directly_below_and_merged_long_text() -> None:
    """空の行の無い注記の行（印か、横に結合した長い文）も注記にする。

    短い 1 つのセルの行は表の行のまま。
    """
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    _table(worksheet)
    worksheet.append(["小計"])
    worksheet["A6"] = _NOTE_2
    _put(
        worksheet, "A7:C7", "印の無い合成の注記の文で、表の全体の幅に横に結合した長い説明の行です。"
    )
    (sheet,) = _sheets(_bytes(workbook))
    assert [block.row_start for block in sheet.blocks] == [2, 3, 4, 5]
    assert sheet.blocks[-1].values == {"項目": "小計"}
    assert [(note.row_number, note.marker, note.cell_range) for note in sheet.notes] == [
        (6, "※2", "A6:A6"),
        (7, None, "A7:C7"),
    ]


def test_marked_note_between_tables_is_not_the_next_table_title() -> None:
    """表と表の間の注記の 1 行は、下の表の題名にせず上の表の注記にする。

    番号だけの行は下の表の題名のまま。
    """
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    _table(worksheet)
    worksheet["A6"] = _NOTE_2
    worksheet["A8"] = "(2) 表の題名"
    worksheet["A9"], worksheet["B9"] = "列X", "列Y"
    worksheet["A10"], worksheet["B10"] = "x1", "y1"
    sheets = _sheets(_bytes(workbook))
    assert [(sheet.region, sheet.title) for sheet in sheets] == [
        ("A1:C6", None),
        ("A8:B10", "(2) 表の題名"),
    ]
    assert [note.text for note in sheets[0].notes] == [_NOTE_2]
    assert sheets[1].notes == []
    assert [column.name for column in sheets[1].columns] == ["列X", "列Y"]


def test_note_markers_are_configurable_and_one_column_lists_are_not_notes() -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    _table(worksheet)
    worksheet["A6"] = "† 合成の注記"
    (default,) = _sheets(_bytes(workbook))
    assert default.notes == []
    (custom,) = _sheets(_bytes(workbook), note_markers=["†"])
    assert [(note.marker, note.text) for note in custom.notes] == [("†", "† 合成の注記")]
    with pytest.raises(ValidationError):
        ExcelOptions(note_markers=["("])

    listing = openpyxl.Workbook()
    sheet = listing.active
    assert sheet is not None
    for value in ["手順", "1) 合成の手順の一つ目", "2) 合成の手順の二つ目"]:
        sheet.append([value])
    (single,) = _sheets(_bytes(listing))
    assert single.notes == []
    assert len(single.blocks) == 2


def test_note_reference_does_not_match_inside_words_or_longer_numbers() -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(["項目", "値"])
    worksheet.append(["受注1", "※10"])
    worksheet.append(["注1 の行", "※1"])
    worksheet.append([])
    worksheet.append(["注1: 合成の注記"])
    worksheet.append(["※1 合成の注記"])
    (sheet,) = _sheets(_bytes(workbook))
    assert [note.marker for note in sheet.notes] == ["注1:", "※1"]
    assert [block.notes for block in sheet.blocks] == [[], ["注1: 合成の注記", "※1 合成の注記"]]


def test_procedure_sheet_keeps_trailing_single_cell_rows_in_the_last_step() -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(["", "作業項目", "作業内容"])
    worksheet.append([1, "手順A", "内容A"])
    worksheet.append([2, "手順B", "内容B"])
    worksheet.append(["", "", "※ 合成の補足"])
    (sheet,) = _sheets(_bytes(workbook))
    assert sheet.mode == "procedure"
    assert sheet.notes == []
    assert "※ 合成の補足" in sheet.blocks[-1].lines[-1]
