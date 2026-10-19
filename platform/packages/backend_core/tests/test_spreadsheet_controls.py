"""帳票のフォームコントロールのチェックの状態と、大きさで置く図形の範囲の検証（#2537）。

テストのデータはすべて合成（ラベル・値は匿名の「項目A」「甲」「乙」）。ブックは openpyxl で作り、
図形・コントロールは xlsx の zip に最小のシートの ``<controls>``・ctrlProp・drawing の XML を
書き足して作る。
"""

from __future__ import annotations

import io
import zipfile
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from pr_backend_core.spreadsheet import SheetRecords, read_excel
from pr_backend_core.spreadsheet.grid import Grid

_XDR = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_REL_TYPES = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
# 列幅 2.5（文字数）は 17 ピクセル、既定の行の高さ 15 ポイント。
_COLUMN_EMU = 17 * 9525
_ROW_EMU = 15 * 12700


def _bytes(workbook: openpyxl.Workbook) -> bytes:
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _put(worksheet: Worksheet, ref: str, value: Any) -> None:
    worksheet[ref.split(":", 1)[0]] = value
    if ":" in ref:
        worksheet.merge_cells(ref)


def _form_workbook() -> openpyxl.Workbook:
    """方眼紙（狭い列の格子）の帳票。"""
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "様式"
    for col in range(1, 31):
        worksheet.column_dimensions[get_column_letter(col)].width = 2.5
    _put(worksheet, "B1:AC1", "様式A")
    _put(worksheet, "B3:E3", "項目A")
    _put(worksheet, "F3:N3", "値A")
    _put(worksheet, "B4:E4", "項目B")
    _put(worksheet, "B5:E5", "項目C")
    _put(worksheet, "B6:E6", "項目D")
    _put(worksheet, "F6:N6", "値D")
    return workbook


def _marker(tag: str, row: int, col: int, *, prefix: str = "xdr:") -> str:
    """アンカーの位置（シートの XML の controlPr の中は ``prefix=""``）。"""
    return (
        f"<{prefix}{tag}><xdr:col>{col}</xdr:col><xdr:colOff>0</xdr:colOff>"
        f"<xdr:row>{row}</xdr:row><xdr:rowOff>0</xdr:rowOff></{prefix}{tag}>"
    )


def _shape(shape_id: int, text: str, *, hidden: bool = False) -> str:
    body = f"<a:p><a:r><a:t>{text}</a:t></a:r></a:p>" if text else "<a:p/>"
    flag = ' hidden="1"' if hidden else ""
    return (
        f'<xdr:sp><xdr:nvSpPr><xdr:cNvPr id="{shape_id}" name="Shape {shape_id}"{flag}/>'
        f"<xdr:cNvSpPr/></xdr:nvSpPr><xdr:spPr/><xdr:txBody><a:bodyPr/>{body}</xdr:txBody></xdr:sp>"
    )


def _two_cell(start: tuple[int, int], end: tuple[int, int], body: str) -> str:
    """0 始まりの from（行・列）と to（行・列。オフセット 0）のアンカー。"""
    return (
        "<xdr:twoCellAnchor>"
        + _marker("from", *start)
        + _marker("to", *end)
        + body
        + "<xdr:clientData/></xdr:twoCellAnchor>"
    )


def _alternate(anchor: str) -> str:
    """Excel がフォームコントロールのアンカーを書く形（wsDr の直下の mc:AlternateContent）。"""
    return (
        f'<mc:AlternateContent><mc:Choice Requires="a14">{anchor}</mc:Choice>'
        "<mc:Fallback/></mc:AlternateContent>"
    )


def _control(shape_id: int, rid: str, box: tuple[int, int, int, int] | None = None) -> str:
    anchor = ""
    if box is not None:
        top, left, bottom, right = box
        anchor = (
            '<controlPr><anchor moveWithCells="1">'
            + _marker("from", top, left, prefix="")
            + _marker("to", bottom, right, prefix="")
            + "</anchor></controlPr>"
        )
    return (
        f'<control shapeId="{shape_id}" r:id="{rid}" name="Control {shape_id}">{anchor}</control>'
    )


def _with_parts(
    source: bytes,
    *,
    anchors: list[str],
    controls: list[str] | None = None,
    props: dict[str, tuple[str, bool]] | None = None,
) -> bytes:
    """xlsx の zip に、シート 1 の drawing・コントロール・ctrlProp を書き足す。

    ``props`` は r:id → （コントロールの種類、オンか）。
    """
    props = props or {}
    drawing = (
        f'<?xml version="1.0" encoding="UTF-8"?><xdr:wsDr xmlns:xdr="{_XDR}" xmlns:a="{_A}" '
        f'xmlns:mc="{_MC}">' + "".join(anchors) + "</xdr:wsDr>"
    )
    relations = [
        f'<Relationship Id="rIdDrawing" Type="{_REL_TYPES}/drawing" '
        'Target="../drawings/drawing1.xml"/>'
    ]
    relations += [
        f'<Relationship Id="{rid}" Type="{_REL_TYPES}/ctrlProp" '
        f'Target="../ctrlProps/ctrlProp{index}.xml"/>'
        for index, rid in enumerate(props, start=1)
    ]
    rels = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        + "".join(relations)
        + "</Relationships>"
    )
    sheet_tail = f'<drawing xmlns:r="{_R}" r:id="rIdDrawing"/>'
    controls = controls or []
    if controls:
        sheet_tail += (
            f'<mc:AlternateContent xmlns:mc="{_MC}"><mc:Choice Requires="x14">'
            f'<controls xmlns:r="{_R}" xmlns:xdr="{_XDR}">'
            + "".join(controls)
            + "</controls></mc:Choice></mc:AlternateContent>"
        )
    package = zipfile.ZipFile(io.BytesIO(source))
    buffer = io.BytesIO()
    with package, zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
        for info in package.infolist():
            payload = package.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                payload = payload.replace(b"</worksheet>", sheet_tail.encode() + b"</worksheet>")
            target.writestr(info, payload)
        target.writestr("xl/drawings/drawing1.xml", drawing)
        target.writestr("xl/worksheets/_rels/sheet1.xml.rels", rels)
        for index, (kind, checked) in enumerate(props.values(), start=1):
            state = ' checked="Checked"' if checked else ""
            target.writestr(
                f"xl/ctrlProps/ctrlProp{index}.xml",
                '<?xml version="1.0" encoding="UTF-8"?><formControlPr '
                'xmlns="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main" '
                f'objectType="{kind}"{state} lockText="1"/>',
            )
    return buffer.getvalue()


def _sheet(source: bytes, **options: Any) -> SheetRecords:
    result = read_excel(source, options=options)
    assert result.document is not None, result.failure
    return result.document.sheets[0]


def _fields(sheet: SheetRecords) -> dict[str | None, Any]:
    return {block.label: block for block in sheet.blocks}


def test_check_boxes_with_labels_become_selected_value() -> None:
    # 項目B の右に「甲」（オフ）と「乙」（オン）のチェックボックス。図形は非表示（表示は VML）。
    source = _with_parts(
        _bytes(_form_workbook()),
        anchors=[
            _alternate(_two_cell((3, 5), (4, 9), _shape(1025, "甲", hidden=True))),
            _alternate(_two_cell((3, 9), (4, 13), _shape(1026, "乙", hidden=True))),
        ],
        controls=[_control(1025, "rIdC1"), _control(1026, "rIdC2")],
        props={"rIdC1": ("CheckBox", False), "rIdC2": ("CheckBox", True)},
    )
    fields = _fields(_sheet(source, mode="form"))
    field = fields["項目B"]
    assert (field.value, field.value_original) == ("乙", "☐ 甲 ☑ 乙")
    assert field.shape_cells == ["F4:I4", "J4:M4"]


def test_option_buttons_and_check_boxes_without_labels() -> None:
    workbook = _form_workbook()
    worksheet = workbook.active
    assert worksheet is not None
    # 項目C の右に、ラベルの無いオプションボタンを置いたセル F5・I5 と、選択肢の文字のセル。
    worksheet["G5"] = "甲"
    worksheet["J5"] = "乙"
    source = _with_parts(
        _bytes(workbook),
        anchors=[],
        # 図形の無いコントロール: controlPr の anchor の位置に記号だけを置く。
        controls=[_control(1030, "rIdR1", (4, 5, 5, 6)), _control(1031, "rIdR2", (4, 8, 5, 9))],
        props={"rIdR1": ("Radio", True), "rIdR2": ("Radio", False)},
    )
    field = _fields(_sheet(source, mode="form"))["項目C"]
    assert (field.value, field.value_original) == ("甲", "● 甲 ○ 乙")


def test_hidden_shapes_other_than_controls_are_not_read() -> None:
    source = _with_parts(
        _bytes(_form_workbook()),
        anchors=[
            _two_cell((3, 5), (4, 13), _shape(2, "見えない文", hidden=True)),
            _two_cell((7, 1), (9, 13), _shape(3, "見える文")),
        ],
    )
    values = [block.value for block in _sheet(source, mode="form").blocks]
    assert "見える文" in values
    assert "見えない文" not in values


def _one_cell(row: int, col: int, width: int, height: int, body: str) -> str:
    return (
        "<xdr:oneCellAnchor>"
        + _marker("from", row, col)
        + f'<xdr:ext cx="{width}" cy="{height}"/>'
        + body
        + "<xdr:clientData/></xdr:oneCellAnchor>"
    )


def _absolute(x: int, y: int, width: int, height: int, body: str) -> str:
    return (
        f'<xdr:absoluteAnchor><xdr:pos x="{x}" y="{y}"/><xdr:ext cx="{width}" cy="{height}"/>'
        + body
        + "<xdr:clientData/></xdr:absoluteAnchor>"
    )


def test_one_cell_and_absolute_anchors_cover_cells_by_size() -> None:
    source = _with_parts(
        _bytes(_form_workbook()),
        anchors=[
            # F4 から 9 列ぶんの幅（F4:N4。項目B の空欄に重ねる）。
            _one_cell(3, 5, 9 * _COLUMN_EMU, _ROW_EMU // 2, _shape(4, "大きさ固定の文")),
            # シートの左上から B9 の位置に、5 列 × 2 行の大きさ（B9:F10）。
            _absolute(
                _COLUMN_EMU + 10,
                8 * _ROW_EMU + 10,
                5 * _COLUMN_EMU - 20,
                2 * _ROW_EMU - 20,
                _shape(5, "絶対位置の文"),
            ),
        ],
    )
    sheet = _sheet(source, mode="form")
    fields = {block.value: block for block in sheet.blocks}
    assert fields["大きさ固定の文"].label == "項目B"
    assert fields["大きさ固定の文"].shape_cells == ["F4:N4"]
    assert fields["絶対位置の文"].shape_cells == ["B9:F10"]
    assert "shape_text_absolute_anchor_skipped" not in {item.code for item in sheet.diagnostics}


def test_anchor_box_uses_widths_heights_and_skips_hidden() -> None:
    grid = Grid(name="様式", hidden=False)
    pixels = 64 * 9525
    # 既定の列幅（64 ピクセル）・行の高さ（15 ポイント）。ちょうど境目で終わる図形は
    # 次の列を入れない。
    assert grid.anchor_box(1, 1, 0, 0, _ROW_EMU, 2 * pixels) == (1, 1, 1, 2)
    assert grid.anchor_box(1, 1, 0, 0, _ROW_EMU + 1, 2 * pixels + 1) == (1, 1, 2, 3)
    # オフセットがセルの幅を超えれば、次のセルから始まる。
    assert grid.anchor_box(1, 1, 0, pixels + 5, 1, 1) == (1, 2, 1, 2)
    # 非表示の列（幅 0）は飛ばし、指定の列幅・行の高さを使う。
    grid.hidden_columns = {2}
    grid.column_widths[3] = 2.5
    grid.row_heights[1] = 30.0
    assert grid.anchor_box(1, 1, 0, 0, 2 * _ROW_EMU, pixels + _COLUMN_EMU) == (1, 1, 1, 3)
