"""会話の添付の Excel を、RAG の取込と同じ読み方で読む（#2544）。

テストの Excel は openpyxl で合成する（業務の内容は書かない）。
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import UTC, datetime
from types import SimpleNamespace
from typing import Any

import anyio
import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from pr_system_settings.chat_attachments import ChatAttachmentService, InMemoryChatAttachmentStore
from pr_system_settings.chat_attachments.models import (
    EXCEL_PARSER_VERSION,
    ParsedAttachment,
    ParsedSheet,
)
from pr_system_settings.chat_attachments.parsing import FORM_COLUMNS, parse_document
from pr_system_settings.chat_attachments.service import WARNING_EXCEL_REPARSE_FAILED


def _bytes(workbook: openpyxl.Workbook) -> bytes:
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _put(worksheet: Worksheet, ref: str, value: Any) -> None:
    worksheet[ref.split(":", 1)[0]] = value
    if ":" in ref:
        worksheet.merge_cells(ref)


def _form_bytes() -> bytes:
    """方眼紙（狭い列の格子）の帳票。"""
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "様式"
    for col in range(1, 31):
        worksheet.column_dimensions[get_column_letter(col)].width = 2.5
    _put(worksheet, "B1:AC1", "様式A")
    _put(worksheet, "B3:E3", "項目A")
    _put(worksheet, "F3:N3", "値A")
    _put(worksheet, "B4:E4", "項目B")
    _put(worksheet, "F4:N4", "値B")
    _put(worksheet, "B5:E5", "項目C")
    _put(worksheet, "F5:N5", "値C")
    _put(worksheet, "B6:E6", "項目D")
    _put(worksheet, "F6:N6", "値D")
    return _bytes(workbook)


def _titled_tables_bytes() -> bytes:
    """表題の行の後に表があり、空の行を挟んで 2 つ目の表がある 1 枚のシート。"""
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "一覧"
    worksheet.append(["表の説明の行"])
    worksheet.append([])
    worksheet.append(["コード", "数量", "単価"])
    worksheet.append(["A-1", 3, 1200])
    worksheet["C4"].number_format = "#,##0"
    worksheet.append(["A-2", 5, 800])
    worksheet.append([])
    worksheet.append([])
    worksheet.append(["区分", "件数"])
    worksheet.append(["X", 10])
    worksheet.append(["Y", 20])
    return _bytes(workbook)


def _sheet(parsed: ParsedAttachment, name: str) -> ParsedSheet:
    return next(item for item in parsed.sheets if item.sheet == name)


def test_form_is_read_as_item_value_rows_with_cells() -> None:
    parsed = parse_document("form.xlsx", _form_bytes(), "excel")

    [sheet] = parsed.sheets
    assert (sheet.layout, sheet.columns) == ("form", list(FORM_COLUMNS))
    assert sheet.rows[:2] == [["項目A", "値A"], ["項目B", "値B"]]
    first = (sheet.locations or [])[0]
    assert (first.row_start, first.label_cell, first.value_cell) == (3, "B3:E3", "F3:N3")
    assert parsed.parser_version == EXCEL_PARSER_VERSION


def test_title_rows_and_multiple_tables_are_separated() -> None:
    parsed = parse_document("tables.xlsx", _titled_tables_bytes(), "excel")

    names = [sheet.sheet for sheet in parsed.sheets]
    assert len(names) == 2 and all(name.startswith("一覧（表 ") for name in names)
    first, second = parsed.sheets
    # 表頭は表題の行ではなく、表の列名の行。表示の書式（桁区切り）を保つ。
    assert first.columns == ["コード", "数量", "単価"]
    assert first.rows[0] == ["A-1", "3", "1,200"]
    assert [(item.row_start, item.cell_range) for item in first.locations or []] == [
        (4, "A4:C4"),
        (5, "A5:C5"),
    ]
    assert second.columns == ["区分", "件数"]
    assert second.rows == [["X", "10"], ["Y", "20"]]
    assert {sheet.source_sheet for sheet in parsed.sheets} == {"一覧"}
    # 表頭より上の説明の行は節にし、元の場所を持たせる。
    [section] = parsed.sections
    assert "表の説明の行" in section.text
    assert section.sheet == first.sheet
    assert section.location is not None and section.location.row_start == 1


def test_note_rows_below_a_table_are_a_section_not_a_table() -> None:
    """表の下の注記の行は別の表・列名にせず、注記の節にする（#3566）。文は合成。"""
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "一覧"
    worksheet.append(["コード", "数量", "単価"])
    worksheet.append(["A-1", 3, "1200※1"])
    worksheet.append(["A-2", 5, 800])
    worksheet.append([])
    worksheet.append(["※1 合成の注記の一つ目の文です。"])
    worksheet.append(["※2 合成の注記の二つ目の文です。"])
    parsed = parse_document("notes.xlsx", _bytes(workbook), "excel")

    [sheet] = parsed.sheets
    assert sheet.columns == ["コード", "数量", "単価"]
    assert len(sheet.rows) == 2
    [section] = parsed.sections
    assert section.title == "一覧 の注記"
    assert section.text == "※1 合成の注記の一つ目の文です。\n※2 合成の注記の二つ目の文です。"
    assert section.location is not None
    assert (section.location.row_start, section.location.row_end) == (5, 6)
    assert section.location.cell_range == "A5:A6"


@dataclass(frozen=True)
class _Principal:
    user_uuid: str
    is_system_admin: bool = False


ALICE = _Principal("user-a")


def _service(store: InMemoryChatAttachmentStore) -> ChatAttachmentService:
    return ChatAttachmentService(
        store,
        product="agent",
        settings=lambda: SimpleNamespace(),
        clock=lambda: datetime(2026, 10, 19, tzinfo=UTC),
    )


def _upload(service: ChatAttachmentService, name: str, data: bytes) -> str:
    async def run() -> Any:
        return await service.upload(ALICE, file_name=name, data=data)

    attachment_id = str(anyio.run(run).attachment_id)
    service.attach_to_conversation(
        ALICE, [attachment_id], conversation_id="c1", owner_user_uuid="user-a"
    )
    return attachment_id


def test_tools_return_cells_of_form_items() -> None:
    store = InMemoryChatAttachmentStore()
    service = _service(store)
    attachment_id = _upload(service, "form.xlsx", _form_bytes())
    toolset = service.toolset("c1")

    rows = toolset.call(
        "chat_attachment_read_rows",
        {"attachment_id": attachment_id, "sheet": "様式", "start_row": 1, "limit": 1},
    )
    found = toolset.call(
        "chat_attachment_find", {"text": "値C", "attachment_id": attachment_id, "limit": 5}
    )

    assert rows["rows"][0] == {
        "row_number": 1,
        "row_start": 3,
        "row_end": 3,
        "cell_range": rows["rows"][0]["cell_range"],
        "label_cell": "B3:E3",
        "value_cell": "F3:N3",
        "values": {"項目": "項目A", "値": "値A"},
    }
    [match] = found["matches"]
    assert (match["sheet"], match["row_start"], match["value_cell"]) == ("様式", 5, "F5:N5")


def _old_parsed() -> ParsedAttachment:
    """#2544 より前の形の結果（読み方の版・元の場所が無い）。"""
    return ParsedAttachment.model_validate(
        {
            "sections": [],
            "sheets": [{"sheet": "様式", "columns": ["様式A"], "rows": [["項目A"]]}],
            "warnings": [],
        }
    )


def _as_legacy(store: InMemoryChatAttachmentStore, attachment_id: str) -> None:
    """#2652 より前の行にする（アップロードの時点の古い読み方の結果。`parse_state` が無い）。"""
    old = _old_parsed()
    record = store.get(attachment_id)
    assert record is not None
    store._records[attachment_id] = record.model_copy(update={"outline": old.outline()})  # noqa: SLF001
    store._parsed[attachment_id] = old  # noqa: SLF001


def test_old_excel_results_are_reparsed_from_original() -> None:
    store = InMemoryChatAttachmentStore()
    service = _service(store)
    attachment_id = _upload(service, "form.xlsx", _form_bytes())
    _as_legacy(store, attachment_id)
    assert store.get(attachment_id).outline.parser_version is None  # type: ignore[union-attr]

    toolset = service.toolset("c1")
    parsed = toolset.parsed(attachment_id)

    assert parsed.parser_version == EXCEL_PARSER_VERSION and parsed.sheets[0].layout == "form"
    # 保存の結果も新しい読み方の結果に置き換わる。
    saved = store.parsed(attachment_id)
    assert saved is not None and saved.parser_version == EXCEL_PARSER_VERSION
    record = store.get(attachment_id)
    assert record is not None and record.outline.parse_status == "ready"


def test_reparse_failure_keeps_old_result_with_warning() -> None:
    store = InMemoryChatAttachmentStore()
    service = _service(store)
    attachment_id = _upload(service, "form.xlsx", _form_bytes())
    _as_legacy(store, attachment_id)
    store._content[attachment_id] = b"broken"  # noqa: SLF001 - 原本が読めない状態を作る

    parsed = service.toolset("c1").parsed(attachment_id)

    assert WARNING_EXCEL_REPARSE_FAILED in parsed.warnings
    assert parsed.sheets[0].columns == ["様式A"]
    # 保存の結果は変えない（失敗の状態だけを残す）。
    saved = store.parsed(attachment_id)
    assert saved is not None and saved.parser_version is None
    record = store.get(attachment_id)
    assert record is not None and record.outline.parse_status == "failed"
    # 2 回目も同じ（入力のせいの失敗は同じ原本と版では試し直さない）。
    again = service.toolset("c1").parsed(attachment_id)
    assert WARNING_EXCEL_REPARSE_FAILED in again.warnings
