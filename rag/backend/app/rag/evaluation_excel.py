"""品質評価の評価リクエストの Excel（#2065）。

画面の「評価リクエスト」は、テンプレートの Excel（1 行 = 1 つの評価のケース）を読み、評価の job の
API（`POST /api/evaluation/jobs/run`）のケース（`EvaluationCase`）にする。列は次の 3 つ。

- 質問（必須）: 利用者が尋ねる文（Ragas の `user_input`、Azure AI の `query`）。
- 期待する根拠（任意）: 回答の根拠になるべき資料（Ragas の `reference_contexts`、TREC の qrels）。
  1 行に 1 つ、`<ファイル名>`（か `文書ID: <ID>`）の後に `|` で区切って、任意で位置と本文を書く。
  位置は文書の形式ごとに、ページ（`p.12` / `p.12-14`。PDF にした Word・PowerPoint も）・スライド
  （`スライド: 3`。PDF にしたスライドのページ）・シートと行（`シート: <名前>` / `行: 10-12`。
  表計算）・
  章節（`章節: <見出し>`）。本文の抜粋（`抜粋: <本文の語句>`）はどの形式にも使える（#2086。W3C Web
  Annotation の selector〔ページ・範囲・引用の文〕、RFC 7111 の行の指定にならう）。ファイル名は
  選んだナレッジベースの文書の名前と完全に一致させる（NFKC・大小文字・前後の空白は無視）。
- 標準回答（任意）: 期待する回答（Ragas の `reference`、Azure AI の `ground_truth`）。

期待する根拠と標準回答の両方が空の行は、資料に答えが無い質問（回答を控えるのが正しい）として扱う
（`EvaluationCase.expects_answer`）。

照合の規則:
- ファイル名だけの根拠は正解の文書（`relevant_document_ids`）にし、正解文書の再現率・順位で測る。
- ページ・章節・抜粋のある根拠は、正解の文書に加えて必要な根拠（`required_evidence`）にし、回答の
  引用にその条件をすべて満たす chunk があるかで測る（`evaluation.evidence_matches_chunk`）。

読み込みは 2 段で、構造の検証（`parse_evaluation_workbook`。ナレッジベースを見ない）と、文書の
解決（`build_evaluation_cases`。選んだナレッジベースの文書を、文書 ID があれば ID で、無ければ
ファイル名で引く）に分ける。どちらの誤りも行・列・理由の形（`EvaluationImportIssue`）で返し、
1 つでもあればケースを作らない。文書の形式に合わない位置の条件（表計算にページなど）と、同じ名前の
文書が複数あるファイル名は警告にする（評価は実行できる。同じ名前の文書は、どれが引用されても一致とする）。

テンプレート（`template_xlsx`）は「評価リクエスト」「記入の説明」「記入例」「文書一覧」の 4 つの
シートを持つ。記入例は別のシートに置き、取込は「評価リクエスト」のシートだけを読む（例の行を評価しない）。
「文書一覧」は、選んだナレッジベースの文書（ファイル名・文書 ID・形式・ページ数 / シート名・
アップロード日）で、「期待する根拠」の列の入力規則の選択肢にする。
"""

from __future__ import annotations

import re
import unicodedata
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass, field
from datetime import datetime
from io import BytesIO
from pathlib import PurePath
from typing import Any, Literal

# openpyxl は rag-engine[excel] が入れる（新しい依存は足さない）。
from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]
from openpyxl.comments import Comment  # type: ignore[import-untyped]
from openpyxl.styles import Alignment, Font, PatternFill  # type: ignore[import-untyped]
from openpyxl.worksheet.datavalidation import DataValidation  # type: ignore[import-untyped]
from pydantic import ValidationError

from app.schemas.evaluation import (
    EVIDENCE_LABEL_MAX_CHARS,
    EVIDENCE_SECTION_MAX_CHARS,
    EVIDENCE_TEXT_MAX_CHARS,
    STANDARD_ANSWER_MAX_CHARS,
    EvaluationCase,
    EvaluationEvidence,
    EvaluationImportIssue,
)

# アップロードの上限。評価は 1 行ごとに回答を生成するため、行数は job の時間の上限
# （`evaluation_job_time_limit_seconds`）に収まる量にする。
EVALUATION_EXCEL_MAX_BYTES = 10 * 1024 * 1024
EVALUATION_EXCEL_MAX_ROWS = 200
# 1 行の質問の長さの上限（文字）。
EVALUATION_QUESTION_MAX_CHARS = 2000
# 1 つのセルの「期待する根拠」の行数の上限（`EvaluationCase.required_evidence` の上限と同じ）。
EVALUATION_EVIDENCE_MAX_LINES = 30
# 画面に返す誤りの件数の上限（それより多い分は件数だけを返す）。
EVALUATION_IMPORT_MAX_ISSUES = 100
# ページの番号の上限（Excel の書き間違いを止める）。
_MAX_PAGE = 100_000

CASES_SHEET_TITLE = "評価リクエスト"
GUIDE_SHEET_TITLE = "記入の説明"
EXAMPLES_SHEET_TITLE = "記入例"
DOCUMENTS_SHEET_TITLE = "文書一覧"
# 取込で読まないシート（ケースのシートが見つからないときの候補から外す）。
_AUXILIARY_SHEET_TITLES = frozenset(
    {GUIDE_SHEET_TITLE, EXAMPLES_SHEET_TITLE, DOCUMENTS_SHEET_TITLE}
)
QUESTION_HEADER = "質問"
EVIDENCE_HEADER = "期待する根拠"
STANDARD_ANSWER_HEADER = "標準回答"
HEADERS = (QUESTION_HEADER, EVIDENCE_HEADER, STANDARD_ANSWER_HEADER)
TEMPLATE_FILE_NAME = "evaluation_request_template.xlsx"
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# 見出しの別名（業界の評価データの列名でも読めるようにする。比較は NFKC・小文字・空白なし）。
_HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "question": (QUESTION_HEADER, "question", "query", "user_input", "質問文"),
    "evidence": (
        EVIDENCE_HEADER,
        "根拠",
        "正解の根拠",
        "reference_contexts",
        "contexts",
        "context",
        "evidence",
    ),
    "standard_answer": (
        STANDARD_ANSWER_HEADER,
        "回答",
        "期待する回答",
        "reference",
        "ground_truth",
        "expected_answer",
        "standard_answer",
    ),
}
_COLUMN_LABELS = {
    "question": QUESTION_HEADER,
    "evidence": EVIDENCE_HEADER,
    "standard_answer": STANDARD_ANSWER_HEADER,
}
# 見出しの末尾の注記（「質問（必須）」など）は無視する。
_HEADER_NOTE = re.compile(r"[（(][^）)]*[）)]\s*$")
_PART_SEPARATOR = re.compile(r"\s*[|｜]\s*")
_PAGE_PATTERN = re.compile(
    r"(?:p\.?|pp\.?|page|ページ|頁)?\s*(\d+)\s*(?:[-~〜～–—]\s*(\d+))?\s*(?:ページ|頁)?",
    re.IGNORECASE,
)
_NUMBER_RANGE_PATTERN = re.compile(r"(\d+)\s*(?:[-~〜～–—]\s*(\d+))?")
_SECTION_PREFIXES = ("章節", "見出し", "section")
_EXCERPT_PREFIXES = ("抜粋", "本文", "excerpt", "text")
_DOCUMENT_ID_PREFIXES = ("文書ID", "文書 ID", "document_id", "document id", "doc_id")
_SHEET_PREFIXES = ("シート", "sheet")
_ROW_PREFIXES = ("行", "rows", "row")
_SLIDE_PREFIXES = ("スライド", "slide")
# 行の番号の上限（Excel の行数の上限）。
_MAX_ROW = 1_048_576
_LABELLED_PREFIXES = (
    _SECTION_PREFIXES
    + _EXCERPT_PREFIXES
    + _DOCUMENT_ID_PREFIXES
    + _SHEET_PREFIXES
    + _ROW_PREFIXES
    + _SLIDE_PREFIXES
)

# 文書の形式（位置の書き方を決める。#2086）。ページのある形式（PDF・PDF にした Word と画像）、
# スライド（PDF にした PowerPoint。スライドの番号がページの番号）、表計算（シート・行）、
# ページの無い文書（Markdown・テキスト・HTML・メール）、不明（警告しない）。
DocumentFormat = Literal["paged", "slides", "sheet", "text", "unknown"]
_FORMAT_BY_EXTENSION: dict[str, DocumentFormat] = {
    **dict.fromkeys((".pdf", ".docx", ".doc", ".rtf", ".odt"), "paged"),
    **dict.fromkeys(
        (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff", ".webp", ".heic"), "paged"
    ),
    **dict.fromkeys((".pptx", ".ppt", ".odp"), "slides"),
    **dict.fromkeys((".xlsx", ".xlsm", ".xls", ".ods", ".csv", ".tsv"), "sheet"),
    **dict.fromkeys(
        (".md", ".markdown", ".txt", ".text", ".html", ".htm", ".eml", ".msg", ".json", ".xml"),
        "text",
    ),
}
_FORMAT_LABELS: dict[str, str] = {
    ".pdf": "PDF",
    ".docx": "Word",
    ".doc": "Word",
    ".pptx": "PowerPoint",
    ".ppt": "PowerPoint",
    ".xlsx": "表計算（Excel）",
    ".xlsm": "表計算（Excel）",
    ".xls": "表計算（Excel）",
    ".csv": "表計算（CSV）",
    ".tsv": "表計算（TSV）",
    ".md": "Markdown",
    ".markdown": "Markdown",
    ".txt": "テキスト",
    ".html": "HTML",
    ".htm": "HTML",
    ".eml": "メール",
    ".msg": "メール",
}


@dataclass(frozen=True)
class EvidenceSpec:
    """「期待する根拠」の 1 行を読んだもの（文書はまだ解決していない）。

    文書は `file_name` か `document_id`（両方なら ID で引く）。`slide` はページの番号をスライドの
    番号として書いたか（形式の警告に使う）。
    """

    line: str
    file_name: str | None = None
    document_id: str | None = None
    page_start: int | None = None
    page_end: int | None = None
    slide: bool = False
    sheet: str | None = None
    row_start: int | None = None
    row_end: int | None = None
    section: str | None = None
    text: str | None = None

    @property
    def has_conditions(self) -> bool:
        return any(
            value is not None
            for value in (self.page_start, self.sheet, self.row_start, self.section, self.text)
        )

    @property
    def document_label(self) -> str:
        """誤り・警告の文に出す文書の名前（ファイル名、無ければ文書 ID）。"""
        return self.file_name or f"文書ID: {self.document_id}"


@dataclass(frozen=True)
class EvaluationExcelRow:
    """Excel の 1 行（構造の検証を通ったもの）。"""

    row: int
    question: str
    evidence: tuple[EvidenceSpec, ...] = ()
    standard_answer: str | None = None

    @property
    def case_id(self) -> str:
        return evaluation_case_id(self.row)


@dataclass(frozen=True)
class EvaluationDocumentRef:
    """文書の解決とテンプレートの「文書一覧」に使う、ナレッジベースの文書の参照。"""

    document_id: str
    file_name: str
    superseded_by_file_name: str | None = None
    superseded: bool = False
    content_type: str | None = None
    uploaded_at: datetime | None = None
    # 「文書一覧」に出すページ数・シート名（chunk から求める。分からなければ None / 空）。
    page_count: int | None = None
    sheet_names: tuple[str, ...] = ()

    @property
    def format(self) -> DocumentFormat:
        return document_format(self.file_name, self.content_type)


def document_format(file_name: str, content_type: str | None = None) -> DocumentFormat:
    """文書の形式（拡張子、無ければ MIME type で決める。分からなければ unknown）。"""
    found = _FORMAT_BY_EXTENSION.get(PurePath(file_name).suffix.lower())
    if found is not None:
        return found
    media = (content_type or "").split(";", 1)[0].strip().lower()
    if media == "application/pdf" or media.startswith("image/"):
        return "paged"
    if media in {"text/csv", "text/tab-separated-values"} or "spreadsheet" in media:
        return "sheet"
    if "presentation" in media:
        return "slides"
    if media.startswith("text/") or media == "message/rfc822":
        return "text"
    return "unknown"


def document_format_label(file_name: str, content_type: str | None = None) -> str:
    """「文書一覧」に出す形式の名前。"""
    label = _FORMAT_LABELS.get(PurePath(file_name).suffix.lower())
    if label is not None:
        return label
    return {
        "paged": "PDF・画像",
        "slides": "PowerPoint",
        "sheet": "表計算",
        "text": "テキスト",
    }.get(document_format(file_name, content_type), "その他")


@dataclass
class EvaluationExcelParse:
    """Excel の構造の検証の結果。"""

    rows: list[EvaluationExcelRow] = field(default_factory=list)
    issues: list[EvaluationImportIssue] = field(default_factory=list)


def evaluation_case_id(row: int) -> str:
    """Excel の行番号から作るケースの id（結果の表で行に戻せるようにする）。"""
    return f"row-{row}"


def normalize_file_name(value: str) -> str:
    """ファイル名の比較の形（NFKC・前後の空白なし・大小文字を区別しない）。"""
    return unicodedata.normalize("NFKC", value).strip().casefold()


def _normalized_header(value: object) -> str:
    text = unicodedata.normalize("NFKC", str(value or "")).strip()
    text = _HEADER_NOTE.sub("", text).strip()
    return re.sub(r"\s+", "", text).casefold()


_ALIAS_LOOKUP = {
    _normalized_header(alias): key for key, aliases in _HEADER_ALIASES.items() for alias in aliases
}


def _cell_text(value: object) -> str:
    """セルの値を文字にする（数値・日付も文字にする。改行は LF にそろえる）。"""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).replace("\r\n", "\n").replace("\r", "\n").strip()


def parse_evaluation_workbook(data: bytes) -> EvaluationExcelParse:
    """Excel の構造（見出し・必須・長さ・行数・根拠の書き方）を検証して行を読む。

    シートは「評価リクエスト」（無ければ最初のシート）。1 行目が見出し、2 行目からがケース。
    すべての列が空の行は飛ばす。誤りは行・列・理由で返す（`rows` は誤りの無い行だけ）。
    """
    result = EvaluationExcelParse()
    if len(data) > EVALUATION_EXCEL_MAX_BYTES:
        result.issues.append(
            _file_issue(
                f"ファイルが大きすぎます（上限 {EVALUATION_EXCEL_MAX_BYTES // (1024 * 1024)}MB）。"
                "行を分けて複数のファイルにしてください。"
            )
        )
        return result
    try:
        book = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 - 形式の誤りは利用者へ返す
        result.issues.append(
            _file_issue("Excel（.xlsx）として読めません。テンプレートの形式で保存してください。")
        )
        return result
    try:
        sheet = _cases_sheet(book)
        if sheet is None:
            result.issues.append(_file_issue("シートがありません。"))
            return result
        rows = sheet.iter_rows(values_only=True)
        first: tuple[object, ...] | None = next(rows, None)
        header = tuple(first or ())
        columns = _header_columns(header)
        if "question" not in columns:
            result.issues.append(
                EvaluationImportIssue(
                    row=1,
                    column=None,
                    message=(
                        f"1 行目に「{QUESTION_HEADER}」の見出しがありません。"
                        "テンプレートの見出し（質問・期待する根拠・標準回答）を使ってください。"
                    ),
                )
            )
            return result
        data_rows = 0
        for number, values in enumerate(rows, start=2):
            cells = {key: _cell_text(_at(values, index)) for key, index in columns.items()}
            if not any(cells.values()):
                continue
            data_rows += 1
            if data_rows > EVALUATION_EXCEL_MAX_ROWS:
                result.issues.append(
                    _file_issue(
                        f"評価のケースは {EVALUATION_EXCEL_MAX_ROWS} 行までです。"
                        "ファイルを分けて評価してください。"
                    )
                )
                result.rows.clear()
                return result
            row, issues = _parse_row(number, cells)
            result.issues.extend(issues)
            if row is not None:
                result.rows.append(row)
        if data_rows == 0:
            result.issues.append(
                _file_issue("評価のケースが 1 行もありません。2 行目から質問を書いてください。")
            )
        return result
    finally:
        book.close()


def _cases_sheet(book: Any) -> Any:
    for sheet in book.worksheets:
        if str(sheet.title).strip() == CASES_SHEET_TITLE:
            return sheet
    for sheet in book.worksheets:
        if str(sheet.title).strip() not in _AUXILIARY_SHEET_TITLES:
            return sheet
    return None


def _header_columns(header: Sequence[object]) -> dict[str, int]:
    found: dict[str, int] = {}
    for index, value in enumerate(header):
        key = _ALIAS_LOOKUP.get(_normalized_header(value))
        if key and key not in found:
            found[key] = index
    return found


def _at(values: Sequence[object], index: int) -> object:
    return values[index] if index < len(values) else None


def _file_issue(message: str) -> EvaluationImportIssue:
    return EvaluationImportIssue(row=None, column=None, message=message)


def _parse_row(
    number: int, cells: Mapping[str, str]
) -> tuple[EvaluationExcelRow | None, list[EvaluationImportIssue]]:
    issues: list[EvaluationImportIssue] = []

    def issue(column: str, message: str) -> None:
        issues.append(
            EvaluationImportIssue(row=number, column=_COLUMN_LABELS[column], message=message)
        )

    question = cells.get("question", "")
    if not question:
        issue("question", "質問を入力してください。")
    elif len(question) > EVALUATION_QUESTION_MAX_CHARS:
        issue(
            "question",
            f"質問は {EVALUATION_QUESTION_MAX_CHARS} 文字以内にしてください"
            f"（{len(question)} 文字）。",
        )
    standard_answer = cells.get("standard_answer", "") or None
    if standard_answer is not None and len(standard_answer) > STANDARD_ANSWER_MAX_CHARS:
        issue(
            "standard_answer",
            f"標準回答は {STANDARD_ANSWER_MAX_CHARS} 文字以内にしてください"
            f"（{len(standard_answer)} 文字）。",
        )
    lines = [line.strip() for line in cells.get("evidence", "").split("\n") if line.strip()]
    if len(lines) > EVALUATION_EVIDENCE_MAX_LINES:
        issue(
            "evidence",
            f"期待する根拠は {EVALUATION_EVIDENCE_MAX_LINES} 行までにしてください"
            f"（{len(lines)} 行）。",
        )
        lines = []
    evidence: list[EvidenceSpec] = []
    for index, line in enumerate(lines, start=1):
        parsed = parse_evidence_line(line)
        if isinstance(parsed, str):
            issue("evidence", f"{index} 行目: {parsed}")
        else:
            evidence.append(parsed)
    if issues:
        return None, issues
    return (
        EvaluationExcelRow(
            row=number,
            question=question,
            evidence=tuple(evidence),
            standard_answer=standard_answer,
        ),
        [],
    )


def parse_evidence_line(line: str) -> EvidenceSpec | str:
    """「期待する根拠」の 1 行を読む。読めなければ理由の文を返す。

    形: `<ファイル名> | p.12 | 章節: <見出し> | 抜粋: <本文の語句>`
    （ファイル名の後は任意・順不同）。
    ファイル名の代わりに `文書ID: <ID>` も書ける（ファイル名の後に書けば、ID で引く）。位置は
    ページ（`p.12`）・スライド（`スライド: 3`）・シート（`シート: <名前>`）・行（`行: 10-12`）。
    """
    if len(line) > EVIDENCE_LABEL_MAX_CHARS:
        return f"1 行は {EVIDENCE_LABEL_MAX_CHARS} 文字以内にしてください。"
    head, *parts = _PART_SEPARATOR.split(line.strip())
    head = head.strip()
    file_name: str | None = head
    document_id: str | None = None
    if (value := _labelled_value(head, _DOCUMENT_ID_PREFIXES)) is not None:
        if not value:
            return "「文書ID:」の後に文書の ID を書いてください。"
        file_name, document_id = None, value
    elif not head:
        return (
            "先頭に文書のファイル名か「文書ID: 〈ID〉」を書いてください"
            "（例: 製品マニュアル.pdf | p.3）。"
        )
    elif _labelled_value(head, _LABELLED_PREFIXES) is not None or _page_range(head) is not None:
        return (
            "先頭に文書のファイル名か「文書ID: 〈ID〉」を書いてください"
            "（ページ・シート・章節・抜粋はファイル名の後に書きます）。"
        )
    values: dict[str, str] = {}
    page: tuple[int, int | None] | None = None
    slide = False
    rows: tuple[int, int | None] | None = None
    for raw in parts:
        part = raw.strip()
        if not part:
            continue
        labelled = _labelled_part(part)
        if labelled is not None:
            kind, value = labelled
            if kind in {"slide", "row"}:
                parsed = _number_range(value, kind=kind, part=part)
                if isinstance(parsed, str):
                    return parsed
                if kind == "row":
                    if rows is not None:
                        return (
                            "行は 1 行に 1 つにしてください"
                            "（範囲は「行: 10-12」のように書きます）。"
                        )
                    rows = parsed
                    continue
                if page is not None:
                    return "ページとスライドは 1 行にどちらか 1 つにしてください。"
                page, slide = parsed, True
                continue
            error = _labelled_error(kind, value, values)
            if error:
                return error
            values[kind] = value
        elif (pages := _page_range(part)) is not None:
            if page is not None:
                return (
                    "ページとスライドは 1 行に 1 つにしてください"
                    "（範囲は p.12-14 のように書きます）。"
                )
            parsed = _checked_range(pages, limit=_MAX_PAGE, part=part, kind="page")
            if isinstance(parsed, str):
                return parsed
            page = parsed
        else:
            return (
                f"「{part}」を読めません。ページは「p.12」「p.12-14」、スライドは「スライド: 3」、"
                "表計算は「シート: 名前」「行: 10-12」、章節は「章節: 見出し」、"
                "本文は「抜粋: 本文の語句」と書いてください。"
            )
    if "document_id" in values:
        document_id = values["document_id"]
    return EvidenceSpec(
        line=line.strip(),
        file_name=file_name,
        document_id=document_id,
        page_start=page[0] if page else None,
        page_end=page[1] if page else None,
        slide=slide,
        sheet=values.get("sheet"),
        row_start=rows[0] if rows else None,
        row_end=rows[1] if rows else None,
        section=values.get("section"),
        text=values.get("text"),
    )


_LABEL_KINDS: tuple[tuple[str, tuple[str, ...]], ...] = (
    ("document_id", _DOCUMENT_ID_PREFIXES),
    ("section", _SECTION_PREFIXES),
    ("text", _EXCERPT_PREFIXES),
    ("sheet", _SHEET_PREFIXES),
    ("slide", _SLIDE_PREFIXES),
    ("row", _ROW_PREFIXES),
)
_LABEL_NAMES = {
    "document_id": "文書ID",
    "section": "章節",
    "text": "抜粋",
    "sheet": "シート",
}
_LABEL_HINTS = {
    "document_id": "文書の ID",
    "section": "見出し",
    "text": "本文の語句",
    "sheet": "シートの名前",
}
_LABEL_MAX_CHARS = {
    "document_id": 200,
    "section": EVIDENCE_SECTION_MAX_CHARS,
    "text": EVIDENCE_TEXT_MAX_CHARS,
    "sheet": EVIDENCE_SECTION_MAX_CHARS,
}


def _labelled_part(part: str) -> tuple[str, str] | None:
    """`<接頭辞>: <値>` の種類と値（どの接頭辞でもなければ None）。"""
    for kind, prefixes in _LABEL_KINDS:
        value = _labelled_value(part, prefixes)
        if value is not None:
            return kind, value
    return None


def _labelled_error(kind: str, value: str, found: Mapping[str, str]) -> str:
    """文書ID・章節・抜粋・シートの値の誤り（無ければ空）。"""
    name = _LABEL_NAMES[kind]
    if kind in found:
        return f"{name}は 1 行に 1 つにしてください（複数の{name}は行を分けて書きます）。"
    if not value:
        return f"「{name}:」の後に{_LABEL_HINTS[kind]}を書いてください。"
    if len(value) > _LABEL_MAX_CHARS[kind]:
        return f"{name}は {_LABEL_MAX_CHARS[kind]} 文字以内にしてください。"
    return ""


def _number_range(value: str, *, kind: str, part: str) -> tuple[int, int | None] | str:
    """「スライド:」「行:」の番号（`3` / `10-12`）。読めなければ理由の文。"""
    name = "スライド" if kind == "slide" else "行"
    match = _NUMBER_RANGE_PATTERN.fullmatch(value.strip())
    if match is None:
        return (
            f"「{part}」を読めません。"
            f"{name}は「{name}: 3」「{name}: 10-12」のように番号で書いてください。"
        )
    end = match.group(2)
    return _checked_range(
        (int(match.group(1)), int(end) if end is not None else None),
        limit=_MAX_ROW if kind == "row" else _MAX_PAGE,
        part=part,
        kind=kind,
    )


def _checked_range(
    pages: tuple[int, int | None], *, limit: int, part: str, kind: str
) -> tuple[int, int | None] | str:
    """番号の範囲の向きと上限を確かめる（終わりが始まりと同じなら終わりを省く）。"""
    name = {"page": "ページ", "slide": "スライド", "row": "行"}[kind]
    start, end = pages
    if start < 1 or start > limit or (end is not None and end > limit):
        return f"{name}「{part}」は 1 以上の番号にしてください。"
    if end is not None and end < start:
        example = "p.12-14" if kind == "page" else f"{name}: 10-12"
        return f"{name}「{part}」の範囲は、小さい番号から書いてください（例: {example}）。"
    return start, end if end != start else None


def _labelled_value(part: str, prefixes: Iterable[str]) -> str | None:
    """`<接頭辞>: <値>` の値（接頭辞が違えば None）。コロンは全角も受け付ける。"""
    normalized = unicodedata.normalize("NFKC", part)
    for prefix in prefixes:
        if normalized.casefold().startswith(prefix.casefold()):
            rest = normalized[len(prefix) :].lstrip()
            if rest.startswith(":"):
                return rest[1:].strip()
    return None


def _page_range(part: str) -> tuple[int, int | None] | None:
    match = _PAGE_PATTERN.fullmatch(unicodedata.normalize("NFKC", part).strip())
    if match is None:
        return None
    end = match.group(2)
    return int(match.group(1)), int(end) if end is not None else None


@dataclass
class EvaluationCasesBuild:
    """文書を解決したケースと、誤り・警告（#2086）。誤りが 1 つでもあればケースを使わない。"""

    cases: list[EvaluationCase] = field(default_factory=list)
    issues: list[EvaluationImportIssue] = field(default_factory=list)
    warnings: list[EvaluationImportIssue] = field(default_factory=list)


def _evidence_issue(row: int, index: int, message: str) -> EvaluationImportIssue:
    return EvaluationImportIssue(
        row=row, column=EVIDENCE_HEADER, message=f"{index} 行目: {message}"
    )


@dataclass(frozen=True)
class _ResolvedDocument:
    document_id: str
    equivalents: tuple[str, ...] = ()
    document: EvaluationDocumentRef | None = None
    warning: str = ""


def build_evaluation_cases(
    rows: Sequence[EvaluationExcelRow],
    documents: Sequence[EvaluationDocumentRef],
) -> EvaluationCasesBuild:
    """読んだ行を評価のケースにする（根拠の文書を、選んだナレッジベースの文書 ID にする）。

    `文書ID:` があれば ID で、無ければファイル名（完全一致。`normalize_file_name`）で引く。
    新しい版に置き換えた文書（旧版）は検索の対象外のため、今有効な版があればそれを使い、旧版しか無ければ誤りに
    する。同じ名前の有効な文書が複数あるファイル名は警告にし、どれが引用されても一致とする（同じ名前は
    同じ資料の版のことが多い。区別するときは文書 ID で書く。#2086）。文書の形式に合わない位置の条件
    （表計算にページ、ページの無い文書にページ、表計算でない文書にシート・行）も警告にする。
    """
    by_name: dict[str, list[EvaluationDocumentRef]] = {}
    by_id: dict[str, EvaluationDocumentRef] = {}
    for document in documents:
        by_name.setdefault(normalize_file_name(document.file_name), []).append(document)
        by_id[document.document_id] = document
    result = EvaluationCasesBuild()
    for row in rows:
        relevant: list[str] = []
        equivalents: dict[str, list[str]] = {}
        required: list[EvaluationEvidence] = []
        row_issues: list[EvaluationImportIssue] = []
        for index, spec in enumerate(row.evidence, start=1):
            resolved, error = _resolve_document(spec, by_name, by_id)
            if resolved is None:
                row_issues.append(_evidence_issue(row.row, index, error))
                continue
            if resolved.warning:
                result.warnings.append(_evidence_issue(row.row, index, resolved.warning))
            if resolved.document is not None and (
                format_warning := _format_warning(spec, resolved.document)
            ):
                result.warnings.append(_evidence_issue(row.row, index, format_warning))
            if resolved.document_id not in relevant:
                relevant.append(resolved.document_id)
            if resolved.equivalents:
                merged = equivalents.setdefault(resolved.document_id, [])
                merged.extend(item for item in resolved.equivalents if item not in merged)
            if spec.has_conditions:
                required.append(
                    EvaluationEvidence(
                        id=f"e{index}",
                        document_id=resolved.document_id,
                        text=spec.text,
                        page_start=spec.page_start,
                        page_end=spec.page_end,
                        sheet=spec.sheet,
                        row_start=spec.row_start,
                        row_end=spec.row_end,
                        section=spec.section,
                        label=spec.line,
                    )
                )
        if row_issues:
            result.issues.extend(row_issues)
            continue
        try:
            result.cases.append(
                EvaluationCase(
                    id=row.case_id,
                    query=row.question,
                    relevant_document_ids=relevant,
                    equivalent_document_ids=equivalents,
                    required_evidence=required,
                    standard_answer=row.standard_answer,
                )
            )
        except ValidationError as exc:
            result.issues.append(
                EvaluationImportIssue(row=row.row, column=None, message=_validation_message(exc))
            )
    return result


def _resolve_document(
    spec: EvidenceSpec,
    by_name: Mapping[str, Sequence[EvaluationDocumentRef]],
    by_id: Mapping[str, EvaluationDocumentRef],
) -> tuple[_ResolvedDocument | None, str]:
    """根拠の文書と、解決できないときの理由（解決できたら空）。"""
    if spec.document_id is not None:
        document = by_id.get(spec.document_id)
        if document is None:
            return None, (
                f"文書 ID「{spec.document_id}」の文書は、選んだナレッジベースにありません"
                "（ID は「文書一覧」のシートで確かめられます）。"
            )
        if document.superseded:
            newer = document.superseded_by_file_name
            hint = f"新しい版（{newer}）" if newer else "新しい版"
            return None, (
                f"文書 ID「{spec.document_id}」は{hint}に置き換えた文書です。"
                "新しい版の文書 ID を書いてください。"
            )
        warning = ""
        if spec.file_name is not None and normalize_file_name(
            spec.file_name
        ) != normalize_file_name(document.file_name):
            warning = (
                f"文書 ID「{spec.document_id}」の文書のファイル名は「{document.file_name}」です"
                f"（「{spec.file_name}」ではありません）。文書 ID の文書で照合します。"
            )
        return _ResolvedDocument(document.document_id, document=document, warning=warning), ""
    file_name = spec.file_name or ""
    candidates = by_name.get(normalize_file_name(file_name), ())
    if not candidates:
        return None, (
            f"「{file_name}」は選んだナレッジベースの文書にありません"
            "（ファイル名を完全に一致させるか、「文書ID: 〈ID〉」で書いてください）。"
        )
    current = [item for item in candidates if not item.superseded]
    if not current:
        newer = next(
            (item.superseded_by_file_name for item in candidates if item.superseded_by_file_name),
            None,
        )
        hint = f"新しい版（{newer}）" if newer else "新しい版"
        return None, (
            f"「{file_name}」は{hint}に置き換えた文書です。新しい版のファイル名を書いてください。"
        )
    # 同じ名前の文書は、新しくアップロードしたものを正解の文書にし、ほかを同じ文書とみなす。
    ordered = sorted(
        {item.document_id: item for item in current}.values(),
        key=lambda item: -(item.uploaded_at.timestamp() if item.uploaded_at else 0.0),
    )
    first, *others = ordered
    warning = ""
    if others:
        warning = (
            f"「{file_name}」という名前の文書が {len(ordered)} 件あります。"
            "どれか 1 つが引用されれば一致とみなします。"
            "区別するときは「文書ID: 〈ID〉」で書いてください"
            "（ID は「文書一覧」のシートにあります）。"
        )
    return (
        _ResolvedDocument(
            first.document_id,
            equivalents=tuple(item.document_id for item in others),
            document=first,
            warning=warning,
        ),
        "",
    )


def _format_warning(spec: EvidenceSpec, document: EvaluationDocumentRef) -> str:
    """文書の形式に合わない位置の条件の警告（合っていれば空。#2086）。"""
    kind = document.format
    name = document.file_name
    if spec.page_start is not None and kind == "sheet":
        return (
            f"「{name}」は表計算の文書で、ページがありません。"
            "「シート: 名前 | 行: 10-12」か「抜粋:」で指定してください。"
        )
    if spec.page_start is not None and kind == "text":
        return f"「{name}」にはページがありません。「章節: 見出し」か「抜粋:」で指定してください。"
    if (spec.sheet is not None or spec.row_start is not None) and kind in {
        "paged",
        "slides",
        "text",
    }:
        return (
            f"「{name}」は表計算の文書ではないため、シート・行では一致しません。"
            "ページ・章節か「抜粋:」で指定してください。"
        )
    return ""


def _validation_message(exc: ValidationError) -> str:
    messages = [str(error.get("msg", "")).removeprefix("Value error, ") for error in exc.errors()]
    return (
        "".join(message for message in messages if message) or "この行を評価のケースにできません。"
    )


# --- 書き出し（テンプレートと、フィードバックから作る 1 行） ------------------------------------

_HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
_HEADER_NOTES = {
    QUESTION_HEADER: "必須。利用者が尋ねる文を 1 セルに 1 つ書きます。",
    EVIDENCE_HEADER: (
        "任意。回答の根拠になるべき資料を 1 行に 1 つ書きます（セルの中で改行して複数）。"
        "形: ファイル名 | 位置 | 抜粋: 本文の語句。位置は PDF・Word が p.12、PowerPoint が"
        "スライド: 3、表計算がシート: 名前 | 行: 10-12、Markdown などが章節: 見出し。"
        "ファイル名の代わりに文書ID: 〈ID〉も書けます（「文書一覧」のシート）。"
    ),
    STANDARD_ANSWER_HEADER: "任意。期待する回答を書きます。回答をこの文と比べて採点します。",
}
_COLUMN_WIDTHS = (48, 56, 64)
_GUIDE_HEADINGS = frozenset({"照合の規則", "形式ごとの「期待する根拠」の書き方", "上限", "注意"})

# 「記入の説明」のシート（列の説明・書き方・照合の規則・形式ごとの書き方）。例は合成の内容。
_GUIDE_ROWS: tuple[tuple[str, ...], ...] = (
    ("列", "必須 / 任意", "意味", "書き方"),
    (
        QUESTION_HEADER,
        "必須",
        "利用者が尋ねる文。業界の評価データの question / query / user_input にあたります。",
        f"1 セルに 1 つ。{EVALUATION_QUESTION_MAX_CHARS} 文字まで。",
    ),
    (
        EVIDENCE_HEADER,
        "任意",
        "回答の根拠になるべき資料の場所と内容（正解の文書・ページ・シートと行・章節・本文の抜粋）。"
        "reference_contexts / qrels にあたります。",
        f"1 行に 1 つの根拠（セルの中で改行して複数。{EVALUATION_EVIDENCE_MAX_LINES} 行まで）。"
        "先頭に文書のファイル名（選ぶナレッジベースの文書の名前と完全に一致）か「文書ID: 〈ID〉」を"
        "書き、続けて「|」で区切って、任意で位置（下の「形式ごとの書き方」）と、本文の抜粋"
        f"「抜粋: 本文の語句」（{EVIDENCE_TEXT_MAX_CHARS} 文字まで）を書きます。",
    ),
    (
        STANDARD_ANSWER_HEADER,
        "任意",
        "期待する回答。reference / ground_truth にあたります。",
        f"{STANDARD_ANSWER_MAX_CHARS} 文字まで。書くと、回答をこの文と比べて採点します"
        "（標準回答での合格・主張の裏付け・必要な内容の網羅）。",
    ),
    (),
    ("照合の規則",),
    (
        "ファイル名だけ",
        "",
        "その文書が回答に引用されたか（正解文書の再現率・順位）で測ります。",
        "例: 製品マニュアル.pdf",
    ),
    (
        "位置・抜粋つき",
        "",
        "引用の中に、書いた条件をすべて満たす箇所があるか（必要な根拠の再現率）でも測ります。"
        "ページ・行は範囲が重なれば一致、シートは名前が一致、章節は見出しの一部が一致、抜粋は本文に"
        "含まれれば一致です（全角・半角・大小文字・空白の違いは無視）。",
        "例: 製品マニュアル.pdf | p.3-4 | 章節: 初期設定 | 抜粋: 電源を入れる",
    ),
    (
        "根拠も標準回答も空",
        "",
        "資料に答えが無い質問として、回答を控えたか（回答を控える判断の正しさ）を測ります。",
        "例: 来月の天気を教えてください。",
    ),
    (),
    ("形式ごとの「期待する根拠」の書き方",),
    ("形式", "位置の書き方", "説明", "例"),
    (
        "PDF",
        "p.12 / p.12-14、章節: 見出し、抜粋:",
        "ページはページの番号（範囲も可）。「12ページ」とも書けます。",
        "製品マニュアル.pdf | p.12 | 章節: 初期設定",
    ),
    (
        "Word（PDF にしたもの）",
        "p.3、章節: 見出し、抜粋:",
        "取込で PDF にしたページの番号です。",
        "操作手順.docx | p.3 | 抜粋: 申請の前に",
    ),
    (
        "PowerPoint（PDF にしたもの）",
        "スライド: 3（p.3 とも書けます）、抜粋:",
        "スライドの番号は、PDF にしたページの番号と同じです。",
        "説明資料.pptx | スライド: 3",
    ),
    (
        "表計算（Excel・CSV）",
        "シート: 名前 | 行: 10-12、抜粋:",
        "行はシートの行番号（範囲も可）。表計算にはページが無いため、p. は一致しません"
        "（取込で警告します）。",
        "料金表.xlsx | シート: 料金 | 行: 10-12",
    ),
    (
        "Markdown・テキスト・HTML",
        "章節: 見出し、抜粋:",
        "ページが無いため、見出しか本文の語句で指定します。",
        "よくある質問.md | 章節: はじめに | 抜粋: 対象となる方",
    ),
    (
        "文書 ID での指定",
        "文書ID: 〈ID〉",
        "ファイル名の代わりに書けます（ファイル名の後に書いても、ID で引きます）。ID は「"
        f"{DOCUMENTS_SHEET_TITLE}」のシートか、文書の詳細の"
        "「処理の詳細（診断）› 原本情報」にあります。",
        "文書ID: 0123456789abcdef | p.5",
    ),
    (
        "同じ名前の文書があるとき",
        "ファイル名だけ、か文書ID:",
        "ファイル名で書くと、同じ名前の文書のどれか 1 つが引用されれば一致とみなします"
        "（取込で警告します）。"
        "区別するときは文書 ID で書きます。",
        "文書ID: 0123456789abcdef",
    ),
    (),
    (
        "上限",
        "",
        f"1 ファイル {EVALUATION_EXCEL_MAX_ROWS} 行・"
        f"{EVALUATION_EXCEL_MAX_BYTES // (1024 * 1024)}MB まで。",
        "",
    ),
    (
        "注意",
        "",
        f"ケースは「{CASES_SHEET_TITLE}」のシートの 2 行目から書きます。"
        "1 行目の見出しは変えないでください。すべての列が空の行は飛ばします。"
        f"「{EXAMPLES_SHEET_TITLE}」「{DOCUMENTS_SHEET_TITLE}」のシートは評価しません。",
        "",
    ),
)

# 「記入例」のシート（形式ごとの記入の例。合成の内容。取込はこのシートを読まない）。
_EXAMPLE_NOTE_HEADER = "この例の説明"
_EXAMPLE_ROWS: tuple[tuple[str, str, str, str], ...] = (
    (
        "初期設定の手順を教えてください。",
        "製品マニュアル.pdf | p.12 | 章節: 初期設定",
        "電源を入れ、画面の案内に従って言語と時刻を設定します。",
        "PDF のページと章節。標準回答があると回答も採点します。",
    ),
    (
        "申請の前に確かめることは何ですか？",
        "操作手順.docx | p.3 | 抜粋: 申請の前に",
        "",
        "Word（PDF にしたもの）のページと本文の抜粋。",
    ),
    (
        "新しい機能の概要を教えてください。",
        "説明資料.pptx | スライド: 3",
        "",
        "PowerPoint のスライドの番号。",
    ),
    (
        "プラン B の月額はいくらですか？",
        "料金表.xlsx | シート: 料金 | 行: 10-12",
        "",
        "表計算のシートと行。",
    ),
    (
        "対象となる方を教えてください。",
        "よくある質問.md | 章節: はじめに\nよくある質問.md | 抜粋: 対象となる方",
        "",
        "Markdown の章節と抜粋。1 つのセルに 2 つの根拠を改行で書いた例。",
    ),
    (
        "保証の期間はどれくらいですか？",
        "文書ID: 0123456789abcdef | p.5",
        "",
        "文書 ID での指定（同じ名前の文書を区別するとき）。ID は「文書一覧」のシートにあります。",
    ),
    (
        "来月の天気を教えてください。",
        "",
        "",
        "根拠も標準回答も空。資料に答えが無い質問として、回答を控えたかを測ります。",
    ),
)
_DOCUMENT_HEADERS = ("ファイル名", "文書 ID", "形式", "ページ数 / シート名", "アップロード日")


def template_xlsx(documents: Sequence[EvaluationDocumentRef] = ()) -> bytes:
    """評価リクエストのテンプレート（見出しだけのケースのシート・記入の説明・記入例・文書一覧）。

    `documents` は選んだナレッジベースの文書（「文書一覧」のシートと、「期待する根拠」の列の
    入力規則の選択肢にする。#2086）。
    """
    return evaluation_cases_xlsx([], documents=documents)


def evaluation_cases_xlsx(
    rows: Sequence[tuple[str, Sequence[str], str | None]],
    *,
    documents: Sequence[EvaluationDocumentRef] = (),
) -> bytes:
    """評価リクエストの Excel（テンプレートと同じ形）。`rows` は（質問, 根拠の行, 標準回答）。"""
    book = Workbook()
    sheet = book.active or book.create_sheet()
    sheet.title = CASES_SHEET_TITLE
    _write_header(sheet, HEADERS, notes=_HEADER_NOTES)
    for question, evidence, standard_answer in rows:
        sheet.append([question, "\n".join(evidence), standard_answer or ""])
    _set_widths(sheet, _COLUMN_WIDTHS)
    _wrap(sheet, min_row=2)
    sheet.freeze_panes = "A2"
    if documents:
        # ファイル名の選択肢（選ぶと、セルの値がファイル名になる。位置は続けて書き足す）。
        # 選択肢に無い値も書けるよう、誤りの表示はしない。
        validation = DataValidation(
            type="list",
            formula1=f"'{DOCUMENTS_SHEET_TITLE}'!$A$2:$A${len(documents) + 1}",
            allow_blank=True,
            showErrorMessage=False,
            showInputMessage=True,
            promptTitle=EVIDENCE_HEADER,
            prompt="ファイル名を選べます。位置は「 | p.12」のように続けて書きます。",
        )
        sheet.add_data_validation(validation)
        validation.add(f"B2:B{EVALUATION_EXCEL_MAX_ROWS + 1}")

    guide = book.create_sheet(GUIDE_SHEET_TITLE)
    for values in _GUIDE_ROWS:
        guide.append(list(values))
    for cell in guide[1]:
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
    _wrap(guide)
    for row in guide.iter_rows():
        if row[0].value in _GUIDE_HEADINGS:
            row[0].font = Font(bold=True)
        if row[0].value == "形式":
            for cell in row:
                cell.font = Font(bold=True)
                cell.fill = _HEADER_FILL
    _set_widths(guide, (28, 32, 64, 56))

    examples = book.create_sheet(EXAMPLES_SHEET_TITLE)
    _write_header(examples, (*HEADERS, _EXAMPLE_NOTE_HEADER))
    for example in _EXAMPLE_ROWS:
        examples.append(list(example))
    _set_widths(examples, (*_COLUMN_WIDTHS, 56))
    _wrap(examples, min_row=2)
    examples.freeze_panes = "A2"

    listing = book.create_sheet(DOCUMENTS_SHEET_TITLE)
    _write_header(listing, _DOCUMENT_HEADERS)
    if documents:
        for document in documents:
            listing.append(_document_row(document))
    else:
        listing.append(
            [
                "品質評価の画面でナレッジベースを選んでからテンプレートをダウンロードすると、"
                "その文書の一覧（ファイル名・文書 ID）が入ります。"
            ]
        )
    _set_widths(listing, (48, 36, 18, 36, 18))
    _wrap(listing, min_row=2)
    listing.freeze_panes = "A2"

    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _write_header(
    sheet: Any, headers: Sequence[str], *, notes: Mapping[str, str] | None = None
) -> None:
    sheet.append(list(headers))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
        if notes and (note := notes.get(str(cell.value))):
            cell.comment = Comment(note, "RAG")


def _set_widths(sheet: Any, widths: Sequence[int]) -> None:
    for column, width in zip("ABCDEFGH", widths, strict=False):
        sheet.column_dimensions[column].width = width


def _wrap(sheet: Any, *, min_row: int = 1) -> None:
    for row in sheet.iter_rows(min_row=min_row):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def _document_row(document: EvaluationDocumentRef) -> list[str]:
    """「文書一覧」の 1 行（ファイル名・文書 ID・形式・ページ数 / シート名・アップロード日）。"""
    if document.sheet_names:
        location = "シート: " + "、".join(document.sheet_names)
    elif document.page_count:
        unit = "スライド" if document.format == "slides" else "ページ"
        location = f"{document.page_count} {unit}"
    else:
        location = ""
    uploaded = document.uploaded_at.strftime("%Y-%m-%d") if document.uploaded_at else ""
    return [
        document.file_name,
        document.document_id,
        document_format_label(document.file_name, document.content_type),
        location,
        uploaded,
    ]


def evidence_line(file_name: str, page: int | None = None) -> str:
    """根拠の 1 行（フィードバックの引用から作る）。"""
    return f"{file_name} | p.{page}" if page else file_name
