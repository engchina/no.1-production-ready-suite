"""品質評価の完了結果を、実行時の入力と合わせて Excel に出力する。"""

from io import BytesIO
from typing import Any

from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE  # type: ignore[import-untyped]
from openpyxl.styles import Alignment, Font, PatternFill  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]

from app.schemas.evaluation import EvaluationMetrics, EvaluationResultExportRequest

METRICS = (
    ("required_evidence_recall", "required_evidence_recall", "必要な根拠の取得率"),
    ("mrr", "mrr", "正解文書の順位（MRR）"),
    ("answer_correctness", "answer_correctness", "回答の正確性"),
    ("requirement_coverage", "requirement_coverage", "回答の網羅性"),
    ("claim_support_rate", "claim_support_rate", "根拠への忠実性"),
    ("response_decision_accuracy", "response_decision_accuracy", "回答・確認・保留の判断の正しさ"),
)
DECISIONS = {"answer": "回答", "clarify": "確認", "hold": "保留"}


def _append(sheet: Any, values: list[Any]) -> None:
    sheet.append([ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v for v in values])
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.value = ILLEGAL_CHARACTERS_RE.sub("", cell.value)
            cell.data_type = "s"
        cell.alignment = Alignment(vertical="top", wrap_text=True)


def _format(sheet: Any, widths: list[int]) -> None:
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "1:1"
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def evaluation_results_xlsx(
    metrics: EvaluationMetrics, inputs: EvaluationResultExportRequest
) -> bytes:
    """全ページ分の結果を出す。文字列は式として評価させない。"""
    workbook = Workbook()
    overview = workbook.active
    assert overview is not None
    overview.title = "概要"
    groups = [("主条件", inputs.profile_name or "主条件", metrics)]
    if metrics.comparison:
        groups.append(
            (
                "比較対象",
                inputs.compare_profile_name or metrics.comparison.search_answer_profile_id,
                metrics.comparison.metrics,
            )
        )
    _append(overview, ["指標", *[title + ": " + profile for title, profile, _ in groups]])
    _append(overview, ["ケース数", *[result.case_count for _, _, result in groups]])
    for key, _, title in METRICS:
        _append(
            overview,
            [
                title,
                *[
                    getattr(result, key) if getattr(result, key) is not None else "未評価"
                    for _, _, result in groups
                ],
            ],
        )
    rows = {row.case_id: row for row in inputs.rows}
    cases = {case.id: case for case in inputs.cases}
    for title, _profile, result in groups:
        sheet = workbook.create_sheet(title + "の結果")
        _append(
            sheet,
            [
                "ケース ID",
                "質問",
                "期待する根拠",
                "標準回答",
                "必須の回答要点",
                "期待する応答の判断",
                "回答の過程",
                *[label for _, _, label in METRICS],
                "状態",
                "失敗理由",
                "Trace",
                "エラー",
            ],
        )
        for item in result.case_results:
            row = rows.get(item.case_id)
            case = cases.get(item.case_id)
            _append(
                sheet,
                [
                    item.case_id,
                    row.question if row else case.query if case else "",
                    "\n".join(row.evidence) if row else "\n".join(item.relevant_document_ids),
                    (case.standard_answer if case else row.standard_answer_preview if row else "")
                    or "",
                    "\n".join(
                        case.required_answer_points
                        if case
                        else row.required_answer_points or []
                        if row
                        else []
                    ),
                    DECISIONS.get(
                        (
                            case.expected_response_decision
                            if case
                            else row.expected_response_decision
                            if row
                            else None
                        )
                        or "",
                        "未評価",
                    ),
                    row.answer_process or "" if row else "",
                    *[
                        getattr(item, key) if getattr(item, key) is not None else "未評価"
                        for _, key, _ in METRICS
                    ],
                    item.status,
                    "\n".join(item.failure_reasons),
                    item.trace_id,
                    item.error_message or "",
                ],
            )
            for cell in sheet[sheet.max_row][7:13]:
                cell.number_format = "0.0%"
        _format(sheet, [24, 48, 40, 60, 40, 20, 60, *([20] * 6), 16, 36, 36, 48])
    _format(overview, [36, *([36] * len(groups))])
    for row in overview.iter_rows(min_row=3, min_col=2):
        for cell in row:
            cell.number_format = "0.0%"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
