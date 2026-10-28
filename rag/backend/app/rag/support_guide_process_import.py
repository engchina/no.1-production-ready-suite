"""回答の過程を書いた表から業務ガイドの下書きとルールの候補を作る（#2397）。

利用者の担当者は、問ごとの「回答の過程」を表（Excel）の 1 セルに、決まった見出しの半構造の文で
書いている（回答の手順・必要な確認事項と分岐・検索する資料と頁・検索の語・補足）。この表の行を、
OCI Enterprise AI で業務ガイドの内容（``SupportGuideContent``）の下書きに直す。

1. Excel の読み込み（``read_process_workbook``）: シートと列（1 行目の見出し）を返し、利用者が
   選んだ列（質問・回答の過程・任意で標準回答）で 2 行目からの行を読む。モデルは呼ばない。
2. 節に分ける（``split_process_sections``）: 回答の過程の文を、言語資源の見出しの語
   （``guide_import_cues.process_headings``。「見出し=項目」）で節に分ける。見出しの語はコードに
   書かない。
3. 構造化（``draft_guides``）: 1 行で LLM を 1 回呼ぶ（共通の上限つき並行の部品。#2004）。応答は
   ``GuideDraftOutput``（Pydantic）で検証し、業務ガイドの内容に直して ``SupportGuideContent`` と
   ``validate_content`` で検証する。
4. 資料名の解決（``resolve_document_name``）: 回答の過程に書いた資料の名前を、検索・回答
   プロファイルのナレッジベースの文書の名前と照合して ``document_id`` にする。同じ正本を指す重複の
   アップロード（``duplicate_of_document_id``）は正本にまとめてから決める（#2662）。解決できない・
   複数に当たる資料は検証のエラー（KB への取り込みを促す）。
5. 資料の位置（``_location``）: PDF などの頁は頁の範囲に、表計算の資料のシート・行・セルはシートの
   行の範囲にする。表計算の資料の頁・頁の資料の行、頁として読めない頁は変換せず検証のエラーにする
   （行の番号を頁にしない。#2663）。
6. 分岐の検証（``merged_branch_issues``）: 同じ条件の別の値が、回答の過程では別の操作なのに同じ
   手順へ進む下書きは、検証のエラーにする（値ごとの手順と資料の絞り込みを失うため。#2666）。

下書きは保存しない（画面で行ごとに確かめ、選んだものだけを取り込みの API で下書きにする）。
"""

from __future__ import annotations

import json
import logging
import re
import unicodedata
from collections.abc import Awaitable, Callable, Mapping, Sequence
from dataclasses import dataclass
from io import BytesIO
from typing import Any, Literal, Protocol

from openpyxl import load_workbook  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]
from pr_backend_core.bounded_calls import Attempt
from pr_backend_core.observability.model_purpose import async_call_with_model_purpose
from pydantic import BaseModel, ConfigDict, Field, ValidationError
from rag_engine.resources.runtime import resolve_prompt
from rag_parser_core.language import LanguagePack
from rag_parser_core.language.markers import marker_patterns

from app.rag.cross_references import document_name_match, title_key
from app.rag.parser_source_guard import is_spreadsheet_file
from app.rag.support_guide import has_errors, validate_content
from app.schemas.support_guide import (
    MAX_PAGE_NUMBER,
    MAX_REFERENCE_PAGE_RANGES,
    MAX_SHEET_ROW,
    MAX_STEP_REFERENCES,
    PROCESS_IMPORT_QUESTION_MAX_CHARS,
    PROCESS_IMPORT_TEXT_MAX_CHARS,
    SupportGuideContent,
    SupportGuideIssue,
    SupportGuidePageRange,
    SupportGuideProcessColumn,
    SupportGuideProcessItem,
    SupportGuideProcessReference,
    SupportGuideProcessRow,
    SupportGuideProcessSheet,
    SupportGuideProcessWorkbookData,
    SupportGuideRuleCandidate,
    SupportGuideSheetRange,
)

logger = logging.getLogger(__name__)

PROCESS_WORKBOOK_MAX_BYTES = 10 * 1024 * 1024
# Excel の読み込みで見る行の上限（上限を超えた分は読まない。行の数の数えすぎを止める）。
_WORKBOOK_SCAN_ROWS = 5000
_MAX_ISSUES = 50
ProcessField = Literal["steps", "conditions", "references", "search_terms", "notes"]
PROCESS_FIELDS: tuple[str, ...] = ("steps", "conditions", "references", "search_terms", "notes")


# ---- Excel ----------------------------------------------------------------------------------


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).replace("\r\n", "\n").strip()


def read_process_workbook(
    data: bytes,
    file_name: str,
    *,
    max_rows: int,
    sheet: str | None = None,
    question_column: int | None = None,
    process_column: int | None = None,
    answer_column: int | None = None,
    start_row: int = 2,
) -> SupportGuideProcessWorkbookData:
    """シートと列を返し、列を選んでいれば ``start_row`` から ``max_rows`` 行を読む。

    1 行目は見出し。質問と回答の過程の両方が空の行は飛ばす。片方だけの行は読めない行として
    ``issues`` に出す。形式の誤り（.xlsx でない・大きすぎる）は ``ValueError``（利用者向けの文）。
    """
    if len(data) > PROCESS_WORKBOOK_MAX_BYTES:
        raise ValueError(
            f"ファイルが大きすぎます（上限 {PROCESS_WORKBOOK_MAX_BYTES // (1024 * 1024)}MB）。"
        )
    try:
        book = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - 形式の誤りは利用者へ返す
        raise ValueError(
            "Excel（.xlsx）として読めません。.xlsx の形式で保存してください。"
        ) from exc
    try:
        sheets: list[SupportGuideProcessSheet] = []
        selected: Any = None
        for worksheet in book.worksheets:
            rows = worksheet.iter_rows(values_only=True, max_row=_WORKBOOK_SCAN_ROWS)
            header = list(next(rows, None) or ())
            count = sum(1 for values in rows if any(_cell_text(value) for value in values))
            sheets.append(
                SupportGuideProcessSheet(
                    name=str(worksheet.title),
                    columns=[
                        SupportGuideProcessColumn(
                            index=index,
                            letter=get_column_letter(index + 1),
                            header=_cell_text(value)[:100],
                        )
                        for index, value in enumerate(header)
                    ],
                    row_count=count,
                )
            )
            if sheet is not None and str(worksheet.title) == sheet:
                selected = worksheet
        result = SupportGuideProcessWorkbookData(
            file_name=file_name, sheets=sheets, max_rows=max_rows
        )
        if question_column is None or process_column is None:
            return result
        if selected is None:
            raise ValueError("選んだシートがありません。シートを選び直してください。")
        result.sheet = str(selected.title)
        _read_rows(
            selected,
            result,
            question_column=question_column,
            process_column=process_column,
            answer_column=answer_column,
            start_row=max(2, start_row),
            max_rows=max_rows,
        )
        return result
    finally:
        book.close()


def _read_rows(
    worksheet: Any,
    result: SupportGuideProcessWorkbookData,
    *,
    question_column: int,
    process_column: int,
    answer_column: int | None,
    start_row: int,
    max_rows: int,
) -> None:
    def at(values: Sequence[object], index: int | None) -> str:
        if index is None or index < 0 or index >= len(values):
            return ""
        return _cell_text(values[index])

    for number, values in enumerate(
        worksheet.iter_rows(min_row=start_row, max_row=_WORKBOOK_SCAN_ROWS, values_only=True),
        start=start_row,
    ):
        question = at(values, question_column)
        process = at(values, process_column)
        if not question and not process:
            continue
        if len(result.rows) >= max_rows:
            result.truncated = True
            result.next_start_row = number
            return
        problem = _row_problem(question, process)
        if problem:
            if len(result.issues) < _MAX_ISSUES:
                result.issues.append(f"{number} 行目: {problem}")
            continue
        result.rows.append(
            SupportGuideProcessRow(
                row=number,
                question=question,
                process=process,
                standard_answer=at(values, answer_column)[:PROCESS_IMPORT_TEXT_MAX_CHARS],
            )
        )


def _row_problem(question: str, process: str) -> str:
    if not question:
        return "質問が空です。"
    if not process:
        return "回答の過程が空です。"
    if len(question) > PROCESS_IMPORT_QUESTION_MAX_CHARS:
        return f"質問が長すぎます（{PROCESS_IMPORT_QUESTION_MAX_CHARS} 字まで）。"
    if len(process) > PROCESS_IMPORT_TEXT_MAX_CHARS:
        return f"回答の過程が長すぎます（{PROCESS_IMPORT_TEXT_MAX_CHARS} 字まで）。"
    return ""


# ---- 節に分ける -----------------------------------------------------------------------------


@dataclass(frozen=True)
class ProcessSection:
    """回答の過程の 1 つの節（見出しの無い先頭の文は ``field`` が None）。"""

    field: str | None
    heading: str
    text: str


def process_headings(pack: LanguagePack) -> dict[str, str]:
    """言語資源の見出しの語 → 業務ガイドの項目（``guide_import_cues.process_headings``）。"""
    mapping = pack.mapping("guide_import_cues", "process_headings")
    return {word: field for word, field in mapping.items() if field in PROCESS_FIELDS}


def _is_mark(char: str) -> bool:
    """見出しの飾り（記号・句読点・空白）か。言語に依らない Unicode の分類で決める。"""
    return char.isspace() or unicodedata.category(char)[0] in {"P", "S"}


def _heading(line: str, keys: Sequence[tuple[str, str, str]]) -> tuple[str, str, str] | None:
    """行が見出しなら (項目, 見出しの語, 見出しの後ろの文)。

    行の前後の飾り（【】・■・「：」など）を外した文が見出しの語で始まり、その後ろが空か飾りで
    始まるときだけ見出しとみなす（「手順書を開く」の「手順」を見出しにしない）。
    """
    text = unicodedata.normalize("NFKC", line).strip()
    start = 0
    while start < len(text) and _is_mark(text[start]):
        start += 1
    body = text[start:]
    folded = body.casefold()
    for key, word, field in keys:
        if not folded.startswith(key):
            continue
        rest = body[len(key) :]
        if rest and not _is_mark(rest[0]):
            continue
        # 見出しの後ろの閉じ括弧・区切り（「】」「：」）を外す。開き括弧から始まる内容は残す。
        index = 0
        while index < len(rest) and (
            rest[index].isspace() or unicodedata.category(rest[index]) in _HEADING_TAIL
        ):
            index += 1
        return field, word, rest[index:].strip()
    return None


# 見出しの後ろに続く飾りの分類（閉じ括弧・区切りの記号・ダッシュ）。
_HEADING_TAIL = frozenset({"Pe", "Pf", "Po", "Pd"})


def split_process_sections(text: str, headings: Mapping[str, str]) -> list[ProcessSection]:
    """回答の過程の文を、見出しの語の行で節に分ける（見出しが無ければ全体を 1 つの節にする）。"""
    keys = sorted(
        (
            (unicodedata.normalize("NFKC", word).casefold().strip(), word, field)
            for word, field in headings.items()
            if word.strip()
        ),
        key=lambda item: -len(item[0]),
    )
    sections: list[ProcessSection] = []
    field: str | None = None
    heading = ""
    lines: list[str] = []

    def flush() -> None:
        body = "\n".join(lines).strip()
        if body or field is not None:
            sections.append(ProcessSection(field, heading, body))

    for line in text.splitlines():
        found = _heading(line, keys) if keys else None
        if found is None:
            lines.append(line)
            continue
        flush()
        field, heading, rest = found
        lines = [rest] if rest else []
    flush()
    return [section for section in sections if section.text or section.field]


# ---- LLM の応答の形 -------------------------------------------------------------------------


class _Out(BaseModel):
    model_config = ConfigDict(extra="ignore", str_strip_whitespace=True)


class DraftStep(_Out):
    id: str = Field(description="手順の id（s1・s2…）。")
    title: str = Field(description="手順の題名（画面・操作が分かる短い文）。")
    purpose: str = Field(default="", description="手順の目的・内容。")
    depends_on: list[str] = Field(default_factory=list, description="先に済ませる手順の id。")
    retrieval_hints: list[str] = Field(
        default_factory=list, description="この手順の資料を検索する語（画面名・機能名）。"
    )
    note: str = Field(default="", description="この手順の注意（補足のうち手順に係るもの）。")


class DraftCondition(_Out):
    id: str = Field(description="条件の id（c1・c2…）。")
    label: str = Field(description="条件の名前。")
    type: Literal["enum", "boolean", "text"] = "enum"
    allowed_values: list[str] = Field(default_factory=list, description="選択肢（enum のとき）。")
    question: str = Field(default="", description="利用者に確かめる問い。")
    handling: Literal["ask", "branch", "handoff"] = Field(
        default="ask", description="分からないとき: ask=確かめる / branch=分岐 / handoff=窓口へ。"
    )


class DraftBranch(_Out):
    condition_id: str
    values: list[str] = Field(
        default_factory=list, description="この値のとき（空なら条件が分からないとき）。"
    )
    goto_step: str = Field(description="進む手順の id。")
    operation: str = Field(
        default="", description="この値のときに回答の過程が案内する操作（書いたまま）。"
    )
    note: str = ""


class DraftReference(_Out):
    document_name: str = Field(description="資料の名前（回答の過程に書いたまま）。")
    pages: str = Field(
        default="", description="頁（書いたまま。例: 12-14）。表計算の行・セルは入れない。"
    )
    sheet_name: str = Field(default="", description="表計算の資料のシート名（書いたまま）。")
    rows: str = Field(default="", description="表計算の資料の行（書いたまま。例: 42-45）。")
    cells: str = Field(
        default="", description="表計算の資料のセルの範囲（書いたまま。例: B3:D10）。"
    )
    section: str = Field(default="", description="章節（書いたまま）。")
    step_ids: list[str] = Field(
        default_factory=list,
        description="この資料を根拠にする手順の id（全手順に共通の資料なら空）。",
    )


class DraftRule(_Out):
    title: str = Field(description="ルールの題名。")
    triggers: list[str] = Field(default_factory=list, description="ルールを当てる質問の語。")
    content: str = Field(description="ルールの内容（回答のときに守ること）。")


class GuideDraftOutput(_Out):
    """LLM の応答（業務ガイドの下書きとルールの候補）。"""

    title: str = Field(description="業務ガイドの題名（利用者がしたいこと）。")
    expected_result: str = Field(description="案内のとおり進めたときの結果。")
    match_terms: list[str] = Field(default_factory=list, description="照合の語（検索の語）。")
    steps: list[DraftStep] = Field(default_factory=list)
    conditions: list[DraftCondition] = Field(default_factory=list)
    branches: list[DraftBranch] = Field(default_factory=list)
    references: list[DraftReference] = Field(default_factory=list)
    handoff_contact: str = Field(default="", description="窓口へ連絡するときの窓口の名前。")
    handoff_conditions: list[str] = Field(default_factory=list, description="窓口へ連絡する条件。")
    rule_candidates: list[DraftRule] = Field(default_factory=list)


GUIDE_DRAFT_RESPONSE_SCHEMA = GuideDraftOutput.model_json_schema()

GUIDE_DRAFT_SYSTEM_PROMPT = (
    "## Step 1: 業務ガイドの下書きを作る\n"
    "あなたは、問い合わせ対応の「回答の過程」を、業務ガイドの下書き（JSO"
    "N）に直す担当です。\n"
    "\n"
    "## Step 2: 入力の範囲を守る\n"
    "- 回答の過程に書いてあることだけを使い、書いていない手順・条件・資料"
    "・頁を足さないでください。\n"
)

GUIDE_DRAFT_PROMPT = (
    "## Step 1: 入力の項目を確認する\n"
    "\n"
    "入力の JSON は、利用者の質問（question）、回答の過程を見"
    "出しで分けた節（sections）、任意の標準回答\n"
    "（standard_answer）です。節の field は次の項目を"
    "表します（null は見出しの無い文）。\n"
    "\n"
    "## Step 2: 手順・条件・根拠を対応付ける\n"
    "- steps（回答の手順）: steps に順に入れる。前の手順を済"
    "ませてから行う手順は depends_on に前の手順の\n"
    "  id を入れる。手順の画面名・機能名は retrieval_hin"
    "ts に入れる。\n"
    "- conditions（必要な確認事項）: 確かめる事項を cond"
    "itions に、「X なら手順 P」を branches に入れる。"
    "\n"
    "  選択肢が決まっていれば type=enum（2 つ以上）、はい /"
    " いいえなら boolean、それ以外は text。\n"
    "  値ごとに案内する操作が違えば、値ごとに別の手順にして goto_s"
    "tep を分け、その値の操作を\n"
    "  operation に入れる（同じ手順に進めるのは、値に関わらず同"
    "じ操作を案内するときだけ）。\n"
    "  「窓口へ連絡」のような人へ引き継ぐ場合は handling=han"
    "doff にし、handoff_contact に窓口の名前、\n"
    "  handoff_conditions にその場合を入れる。利用者に"
    "聞けば決まる事項は handling=ask と question を"
    "\n"
    "  入れる。\n"
    "- references（検索する資料と根拠の頁）: 資料の名前は書い"
    "たまま document_name に、頁は pages、\n"
    "  表計算のシート・行・セルの範囲は sheet_name・rows・"
    "cells、章節は section に入れる。どの手順の\n"
    "  根拠の資料かが書いてあれば、step_ids にその手順の id "
    "を入れる（全手順に共通の資料なら空）。\n"
    "- search_terms（検索の語）: match_terms と"
    "、関係する手順の retrieval_hints に入れる。\n"
    "- notes（補足）: 1 つの手順に係る注意はその手順の note"
    " に入れる。この質問に限らず回答の仕方を決める\n"
    "  規則（「この種類の質問はこの画面を案内する」「言い切らず可能性にと"
    "どめる」「Aと Bを区別する」\n"
    "  「範囲外は窓口へ」など）は rule_candidates に入れ"
    "る（triggers はその規則を当てる質問の語）。\n"
    "\n"
    "## Step 3: 名前と識別子を設定する\n"
    "title は利用者がしたいことを表す短い名前、expected_re"
    "sult は手順のとおり進めたときの結果にします。\n"
    "id は手順が s1・s2…、条件が c1・c2… の順の番号にします"
    "。\n"
)


class GuideDraftLlm(Protocol):
    async def generate(
        self,
        prompt: str,
        context: str,
        *,
        system_prompt: str | None = None,
        response_schema: Mapping[str, Any] | None = None,
        response_schema_name: str = ...,
    ) -> str: ...


_JSON_OBJECT = re.compile(r"\{.*\}", re.DOTALL)


def parse_draft_output(raw: str) -> GuideDraftOutput:
    """LLM の応答を検証する（JSON の前後の文は外す）。形が合わなければ ``ValueError``。"""
    match = _JSON_OBJECT.search(str(raw or ""))
    if match is None:
        raise ValueError("応答に JSON がありません。")
    try:
        return GuideDraftOutput.model_validate(json.loads(match.group(0)))
    except (json.JSONDecodeError, ValidationError) as exc:
        raise ValueError("応答の JSON が業務ガイドの下書きの形ではありません。") from exc


def draft_context(row: SupportGuideProcessRow, headings: Mapping[str, str]) -> str:
    """LLM に渡す 1 行の入力（質問・節に分けた回答の過程・標準回答）。"""
    sections = split_process_sections(row.process, headings)
    payload = {
        "question": row.question,
        "sections": [
            {"field": section.field, "heading": section.heading, "text": section.text}
            for section in sections
        ],
        "standard_answer": row.standard_answer,
    }
    return json.dumps(payload, ensure_ascii=False)


# ---- 資料名の解決 ---------------------------------------------------------------------------


@dataclass(frozen=True)
class KnowledgeDocument:
    """検索・回答プロファイルのナレッジベースの文書（照合に使う名前）。"""

    document_id: str
    file_name: str
    # 同じ内容の重複のアップロードなら正本の文書の id（``duplicate_of_document_id``。#2662）。
    duplicate_of: str = ""
    # 表計算の資料か（頁ではなくシート・行で位置を書く。#2663）。
    spreadsheet: bool = False

    @property
    def identity(self) -> str:
        """同じ内容の文書をまとめる id（重複なら正本の id）。"""
        return self.duplicate_of or self.document_id


async def list_knowledge_documents(
    oracle: Any, knowledge_base_ids: Sequence[str]
) -> list[KnowledgeDocument]:
    """ナレッジベースの文書（利用者が見られるもの。新しい版に置き換えた文書は除く）。

    重複のアップロードは正本の id（``duplicate_of_document_id``）を持たせ、資料名の解決で正本に
    まとめる（#2662）。表計算の資料かは、頁ではなくシート・行で位置を書く資料かの判定に使う（#2663）。
    """
    documents: dict[str, KnowledgeDocument] = {}
    for knowledge_base_id in knowledge_base_ids:
        for summary in await oracle.list_documents(knowledge_base_id=knowledge_base_id):
            if summary.is_superseded:
                continue
            documents[summary.id] = KnowledgeDocument(
                summary.id,
                summary.file_name,
                duplicate_of=summary.duplicate_of_document_id or "",
                spreadsheet=is_spreadsheet_file(summary.file_name, summary.content_type),
            )
    return list(documents.values())


_EXTENSION = re.compile(r"\.[A-Za-z0-9]{1,5}$")


def _stem_key(file_name: str) -> str:
    return title_key(_EXTENSION.sub("", unicodedata.normalize("NFKC", file_name).strip()))


def _settle(
    candidates: Sequence[KnowledgeDocument], documents: Sequence[KnowledgeDocument]
) -> tuple[KnowledgeDocument | None, list[KnowledgeDocument]]:
    """候補を同じ内容の文書（正本と、その正本を指す重複）ごとにまとめて決める（#2662）。

    正本が候補の文書の中に見えるとき（このプロファイルのナレッジベースにあり、利用者が見られる）だけ、
    その正本と重複を 1 つにまとめ、正本の文書にする。正本が見えない（範囲の外・見られない・無い）
    重複は、まとめずにそれぞれの文書のまま扱う（見えない正本の id を参照にしない）。
    """
    visible = {item.document_id: item for item in documents}
    groups: dict[str, KnowledgeDocument] = {}
    for item in candidates:
        canonical = visible.get(item.identity)
        key = item.identity if canonical is not None else item.document_id
        groups.setdefault(key, canonical or item)
    settled = list(groups.values())
    return (settled[0], []) if len(settled) == 1 else (None, settled)


def resolve_document_name(
    name: str, documents: Sequence[KnowledgeDocument]
) -> tuple[KnowledgeDocument | None, list[KnowledgeDocument]]:
    """資料の名前に当たる文書（1 つに決まれば文書、決まらなければ None と当たった文書）。

    確かな順に照合し、最初に当たった順位の文書だけを候補にする。
    1. ファイル名（拡張子を除く形を含む）が同じ。
    2. 文書のタイトル・ファイル名の別名と同じ、またはファイル名が資料の名前を含む
       （``document_name_match``。他の文書への参照の解決と同じ規則）。
    3. 資料の名前がファイル名（拡張子を除く）を含む（利用者のファイル名の接頭辞・接尾辞）。最も
       長いファイル名だけを候補にする。
    候補は同じ正本の重複をまとめてから数える（``_settle``）。
    """
    key = title_key(name)
    stem = _stem_key(name)
    if len(key) < 2:
        return None, []
    exact = [
        item
        for item in documents
        if title_key(item.file_name) == key or (stem and _stem_key(item.file_name) == stem)
    ]
    if exact:
        return _settle(exact, documents)
    ranked: dict[int, list[KnowledgeDocument]] = {}
    for item in documents:
        rank = document_name_match(name, item.file_name)
        if rank is not None:
            ranked.setdefault(rank, []).append(item)
    if ranked:
        return _settle(ranked[min(ranked)], documents)
    contained = [
        (len(_stem_key(item.file_name)), item)
        for item in documents
        if len(_stem_key(item.file_name)) >= 2 and _stem_key(item.file_name) in key
    ]
    if not contained:
        return None, []
    longest = max(length for length, _ in contained)
    return _settle([item for length, item in contained if length == longest], documents)


# ---- 下書きにする ---------------------------------------------------------------------------


# 頁の書き方（「12-14」「p.12〜14、20」）から頁の範囲を読む（数字と範囲の記号だけを見る。#2395）。
_PAGE_RANGE = re.compile(r"(\d+)\s*[-~〜–—]\s*(\d+)|(\d+)")
# 頁の書き方の略記（p. / pp.）。頁の語そのもの（「頁」「ページ」）は言語資源の markers.page_words。
_PAGE_ABBREVIATION = re.compile(r"(?<![A-Za-z])pp?(?![A-Za-z])", re.IGNORECASE)
# 表計算のセルの範囲（A1 形式）。
_CELL_RANGE = re.compile(r"[A-Za-z]{1,3}[1-9]\d{0,6}(?:\s*:\s*[A-Za-z]{1,3}[1-9]\d{0,6})?")


def _numbered_ranges(text: str, limit: int) -> list[tuple[int, int]]:
    ranges: list[tuple[int, int]] = []
    for match in _PAGE_RANGE.finditer(unicodedata.normalize("NFKC", text)):
        start = int(match.group(1) or match.group(3))
        end = int(match.group(2) or start)
        if 1 <= start <= end <= limit and (start, end) not in ranges:
            ranges.append((start, end))
    return ranges[:MAX_REFERENCE_PAGE_RANGES]


def page_ranges_from_text(text: str) -> list[dict[str, int]]:
    """回答の過程に書いた頁を、業務ガイドの資料の頁の範囲にする（読めない・逆の範囲は除く）。"""
    return [
        {"page_start": start, "page_end": end}
        for start, end in _numbered_ranges(text, MAX_PAGE_NUMBER)
    ]


def readable_as_pages(text: str) -> bool:
    """頁の書き方として読めるか（数字・範囲・区切りと、頁の語・略記だけ）。

    「行42」「A1:C3」のように、頁の語でない文字（行・セル・シートの名前など）を含む書き方は頁として
    読まない（行の番号を頁にしない。#2663）。頁の語は言語資源（``markers.page_words``）。
    """
    normalized = unicodedata.normalize("NFKC", text)
    stripped = re.sub(marker_patterns().page_words, " ", normalized, flags=re.IGNORECASE)
    stripped = _PAGE_ABBREVIATION.sub(" ", stripped)
    return not any(unicodedata.category(char)[0] == "L" for char in stripped)


def _location_issue(index: int, code: str, message: str) -> SupportGuideIssue:
    return SupportGuideIssue(
        severity="error", code=code, path=f"references[{index}]", message=message
    )


def _sheet_ranges(item: DraftReference) -> list[SupportGuideSheetRange] | None:
    """シート・行・セルからシートの行の範囲（読めなければ None）。

    行は書いた数字（範囲・列挙）を行の番号として読み、セルは A1 形式の範囲として読む。シートの名前の
    無い行・セルは、どのシートの行か決まらないので読まない。
    """
    sheet = item.sheet_name.strip()[:200]
    if not sheet:
        return None
    rows = _numbered_ranges(item.rows, MAX_SHEET_ROW) if item.rows.strip() else []
    cells = [
        re.sub(r"\s+", "", match.group(0)).upper()
        for match in _CELL_RANGE.finditer(unicodedata.normalize("NFKC", item.cells))
    ]
    if (item.rows.strip() and not rows) or (item.cells.strip() and not cells):
        return None
    try:
        if cells:
            return [
                SupportGuideSheetRange(sheet_name=sheet, cell_range=cell)
                for cell in cells[:MAX_REFERENCE_PAGE_RANGES]
            ]
        if rows:
            return [
                SupportGuideSheetRange(sheet_name=sheet, row_start=start, row_end=end)
                for start, end in rows
            ]
        return [SupportGuideSheetRange(sheet_name=sheet)]
    except ValidationError:
        return None


def _location(
    index: int, item: DraftReference, document: KnowledgeDocument
) -> tuple[list[SupportGuidePageRange], list[SupportGuideSheetRange], list[SupportGuideIssue]]:
    """解決した資料の根拠の位置（頁の範囲・表計算のシートの行の範囲）と、位置の検証の誤り（#2663）。

    表計算の資料はシート・行・セルで、それ以外の資料は頁で位置を書く。書き方が資料の種類に合わない・
    読めない位置は変換せず、誤りにする（行の番号を頁の範囲にしない）。
    """
    name = item.document_name.strip()
    pages = item.pages.strip()
    has_pages = bool(re.search(r"\d", unicodedata.normalize("NFKC", pages)))
    has_sheet = bool(item.sheet_name.strip() or item.rows.strip() or item.cells.strip())
    if document.spreadsheet:
        issues: list[SupportGuideIssue] = []
        if has_pages and not has_sheet:
            issues.append(
                _location_issue(
                    index,
                    "reference_location_mismatch",
                    f"資料「{name}」は表計算の資料のため、頁「{pages}」を根拠の位置にできません。"
                    "シート・行・セルの範囲で指定してください。",
                )
            )
        if not has_sheet:
            return [], [], issues
        sheets = _sheet_ranges(item)
        if sheets is None:
            location = " ".join(
                part for part in (item.sheet_name, item.rows, item.cells) if part.strip()
            )
            issues.append(
                _location_issue(
                    index,
                    "reference_location_unreadable",
                    f"資料「{name}」のシート・行・セル「{location.strip()}」を読めません。"
                    "シートの名前と、行の番号（例: 42-45）かセルの範囲（例: B3:D10）で"
                    "指定してください。",
                )
            )
            return [], [], issues
        return [], sheets, issues
    if has_sheet:
        return (
            [],
            [],
            [
                _location_issue(
                    index,
                    "reference_location_mismatch",
                    f"資料「{name}」は表計算の資料ではないため、シート・行・セルを根拠の位置に"
                    "できません。頁か章節で指定してください。",
                )
            ],
        )
    if pages and not readable_as_pages(pages):
        return (
            [],
            [],
            [
                _location_issue(
                    index,
                    "reference_location_unreadable",
                    f"資料「{name}」の頁「{pages}」を頁の範囲として読めません。"
                    "頁の番号と範囲（例: 12-14）で指定してください。",
                )
            ],
        )
    return (
        [SupportGuidePageRange.model_validate(page) for page in page_ranges_from_text(pages)],
        [],
        [],
    )


def _clean(values: Sequence[str], limit: int, max_chars: int = 200) -> list[str]:
    return list(dict.fromkeys(value.strip()[:max_chars] for value in values if value.strip()))[
        :limit
    ]


def _references(
    output: GuideDraftOutput, documents: Sequence[KnowledgeDocument]
) -> tuple[list[SupportGuideProcessReference], list[SupportGuideIssue]]:
    references: list[SupportGuideProcessReference] = []
    issues: list[SupportGuideIssue] = []
    for index, item in enumerate(output.references[:30]):
        name = item.document_name.strip()
        if not name:
            continue
        written = {
            "name": name,
            "pages": item.pages[:64],
            "sheet_name": item.sheet_name[:200],
            "rows": item.rows[:64],
            "cells": item.cells[:64],
            "section": item.section[:200],
            "step_ids": _clean(item.step_ids, 20, 64),
        }
        found, candidates = resolve_document_name(name, documents)
        if found is not None:
            page_ranges, sheet_ranges, location_issues = _location(index, item, found)
            issues.extend(location_issues)
            references.append(
                SupportGuideProcessReference(
                    **written,
                    status="resolved",
                    document_id=found.document_id,
                    file_name=found.file_name,
                    page_ranges=page_ranges,
                    sheet_ranges=sheet_ranges,
                )
            )
            continue
        status: Literal["unresolved", "ambiguous"] = "ambiguous" if candidates else "unresolved"
        references.append(
            SupportGuideProcessReference(
                **written,
                status=status,
                candidates=[candidate.file_name for candidate in candidates[:10]],
            )
        )
        message = (
            f"資料「{name}」に当たる文書が複数あります（"
            + "、".join(candidate.file_name for candidate in candidates[:5])
            + "）。回答の過程の資料の名前を文書の名前に合わせてください。"
            if candidates
            else f"資料「{name}」は、このプロファイルのナレッジベースの文書にありません。"
            "資料をナレッジベースに取り込んでから、もう一度下書きを作ってください。"
        )
        issues.append(
            SupportGuideIssue(
                severity="error",
                code="reference_ambiguous" if candidates else "reference_unresolved",
                path=f"references[{index}]",
                message=message,
            )
        )
    return references, issues


def _content_payload(
    output: GuideDraftOutput,
    question: str,
    references: Sequence[SupportGuideProcessReference],
) -> dict[str, Any]:
    step_ids = {step.id for step in output.steps}
    # 解決できた資料を、手順の資料（step_ids が手順を指す）と全手順に共通の資料に分ける（#2395）。
    step_references: dict[str, list[dict[str, Any]]] = {}
    common_references: list[dict[str, Any]] = []
    for item in references:
        if not item.document_id:
            continue
        reference = {
            "document_id": item.document_id,
            "title": (item.file_name or item.name)[:200],
            "section_path": [item.section[:200]] if item.section else [],
            "page_ranges": [
                page.model_dump() for page in item.page_ranges[:MAX_REFERENCE_PAGE_RANGES]
            ],
            "sheet_ranges": [
                sheet.model_dump() for sheet in item.sheet_ranges[:MAX_REFERENCE_PAGE_RANGES]
            ],
        }
        owners = [step_id for step_id in item.step_ids if step_id in step_ids]
        for step_id in owners:
            step_references.setdefault(step_id, []).append(reference)
        if not owners:
            common_references.append(reference)
    steps = []
    for step in output.steps[:50]:
        purpose = step.purpose.strip()
        if step.note.strip():
            purpose = f"{purpose}\n注意: {step.note.strip()}" if purpose else f"注意: {step.note}"
        steps.append(
            {
                "id": step.id,
                "title": step.title[:200],
                "purpose": purpose[:2000],
                "depends_on": [item for item in step.depends_on if item in step_ids][:20],
                "retrieval_hints": _clean(step.retrieval_hints, 10),
                "references": step_references.get(step.id, [])[:MAX_STEP_REFERENCES],
            }
        )
    conditions = []
    for condition in output.conditions[:30]:
        values = _clean(condition.allowed_values, 20)
        kind = condition.type
        if kind == "enum" and len(values) < 2:
            kind = "text"
        question_text = condition.question.strip()
        if condition.handling == "ask" and not question_text:
            question_text = f"{condition.label}を教えてください。"
        conditions.append(
            {
                "id": condition.id,
                "label": condition.label[:200],
                "type": kind,
                "allowed_values": values if kind == "enum" else [],
                "unknown_handling": condition.handling,
                "question": question_text[:200],
            }
        )
    branches = []
    for index, branch in enumerate(output.branches[:50], start=1):
        values = _clean(branch.values, 20)
        operator = "unknown" if not values else "equals" if len(values) == 1 else "in"
        branches.append(
            {
                "id": f"b{index}",
                "when": {
                    "condition_id": branch.condition_id,
                    "operator": operator,
                    "values": values,
                },
                "goto_step": branch.goto_step,
                "note": (branch.note or branch.operation)[:200],
            }
        )
    title = output.title.strip() or question
    return {
        "title": title[:200],
        "goal": {
            "expected_result": (output.expected_result.strip() or title)[:2000],
            "intent_examples": [question[:2000]],
            "match_terms": _clean(output.match_terms, 30),
        },
        "conditions": conditions,
        "steps": steps,
        "branches": branches,
        "references": common_references[:30],
        "handoff": {
            "contact": output.handoff_contact[:200],
            "conditions": _clean(output.handoff_conditions, 20),
        },
    }


def _operation_key(text: str) -> str:
    """操作の文の比べる形（NFKC・大小・記号と空白を除く）。"""
    folded = unicodedata.normalize("NFKC", text).casefold()
    return "".join(char for char in folded if unicodedata.category(char)[0] in {"L", "N"})


def merged_branch_issues(output: GuideDraftOutput) -> list[SupportGuideIssue]:
    """同じ条件の別の値が、別の操作なのに同じ手順へ進む分岐（#2666）。

    回答の過程が値ごとに別の操作を案内するのに、下書きがその値の分岐を同じ手順にまとめると、値を
    選んでもその操作の手順と資料の絞り込みに分けられない。値ごとの操作（``operation``）が違う分岐が
    同じ手順へ進むときを誤りにする。値に関わらず同じ操作を案内する合流（操作が同じ・書いていない）と、
    条件が分からないときの分岐は対象にしない。
    """
    titles = {step.id: step.title for step in output.steps}
    labels = {condition.id: condition.label for condition in output.conditions}
    groups: dict[tuple[str, str], list[tuple[int, DraftBranch]]] = {}
    for index, branch in enumerate(output.branches[:50]):
        if branch.values and _operation_key(branch.operation):
            groups.setdefault((branch.condition_id, branch.goto_step), []).append((index, branch))
    issues: list[SupportGuideIssue] = []
    for (condition_id, goto_step), members in groups.items():
        operations = {_operation_key(branch.operation): branch for _, branch in members}
        if len(operations) < 2:
            continue
        values = "」「".join("／".join(branch.values) for branch in operations.values())
        issues.append(
            SupportGuideIssue(
                severity="error",
                code="branch_operations_merged",
                path=f"branches[{members[0][0]}]",
                message=(
                    f"確認事項「{labels.get(condition_id, condition_id)}」の値「{values}」は"
                    "回答の過程で別の操作ですが、同じ手順"
                    f"「{titles.get(goto_step, goto_step)}」に進みます。"
                    "値ごとに手順と資料を分けて、もう一度下書きを作ってください。"
                ),
            )
        )
    return issues


def build_item(
    row: SupportGuideProcessRow,
    output: GuideDraftOutput,
    documents: Sequence[KnowledgeDocument],
) -> SupportGuideProcessItem:
    """LLM の応答を業務ガイドの下書きにし、資料の対応と検証の結果を付ける。"""
    references, issues = _references(output, documents)
    issues.extend(merged_branch_issues(output))
    rules = [
        SupportGuideRuleCandidate(
            title=rule.title[:200] or rule.content[:40],
            triggers=_clean(rule.triggers, 32, 160),
            content=rule.content[:2000],
        )
        for rule in output.rule_candidates[:10]
        if rule.content.strip()
    ]
    content: SupportGuideContent | None = None
    try:
        content = SupportGuideContent.model_validate(
            _content_payload(output, row.question, references)
        )
    except ValidationError as error:
        issues.extend(
            SupportGuideIssue(
                severity="error",
                code="invalid_content",
                path=".".join(str(part) for part in item["loc"]),
                message=str(item["msg"]),
            )
            for item in error.errors()[:20]
        )
    if content is not None:
        issues = [*validate_content(content), *issues]
    return SupportGuideProcessItem(
        row=row.row,
        question=row.question,
        valid=content is not None and not has_errors(issues),
        content=content,
        issues=issues,
        references=references,
        rule_candidates=rules,
    )


def failed_item(row: SupportGuideProcessRow, message: str) -> SupportGuideProcessItem:
    """下書きを作れなかった行（LLM の失敗・応答の形の誤り）。"""
    return SupportGuideProcessItem(
        row=row.row,
        question=row.question,
        valid=False,
        issues=[SupportGuideIssue(severity="error", code="draft_failed", path="", message=message)],
    )


DraftCaller = Callable[
    [Sequence[SupportGuideProcessRow], Callable[[SupportGuideProcessRow, Attempt], Awaitable[str]]],
    Awaitable[list[tuple[str | None, str | None]]],
]


async def draft_guides(
    rows: Sequence[SupportGuideProcessRow],
    documents: Sequence[KnowledgeDocument],
    *,
    llm: GuideDraftLlm,
    headings: Mapping[str, str],
    run: DraftCaller,
) -> list[SupportGuideProcessItem]:
    """行ごとに LLM で下書きを作る。

    ``run`` は上限つき並行で呼び、行の順に (応答, 失敗の文) を返す。
    """

    async def once(row: SupportGuideProcessRow, _: Attempt) -> str:
        return await async_call_with_model_purpose(
            "guide.draft",
            llm.generate,
            resolve_prompt("support_guide_draft.instruction", GUIDE_DRAFT_PROMPT),
            draft_context(row, headings),
            system_prompt=resolve_prompt("support_guide_draft.system", GUIDE_DRAFT_SYSTEM_PROMPT),
            response_schema=GUIDE_DRAFT_RESPONSE_SCHEMA,
            response_schema_name="support_guide_draft",
        )

    outcomes = await run(rows, once)
    items: list[SupportGuideProcessItem] = []
    for row, (raw, error) in zip(rows, outcomes, strict=True):
        if error is not None or raw is None:
            items.append(failed_item(row, error or "下書きを作れませんでした。"))
            continue
        try:
            output = parse_draft_output(raw)
        except ValueError as exc:
            logger.warning("support guide draft invalid", extra={"row": row.row})
            items.append(failed_item(row, f"下書きを作れませんでした（{exc}）。"))
            continue
        items.append(build_item(row, output, documents))
    return items


__all__ = [
    "GUIDE_DRAFT_PROMPT",
    "GUIDE_DRAFT_RESPONSE_SCHEMA",
    "GUIDE_DRAFT_SYSTEM_PROMPT",
    "PROCESS_WORKBOOK_MAX_BYTES",
    "GuideDraftOutput",
    "KnowledgeDocument",
    "ProcessSection",
    "build_item",
    "draft_context",
    "draft_guides",
    "failed_item",
    "list_knowledge_documents",
    "parse_draft_output",
    "merged_branch_issues",
    "process_headings",
    "read_process_workbook",
    "readable_as_pages",
    "resolve_document_name",
    "split_process_sections",
]
