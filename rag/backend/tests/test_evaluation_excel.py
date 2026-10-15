"""品質評価の評価リクエストの Excel（#2065）: テンプレート・構造の検証・ファイル名の解決・API。"""

from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]

from app.api.routes import evaluation as evaluation_route
from app.main import app
from app.rag.evaluation import evidence_matches_chunk
from app.rag.evaluation_excel import (
    CASES_SHEET_TITLE,
    EVALUATION_EXCEL_MAX_BYTES,
    EVALUATION_EXCEL_MAX_ROWS,
    GUIDE_SHEET_TITLE,
    HEADERS,
    EvaluationDocumentRef,
    build_evaluation_cases,
    evaluation_cases_xlsx,
    parse_evaluation_workbook,
    parse_evidence_line,
    template_xlsx,
)
from app.schemas.document import DocumentSummary
from app.schemas.evaluation import EvaluationEvidence
from app.schemas.search import RetrievedChunk
from tests.support import AsgiTestClient

client = AsgiTestClient(app)

DOCUMENTS = [
    EvaluationDocumentRef(document_id="doc-expense", file_name="経費精算規程.pdf"),
    EvaluationDocumentRef(document_id="doc-travel", file_name="出張旅費規程.docx"),
    EvaluationDocumentRef(
        document_id="doc-travel-old",
        file_name="出張旅費規程_旧版.docx",
        superseded=True,
        superseded_by_file_name="出張旅費規程.docx",
    ),
]


def _workbook(rows: list[list[object]], *, header: list[object] | None = None) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = CASES_SHEET_TITLE
    sheet.append(header if header is not None else list(HEADERS))
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def test_template_has_headers_and_guide_sheet() -> None:
    book = load_workbook(BytesIO(template_xlsx()))

    assert book.sheetnames == [CASES_SHEET_TITLE, GUIDE_SHEET_TITLE]
    cases = book[CASES_SHEET_TITLE]
    assert [cell.value for cell in cases[1]] == ["質問", "期待する根拠", "標準回答"]
    # 見出しのセルに意味・必須か任意か・書き方を注記する。
    assert "必須" in cases["A1"].comment.text
    assert "任意" in cases["B1"].comment.text
    assert cases.max_row == 1
    guide_text = "\n".join(
        str(cell.value) for row in book[GUIDE_SHEET_TITLE].iter_rows() for cell in row if cell.value
    )
    for phrase in (
        "必須",
        "任意",
        "p.12-14",
        "章節:",
        "抜粋:",
        "記入例",
        f"{EVALUATION_EXCEL_MAX_ROWS} 行",
    ):
        assert phrase in guide_text


def test_template_round_trips_through_the_parser() -> None:
    # 書き出した Excel（フィードバックの 1 行と同じ形）をそのまま読み込める。
    data = evaluation_cases_xlsx(
        [("経費精算の締め日はいつですか？", ["経費精算規程.pdf | p.3"], "毎月 25 日です。")]
    )

    parsed = parse_evaluation_workbook(data)

    assert parsed.issues == []
    assert [(row.row, row.question, row.standard_answer) for row in parsed.rows] == [
        (2, "経費精算の締め日はいつですか？", "毎月 25 日です。")
    ]
    assert parsed.rows[0].evidence[0].page_start == 3


def test_parse_reads_rows_and_skips_empty_rows() -> None:
    data = _workbook(
        [
            ["経費精算の締め日はいつですか？", "経費精算規程.pdf", "毎月 25 日です。"],
            [None, None, None],
            [
                "出張の宿泊費の上限は？",
                "出張旅費規程.docx | p.4-5 | 章節: 宿泊費\n"
                "出張旅費規程.docx｜抜粋：1 泊あたりの上限",
                None,
            ],
            ["社員食堂の来月の献立は？", "", ""],
        ],
        # 見出しの注記（「（必須）」）と英語の列名でも読める。
        header=["質問（必須）", "reference_contexts", "ground_truth"],
    )

    parsed = parse_evaluation_workbook(data)

    assert parsed.issues == []
    assert [row.row for row in parsed.rows] == [2, 4, 5]
    travel = parsed.rows[1]
    assert travel.case_id == "row-4"
    assert [
        (spec.file_name, spec.page_start, spec.page_end, spec.section, spec.text)
        for spec in travel.evidence
    ] == [
        ("出張旅費規程.docx", 4, 5, "宿泊費", None),
        ("出張旅費規程.docx", None, None, None, "1 泊あたりの上限"),
    ]
    assert parsed.rows[2].evidence == ()
    assert parsed.rows[2].standard_answer is None


@pytest.mark.parametrize(
    ("line", "expected"),
    [
        ("経費精算規程.pdf", ("経費精算規程.pdf", None, None, None, None)),
        ("経費精算規程.pdf | p.3", ("経費精算規程.pdf", 3, None, None, None)),
        ("経費精算規程.pdf | P12-14", ("経費精算規程.pdf", 12, 14, None, None)),
        ("経費精算規程.pdf | １２〜１４ページ", ("経費精算規程.pdf", 12, 14, None, None)),
        ("経費精算規程.pdf | p.5-5", ("経費精算規程.pdf", 5, None, None, None)),
        (
            "経費精算規程.pdf | 抜粋: 部門長が承認する | 章節：申請の手続き",
            ("経費精算規程.pdf", None, None, "申請の手続き", "部門長が承認する"),
        ),
    ],
    ids=["file-only", "page", "page-range", "fullwidth-page", "same-page", "excerpt-and-section"],
)
def test_parse_evidence_line(line: str, expected: tuple[object, ...]) -> None:
    spec = parse_evidence_line(line)

    assert not isinstance(spec, str)
    assert (spec.file_name, spec.page_start, spec.page_end, spec.section, spec.text) == expected


@pytest.mark.parametrize(
    ("line", "message"),
    [
        ("p.3 | 経費精算規程.pdf", "先頭に文書のファイル名"),
        ("経費精算規程.pdf | 3 ページ目あたり", "を読めません"),
        ("経費精算規程.pdf | p.14-12", "小さい番号から"),
        ("経費精算規程.pdf | p.0", "1 以上の番号"),
        ("経費精算規程.pdf | p.3 | p.4", "ページは 1 行に 1 つ"),
        ("経費精算規程.pdf | 章節:", "見出しを書いて"),
        ("経費精算規程.pdf | 抜粋: " + "あ" * 301, "300 文字以内"),
    ],
    ids=[
        "no-file-name",
        "unknown-part",
        "reversed-pages",
        "page-zero",
        "two-pages",
        "empty-section",
        "long-excerpt",
    ],
)
def test_parse_evidence_line_errors(line: str, message: str) -> None:
    result = parse_evidence_line(line)

    assert isinstance(result, str)
    assert message in result


def test_parse_reports_row_and_column_errors() -> None:
    data = _workbook(
        [
            [None, "経費精算規程.pdf", "回答"],
            ["質" * 2001, None, None],
            ["正しい行", "経費精算規程.pdf\n経費精算規程.pdf | p.9-1", None],
            ["正しい行 2", None, "長" * 8001],
        ]
    )

    parsed = parse_evaluation_workbook(data)

    assert [(issue.row, issue.column) for issue in parsed.issues] == [
        (2, "質問"),
        (3, "質問"),
        (4, "期待する根拠"),
        (5, "標準回答"),
    ]
    assert "質問を入力してください" in parsed.issues[0].message
    assert "2000 文字以内" in parsed.issues[1].message
    # セルの中の何行目の根拠かも示す。
    assert parsed.issues[2].message.startswith("2 行目: ")
    assert parsed.rows == []


@pytest.mark.parametrize(
    ("data", "message"),
    [
        (b"not an excel file", "Excel（.xlsx）として読めません"),
        (b"0" * (EVALUATION_EXCEL_MAX_BYTES + 1), "ファイルが大きすぎます"),
        (_workbook([], header=["問い", "答え"]), "「質問」の見出しがありません"),
        (_workbook([]), "1 行もありません"),
        (
            _workbook([[f"質問 {index}"] for index in range(EVALUATION_EXCEL_MAX_ROWS + 1)]),
            "200 行まで",
        ),
    ],
    ids=["not-xlsx", "too-large", "no-header", "no-rows", "too-many-rows"],
)
def test_parse_reports_file_errors(data: bytes, message: str) -> None:
    parsed = parse_evaluation_workbook(data)

    assert len(parsed.issues) == 1
    assert message in parsed.issues[0].message
    assert parsed.rows == []


def test_build_cases_resolves_file_names_in_the_knowledge_bases() -> None:
    parsed = parse_evaluation_workbook(
        _workbook(
            [
                ["締め日は？", "ｹｲﾋﾁｮｳ.PDF", None],
                [
                    "宿泊費の上限は？",
                    "経費精算規程.PDF\n出張旅費規程.docx | p.4 | 抜粋: 1 泊あたりの上限",
                    "1 泊 1 万円までです。",
                ],
                ["献立は？", None, None],
            ]
        )
    )
    documents = [
        *DOCUMENTS,
        EvaluationDocumentRef(document_id="doc-kana", file_name="ケイヒチョウ.pdf"),
    ]

    cases, issues = build_evaluation_cases(parsed.rows, documents)

    assert issues == []
    # NFKC（半角カナ）と大小文字の違いは無視する。
    assert cases[0].relevant_document_ids == ["doc-kana"]
    assert cases[0].required_evidence == []
    lodging = cases[1]
    assert lodging.id == "row-3"
    assert lodging.relevant_document_ids == ["doc-expense", "doc-travel"]
    assert [evidence.model_dump(exclude_none=True) for evidence in lodging.required_evidence] == [
        {
            "id": "e2",
            "document_id": "doc-travel",
            "text": "1 泊あたりの上限",
            "page_start": 4,
            "label": "出張旅費規程.docx | p.4 | 抜粋: 1 泊あたりの上限",
        }
    ]
    assert lodging.standard_answer == "1 泊 1 万円までです。"
    # 根拠も標準回答も無い行は、資料に答えが無い質問（拒答の正しさを測る）。
    assert cases[2].expects_answer is False


def test_build_cases_reports_unknown_ambiguous_and_superseded_file_names() -> None:
    parsed = parse_evaluation_workbook(
        _workbook(
            [
                ["q1", "存在しない規程.pdf", None],
                ["q2", "就業規則.pdf", None],
                ["q3", "出張旅費規程_旧版.docx", None],
                ["q4", "経費精算規程.pdf", None],
            ]
        )
    )
    documents = [
        *DOCUMENTS,
        EvaluationDocumentRef(document_id="doc-a", file_name="就業規則.pdf"),
        EvaluationDocumentRef(document_id="doc-b", file_name="就業規則.pdf"),
    ]

    cases, issues = build_evaluation_cases(parsed.rows, documents)

    assert [(issue.row, issue.column) for issue in issues] == [
        (2, "期待する根拠"),
        (3, "期待する根拠"),
        (4, "期待する根拠"),
    ]
    assert "選んだナレッジベースの文書にありません" in issues[0].message
    assert "2 件あります" in issues[1].message
    assert "新しい版（出張旅費規程.docx）" in issues[2].message
    assert [case.id for case in cases] == ["row-5"]


def _chunk(**metadata: Any) -> RetrievedChunk:
    return RetrievedChunk(
        document_id=metadata.pop("document_id", "doc-1"),
        chunk_id="doc-1:c1",
        text=metadata.pop("text", "申請は部門長が承認する。"),
        score=1.0,
        metadata=metadata,
    )


@pytest.mark.parametrize(
    ("evidence", "chunk", "matched"),
    [
        ({"page_start": 3}, _chunk(page_start=2, page_end=4), True),
        ({"page_start": 5, "page_end": 6}, _chunk(page_start=2, page_end=4), False),
        ({"page_start": 3}, _chunk(page_number="3"), True),
        ({"page_start": 3}, _chunk(), False),
        ({"section": "申請の手続き"}, _chunk(section_path="第 2 章 > 申請の手続き"), True),
        ({"section": "宿泊費"}, _chunk(section_path="第 2 章 > 申請の手続き"), False),
        ({"text": "部門長が承認"}, _chunk(), True),
        ({"text": "部門長が承認", "page_start": 3}, _chunk(page_start=3), True),
        ({"text": "部門長が承認"}, _chunk(document_id="doc-2"), False),
    ],
    ids=[
        "page-overlap",
        "page-outside",
        "page-number",
        "no-page-info",
        "section-partial",
        "section-other",
        "text",
        "text-and-page",
        "other-document",
    ],
)
def test_evidence_matches_chunk(
    evidence: dict[str, Any], chunk: RetrievedChunk, matched: bool
) -> None:
    item = EvaluationEvidence(id="e1", document_id="doc-1", **evidence)

    assert evidence_matches_chunk(item, chunk) is matched


def test_evidence_requires_a_condition_and_ordered_pages() -> None:
    with pytest.raises(ValueError, match="どれかを指定"):
        EvaluationEvidence(id="e1", document_id="doc-1")
    with pytest.raises(ValueError, match="page_start 以上"):
        EvaluationEvidence(id="e1", page_start=5, page_end=3)


class _FakeOracle:
    def __init__(self) -> None:
        self.knowledge_base_ids: list[str | None] = []

    async def list_documents(self, **kwargs: Any) -> list[DocumentSummary]:
        self.knowledge_base_ids.append(kwargs.get("knowledge_base_id"))
        now = datetime(2026, 10, 1, tzinfo=UTC)
        return [
            DocumentSummary(
                id="doc-expense", file_name="経費精算規程.pdf", status="INDEXED", uploaded_at=now
            ),
            DocumentSummary(
                id="doc-other", file_name="その他.pdf", status="INDEXED", uploaded_at=now
            ),
        ]


def _upload(
    data: bytes, *, name: str = "cases.xlsx", knowledge_base_ids: list[str] | None = None
) -> Any:
    return client.post(
        "/api/evaluation/cases/import",
        files={
            "file": (
                name,
                data,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        data={
            "knowledge_base_ids": knowledge_base_ids if knowledge_base_ids is not None else ["kb-1"]
        },
    )


def test_import_api_returns_cases_preview_and_document_names(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    fake = _FakeOracle()
    monkeypatch.setattr(evaluation_route, "OracleClient", lambda: fake)

    response = _upload(
        _workbook(
            [
                ["締め日は？", "経費精算規程.pdf | p.3", "毎月 25 日です。" * 30],
                ["献立は？", None, None],
            ]
        ),
        knowledge_base_ids=["kb-1", "kb-2", "kb-1"],
    )

    assert response.status_code == 200, response.text
    data = response.json()["data"]
    assert data["errors"] == []
    assert data["file_name"] == "cases.xlsx"
    assert data["row_count"] == 2
    assert data["knowledge_base_ids"] == ["kb-1", "kb-2"]
    assert fake.knowledge_base_ids == ["kb-1", "kb-2"]
    assert [case["id"] for case in data["cases"]] == ["row-2", "row-3"]
    assert data["cases"][0]["relevant_document_ids"] == ["doc-expense"]
    assert data["document_names"] == {"doc-expense": "経費精算規程.pdf"}
    preview = data["rows"][0]
    assert preview["evidence"] == ["経費精算規程.pdf | p.3"]
    assert preview["has_standard_answer"] is True
    assert preview["standard_answer_preview"].endswith("…")
    assert data["rows"][1]["expects_answer"] is False
    assert data["max_rows"] == EVALUATION_EXCEL_MAX_ROWS


def test_import_api_returns_row_errors_without_cases(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(evaluation_route, "OracleClient", _FakeOracle)

    response = _upload(
        _workbook([["締め日は？", "無い規程.pdf", None], [None, "経費精算規程.pdf", None]])
    )

    assert response.status_code == 200
    data = response.json()["data"]
    assert data["cases"] == []
    assert data["error_count"] == 2
    assert [(issue["row"], issue["column"]) for issue in data["errors"]] == [
        (2, "期待する根拠"),
        (3, "質問"),
    ]


def test_import_api_does_not_read_documents_without_evidence(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    fake = _FakeOracle()
    monkeypatch.setattr(evaluation_route, "OracleClient", lambda: fake)

    response = _upload(_workbook([["献立は？", None, None]]))

    assert response.json()["data"]["cases"][0]["id"] == "row-2"
    assert fake.knowledge_base_ids == []


def test_import_api_requires_knowledge_base_and_xlsx() -> None:
    no_kb = _upload(_workbook([["q", None, None]]), knowledge_base_ids=[" "])
    assert no_kb.status_code == 422
    assert "ナレッジベースを選んでください" in no_kb.json()["error_messages"][0]

    csv = _upload(b"q,a\n", name="cases.csv")
    assert csv.status_code == 200
    assert "Excel（.xlsx）" in csv.json()["data"]["errors"][0]["message"]


def test_template_api_returns_xlsx() -> None:
    response = client.get("/api/evaluation/cases/template")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "evaluation_request_template.xlsx" in response.headers["content-disposition"]
    assert load_workbook(BytesIO(response.content)).sheetnames == [
        CASES_SHEET_TITLE,
        GUIDE_SHEET_TITLE,
    ]
