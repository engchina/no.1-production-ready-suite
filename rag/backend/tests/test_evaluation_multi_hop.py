"""多段の質問（multi-hop）の評価セットと指標のテスト（#1335・#1352）。

- 評価セット（`rag/evaluation/multi-hop/multi-hop.json`）が今の評価のスキーマで読め、必要な根拠の
  語句が、指した資料の原稿（HTML / 表の原稿）と同梱の PDF / xlsx に実際に書いてある。
- 根拠の連鎖の完全率（`evidence_chain_complete_rate`）と、種類別・段の数別の内訳。
- 資料の生成の script（`rag/scripts/generate_evaluation_corpus.py`）が、表の原稿から同梱と同じ
  xlsx を作る。多段の資料の原稿と評価セットは、実体のデータ（`rag/scripts/multi_hop_corpus.py`）から
  同梱と同じものができる（#1352）。
"""

import importlib.util
import json
import re
from collections.abc import Sequence
from pathlib import Path
from types import ModuleType
from typing import Any

import pytest
from pydantic import ValidationError

from app.rag.evaluation import case_error_result, score_case_answers, summarize_case_results
from app.rag.evaluation_handling import contains_normalized
from app.schemas.evaluation import (
    EvaluationCase,
    EvaluationCaseResult,
    EvaluationMetrics,
    EvaluationRunRequest,
)
from app.schemas.search import RetrievedChunk, SearchDiagnostics, SearchResponse

RAG_DIR = Path(__file__).resolve().parents[2]
CORPUS = RAG_DIR / "evaluation" / "multi-hop"
GENERATOR = RAG_DIR / "scripts" / "generate_evaluation_corpus.py"
SOURCE_BUILDER = RAG_DIR / "scripts" / "multi_hop_corpus.py"
# #1335 の 33 問（#1352 で資料を増やしても ID・質問・期待する語を変えない）。
ORIGINAL_CASE_COUNT = 33


def _generator() -> ModuleType:
    spec = importlib.util.spec_from_file_location("generate_evaluation_corpus_1335", GENERATOR)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _source_builder() -> ModuleType:
    module: ModuleType | None = _generator().source_builder(CORPUS)
    assert module is not None
    return module


def _load_set() -> tuple[dict[str, Any], EvaluationRunRequest]:
    payload = json.loads((CORPUS / "multi-hop.json").read_text(encoding="utf-8"))
    return payload, EvaluationRunRequest.model_validate(payload)


def _html_text(name: str) -> str:
    source = CORPUS / "sources" / f"{Path(name).stem}.html"
    return re.sub(r"<[^>]+>", "", source.read_text(encoding="utf-8"))


def _workbook_cells(spec_path: Path) -> list[str]:
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    return [
        str(value)
        for sheet in spec["sheets"]
        for line in [*([item] for item in sheet.get("preamble", [])), sheet["header"]]
        + sheet["rows"]
        for value in line
    ]


def _xlsx_cells(path: Path) -> list[tuple[str, list[list[tuple[Any, str]]]]]:
    from openpyxl import load_workbook  # type: ignore[import-untyped]

    workbook = load_workbook(path)
    return [
        (
            sheet.title,
            [[(cell.value, cell.number_format) for cell in row] for row in sheet.iter_rows()],
        )
        for sheet in workbook.worksheets
    ]


# ---- 評価セット ----------------------------------------------------------------------------


def test_multi_hop_set_has_every_reasoning_type_and_both_splits() -> None:
    """60 問以上で、種類（対照の 1 段を含む）と区分がそろい、各ケースが段ごとの根拠を持つ。"""
    _payload, request = _load_set()
    cases = request.cases
    assert 60 <= len(cases) <= 90
    assert len({case.id for case in cases}) == len(cases)
    assert {case.reasoning_type for case in cases} == {
        "single_hop",
        "bridge",
        "comparison",
        "intra_document_reference",
        "table_lookup",
    }
    for split in ("dev", "holdout"):
        members = [case for case in cases if case.split == split]
        # 区分のどちらにも、すべての種類がある（dev の結果だけで基準を決められるように）。
        assert {case.reasoning_type for case in members} == {
            case.reasoning_type for case in cases
        }, split
    holdout = [case for case in cases if case.split == "holdout"]
    assert 0.3 <= len(holdout) / len(cases) <= 0.45
    # 多段の質問が中心で、1 段の対照は数問。3 段以上の橋渡しもある。
    assert sum(case.reasoning_type == "single_hop" for case in cases) <= 6
    assert any(case.reasoning_type == "bridge" and (case.hops or 0) >= 3 for case in cases)
    for case in cases:
        assert case.hops is not None and case.required_evidence, case.id
        assert case.expected_answer_keywords and case.expected_outcomes, case.id
        # 正解の文書は、必要な根拠の文書と同じ（多段は複数の文書をまたぐ）。
        assert set(case.relevant_document_ids) == {
            item.document_id for item in case.required_evidence
        }, case.id
    assert any(len(case.relevant_document_ids) >= 3 for case in cases)


def test_multi_hop_evidence_is_written_in_its_sources_and_corpus() -> None:
    """必要な根拠の語句は、指した資料の原稿と、同梱の PDF / xlsx に実際に書いてある。"""
    from pypdf import PdfReader

    from app.rag.evaluation_corpus_cli import referenced_files

    payload, request = _load_set()
    names = referenced_files(payload)
    # 紛らわしい資料（旧版・別の会社・似た承認の規程）も、評価の CLI が取り込むよう 1 問は参照する。
    assert len(names) == 15
    assert all((CORPUS / name).is_file() for name in names)
    pdf_text = {
        name: "\n".join(page.extract_text() for page in PdfReader(CORPUS / name).pages)
        for name in names
        if name.endswith(".pdf")
    }
    workbook_cells = {
        name: _workbook_cells(CORPUS / "sources" / f"{Path(name).stem}.workbook.json")
        for name in names
        if name.endswith(".xlsx")
    }
    xlsx_cells = {
        name: [
            str(value)
            for _title, rows in _xlsx_cells(CORPUS / name)
            for row in rows
            for value, _format in row
            if value is not None
        ]
        for name in workbook_cells
    }
    for case in request.cases:
        for item in case.required_evidence:
            assert item.document_id is not None and item.document_id.startswith("file:")
            name = item.document_id.removeprefix("file:")
            if name.endswith(".pdf"):
                assert contains_normalized(_html_text(name), item.text), (case.id, item.id)
                assert contains_normalized(pdf_text[name], item.text), (case.id, item.id)
            else:
                # 表は 1 行 1 記録で取り込む。根拠は行を指すセルの値（システム ID など）にする。
                assert item.text in workbook_cells[name], (case.id, item.id)
                assert item.text in xlsx_cells[name], (case.id, item.id)


def test_multi_hop_corpus_has_notation_variants_and_references() -> None:
    """実体の表記ゆれ（全角 / 半角・カナ / 英字・略称 / 正式名・1 文字の部署名）と文書内の参照。"""
    ledger = _workbook_cells(CORPUS / "sources" / "system-ledger.workbook.json")
    maintenance = _html_text("maintenance-plan.pdf")
    incident = _html_text("incident-contact-rules.pdf")
    organization = _html_text("organization-rules.pdf")
    # 全角 / 半角: 台帳は半角、保守計画・障害連絡は全角。
    assert "HRM" in ledger and "ＨＲＭ" in maintenance and "HRM" not in maintenance
    assert "SYS-104" in ledger and "ＳＹＳ－１０４" in incident
    # カナ / 英字: 台帳の正式名はカナ、保守計画は英字。
    assert "ドキュメントポータル" in ledger and "Document Portal" in maintenance
    # 略称 / 正式名: 部署の通称（情シス）と、システムの別表記（経費 Portal）。
    assert "通称「情シス」" in organization and "情シス部長" in incident
    assert "経費 Portal" in ledger and "経費 Portal" in _html_text("change-procedure.pdf")
    # 1 文字の漢字の部署名（台帳の担当部署の列）。
    assert {"情", "経", "人", "総", "営"} <= set(ledger)
    # 文書の中の参照。
    assert "第 2 章の共通の保守枠" in maintenance
    assert "第 4 章を参照してください" in organization


# ---- スキーマ ------------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("reasoning_type", "hops", "evidence_count"),
    [
        ("single_hop", 1, 1),
        ("bridge", 3, 3),
        ("comparison", 1, 2),
        ("table_lookup", 1, 1),
        ("intra_document_reference", 2, 2),
        ("bridge", None, 1),
    ],
    ids=["single", "bridge", "comparison-1", "table-1", "intra", "type-only"],
)
def test_case_accepts_reasoning_type_and_hops(
    reasoning_type: str, hops: int | None, evidence_count: int
) -> None:
    case = EvaluationCase.model_validate(
        {
            "id": "c",
            "query": "q",
            "reasoning_type": reasoning_type,
            "hops": hops,
            "required_evidence": [{"id": f"e{i}", "text": "t"} for i in range(evidence_count)],
        }
    )
    assert (case.reasoning_type, case.hops) == (reasoning_type, hops)


@pytest.mark.parametrize(
    ("overrides", "message"),
    [
        ({"hops": 2}, "reasoning_type と一緒に"),
        ({"reasoning_type": "single_hop", "hops": 2}, "single_hop"),
        ({"reasoning_type": "bridge", "hops": 1}, "2 以上"),
        ({"reasoning_type": "intra_document_reference", "hops": 1}, "2 以上"),
        ({"reasoning_type": "bridge", "hops": 3}, "段ごとに"),
        ({"reasoning_type": "bridge", "hops": 0}, "greater than or equal"),
        ({"reasoning_type": "bridge", "hops": 11}, "less than or equal"),
        ({"reasoning_type": "two_hop"}, "reasoning_type"),
    ],
    ids=[
        "hops-without-type",
        "single-hop-2",
        "bridge-1",
        "intra-1",
        "fewer-evidence",
        "zero",
        "too-many",
        "unknown-type",
    ],
)
def test_case_rejects_inconsistent_reasoning(overrides: dict[str, Any], message: str) -> None:
    payload: dict[str, Any] = {
        "id": "c",
        "query": "q",
        "required_evidence": [{"id": "e1", "text": "a"}, {"id": "e2", "text": "b"}],
        **overrides,
    }
    with pytest.raises(ValidationError, match=message):
        EvaluationCase.model_validate(payload)


# ---- 指標 ----------------------------------------------------------------------------------


def _response(chunks: Sequence[tuple[str, str]], answer: str = "回答", **details: Any) -> Any:
    return SearchResponse(
        answer=answer,
        citations=[
            RetrievedChunk(document_id=doc, chunk_id=f"{doc}:{index}", text=text, score=1.0)
            for index, (doc, text) in enumerate(chunks)
        ],
        trace_id="trace",
        guardrail_warnings=[],
        elapsed_ms=1.0,
        diagnostics=SearchDiagnostics(answer=details or None),
    )


def _chain_case(case_id: str, reasoning_type: str, hops: int, **extra: Any) -> EvaluationCase:
    return EvaluationCase.model_validate(
        {
            "id": case_id,
            "query": "q",
            "split": "dev",
            "reasoning_type": reasoning_type,
            "hops": hops,
            "expected_answer_keywords": ["管理本部長"],
            "required_evidence": [
                {"id": "ledger", "document_id": "ledger", "text": "SYS-101"},
                {"id": "org", "document_id": "org", "text": "経理部の承認者は管理本部長です"},
            ],
            **extra,
        }
    )


def _scored(case: EvaluationCase, response: SearchResponse) -> EvaluationCaseResult:
    return score_case_answers(case, [(response, {})])


def test_evidence_chain_complete_rate_counts_cases_with_all_evidence() -> None:
    """全部そろう・1 つ欠ける・対象外（根拠の無いケース・検索をしない回答）・失敗のケース。"""
    complete = _scored(
        _chain_case("all", "bridge", 2),
        _response(
            [
                ("ledger", "SYS-101 経費精算ポータル 経"),
                ("org", "経理部の承認者は管理本部長です。"),
            ],
            answer="管理本部長が承認します。",
        ),
    )
    partial = _scored(
        _chain_case("one-missing", "bridge", 2),
        _response([("ledger", "SYS-101")], answer="経理部長が承認します。"),
    )
    no_evidence = _scored(
        EvaluationCase(id="none", query="q", expected_answer_keywords=["x"]),
        _response([("ledger", "x")], answer="x"),
    )
    clarification = _scored(
        _chain_case("clarify", "table_lookup", 1),
        _response([], answer="どのシステムですか？", outcome="needs_clarification"),
    )
    failed = case_error_result(
        _chain_case("error", "bridge", 2),
        trace_id="t",
        elapsed_ms=0.0,
        error_type="RuntimeError",
        error_message="失敗",
    )

    assert (complete.evidence_recall, complete.evidence_chain_complete) == (1.0, True)
    assert (partial.evidence_recall, partial.evidence_chain_complete) == (0.5, False)
    assert partial.missing_evidence == ["org"]
    assert no_evidence.evidence_chain_complete is None
    assert clarification.evidence_chain_complete is None
    assert (failed.reasoning_type, failed.hops) == ("bridge", 2)

    metrics = summarize_case_results([complete, partial, no_evidence, clarification, failed])
    # 再現率の平均（0.75）とは別に、全部そろったケースの割合（0.5）を見る。
    assert metrics.required_evidence_recall == 0.75
    assert metrics.evidence_chain_complete_rate == 0.5
    assert metrics.metric_case_counts["evidence_chain_complete_rate"] == 2
    assert metrics.split_breakdown["dev"].metrics["evidence_chain_complete_rate"] == 0.5


def test_required_evidence_matches_only_the_document_of_the_evaluation_set() -> None:
    """必要な根拠は評価セットの文書 ID で照合し、同じ内容の別の文書の引用では数えない(#1381)。

    評価のナレッジベースの文書は、前の評価で取り込んだ同じ内容の文書と別の文書になる。検索の範囲は
    そのナレッジベースの文書だけ(`_oracle_retrieval_where`)なので、別の文書の引用は範囲の漏れであり、
    数えずに欠けとして出す(ファイル名・内容のハッシュで同じとみなすと、旧版の登録の無い文書の漏れが
    隠れる)。
    """
    case = _chain_case("other-kb", "bridge", 2)
    leaked = _scored(
        case,
        _response(
            [
                ("ledger", "SYS-101 経費精算ポータル 経"),
                ("org-earlier-kb", "経理部の承認者は管理本部長です。"),
            ],
            answer="管理本部長が承認します。",
        ),
    )

    assert (leaked.evidence_recall, leaked.missing_evidence) == (0.5, ["org"])
    assert leaked.evidence_chain_complete is False


def test_reasoning_and_hops_breakdowns() -> None:
    """種類別・段の数別に、根拠の再現率・連鎖の完全率・期待する語の一致率と件数を出す。"""
    bridge_ok = _scored(
        _chain_case("b1", "bridge", 2),
        _response(
            [("ledger", "SYS-101"), ("org", "経理部の承認者は管理本部長です")],
            answer="管理本部長です。",
        ),
    )
    bridge_miss = _scored(
        _chain_case("b2", "bridge", 2), _response([("ledger", "SYS-101")], answer="わかりません")
    )
    comparison = _scored(
        _chain_case("c1", "comparison", 1, split="holdout"),
        _response(
            [("ledger", "SYS-101"), ("org", "経理部の承認者は管理本部長です")],
            answer="管理本部長",
        ),
    )
    single = _scored(
        EvaluationCase(id="s1", query="q", expected_answer_keywords=["x"]),
        _response([("ledger", "x")], answer="x"),
    )
    failed = case_error_result(
        _chain_case(
            "b3", "bridge", 3, required_evidence=[{"id": f"e{i}", "text": "t"} for i in range(3)]
        ),
        trace_id="t",
        elapsed_ms=0.0,
        error_type="RuntimeError",
        error_message="失敗",
    )
    metrics = summarize_case_results([bridge_ok, bridge_miss, comparison, single, failed])

    assert list(metrics.reasoning_type_breakdown) == ["bridge", "comparison", "unspecified"]
    bridge = metrics.reasoning_type_breakdown["bridge"]
    assert (bridge.case_count, bridge.error_count) == (3, 1)
    assert bridge.metrics == {
        "required_evidence_recall": 0.75,
        "evidence_chain_complete_rate": 0.5,
        "answer_keyword_hit_rate": 0.5,
    }
    assert bridge.metric_case_counts["evidence_chain_complete_rate"] == 2
    assert (
        metrics.reasoning_type_breakdown["comparison"].metrics["evidence_chain_complete_rate"]
        == 1.0
    )
    unspecified = metrics.reasoning_type_breakdown["unspecified"]
    assert unspecified.metrics["evidence_chain_complete_rate"] is None
    assert unspecified.metrics["answer_keyword_hit_rate"] == 1.0
    # 段の数の順（無いケースは最後）。
    assert list(metrics.hops_breakdown) == ["1", "2", "3", "unspecified"]
    assert metrics.hops_breakdown["2"].metrics["evidence_chain_complete_rate"] == 0.5
    assert metrics.hops_breakdown["3"].error_count == 1
    # 区分の中の内訳。
    dev = metrics.split_breakdown["dev"]
    assert set(dev.reasoning_type_breakdown) == {"bridge"}
    assert metrics.split_breakdown["holdout"].hops_breakdown["1"].case_count == 1


def test_breakdowns_are_empty_without_reasoning_fields() -> None:
    result = _scored(
        EvaluationCase(id="s", query="q", expected_answer_keywords=["x"]),
        _response([("doc", "x")], answer="x"),
    )
    metrics = summarize_case_results([result])
    assert metrics.reasoning_type_breakdown == {}
    assert metrics.hops_breakdown == {}
    assert metrics.evidence_chain_complete_rate is None


@pytest.mark.parametrize(
    ("answer", "keywords", "hit"),
    [
        ("ＨＲＭ の保守が先です。", ["HRM"], True),
        ("保管は 5 年です。", ["5年"], True),
        ("Approved", ["approved"], True),
        ("経理部長が承認します。", ["管理本部長"], False),
        ("回答", ["回答", " "], True),
    ],
    ids=["fullwidth", "space", "case", "miss", "blank-keyword"],
)
def test_answer_keywords_ignore_width_case_and_spaces(
    answer: str, keywords: list[str], hit: bool
) -> None:
    case = EvaluationCase(id="k", query="q", expected_answer_keywords=keywords)
    assert _scored(case, _response([("doc", "x")], answer=answer)).answer_keyword_hit is hit


# ---- CLI -----------------------------------------------------------------------------------


def _metrics_with_breakdowns() -> EvaluationMetrics:
    bridge = _scored(
        _chain_case("b1", "bridge", 2),
        _response(
            [("ledger", "SYS-101"), ("org", "経理部の承認者は管理本部長です")],
            answer="管理本部長です。",
        ),
    )
    single = _scored(
        _chain_case("s1", "comparison", 1, split="holdout"),
        _response([("ledger", "SYS-101")], answer="不明"),
    )
    return summarize_case_results([bridge, single])


def test_evaluation_cli_surfaces_chain_rate_and_breakdowns(tmp_path: Path) -> None:
    from app.rag.evaluation_cli import (
        GateEvaluation,
        _gate_summary,
        _load_evaluation_request,
        _metrics_trend,
    )

    metrics = _metrics_with_breakdowns()
    trend = _metrics_trend(metrics)
    assert trend["evidence_chain_complete_rate"] == 0.5
    assert (
        trend["reasoning_type_breakdown"]["bridge"]["metrics"]["evidence_chain_complete_rate"]
        == 1.0
    )
    assert trend["hops_breakdown"]["1"]["case_count"] == 1
    summary = _gate_summary(GateEvaluation(metrics=metrics), passed=False)
    assert "evidence_chain_complete_rate=0.5" in summary
    assert "reasoning_type=bridge: cases=1, errors=0" in summary
    assert "hops=2: cases=1" in summary

    golden = tmp_path / "set.json"
    golden.write_text((CORPUS / "multi-hop.json").read_text(encoding="utf-8"), encoding="utf-8")
    loaded = _load_evaluation_request(golden, split="dev")
    assert loaded.payload["cases"]
    assert {case["split"] for case in loaded.payload["cases"]} == {"dev"}
    assert {case["reasoning_type"] for case in loaded.payload["cases"]} >= {"bridge"}


def test_agent_evaluation_summary_shows_reasoning_rows() -> None:
    from app.rag.agent_evaluation_cli import summarize_result, summary_markdown

    data = _metrics_with_breakdowns().model_dump(mode="json")
    summary = summarize_result(data)
    assert summary["evidence_chain_complete_rate"] == 0.5
    assert summary["reasoning_types"]["bridge"]["evidence_chain_complete_rate"] == 1.0
    assert summary["hops"]["1"]["evidence_chain_complete_rate"] == 0.0
    assert summary["splits"]["dev"]["reasoning_types"]["bridge"]["case_count"] == 1
    table = summary_markdown({"A": summary, "D": summary})
    assert "| 根拠の連鎖の完全率 | 0.50 | 0.50 |" in table
    assert "| bridge: 根拠の連鎖の完全率 | 1.00 | 1.00 |" in table
    assert "| 2 段: 必要な根拠の再現率 | 1.00 | 1.00 |" in table
    assert "| dev / bridge: 根拠の連鎖の完全率 | 1.00 | 1.00 |" in table


# ---- 資料の生成 ----------------------------------------------------------------------------


@pytest.mark.parametrize(
    "corpus",
    ["multi-hop", "business-support"],
    ids=["multi-hop", "business-support"],
)
def test_generator_rebuilds_committed_workbooks(corpus: str, tmp_path: Path) -> None:
    """表の原稿から作った xlsx は、同梱の xlsx と同じセルの値・書式になる（作り直しても同じ）。"""
    generator = _generator()
    corpus_dir = RAG_DIR / "evaluation" / corpus
    specs = sorted((corpus_dir / "sources").glob("*.workbook.json"))
    assert specs
    for spec_path in specs:
        name = generator.workbook_name(spec_path)
        spec = generator.load_workbook_spec(spec_path)
        first, second = tmp_path / f"1-{name}", tmp_path / f"2-{name}"
        generator.write_workbook(spec, first)
        generator.write_workbook(spec, second)
        assert _xlsx_cells(first) == _xlsx_cells(corpus_dir / name)
        assert _xlsx_cells(first) == _xlsx_cells(second)


def test_generator_builds_only_workbooks_without_libreoffice(tmp_path: Path) -> None:
    generator = _generator()
    sources = tmp_path / "sources"
    sources.mkdir()
    (sources / "doc.html").write_text("<p>本文</p>", encoding="utf-8")
    (sources / "table.workbook.json").write_text(
        json.dumps(
            {
                "sheets": [
                    {"title": "表", "header": ["列"], "rows": [["030"]], "text_columns": ["A"]}
                ]
            }
        ),
        encoding="utf-8",
    )
    written = generator.build_corpus(tmp_path, xlsx_only=True)
    assert [path.name for path in written] == ["table.xlsx"]
    assert _xlsx_cells(tmp_path / "table.xlsx") == [("表", [[("列", "General")], [("030", "@")]])]


def test_generator_rejects_rows_that_do_not_match_the_header(tmp_path: Path) -> None:
    generator = _generator()
    spec_path = tmp_path / "bad.workbook.json"
    spec_path.write_text(
        json.dumps({"sheets": [{"title": "表", "header": ["a", "b"], "rows": [["1"]]}]}),
        encoding="utf-8",
    )
    with pytest.raises(ValueError, match="列の数"):
        generator.load_workbook_spec(spec_path)


# ---- 多段の資料の原稿と評価セットの生成（#1352） ------------------------------------------


def test_source_builder_rebuilds_committed_sources_and_set(tmp_path: Path) -> None:
    """実体のデータから作った原稿と評価セットは、同梱のものと同じ（作り直しても同じ）。"""
    builder = _source_builder()
    corpus_dir = tmp_path / "multi-hop"
    first = builder.write_sources(corpus_dir)
    second = builder.write_sources(corpus_dir)
    assert [path.name for path in first] == [path.name for path in second]
    committed = sorted(path.name for path in (CORPUS / "sources").iterdir())
    assert sorted(path.name for path in (corpus_dir / "sources").iterdir()) == committed
    for path in first:
        relative = path.relative_to(corpus_dir)
        assert path.read_bytes() == (CORPUS / relative).read_bytes(), relative


def test_generator_sources_only_writes_sources_and_set(tmp_path: Path) -> None:
    """`multi-hop` は変換の前に原稿と評価セットを作る（`--sources-only` はそこまで）。"""
    corpus_dir = tmp_path / "multi-hop"
    written = _generator().build_corpus(corpus_dir, sources_only=True)
    assert written[-1] == corpus_dir / "multi-hop.json"
    assert all(path.is_file() for path in written)
    assert not list(corpus_dir.glob("*.pdf")) and not list(corpus_dir.glob("*.xlsx"))


def test_multi_hop_set_keeps_the_original_cases() -> None:
    """#1335 の 33 問が先頭に同じ順で残り、#1352 の問は dev / holdout の両方にある。"""
    builder = _source_builder()
    _payload, request = _load_set()
    original = [spec.id for spec in builder.ORIGINAL_CASES]
    assert len(original) == ORIGINAL_CASE_COUNT
    assert [case.id for case in request.cases[:ORIGINAL_CASE_COUNT]] == original
    for spec, case in zip(builder.ORIGINAL_CASES, request.cases, strict=False):
        assert (case.query, case.expected_answer_keywords) == (spec.query, list(spec.keywords))
    added = request.cases[ORIGINAL_CASE_COUNT:]
    assert len(added) >= 30
    assert {case.split for case in added} == {"dev", "holdout"}


def test_multi_hop_corpus_is_larger_than_the_default_context() -> None:
    """既定の top_k で資料の全部が文脈に入らない大きさ（子 chunk の見積もりが 200 以上）。"""
    builder = _source_builder()
    payload, _request = _load_set()
    counts = builder.estimate_chunks(builder.build_documents())
    assert sum(counts.values()) >= 200
    assert sum(counts.values()) >= 10 * payload["top_k"]
    # システム台帳は 50〜100 行。
    assert 50 <= len(builder.SYSTEMS) <= 100


def test_multi_hop_corpus_has_confusing_entities_and_documents() -> None:
    """名前の似たシステム・部署、略称の衝突、旧版・別の会社・似た承認の規程がある。"""
    ledger = _workbook_cells(CORPUS / "sources" / "system-ledger.workbook.json")
    logistics = _workbook_cells(CORPUS / "sources" / "logistics-system-ledger.workbook.json")
    organization = _html_text("organization-rules.pdf")
    # 名前の似たシステム・部署。
    similar = {"経費精算ポータル", "交通費精算ポータル", "人事評価システム", "人材評価分析システム"}
    assert similar <= set(ledger)
    assert "経理部" in organization and "財務部" in organization
    assert "人事部" in organization and "労務部" in organization
    # 略称の衝突（同じ会社の中・別の会社との間）。
    assert ledger.count("PMS") == 2 and ledger.count("TMS") == 2
    assert "OMS" in logistics and "OMS、受発注" in ledger
    # 旧版・似た承認の規程・別の会社の組織規程。
    assert "廃止: 2026年3月31日" in _html_text("approval-rules-2023.pdf")
    assert "終了: 2026年3月31日" in _html_text("maintenance-plan-2025.pdf")
    assert "システムの変更の申請は承認規程で扱い" in _html_text("purchase-approval-rules.pdf")
    assert "サンプル物流社だけに適用します" in _html_text("logistics-organization-rules.pdf")


def test_multi_hop_set_declares_superseded_versions() -> None:
    """旧版の 2 文書は、新しい版に置き換えた文書として取り込む（#1366）。

    旧版の文を根拠にする問（版の比較）だけが、旧版も検索する（`include_superseded`）。
    """
    payload, request = _load_set()
    assert payload["document_versions"] == [
        {"document_id": "file:approval-rules-2023.pdf", "superseded_by": "file:approval-rules.pdf"},
        {
            "document_id": "file:maintenance-plan-2025.pdf",
            "superseded_by": "file:maintenance-plan.pdf",
        },
    ]
    superseded = {entry["document_id"] for entry in payload["document_versions"]}
    marked = {case.id for case in request.cases if case.include_superseded}
    uses_old = {
        case.id
        for case in request.cases
        if superseded & {item.document_id for item in case.required_evidence}
    }
    assert marked == uses_old == {"cmp-approval-deadline-versions", "cmp-hrm-maintenance-versions"}


def test_source_builder_rejects_a_change_that_breaks_original_answers(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """資料を増やして #1335 の問の正解が変わる（重要度 A かつ極秘が増える）と、作らずに止める。"""
    builder = _source_builder()
    extra = builder.System(999, "評価用の追加システム", "追加", "総", "A", "極秘")
    monkeypatch.setattr(builder, "SYSTEMS", (*builder.SYSTEMS, extra))
    with pytest.raises(ValueError, match="重要度 A かつ極秘"):
        builder.build_golden_set(builder.build_documents())


def test_source_builder_rejects_evidence_missing_from_the_document() -> None:
    builder = _source_builder()
    document = builder.Document("x.pdf", "x", "<p>本文</p>", {"x-1": "書いていない文"})
    with pytest.raises(ValueError, match="根拠の文がありません"):
        builder.evidence_registry([document])
