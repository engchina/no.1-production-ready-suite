"""品質評価の結果 Excel と、所有者を検証する出力 API。"""

from datetime import UTC, datetime
from io import BytesIO
from typing import Any

import pytest
from openpyxl import load_workbook  # type: ignore[import-untyped]

from app.api.routes import evaluation as route
from app.main import app
from app.rag.evaluation_jobs import EvaluationJobNotFoundError
from app.rag.evaluation_results_excel import evaluation_results_xlsx
from app.schemas.evaluation import (
    EvaluationCaseResult,
    EvaluationComparison,
    EvaluationJob,
    EvaluationMetrics,
    EvaluationResultExportRequest,
)
from tests.support import AsgiTestClient


def metrics() -> EvaluationMetrics:
    return EvaluationMetrics(
        case_count=1,
        mrr=0.5,
        case_results=[
            EvaluationCaseResult(
                case_id="row-2",
                trace_id="trace-a",
                elapsed_ms=10,
                mrr=0.5,
                required_evidence_recall=0,
                failure_reasons=["required_evidence_missed"],
            )
        ],
    )


def inputs() -> EvaluationResultExportRequest:
    return EvaluationResultExportRequest.model_validate(
        {
            "profile_name": "プロファイル A",
            "compare_profile_name": "プロファイル B",
            "rows": [
                {
                    "row": 2,
                    "case_id": "row-2",
                    "question": "=2+2",
                    "evidence": ["資料 A.pdf"],
                    "standard_answer_preview": "短いプレビュー",
                    "required_answer_points": ["要点 A"],
                    "expected_response_decision": "answer",
                    "answer_process": "手順 A",
                }
            ],
            "cases": [
                {
                    "id": "row-2",
                    "query": "=2+2",
                    "standard_answer": "合成の回答です。" * 100,
                    "required_answer_points": ["要点 A"],
                    "expected_response_decision": "answer",
                }
            ],
        }
    )


def test_export_keeps_full_inputs_scores_unassessed_and_comparison() -> None:
    result = metrics()
    result.comparison = EvaluationComparison(
        search_answer_profile_id="profile-b", metrics=metrics()
    )
    book = load_workbook(BytesIO(evaluation_results_xlsx(result, inputs())))
    assert book.sheetnames == ["概要", "主条件の結果", "比較対象の結果"]
    sheet = book["主条件の結果"]
    assert sheet["B2"].value == "=2+2"
    assert sheet["B2"].data_type == "s"
    assert sheet["D2"].value == "合成の回答です。" * 100
    assert sheet["G2"].value == "手順 A"
    assert sheet["H2"].value == 0
    assert sheet["I2"].value == 0.5
    assert sheet["I2"].number_format == "0.0%"
    assert sheet["J2"].value == "未評価"
    assert sheet["P2"].value == "trace-a"
    assert sheet.freeze_panes == "B2"
    assert book["概要"]["C1"].value == "比較対象: プロファイル B"


@pytest.mark.parametrize(
    "state,code", [("SUCCEEDED", 200), ("RUNNING", 409), ("FAILED", 409), ("missing", 404)]
)
def test_export_api_uses_authorized_job_and_rejects_unavailable_results(
    monkeypatch: pytest.MonkeyPatch, state: str, code: int
) -> None:
    class Service:
        async def get(self, job_id: str) -> EvaluationJob:
            assert job_id == "job-a"
            if state == "missing":
                raise EvaluationJobNotFoundError("見つかりません。")
            return EvaluationJob.model_validate(
                dict(
                    job_id=job_id,
                    kind="run",
                    status=state,
                    total_cases=1,
                    completed_cases=1,
                    time_limit_seconds=60,
                    created_at=datetime.now(UTC),
                    run_result=metrics() if state == "SUCCEEDED" else None,
                )
            )

    monkeypatch.setattr(route, "get_evaluation_job_service", Service)
    body: dict[str, Any] = inputs().model_dump()
    body["metrics"] = {"mrr": 1}
    response = AsgiTestClient(app).post("/api/evaluation/jobs/job-a/results.xlsx", json=body)
    assert response.status_code == code
    if code == 200:
        book = load_workbook(BytesIO(response.content))
        assert book["主条件の結果"]["I2"].value == 0.5
        assert "evaluation_results.xlsx" in response.headers["content-disposition"]


def test_export_includes_all_cases_beyond_one_page() -> None:
    result = metrics()
    result.case_results = [
        result.case_results[0].model_copy(update={"case_id": f"case-{i}"}) for i in range(40)
    ]
    result.case_count = 40
    sheet = load_workbook(
        BytesIO(evaluation_results_xlsx(result, EvaluationResultExportRequest()))
    )["主条件の結果"]
    assert sheet.max_row == 41
    assert sheet["A41"].value == "case-39"
