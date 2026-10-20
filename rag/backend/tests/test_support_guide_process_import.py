"""回答の過程の表から業務ガイドの下書きとルールの候補を作る（#2397）。

内容は合成の例（架空の画面・資料の名前）だけを使う。LLM はスタブ（録画した応答の JSON を返す）。
"""

from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import pytest
from openpyxl import Workbook  # type: ignore[import-untyped]
from pydantic import ValidationError
from pytest import MonkeyPatch
from rag_parser_core.language import builtin_pack

from app.api.routes import support_guide_process_import as route
from app.clients import ai_calls
from app.main import app
from app.rag.search_answer_profile_config import SearchAnswerProfileConfig
from app.rag.support_guide_process_import import (
    GUIDE_DRAFT_RESPONSE_SCHEMA,
    KnowledgeDocument,
    build_item,
    draft_context,
    page_ranges_from_text,
    parse_draft_output,
    process_headings,
    read_process_workbook,
    readable_as_pages,
    resolve_document_name,
    split_process_sections,
)
from app.schemas.support_guide import (
    SupportGuideProcessRow,
    SupportGuideReference,
    SupportGuideSheetRange,
)
from tests.support import AsgiTestClient

client = AsgiTestClient(app)
BASE = "/api/search-answer-profiles/bv-1/support-guides/import/process"
FIXTURES = Path(__file__).parent / "fixtures" / "support_guide_process_import"
RECORDED = (FIXTURES / "draft_response.json").read_text(encoding="utf-8")
HEADINGS = process_headings(builtin_pack())
PROCESS_TEXT = """【回答の手順】
1. 設定画面Aを開く
2. 通知の設定を有効にする（1 の後。保存を押すまで反映されない）
【必要な確認事項】
利用者が管理者か。管理者でなければ窓口へ連絡。
検索する資料：01_操作手順書A.pdf p.12-14、未登録の資料B.pdf p.3
■検索の語
設定画面A、通知
補足:
通知の質問は設定画面Aを案内する。言い切らない。手順書を開いて確かめる。"""
DOCUMENTS = [
    KnowledgeDocument("doc-a", "操作手順書A.pdf"),
    KnowledgeDocument("doc-c", "別の資料C.xlsx"),
]


def _row(**overrides: Any) -> SupportGuideProcessRow:
    values: dict[str, Any] = {
        "row": 2,
        "question": "画面Aで通知が出ない",
        "process": PROCESS_TEXT,
    }
    values.update(overrides)
    return SupportGuideProcessRow(**values)


# ---- 節に分ける（見出しの語は言語資源）------------------------------------------------------


def test_process_text_is_split_by_language_pack_headings() -> None:
    sections = split_process_sections(PROCESS_TEXT, HEADINGS)
    assert [section.field for section in sections] == [
        "steps",
        "conditions",
        "references",
        "search_terms",
        "notes",
    ]
    # 見出しの後ろの内容（「検索する資料：…」）は節の本文にする。
    assert sections[2].text.startswith("01_操作手順書A.pdf")
    # 本文の中の「手順書を開いて」は見出しにしない。
    assert "手順書を開いて確かめる" in sections[4].text
    context = json.loads(draft_context(_row(), HEADINGS))
    assert context["question"] == "画面Aで通知が出ない"
    assert [item["heading"] for item in context["sections"]][:2] == ["回答の手順", "必要な確認事項"]


def test_text_without_headings_is_one_section_and_headings_are_configurable() -> None:
    sections = split_process_sections("設定画面Aを開いて通知を有効にする", HEADINGS)
    assert [(section.field, section.text) for section in sections] == [
        (None, "設定画面Aを開いて通知を有効にする")
    ]
    custom = split_process_sections("[やること]\n画面を開く", {"やること": "steps"})
    assert [(section.field, section.text) for section in custom] == [("steps", "画面を開く")]


# ---- 資料名の解決 ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("name", "expected"),
    [
        ("操作手順書A.pdf", "doc-a"),
        ("操作手順書A", "doc-a"),
        ("01_操作手順書A.pdf", "doc-a"),
        ("別の資料C", "doc-c"),
    ],
    ids=["same", "no-extension", "prefixed", "other"],
)
def test_reference_names_resolve_to_knowledge_base_documents(name: str, expected: str) -> None:
    found, candidates = resolve_document_name(name, DOCUMENTS)
    assert found is not None and found.document_id == expected
    assert candidates == []


def test_unresolved_and_ambiguous_reference_names() -> None:
    assert resolve_document_name("未登録の資料B.pdf", DOCUMENTS) == (None, [])
    twins = [KnowledgeDocument("d1", "手順書A.pdf"), KnowledgeDocument("d2", "手順書A.docx")]
    found, candidates = resolve_document_name("手順書A", twins)
    assert found is None
    assert [item.document_id for item in candidates] == ["d1", "d2"]


# ---- 構造化（録画した LLM の応答）-------------------------------------------------------------


def test_recorded_response_becomes_a_validated_draft_with_rule_candidates() -> None:
    item = build_item(_row(), parse_draft_output(RECORDED), DOCUMENTS)
    assert item.content is not None
    content = item.content
    assert content.title == "画面Aの通知を表示する"
    assert content.goal.intent_examples == ["画面Aで通知が出ない"]
    assert content.goal.match_terms == ["設定画面A", "通知"]
    assert [step.id for step in content.steps] == ["s1", "s2"]
    assert content.steps[1].depends_on == ["s1"]
    assert "注意: 保存を押すまで反映されない" in content.steps[1].purpose
    assert content.conditions[0].unknown_handling == "handoff"
    assert content.branches[0].when.operator == "equals"
    assert (content.handoff.contact, content.handoff.conditions) == (
        "窓口",
        ["利用者が管理者でない"],
    )
    # 解決できた資料だけを document_id で参照にする。
    assert [(ref.document_id, ref.title) for ref in content.references] == [
        ("doc-a", "操作手順書A.pdf")
    ]
    assert [(ref.status, ref.pages) for ref in item.references] == [
        ("resolved", "12-14"),
        ("unresolved", "3"),
    ]
    # 解決できない資料は検証のエラー（取り込めない）。KB への取り込みを促す。
    assert item.valid is False
    error = next(issue for issue in item.issues if issue.code == "reference_unresolved")
    assert error.severity == "error" and error.path == "references[1]"
    assert "ナレッジベースに取り込んで" in error.message
    assert [rule.title for rule in item.rule_candidates] == ["通知の質問は設定画面Aを案内する"]
    assert item.rule_candidates[0].triggers == ["通知"]


def test_reference_pages_and_step_references_are_filled() -> None:
    """頁を資料の頁の範囲にし、手順を指す資料はその手順の参照にする（#2395・#2397）。"""
    recorded = json.loads(RECORDED)
    common, _missing = recorded["references"]
    recorded["references"] = [
        {**common, "pages": "p.12〜14、20"},
        {**common, "pages": "5", "step_ids": ["s2", "s9"]},
    ]
    item = build_item(_row(), parse_draft_output(json.dumps(recorded)), DOCUMENTS)
    assert item.content is not None
    assert [
        [(r.page_start, r.page_end) for r in ref.page_ranges] for ref in item.content.references
    ] == [[(12, 14), (20, 20)]]
    # 手順を指す資料はその手順の参照にする（無い手順の id は捨てる）。共通の資料には入れない。
    assert item.content.steps[0].references == []
    step_ref = item.content.steps[1].references
    assert [
        (ref.document_id, [(r.page_start, r.page_end) for r in ref.page_ranges]) for ref in step_ref
    ] == [("doc-a", [(5, 5)])]
    assert item.references[1].step_ids == ["s2", "s9"]


def test_page_ranges_from_text_reads_numbers_and_ranges_only() -> None:
    assert page_ranges_from_text("12-14") == [{"page_start": 12, "page_end": 14}]
    assert page_ranges_from_text("ｐ．３～５, 3-5, 9") == [
        {"page_start": 3, "page_end": 5},
        {"page_start": 9, "page_end": 9},
    ]
    assert page_ranges_from_text("") == []
    assert page_ranges_from_text("14-12") == []


def test_draft_is_importable_when_all_references_resolve() -> None:
    recorded = json.loads(RECORDED)
    recorded["references"] = recorded["references"][:1]
    item = build_item(_row(), parse_draft_output(json.dumps(recorded)), DOCUMENTS)
    assert item.valid is True
    assert all(issue.severity != "error" for issue in item.issues)


def test_invalid_structures_are_reported_as_validation_errors() -> None:
    recorded = json.loads(RECORDED)
    recorded["references"] = []
    recorded["steps"][1]["depends_on"] = ["s1"]
    recorded["branches"] = [{"condition_id": "c9", "values": ["x"], "goto_step": "s9"}]
    # enum の選択肢が 1 つなら text にする（選択肢の条件の検証で落とさない）。
    recorded["conditions"].append(
        {"id": "c2", "label": "種類", "type": "enum", "allowed_values": ["A"], "handling": "ask"}
    )
    item = build_item(_row(), parse_draft_output(json.dumps(recorded)), DOCUMENTS)
    assert item.content is not None
    assert item.content.conditions[1].type == "text"
    assert item.content.conditions[1].question == "種類を教えてください。"
    assert item.valid is False
    assert any(issue.path.startswith("branches") for issue in item.issues)


def test_unparsable_response_is_rejected() -> None:
    with pytest.raises(ValueError, match="JSON"):
        parse_draft_output("下書きを作れません")
    with pytest.raises(ValueError, match="形ではありません"):
        parse_draft_output('{"steps": "x"}')
    assert GUIDE_DRAFT_RESPONSE_SCHEMA["title"] == "GuideDraftOutput"


# ---- Excel ----------------------------------------------------------------------------------


def _workbook(rows: list[list[Any]]) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "問い合わせ"
    for values in rows:
        sheet.append(values)
    book.create_sheet("説明").append(["メモ"])
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


WORKBOOK = _workbook(
    [
        ["番号", "質問", "回答の過程", "標準回答"],
        [1, "画面Aで通知が出ない", PROCESS_TEXT, "設定画面Aで有効にします"],
        [2, "", "", ""],
        [3, "質問だけの行", "", ""],
        [4, "画面Bの出力", "【回答の手順】\n画面Bを開く", None],
    ]
)


def test_workbook_columns_then_rows_with_the_chosen_columns() -> None:
    columns = read_process_workbook(WORKBOOK, "a.xlsx", max_rows=30)
    assert [sheet.name for sheet in columns.sheets] == ["問い合わせ", "説明"]
    first = columns.sheets[0]
    assert [(column.letter, column.header) for column in first.columns] == [
        ("A", "番号"),
        ("B", "質問"),
        ("C", "回答の過程"),
        ("D", "標準回答"),
    ]
    assert first.row_count == 4
    assert columns.rows == []

    read = read_process_workbook(
        WORKBOOK,
        "a.xlsx",
        max_rows=30,
        sheet="問い合わせ",
        question_column=1,
        process_column=2,
        answer_column=3,
    )
    assert [(row.row, row.question) for row in read.rows] == [
        (2, "画面Aで通知が出ない"),
        (5, "画面Bの出力"),
    ]
    assert read.rows[0].standard_answer == "設定画面Aで有効にします"
    assert read.issues == ["4 行目: 回答の過程が空です。"]

    limited = read_process_workbook(
        WORKBOOK, "a.xlsx", max_rows=1, sheet="問い合わせ", question_column=1, process_column=2
    )
    assert [row.row for row in limited.rows] == [2]
    assert (limited.truncated, limited.next_start_row) == (True, 4)
    rest = read_process_workbook(
        WORKBOOK,
        "a.xlsx",
        max_rows=1,
        sheet="問い合わせ",
        question_column=1,
        process_column=2,
        start_row=5,
    )
    assert [row.row for row in rest.rows] == [5]


def test_workbook_format_errors() -> None:
    with pytest.raises(ValueError, match="Excel"):
        read_process_workbook(b"not a workbook", "a.xlsx", max_rows=30)
    with pytest.raises(ValueError, match="シート"):
        read_process_workbook(
            WORKBOOK, "a.xlsx", max_rows=30, sheet="無い", question_column=1, process_column=2
        )


# ---- API ------------------------------------------------------------------------------------


class _FakeOracle:
    async def get_search_answer_profile(self, profile_id: str) -> Any:
        if profile_id != "bv-1":
            return None
        return SimpleNamespace(
            id="bv-1", config=SearchAnswerProfileConfig(knowledge_base_ids=["kb-1"])
        )

    async def list_documents(self, *, knowledge_base_id: str) -> list[Any]:
        assert knowledge_base_id == "kb-1"
        return [
            _summary("doc-a", "操作手順書A.pdf"),
            _summary("doc-old", "未登録の資料B.pdf", is_superseded=True),
            # 同じ正本を指す、同じ名前の重複のアップロード（#2662）。
            _summary("doc-a-dup", "操作手順書A.pdf", duplicate_of_document_id="doc-a"),
            _summary("doc-sheet", "設定一覧D.xlsx"),
        ]


def _summary(document_id: str, file_name: str, **values: Any) -> SimpleNamespace:
    return SimpleNamespace(
        **{
            "id": document_id,
            "file_name": file_name,
            "is_superseded": False,
            "duplicate_of_document_id": None,
            "content_type": None,
            **values,
        }
    )


class _Llm:
    def __init__(self, responses: dict[str, str | Exception]) -> None:
        self.responses = responses
        self.calls: list[dict[str, Any]] = []

    async def generate(self, prompt: str, context: str, **kwargs: Any) -> str:
        question = json.loads(context)["question"]
        self.calls.append({"question": question, **kwargs})
        response = self.responses[question]
        if isinstance(response, Exception):
            raise response
        return response


@pytest.fixture
def fake_oracle(monkeypatch: MonkeyPatch) -> _FakeOracle:
    oracle = _FakeOracle()
    monkeypatch.setattr(route, "OracleClient", lambda: oracle)
    return oracle


@pytest.mark.usefixtures("fake_oracle")
def test_workbook_endpoint_reads_sheets_and_rows() -> None:
    files = {"file": ("process.xlsx", WORKBOOK, "application/octet-stream")}
    body = client.post(f"{BASE}/workbook", files=files).json()["data"]
    assert body["sheets"][0]["name"] == "問い合わせ"
    assert body["max_rows"] == 30
    chosen = client.post(
        f"{BASE}/workbook",
        files=files,
        data={"sheet": "問い合わせ", "question_column": "1", "process_column": "2"},
    ).json()["data"]
    assert [row["row"] for row in chosen["rows"]] == [2, 5]
    wrong = client.post(f"{BASE}/workbook", files={"file": ("process.csv", b"a,b", "text/csv")})
    assert wrong.status_code == 422
    missing = client.post(
        "/api/search-answer-profiles/other/support-guides/import/process/workbook", files=files
    )
    assert missing.status_code == 404


@pytest.mark.usefixtures("fake_oracle")
def test_preview_endpoint_drafts_each_row_with_the_recorded_response(
    monkeypatch: MonkeyPatch,
) -> None:
    async def no_sleep(_seconds: float) -> None:
        return None

    monkeypatch.setattr(ai_calls, "_sleep", no_sleep)
    llm = _Llm(
        {"画面Aで通知が出ない": RECORDED, "画面Bの出力": ValueError("回答 text がありません")}
    )
    monkeypatch.setattr(route, "guide_draft_llm", lambda _settings: llm)
    rows = [
        _row().model_dump(),
        _row(row=5, question="画面Bの出力", process="【回答の手順】\n画面Bを開く").model_dump(),
    ]
    data = client.post(f"{BASE}/preview", json={"rows": rows}).json()["data"]
    first, second = data["items"]
    assert first["row"] == 2 and first["content"]["title"] == "画面Aの通知を表示する"
    # 置き換えた旧版の文書には当てない（未解決のまま）。同じ正本の重複は正本にまとめる（#2662）。
    assert [ref["status"] for ref in first["references"]] == ["resolved", "unresolved"]
    assert first["references"][0]["document_id"] == "doc-a"
    assert first["valid"] is False
    assert first["rule_candidates"][0]["triggers"] == ["通知"]
    assert second["valid"] is False
    assert second["issues"][0]["code"] == "draft_failed"
    assert data["importable_count"] == 0
    # 構造化出力の schema と system prompt を渡す。
    assert llm.calls[0]["response_schema"] == GUIDE_DRAFT_RESPONSE_SCHEMA
    assert llm.calls[0]["response_schema_name"] == "support_guide_draft"
    assert llm.calls[0]["system_prompt"]


@pytest.mark.usefixtures("fake_oracle")
def test_preview_endpoint_limits_rows(monkeypatch: MonkeyPatch) -> None:
    monkeypatch.setattr(route, "guide_draft_llm", lambda _settings: _Llm({}))
    rows = [_row(row=index).model_dump() for index in range(2, 34)]
    response = client.post(f"{BASE}/preview", json={"rows": rows})
    assert response.status_code == 422
    assert "30 行まで" in response.json()["error_messages"][0]


# ---- 同じ正本の重複のアップロード（#2662）----------------------------------------------------


def test_duplicate_uploads_of_one_canonical_resolve_to_the_canonical() -> None:
    canonical = KnowledgeDocument("doc-a", "手順書A.pdf")
    duplicate = KnowledgeDocument("doc-b", "手順書A.pdf", duplicate_of="doc-a")
    for documents in ([canonical, duplicate], [duplicate, canonical]):
        found, candidates = resolve_document_name("01_手順書A.pdf", documents)
        assert found == canonical
        assert candidates == []
    # 名前が重複のファイル名にだけ当たっても、見える正本の文書にする（chunk は正本にある）。
    renamed = KnowledgeDocument("doc-b", "手順書A_写し.pdf", duplicate_of="doc-a")
    assert resolve_document_name("手順書A_写し", [canonical, renamed]) == (canonical, [])


def test_same_name_with_other_content_or_invisible_canonical_stays_ambiguous() -> None:
    # 同じ名前でも別の内容（重複でない）の文書は、どちらかに決めない。
    other = [KnowledgeDocument("doc-a", "手順書A.pdf"), KnowledgeDocument("doc-c", "手順書A.pdf")]
    found, candidates = resolve_document_name("手順書A.pdf", other)
    assert found is None
    assert [item.document_id for item in candidates] == ["doc-a", "doc-c"]
    # 正本がこのプロファイルの文書に無い（範囲の外・見られない・無い）重複は、まとめない。
    hidden = [
        KnowledgeDocument("doc-b", "手順書A.pdf", duplicate_of="doc-x"),
        KnowledgeDocument("doc-d", "手順書A.pdf", duplicate_of="doc-x"),
    ]
    found, candidates = resolve_document_name("手順書A.pdf", hidden)
    assert found is None
    assert [item.document_id for item in candidates] == ["doc-b", "doc-d"]
    # 1 件だけなら、その文書のまま（見えない正本の id を参照にしない）。
    found, _ = resolve_document_name("手順書A.pdf", hidden[:1])
    assert found is not None and found.document_id == "doc-b"


def test_draft_with_duplicate_documents_is_importable() -> None:
    recorded = json.loads(RECORDED)
    recorded["references"] = recorded["references"][:1]
    documents = [
        KnowledgeDocument("doc-a", "操作手順書A.pdf"),
        KnowledgeDocument("doc-a-dup", "操作手順書A.pdf", duplicate_of="doc-a"),
    ]
    item = build_item(_row(), parse_draft_output(json.dumps(recorded)), documents)
    assert item.valid is True
    assert [(ref.status, ref.document_id) for ref in item.references] == [("resolved", "doc-a")]
    assert item.content is not None
    assert [ref.document_id for ref in item.content.references] == ["doc-a"]


# ---- 表計算の資料の位置（#2663）---------------------------------------------------------------

SHEET_DOCUMENTS = [
    KnowledgeDocument("doc-a", "操作手順書A.pdf"),
    KnowledgeDocument("doc-s", "設定一覧D.xlsx", spreadsheet=True),
]


def _with_references(*references: dict[str, Any]) -> Any:
    recorded = json.loads(RECORDED)
    recorded["references"] = list(references)
    return parse_draft_output(json.dumps(recorded))


def test_spreadsheet_rows_become_sheet_ranges_not_pages() -> None:
    output = _with_references(
        {"document_name": "設定一覧D.xlsx", "sheet_name": "シートA", "rows": "42"},
        {"document_name": "設定一覧D.xlsx", "sheet_name": "シートB", "rows": "42-45"},
        {"document_name": "設定一覧D.xlsx", "sheet_name": "シートA", "cells": "ｂ３：ｄ１０"},
        {"document_name": "操作手順書A.pdf", "pages": "p.12-14"},
    )
    item = build_item(_row(), output, SHEET_DOCUMENTS)
    assert item.valid is True, item.issues
    assert item.content is not None
    sheets = [
        [(r.sheet_name, r.row_start, r.row_end, r.cell_range) for r in ref.sheet_ranges]
        for ref in item.content.references
    ]
    # 同じ行の番号でもシートで分け、セルの範囲はその行の範囲にする。頁の範囲にはしない。
    assert sheets == [
        [("シートA", 42, 42, "")],
        [("シートB", 42, 45, "")],
        [("シートA", 3, 10, "B3:D10")],
        [],
    ]
    assert [ref.page_ranges for ref in item.content.references[:3]] == [[], [], []]
    # 頁の資料は今までどおり頁の範囲にする。
    pdf_pages = item.content.references[3].page_ranges
    assert [(r.page_start, r.page_end) for r in pdf_pages] == [(12, 14)]


def test_row_written_as_page_is_not_converted_to_a_page_range() -> None:
    # 表計算の資料の頁の欄の行の番号は頁にしない。
    sheet = build_item(
        _row(),
        _with_references({"document_name": "設定一覧D.xlsx", "pages": "42"}),
        SHEET_DOCUMENTS,
    )
    assert sheet.valid is False
    assert [issue.code for issue in sheet.issues] == ["reference_location_mismatch"]
    assert sheet.content is not None
    assert sheet.content.references[0].page_ranges == []
    # 頁の資料でも、行・セルの書き方（頁の語でない文字を含む）は頁として読まない。
    for pages in ("行42", "A1:C3", "42行目"):
        pdf = build_item(
            _row(),
            _with_references({"document_name": "操作手順書A.pdf", "pages": pages}),
            DOCUMENTS,
        )
        assert pdf.valid is False, pages
        assert [issue.code for issue in pdf.issues] == ["reference_location_unreadable"]
        assert pdf.content is not None
        assert pdf.content.references[0].page_ranges == []


def test_sheet_location_errors_are_reported() -> None:
    cases = [
        # シートの名前の無い行は、どのシートの行か決まらない。
        ({"document_name": "設定一覧D.xlsx", "rows": "42"}, "reference_location_unreadable"),
        # 読めないセルの範囲。
        (
            {"document_name": "設定一覧D.xlsx", "sheet_name": "シートA", "cells": "どこか"},
            "reference_location_unreadable",
        ),
        # 頁の資料にシート・行は使えない。
        (
            {"document_name": "操作手順書A.pdf", "sheet_name": "シートA", "rows": "3"},
            "reference_location_mismatch",
        ),
    ]
    for reference, code in cases:
        item = build_item(_row(), _with_references(reference), SHEET_DOCUMENTS)
        assert item.valid is False, reference
        assert [issue.code for issue in item.issues] == [code], reference


def test_readable_as_pages_uses_page_words_only() -> None:
    for text in ("12-14", "p.12〜14、20", "ｐ．３～５", "12頁", "12〜14ページ", "pp. 3-5", ""):
        assert readable_as_pages(text), text
    for text in ("行42", "42行目", "A1:C3", "シートA 3"):
        assert not readable_as_pages(text), text


def test_reference_location_kinds_are_exclusive() -> None:
    reference = SupportGuideReference.model_validate(
        {"document_id": "doc-s", "sheet_ranges": [{"sheet_name": "シートA", "cell_range": "B3"}]}
    )
    assert [(r.row_start, r.row_end, r.cell_range) for r in reference.sheet_ranges] == [
        (3, 3, "B3")
    ]
    with pytest.raises(ValidationError, match="どちらかだけ"):
        SupportGuideReference.model_validate(
            {
                "document_id": "doc-s",
                "page_ranges": [{"page_start": 1}],
                "sheet_ranges": [{"sheet_name": "シートA"}],
            }
        )
    with pytest.raises(ValidationError, match="セルの範囲の行"):
        SupportGuideSheetRange(sheet_name="シートA", row_start=1, cell_range="B3:D10")
    with pytest.raises(ValidationError, match="A1 形式"):
        SupportGuideSheetRange(sheet_name="シートA", cell_range="3B")
    with pytest.raises(ValidationError, match="始まり以降"):
        SupportGuideSheetRange(sheet_name="シートA", row_start=9, row_end=3)


# ---- 別の操作の分岐を同じ手順にまとめない（#2666）----------------------------------------------


def _branching_output(*branches: dict[str, Any]) -> Any:
    recorded = json.loads(RECORDED)
    recorded["references"] = recorded["references"][:1]
    recorded["steps"].append(
        {"id": "s3", "title": "別の画面で値を直す", "depends_on": ["s1"], "retrieval_hints": []}
    )
    recorded["conditions"] = [
        {
            "id": "c1",
            "label": "項目の種類",
            "type": "enum",
            "allowed_values": ["種類A", "種類B"],
            "handling": "ask",
            "question": "項目の種類は？",
        }
    ]
    recorded["branches"] = list(branches)
    recorded["handoff_contact"] = ""
    recorded["handoff_conditions"] = []
    return parse_draft_output(json.dumps(recorded))


def _branch(values: list[str], goto: str, operation: str | None = None) -> dict[str, Any]:
    branch: dict[str, Any] = {"condition_id": "c1", "values": values, "goto_step": goto}
    if operation is not None:
        branch["operation"] = operation
    return branch


def test_branches_with_different_operations_merged_into_one_step_are_rejected() -> None:
    output = _branching_output(
        _branch(["種類A"], "s2", "操作Aをする"), _branch(["種類B"], "s2", "操作Bをする")
    )
    item = build_item(_row(), output, DOCUMENTS)
    assert item.valid is False
    issue = next(issue for issue in item.issues if issue.code == "branch_operations_merged")
    assert issue.severity == "error" and issue.path == "branches[0]"
    assert "種類A" in issue.message and "種類B" in issue.message
    assert "通知の設定を有効にする" in issue.message


def test_faithful_branches_and_common_merge_stay_valid() -> None:
    cases = [
        # 値ごとに別の手順へ進む。
        (_branch(["種類A"], "s2", "操作A"), _branch(["種類B"], "s3", "操作B")),
        # 値に関わらず同じ操作の手順に合流する（操作が同じ・書いていない）。
        (_branch(["種類A"], "s2", "共通の操作"), _branch(["種類B"], "s2", "共通の操作。")),
        (_branch(["種類A"], "s2"), _branch(["種類B"], "s2")),
        # 条件が分からないときの分岐は、値の分岐と同じ手順でも別の操作と数えない。
        (_branch(["種類A"], "s2", "操作A"), _branch([], "s2", "確かめる")),
    ]
    for branches in cases:
        item = build_item(_row(), _branching_output(*branches), DOCUMENTS)
        assert item.valid is True, (branches, item.issues)
        assert all(issue.code != "branch_operations_merged" for issue in item.issues)
