"""評価の合成の資料（PDF / xlsx）を原稿から作る（業務支援 #1231・多段の質問 #1335）。

`<資料のフォルダ>/sources/*.html` を LibreOffice（soffice）で PDF にし、`sources/*.workbook.json`
（表の原稿。シートごとの前書き・表頭・行）を openpyxl で xlsx にして、`<資料のフォルダ>/` に置く。
資料はすべて架空の内容で、実在の顧客・製品の内容は入れない。

作った PDF / xlsx は commit する（評価の CLI はこのファイルをそのまま取り込む）。原稿を直したら、
リポジトリ直下から次で作り直す（PDF の生成には LibreOffice と日本語のフォントが要る）::

    uv run --project rag/backend python rag/scripts/generate_evaluation_corpus.py \\
        rag/evaluation/multi-hop
    uv run --project rag/backend python rag/scripts/generate_evaluation_corpus.py \\
        rag/evaluation/business-support

`--xlsx-only` は Excel だけを作る（LibreOffice の無い環境で表の原稿だけを直したとき）。xlsx は
作成時刻を固定するため、同じ原稿からは同じセルの値・書式の xlsx ができる（テストで照合する）。

原稿を実体のデータから作る資料のフォルダ（`SOURCE_BUILDERS`。多段の質問 `multi-hop` は
`multi_hop_corpus.py`。#1352）は、変換の前に原稿（`sources/`）と評価セットを作り直す。
`--sources-only` は原稿と評価セットだけを作る（PDF / xlsx は作らない）。
"""

from __future__ import annotations

import argparse
import importlib.util
import json
import shutil
import subprocess  # nosec B404 - soffice を固定の引数で呼ぶだけ
import sys
import tempfile
from collections.abc import Mapping, Sequence
from datetime import datetime
from pathlib import Path
from types import ModuleType
from typing import Any

SOURCES_DIRNAME = "sources"
WORKBOOK_SUFFIX = ".workbook.json"
# xlsx の作成・更新の時刻（作り直しても同じ内容にする）。
FIXED_TIMESTAMP = datetime(2026, 4, 1)
# 原稿を実体のデータから作る資料のフォルダ（フォルダ名 → この script と同じフォルダの module）。
SOURCE_BUILDERS = {"multi-hop": "multi_hop_corpus.py"}


def source_builder(corpus_dir: Path) -> ModuleType | None:
    """資料のフォルダの原稿を作る module（`write_sources(corpus_dir)` を持つ）。無ければ None。"""
    module_name = SOURCE_BUILDERS.get(corpus_dir.name)
    if module_name is None:
        return None
    path = Path(__file__).resolve().parent / module_name
    spec = importlib.util.spec_from_file_location(path.stem, path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"原稿を作る module を読めません: {path}")
    module = importlib.util.module_from_spec(spec)
    # dataclass が module を sys.modules から引くため、実行の前に登録する。
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def workbook_name(spec_path: Path) -> str:
    """表の原稿（`<名前>.workbook.json`）から作る xlsx のファイル名。"""
    return spec_path.name.removesuffix(WORKBOOK_SUFFIX) + ".xlsx"


def load_workbook_spec(spec_path: Path) -> dict[str, Any]:
    """表の原稿を読み、形を確かめる（シートごとに title・header・rows が要る）。"""
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    sheets = spec.get("sheets") if isinstance(spec, Mapping) else None
    if not isinstance(sheets, list) or not sheets:
        raise ValueError(f"{spec_path.name}: sheets がありません。")
    for sheet in sheets:
        header = sheet.get("header")
        rows = sheet.get("rows")
        if not sheet.get("title") or not isinstance(header, list) or not isinstance(rows, list):
            raise ValueError(f"{spec_path.name}: シートに title・header・rows が要ります。")
        if any(len(row) != len(header) for row in rows):
            raise ValueError(f"{spec_path.name}: 行の列の数が表頭と違います（{sheet['title']}）。")
    return dict(spec)


def write_workbook(spec: Mapping[str, Any], path: Path) -> None:
    """表の原稿から xlsx を書く。

    各シートは、前書きの行（`preamble`）→ 表頭（`header`）→ 行（`rows`）の順。`text_columns`
    （列の記号）の行のセルは文字列の書式（`@`）にし、`030` のような値を数値に読み替えさせない。
    """
    from openpyxl import Workbook  # type: ignore[import-untyped]

    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_spec in spec["sheets"]:
        sheet = workbook.create_sheet(title=str(sheet_spec["title"]))
        preamble = [str(line) for line in sheet_spec.get("preamble", [])]
        for line in preamble:
            sheet.append([line])
        sheet.append([str(value) for value in sheet_spec["header"]])
        for row in sheet_spec["rows"]:
            sheet.append([str(value) for value in row])
        first_data_row = len(preamble) + 1  # 0 始まりの添字（表頭の次の行）
        for column in sheet_spec.get("text_columns", []):
            for cell in sheet[str(column)][first_data_row:]:
                cell.number_format = "@"
    workbook.properties.created = FIXED_TIMESTAMP
    workbook.properties.modified = FIXED_TIMESTAMP
    workbook.save(path)


def convert_html_to_pdf(soffice: str, sources: Sequence[Path], outdir: Path) -> None:
    """原稿の HTML を PDF にする（利用者の LibreOffice の設定を使わず、毎回同じ結果にする）。"""
    with tempfile.TemporaryDirectory() as work:
        profile = Path(work) / "profile"
        subprocess.run(  # nosec B603 - 固定の実行ファイルと引数
            [
                soffice,
                f"-env:UserInstallation=file://{profile}",
                "--headless",
                "--convert-to",
                "pdf:writer_web_pdf_Export",
                "--outdir",
                str(outdir),
                *[str(source) for source in sources],
            ],
            check=True,
            capture_output=True,
            timeout=300,
        )


def build_corpus(
    corpus_dir: Path, *, xlsx_only: bool = False, sources_only: bool = False
) -> list[Path]:
    """資料のフォルダの原稿から PDF / xlsx を作り、作ったファイルを返す。

    原稿を実体のデータから作るフォルダ（`SOURCE_BUILDERS`）は、先に原稿と評価セットを作り直す
    （`sources_only` ならそこで終える）。
    """
    sources_dir = corpus_dir / SOURCES_DIRNAME
    builder = source_builder(corpus_dir)
    generated: list[Path] = list(builder.write_sources(corpus_dir)) if builder else []
    if sources_only:
        return generated
    html_sources = sorted(sources_dir.glob("*.html"))
    workbook_specs = sorted(sources_dir.glob(f"*{WORKBOOK_SUFFIX}"))
    written: list[Path] = []
    if html_sources and not xlsx_only:
        soffice = shutil.which("soffice") or shutil.which("libreoffice")
        if soffice is None:
            raise RuntimeError("soffice（LibreOffice）が見つかりません。")
        convert_html_to_pdf(soffice, html_sources, corpus_dir)
        written.extend(corpus_dir / f"{source.stem}.pdf" for source in html_sources)
    for spec_path in workbook_specs:
        target = corpus_dir / workbook_name(spec_path)
        write_workbook(load_workbook_spec(spec_path), target)
        written.append(target)
    missing = [path for path in written if not path.is_file()]
    if missing:
        raise RuntimeError("作れませんでした: " + ", ".join(path.name for path in missing))
    return generated + written


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("corpus_dir", type=Path, help="資料のフォルダ（sources/ を持つ）")
    parser.add_argument("--xlsx-only", action="store_true", help="Excel だけを作る")
    parser.add_argument(
        "--sources-only", action="store_true", help="実体のデータから原稿と評価セットだけを作る"
    )
    args = parser.parse_args(argv)
    corpus_dir = args.corpus_dir.resolve()
    if corpus_dir.name not in SOURCE_BUILDERS and not (corpus_dir / SOURCES_DIRNAME).is_dir():
        print(f"原稿のフォルダがありません: {corpus_dir / SOURCES_DIRNAME}", file=sys.stderr)
        return 2
    try:
        written = build_corpus(corpus_dir, xlsx_only=args.xlsx_only, sources_only=args.sources_only)
    except (RuntimeError, ValueError, subprocess.CalledProcessError) as error:
        print(str(error), file=sys.stderr)
        return 1
    for path in written:
        print(f"wrote {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
