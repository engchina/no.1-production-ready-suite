"""Excel→構造化レコード JSONL 変換の検証（#1221・#2457。.xlsx は openpyxl、.xls は xlrd）。

テストのデータはすべて合成（顧客の資料は使わない）。
"""

from __future__ import annotations

import datetime as dt
import io
import json
from types import SimpleNamespace
from typing import Any

import openpyxl
import pytest
from pr_backend_core.spreadsheet import reader
from rag_parser_core.sheet_records import (
    SheetRecordsDocument,
    sheet_records_extraction,
)
from rag_parser_core.structured_records import (
    STRUCTURED_RECORDS_CONTENT_TYPE,
    SheetLine,
    StructuredRecordsDocument,
    parse_structured_records,
    structured_records_extraction,
    structured_records_to_sheet_records,
)

from app.converters import convert


def _workbook(sheets: dict[str, list[list[Any]]]) -> openpyxl.Workbook:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        worksheet = workbook.create_sheet(title=name)
        for row in rows:
            worksheet.append(row)
    return workbook


def _bytes(workbook: openpyxl.Workbook) -> bytes:
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _xlsx_bytes(sheets: dict[str, list[list[Any]]]) -> bytes:
    return _bytes(_workbook(sheets))


def _records(derived: bytes) -> StructuredRecordsDocument:
    document = parse_structured_records(derived)
    assert isinstance(document, StructuredRecordsDocument), document
    return document


def _loaded(derived: bytes) -> dict[str, Any]:
    """JSONL を行の記録の形（シートごとの表頭・記録）に戻して比べる。"""
    payload: dict[str, Any] = structured_records_to_sheet_records(_records(derived)).model_dump(
        mode="json"
    )
    # 表・手順書の記録には帳票の項目の欄（#2516・#2526）を比べない。図形の文字の無い記録には
    # 図形のセル範囲（#2526）を比べない。
    # 表の下の注記（#3566）の無いシート・行には注記の欄を比べない。
    for sheet in payload["sheets"]:
        if not sheet["notes"]:
            sheet.pop("notes")
        for block in sheet["blocks"]:
            if not block["notes"]:
                block.pop("notes")
            if block["kind"] != "form_field":
                for key in ("label", "value", "label_cell", "value_cell", "value_original"):
                    block.pop(key)
        for item in [*sheet["blocks"], *sheet["preamble"]]:
            if not item["shape_cells"]:
                item.pop("shape_cells")
    return payload


def _payload(source: bytes, **options: Any) -> dict[str, Any]:
    outcome = convert(source, "", "excel_to_jsonl", None, options=options)
    assert outcome.converted is True, outcome.warnings
    assert outcome.derived_bytes is not None
    assert outcome.derived_content_type == STRUCTURED_RECORDS_CONTENT_TYPE
    assert outcome.converter_version == "v9"
    return _loaded(outcome.derived_bytes)


def _sheet(payload: dict[str, Any], index: int = 0) -> dict[str, Any]:
    sheet: dict[str, Any] = payload["sheets"][index]
    return sheet


def test_description_row_before_header_is_detected_as_preamble() -> None:
    """1 行目が説明、2 行目が表頭のシート（handoff の再現条件）。"""
    payload = _payload(
        _xlsx_bytes(
            {
                "コード表": [
                    ["このシートは経費の区分コードの一覧です。"],
                    ["コード", "名称", "区分"],
                    ["A01", "交通費", "旅費"],
                    ["A02", "宿泊費", "旅費"],
                ]
            }
        )
    )
    sheet = _sheet(payload)
    assert payload["format"] == "sheet_records"
    assert payload["source_format"] == "xlsx"
    assert sheet["header_row"] == 2
    assert sheet["header_detection"]["method"] == "detected"
    assert sheet["header_detection"]["confidence"] >= reader.HEADER_LOW_CONFIDENCE
    assert [column["name"] for column in sheet["columns"]] == ["コード", "名称", "区分"]
    assert [column["column"] for column in sheet["columns"]] == ["A", "B", "C"]
    assert sheet["preamble"] == [
        {"row_number": 1, "cell_range": "A1:A1", "text": "このシートは経費の区分コードの一覧です。"}
    ]
    assert sheet["mode"] == "table"
    assert sheet["blocks"][0] == {
        "kind": "row",
        "row_start": 3,
        "row_end": 3,
        "cell_range": "A3:C3",
        "section_path": [],
        "values": {"コード": "A01", "名称": "交通費", "区分": "旅費"},
        "lines": [],
    }
    assert sheet["diagnostics"] == []


def test_configured_header_row_wins_over_detection() -> None:
    payload = _payload(
        _xlsx_bytes({"S": [["説明"], ["a", "b"], ["x", "y"], ["1", "2"]]}), header_row=3
    )
    sheet = _sheet(payload)
    assert sheet["header_row"] == 3
    assert sheet["header_detection"]["method"] == "configured"
    assert [column["name"] for column in sheet["columns"]] == ["x", "y"]
    assert [block["row_start"] for block in sheet["blocks"]] == [4]
    assert [row["row_number"] for row in sheet["preamble"]] == [1, 2]


def test_original_row_numbers_survive_blank_rows() -> None:
    payload = _payload(
        _xlsx_bytes({"S": [["名前", "点"], ["甲", 1], [None, None], ["乙", 2], [], ["丙", 3]]})
    )
    blocks = _sheet(payload)["blocks"]
    assert [block["row_start"] for block in blocks] == [2, 4, 6]
    assert [block["cell_range"] for block in blocks] == ["A2:B2", "A4:B4", "A6:B6"]


def test_values_keep_display_formats() -> None:
    workbook = _workbook(
        {"S": [["コード", "文字のコード", "長い番号", "率", "金額", "小数", "日付", "真偽"]]}
    )
    worksheet = workbook["S"]
    worksheet.append(
        [1, "001", 123456789012345, 0.25, 1234567, 0.1 + 0.2, dt.date(2026, 4, 1), True]
    )
    worksheet["A2"].number_format = "000"
    worksheet["D2"].number_format = "0%"
    worksheet["E2"].number_format = "#,##0"
    values = _sheet(_payload(_bytes(workbook)))["blocks"][0]["values"]
    assert values == {
        "コード": "001",
        "文字のコード": "001",
        "長い番号": "123456789012345",
        "率": "25%",
        "金額": "1,234,567",
        "小数": "0.3",
        "日付": "2026-04-01",
        "真偽": "TRUE",
    }


def test_formula_without_cached_value_is_diagnosed_not_blank() -> None:
    """openpyxl が書いたファイルは数式のキャッシュ値を持たない（計算されていない）。

    空にせず数式の文字列にし（no.1-rag と同じ）、セルの位置付きの診断を残す。
    """
    workbook = _workbook({"S": [["単価", "数量", "金額"], [100, 3, "=A2*B2"]]})
    outcome = convert(_bytes(workbook), "", "excel_to_jsonl", None)
    assert "excel_formula_without_cached_value:S:1" in outcome.warnings
    assert outcome.derived_bytes is not None
    sheet = _loaded(outcome.derived_bytes)["sheets"][0]
    assert sheet["blocks"][0]["values"]["金額"] == "=A2*B2"
    assert sheet["diagnostics"] == [
        {"code": "formula_without_cached_value", "cell": "C2", "detail": None}
    ]


def test_merged_header_and_merged_data_cells() -> None:
    workbook = _workbook(
        {
            "S": [
                ["部門", "期間", None, "区分"],
                [None, "開始", "終了", None],
                ["営業", "4月", "6月", "通期"],
                [None, "7月", "9月", None],
            ]
        }
    )
    worksheet = workbook["S"]
    worksheet.merge_cells("B1:C1")  # 表頭の横の結合
    worksheet.merge_cells("A1:A2")  # 表頭の縦の結合
    worksheet.merge_cells("A3:A4")  # データの縦の結合
    worksheet.merge_cells("D3:E3")  # データの横の結合（同じ行）
    sheet = _sheet(_payload(_bytes(workbook), header_row=1, header_row_count=2, mode="table"))
    assert [column["name"] for column in sheet["columns"]] == [
        "部門",
        "期間 / 開始",
        "期間 / 終了",
        "区分",
        "column_E",
    ]
    blocks = sheet["blocks"]
    # データの横の結合の値は左上の列にだけ置き、結合した列に繰り返さない（#2531）。
    assert blocks[0]["values"] == {
        "部門": "営業",
        "期間 / 開始": "4月",
        "期間 / 終了": "6月",
        "区分": "通期",
    }
    # 縦の結合の値は後の行へ漏らさない（no.1-rag と同じ）。
    assert blocks[1]["values"] == {"期間 / 開始": "7月", "期間 / 終了": "9月"}


def test_procedure_workbook_groups_steps_with_sections() -> None:
    """手順書（表頭の無い番号の列と、表頭のある最初の列＝題名の列）は、手順ごとに複数行を
    1 block にする（no.1-rag と同じ判定）。"""
    payload = _payload(
        _xlsx_bytes(
            {
                "手順": [
                    [None, "作業項目", "作業内容", "確認ポイント"],
                    ["1", "事前準備", None, None],
                    ["1.1", "アカウントの登録", "管理画面で利用者を追加する", "一覧に表示される"],
                    [None, None, "権限を付ける", None],
                    ["1.2", "通知の設定", "通知先を入れる", "テスト通知が届く"],
                    ["2", "確認", None, None],
                    ["2.1", "動作確認", "ログインする", "トップ画面が出る"],
                ]
            }
        )
    )
    sheet = _sheet(payload)
    assert sheet["mode"] == "procedure"
    blocks = sheet["blocks"]
    assert [block["kind"] for block in blocks] == ["procedure_step"] * 3
    first = blocks[0]
    assert (first["row_start"], first["row_end"], first["cell_range"]) == (3, 4, "A3:D4")
    assert first["section_path"] == ["事前準備", "アカウントの登録"]
    assert first["lines"] == [
        "セクション: 事前準備",
        "column_A: 1.1 | 作業項目: アカウントの登録 | 作業内容: 管理画面で利用者を追加する"
        " | 確認ポイント: 一覧に表示される",
        "作業内容: 権限を付ける",
    ]
    assert blocks[2]["section_path"] == ["確認", "動作確認"]
    # mode="table" を選べば 1 行ずつ読む。
    table = _sheet(
        _payload(
            _xlsx_bytes({"手順": [[None, "作業項目", "内容"], ["1", "a", "x"], ["2", "b", "y"]]}),
            mode="table",
        )
    )
    assert [block["kind"] for block in table["blocks"]] == ["row", "row"]


def test_hidden_sheets_selection_and_excluded_columns() -> None:
    workbook = _workbook(
        {
            "表示": [["名前", "備考", "点"], ["甲", "内部用", 1]],
            "非表示": [["x"], ["1"]],
            "別": [["y"], ["2"]],
        }
    )
    workbook["非表示"].sheet_state = "hidden"
    source = _bytes(workbook)

    default = _payload(source)
    assert [sheet["name"] for sheet in default["sheets"]] == ["表示", "別"]
    assert default["skipped_sheets"] == [{"name": "非表示", "reason": "hidden"}]

    assert [sheet["name"] for sheet in _payload(source, include_hidden_sheets=True)["sheets"]] == [
        "表示",
        "非表示",
        "別",
    ]

    selected = _payload(source, sheets=["表示"], exclude_columns=["備考", "c"])
    assert [sheet["name"] for sheet in selected["sheets"]] == ["表示"]
    sheet = _sheet(selected)
    # 除外した列は列にも値にも出さない（列名でも列の記号でも指定できる）。
    assert [column["name"] for column in sheet["columns"]] == ["名前"]
    assert sheet["blocks"][0]["values"] == {"名前": "甲"}

    excluded = _payload(source, exclude_sheets=["別"])
    assert [sheet["name"] for sheet in excluded["sheets"]] == ["表示"]
    assert {"name": "別", "reason": "excluded"} in excluded["skipped_sheets"]

    outcome = convert(source, "", "excel_to_jsonl", None, options={"sheets": ["無い"]})
    assert outcome.converted is False
    assert "excel_no_rows" in outcome.warnings


def test_low_confidence_header_is_warned() -> None:
    outcome = convert(_xlsx_bytes({"数値": [[1, 2], [3, 4], [5]]}), "", "excel_to_jsonl", None)
    assert "excel_header_low_confidence:数値" in outcome.warnings
    assert outcome.derived_bytes is not None
    sheet = _loaded(outcome.derived_bytes)["sheets"][0]
    assert sheet["diagnostics"][0]["code"] == "header_low_confidence"


def test_ranges_limit_cells_used_for_header_and_records() -> None:
    """読む範囲の外（表の下の集計・右のメモ）は表頭の推定にも記録にも使わない（#1229）。"""
    workbook = _workbook(
        {
            "費目": [
                ["費目の一覧", None, None, "メモ"],
                ["コード", "名称", "上限", "担当: 経理"],
                ["001", "交通費", 5000, None],
                ["002", "宿泊費", 12000, None],
                [None, None, None, None],
                ["合計", None, 17000, None],
            ],
            "手順": [["番号", "作業"], ["1", "準備"], ["2", "確認"]],
        }
    )
    workbook["費目"].merge_cells("A1:D1")
    source = _bytes(workbook)

    payload = _payload(source, ranges=["費目!A2:C4", "A1:B2"])
    sheet = _sheet(payload)
    assert [column["name"] for column in sheet["columns"]] == ["コード", "名称", "上限"]
    assert [block["row_start"] for block in sheet["blocks"]] == [3, 4]
    assert sheet["blocks"][0]["cell_range"] == "A3:C3"
    assert sheet["preamble"] == []
    # シート名の無い範囲は、シート名を付けた範囲の無いシートに当てる。
    steps = _sheet(payload, 1)
    assert [block["row_start"] for block in steps["blocks"]] == [2]

    # 結合セルの左上が範囲の外でも、切り詰めた結合の左上へ値を移す。
    merged = _sheet(_payload(source, sheets=["費目"], ranges=["B1:C4"], header_row=2))
    assert merged["preamble"][0]["text"] == "費目の一覧"

    outcome = convert(source, "", "excel_to_jsonl", None, options={"ranges": ["無い!A1:B2"]})
    assert "excel_range_sheet_not_found:無い" in outcome.warnings


def test_invalid_ranges_are_rejected() -> None:
    from rag_parser_core.sheet_records import ExcelOptions, parse_excel_range

    assert parse_excel_range("'売上 ''26'!b2:a10").sheet == "売上 '26"
    assert (parse_excel_range("C9:A1").min_row, parse_excel_range("C9:A1").max_col) == (1, 3)
    for text in ("A1", "A0:B2", "1:3", "A1:B2:C3"):
        with pytest.raises(ValueError):
            ExcelOptions(ranges=[text])
    outcome = convert(b"PK", "", "excel_to_jsonl", None, options={"ranges": ["A1"]})
    assert outcome.converted is False
    assert "excel_options_invalid" in outcome.warnings


def test_duplicate_and_blank_header_names() -> None:
    sheet = _sheet(_payload(_xlsx_bytes({"S": [["id", "id", None, "名前"], [1, 2, 3, 4]]})))
    # 空の表頭は column_<列>、重複は <名前>__<列>（no.1-rag と同じ）。
    assert [column["name"] for column in sheet["columns"]] == ["id", "id__B", "column_C", "名前"]


def test_is_deterministic() -> None:
    source = _xlsx_bytes({"S": [["a", "b"], [1, 2]]})
    first = convert(source, "", "excel_to_jsonl", None).derived_bytes
    second = convert(source, "", "excel_to_jsonl", None).derived_bytes
    assert first == second


def test_passthrough_cases() -> None:
    assert convert(_xlsx_bytes({"S": [["a"], ["1"]]}), "", "csv_to_jsonl", None).converted is False
    empty = convert(b"", "", "excel_to_jsonl", None)
    assert (empty.converted, empty.warnings) == (False, ("excel_empty",))
    invalid = convert(
        _xlsx_bytes({"S": [["a"], ["1"]]}), "", "excel_to_jsonl", None, options={"header_row": 0}
    )
    assert (invalid.converted, invalid.warnings) == (False, ("excel_options_invalid",))


# ---- 列の役割（#1281） ------------------------------------------------------------------

# 評価セットの portal-parameters.xlsx と同じ形（説明の 2 行・表頭・パラメータの行）。架空の値。
_PARAMETER_SHEET = {
    "パラメータ一覧": [
        ["サンプル業務ポータルのパラメータの一覧"],
        ["既定値は出荷時の値、設定例は書き方の例です。"],
        ["パラメータ名", "既定値", "設定例", "説明"],
        ["session_timeout_minutes", "30", "60", "操作が無いときにログインが切れるまでの分数"],
        ["max_upload_mb", "20", "100", "1 ファイルのアップロードの上限（MB）"],
    ]
}


def _roles(sheet: dict[str, Any]) -> dict[str, tuple[str | None, str | None]]:
    return {
        column["name"]: (column.get("role"), column.get("role_method"))
        for column in sheet["columns"]
    }


def _extraction_texts(payload: dict[str, Any]) -> list[str]:
    extraction = sheet_records_extraction(
        SheetRecordsDocument.model_validate(payload),
        source_parser="test",
        parser_backend="test",
        parser_version="0",
    )
    return [element.text for element in extraction.elements]


def test_parameter_table_columns_get_roles_and_labels_in_text() -> None:
    """既定値・設定例・説明の列に役割を付け、本文の値の列名の後に役割の表示を付ける。"""
    payload = _payload(_xlsx_bytes(_PARAMETER_SHEET))
    sheet = _sheet(payload)
    assert _roles(sheet) == {
        "パラメータ名": (None, None),
        "既定値": ("default", "detected"),
        "設定例": ("example", "detected"),
        "説明": ("definition", "detected"),
    }
    assert sheet["diagnostics"] == []
    # 行の値（values）は列名のまま。役割の表示は本文だけに付ける。
    assert sheet["blocks"][0]["values"]["既定値"] == "30"
    texts = _extraction_texts(payload)
    assert (
        "パラメータ名: session_timeout_minutes / 既定値［資料の既定値］: 30 / "
        "設定例［例示の値］: 60 / 説明: 操作が無いときにログインが切れるまでの分数"
    ) in texts
    extraction = sheet_records_extraction(
        SheetRecordsDocument.model_validate(payload),
        source_parser="test",
        parser_backend="test",
        parser_version="0",
    )
    row = next(
        item for item in extraction.elements if item.metadata.get("sheet_block_kind") == "row"
    )
    assert row.metadata["sheet_column_roles"] == "既定値=default; 設定例=example; 説明=definition"


@pytest.mark.parametrize(
    ("header", "role"),
    [
        ("既定値", "default"),
        ("既定値（分）", "default"),
        ("デフォルト値", "default"),
        ("初期値", "default"),
        ("Default Value", "default"),
        ("タイムアウトの既定値", "default"),
        ("記入例", "example"),
        ("Example", "example"),
        ("e.g.", "example"),
        ("サンプル値", "example"),
        ("現在値", "current"),
        ("本番設定値", "current"),
        ("Current value", "current"),
        ("推奨値", "recommended"),
        ("設定可能範囲", "allowed"),
        ("Allowed values", "allowed"),
        ("説明", "definition"),
        ("Description", "definition"),
        # 複数行の表頭は下の行（具体的な語）から見る。下の行に語が無ければ上の行で決める。
        ("設定 / 既定値", "default"),
        ("既定値 / Windows", "default"),
    ],
)
def test_detect_column_role_from_header_vocabulary(header: str, role: str) -> None:
    detected, confidence, _, reason = reader.detect_column_role(header)
    assert (detected, reason) == (role, None)
    assert confidence >= reader.COLUMN_ROLE_MIN_CONFIDENCE


@pytest.mark.parametrize(
    ("header", "reason"),
    [
        # 「設定値」は現場の値のことも、選べる値の一覧のこともある。推測しない。
        ("設定値", "uncertain"),
        ("Value", "uncertain"),
        # 2 つの役割の語を含む表頭は 1 つに決めない。
        ("既定値の説明", "ambiguous"),
        ("既定値／推奨値", "ambiguous"),
        # 役割の語を含むが値の列ではない。
        ("既定値の変更履歴", "uncertain"),
        # 役割の語に当たらない。
        ("パラメータ名", None),
        ("事例", None),
        ("適用範囲", None),
        ("作業内容", None),
    ],
)
def test_detect_column_role_does_not_guess(header: str, reason: str | None) -> None:
    detected, _, _, why = reader.detect_column_role(header)
    assert (detected, why) == (None, reason)


def test_uncertain_and_ambiguous_columns_get_no_role_but_diagnostics() -> None:
    payload = _payload(
        _xlsx_bytes(
            {
                "S": [
                    ["項目", "設定値", "既定値の説明", "既定値", "説明"],
                    ["timeout", "45", "出荷時は 30", "30", "分数"],
                ]
            }
        )
    )
    sheet = _sheet(payload)
    assert _roles(sheet) == {
        "項目": (None, None),
        "設定値": (None, None),
        "既定値の説明": (None, None),
        "既定値": ("default", "detected"),
        "説明": ("definition", "detected"),
    }
    assert {(item["code"], item["detail"]) for item in sheet["diagnostics"]} == {
        ("column_role_uncertain", "B:設定値"),
        ("column_role_ambiguous", "C:既定値の説明"),
    }
    assert (
        "設定値: 45 / 既定値の説明: 出荷時は 30 / 既定値［資料の既定値］: 30"
        in (_extraction_texts(payload)[0])
    )


def test_single_role_column_is_not_corroborated() -> None:
    """役割の語に当たる列が 1 つだけの表（検体の「サンプル」）は設定値の表と言い切れない。"""
    payload = _payload(
        _xlsx_bytes({"検体": [["検体ID", "サンプル", "結果"], ["K-1", "血液", "陰性"]]})
    )
    sheet = _sheet(payload)
    assert all(column.get("role") is None for column in sheet["columns"])
    assert [(item["code"], item["detail"]) for item in sheet["diagnostics"]] == [
        ("column_role_not_corroborated", "B:サンプル")
    ]
    assert _extraction_texts(payload) == ["検体ID: K-1 / サンプル: 血液 / 結果: 陰性"]


def test_configured_column_roles_win_and_detection_can_be_turned_off() -> None:
    source = _xlsx_bytes(
        {
            "S": [
                ["項目", "設定値", "設定例", "説明"],
                ["timeout", "45", "60", "分数"],
            ]
        }
    )
    # 指定（列名か列の記号）は判定より優先し、none は役割を付けない。
    configured = _sheet(_payload(source, column_roles={"設定値": "current", "c": "none"}))
    assert _roles(configured) == {
        "項目": (None, None),
        "設定値": ("current", "configured"),
        "設定例": (None, None),
        "説明": ("definition", "detected"),
    }
    # off は指定だけを使う。
    off = _payload(source, column_role_detection="off", column_roles={"B": "current"})
    assert _roles(_sheet(off)) == {
        "項目": (None, None),
        "設定値": ("current", "configured"),
        "設定例": (None, None),
        "説明": (None, None),
    }
    assert _extraction_texts(off) == [
        "項目: timeout / 設定値［記載時点の設定値］: 45 / 設定例: 60 / 説明: 分数"
    ]
    # 無効にすると表の本文は役割の表示の無い元の形になる。
    plain = _payload(source, column_role_detection="off")
    assert _extraction_texts(plain) == ["項目: timeout / 設定値: 45 / 設定例: 60 / 説明: 分数"]


def test_configured_column_role_for_missing_column_warns() -> None:
    outcome = convert(
        _xlsx_bytes(_PARAMETER_SHEET),
        "",
        "excel_to_jsonl",
        None,
        options={"column_roles": {"現在値": "current"}},
    )
    assert outcome.converted is True
    assert "excel_column_role_target_not_found:現在値" in outcome.warnings


def test_invalid_column_role_option_is_rejected() -> None:
    outcome = convert(
        _xlsx_bytes(_PARAMETER_SHEET),
        "",
        "excel_to_jsonl",
        None,
        options={"column_roles": {"B": "actual"}},
    )
    assert (outcome.converted, outcome.warnings) == (False, ("excel_options_invalid",))


def test_low_confidence_header_skips_role_detection(monkeypatch: pytest.MonkeyPatch) -> None:
    """表頭が確かでないシートでは、表頭の語で役割を決めない（指定の役割は使う）。"""
    monkeypatch.setattr(reader, "_detect_header", lambda grid, rows: (3, 0.3, "低い"))
    sheet = _sheet(_payload(_xlsx_bytes(_PARAMETER_SHEET), column_roles={"B": "default"}))
    assert _roles(sheet)["既定値"] == ("default", "configured")
    assert _roles(sheet)["設定例"] == (None, None)
    assert "column_roles_skipped_low_header_confidence" in {
        item["code"] for item in sheet["diagnostics"]
    }


def test_procedure_lines_label_role_columns() -> None:
    payload = _payload(
        _xlsx_bytes(
            {
                "手順": [
                    ["", "作業項目", "既定値", "記入例"],
                    [1, "タイムアウトを設定する", "30", "60"],
                    ["", "", "", "90"],
                    [2, "上限を設定する", "20", "100"],
                ]
            }
        ),
        mode="procedure",
    )
    sheet = _sheet(payload)
    assert sheet["mode"] == "procedure"
    assert sheet["blocks"][0]["lines"][0] == (
        "column_A: 1 | 作業項目: タイムアウトを設定する | 既定値［資料の既定値］: 30 | "
        "記入例［例示の値］: 60"
    )


# ---- .xls（xlrd）の値 -----------------------------------------------------------------

_XLRD = SimpleNamespace(
    XL_CELL_EMPTY=0,
    XL_CELL_TEXT=1,
    XL_CELL_NUMBER=2,
    XL_CELL_DATE=3,
    XL_CELL_BOOLEAN=4,
    XL_CELL_ERROR=5,
    XL_CELL_BLANK=6,
    error_text_from_code={0x07: "#DIV/0!"},
    xldate_as_datetime=lambda value, datemode: dt.datetime(1899, 12, 30) + dt.timedelta(days=value),
)


def _xls_text(ctype: int, value: Any, number_format: str = "General") -> str:
    sheet = SimpleNamespace(
        cell=lambda row, col: SimpleNamespace(ctype=ctype, value=value),
        cell_xf_index=lambda row, col: 0,
    )
    book = SimpleNamespace(
        datemode=0,
        xf_list=[SimpleNamespace(format_key=1)],
        format_map={1: SimpleNamespace(format_str=number_format)},
    )
    return reader._xls_cell_text(book, sheet, 0, 0, _XLRD)


def test_xls_dates_booleans_numbers_and_errors() -> None:
    assert _xls_text(_XLRD.XL_CELL_DATE, 46113.0) == "2026-04-01"
    assert _xls_text(_XLRD.XL_CELL_BOOLEAN, 1) == "TRUE"
    assert _xls_text(_XLRD.XL_CELL_BOOLEAN, 0) == "FALSE"
    assert _xls_text(_XLRD.XL_CELL_NUMBER, 7.0, "0000") == "0007"
    assert _xls_text(_XLRD.XL_CELL_NUMBER, 2.5) == "2.5"
    assert _xls_text(_XLRD.XL_CELL_TEXT, "001") == "001"
    assert _xls_text(_XLRD.XL_CELL_ERROR, 0x07) == "#DIV/0!"
    assert _xls_text(_XLRD.XL_CELL_BLANK, "") == ""


def test_xls_round_trip_when_xlwt_is_available() -> None:
    xlwt = pytest.importorskip("xlwt")  # .xls 書き込みで往復検証(無ければ skip)
    workbook = xlwt.Workbook()
    worksheet = workbook.add_sheet("Sheet1")
    for r, row in enumerate([["name", "score"], ["Carol", 99]]):
        for c, value in enumerate(row):
            worksheet.write(r, c, value)
    buffer = io.BytesIO()
    workbook.save(buffer)
    payload = _payload(buffer.getvalue())
    assert payload["source_format"] == "xls"
    assert _sheet(payload)["blocks"][0]["values"] == {"name": "Carol", "score": "99"}


def test_system_ledger_of_the_evaluation_set_is_row_records() -> None:
    """評価セットの台帳（multi-hop の system-ledger.xlsx）は前書きと 1 行ずつの記録（#1349）。

    前書きの行数・行の数とシステム ID は台帳の原稿（`sources/system-ledger.workbook.json`。#1352
    で 80 行。#1406 で前書きに重要度の順位を足して 3 行）から決める。
    """
    import json
    from pathlib import Path

    rag_dir = Path(__file__).resolve().parents[4]
    source = rag_dir / "evaluation" / "multi-hop" / "system-ledger.xlsx"
    workbook = json.loads(
        (source.parent / "sources" / "system-ledger.workbook.json").read_text(encoding="utf-8")
    )
    ledger_ids = [row[0] for row in workbook["sheets"][0]["rows"]]
    preamble_rows = len(workbook["sheets"][0]["preamble"])
    header_row = preamble_rows + 1
    sheet = _sheet(_payload(source.read_bytes()))
    assert sheet["name"] == "システム台帳"
    assert (sheet["header_row"], sheet["mode"]) == (header_row, "table")
    assert [row["cell_range"] for row in sheet["preamble"]] == [
        f"A{row}:A{row}" for row in range(1, header_row)
    ]
    blocks = sheet["blocks"]
    assert [block["cell_range"] for block in blocks] == [
        f"A{row}:F{row}" for row in range(header_row + 1, header_row + 1 + len(ledger_ids))
    ]
    assert all(block["kind"] == "row" for block in blocks)
    assert [block["values"]["システムID"] for block in blocks] == ledger_ids
    assert list(blocks[0]["values"]) == [
        "システムID",
        "正式名",
        "略称・別表記",
        "担当部署",
        "重要度",
        "機密区分",
    ]


def test_jsonl_lines_carry_schema_ids_and_sheet_context() -> None:
    """1 物理行 1 オブジェクト。document・sheet・記録の順で、記録は sheet を親に持つ（#2457）。

    データは合成。
    """
    outcome = convert(
        _xlsx_bytes({"S": [["説明の行"], ["名前", "値"], ["a", "1\n2"], ["b", "3"]]}),
        "",
        "excel_to_jsonl",
        None,
    )
    assert outcome.derived_bytes is not None
    text = outcome.derived_bytes.decode("utf-8")
    assert text.endswith("\n") and "\n\n" not in text
    raw = [json.loads(line) for line in text.split("\n")[:-1]]
    assert [item["record_type"] for item in raw] == ["document", "sheet", "preamble", "row", "row"]
    assert all(item["schema"] == "production-ready.structured-records" for item in raw)
    assert all(item["schema_version"] == 1 and item["source_type"] == "xlsx" for item in raw)
    assert [item["record_id"] for item in raw[1:]] == [
        "sheet-000",
        "sheet-000-r1",
        "sheet-000-r3",
        "sheet-000-r4",
    ]
    assert raw[0]["converter"] == "excel_to_jsonl"
    row = raw[3]
    assert row["parent_id"] == "sheet-000"
    assert row["location"] == {"row_start": 3, "row_end": 3, "cell_range": "A3:B3"}
    assert row["fields"] == [
        {"name": "名前", "position": 1, "column": "A", "value": "a"},
        {"name": "値", "position": 2, "column": "B", "value": "1\n2"},
    ]
    sheet = _records(outcome.derived_bytes).lines[0]
    assert isinstance(sheet, SheetLine) and sheet.header_row == 2


def _as_macro_enabled(data: bytes) -> bytes:
    """.xlsx をマクロ付きの .xlsm の形（本体の content type と vbaProject.bin）にする（合成）。"""
    import zipfile

    source = zipfile.ZipFile(io.BytesIO(data))
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            payload = source.read(info.filename)
            if info.filename == "[Content_Types].xml":
                payload = payload.replace(
                    b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
                    b"application/vnd.ms-excel.sheet.macroEnabled.main+xml",
                )
                assert b"macroEnabled" in payload
            target.writestr(info, payload)
        target.writestr("xl/vbaProject.bin", b"synthetic-macro")
    return buffer.getvalue()


def test_macro_enabled_workbook_is_read_without_macros() -> None:
    # .xlsm も .xlsx と同じく読み、マクロ（vbaProject.bin）は読まない・実行しない（#2470）。
    data = _as_macro_enabled(_xlsx_bytes({"一覧": [["名前", "値"], ["行 A", 1]]}))
    outcome = convert(
        data,
        "application/vnd.ms-excel.sheet.macroEnabled.12",
        "excel_to_jsonl",
        SimpleNamespace(extension=".xlsm"),  # type: ignore[arg-type]
    )
    assert outcome.converted is True
    assert outcome.derived_content_type == STRUCTURED_RECORDS_CONTENT_TYPE
    assert b"synthetic-macro" not in (outcome.derived_bytes or b"")
    loaded = _loaded(outcome.derived_bytes)
    assert loaded["sheets"][0]["blocks"][0]["values"] == {"名前": "行 A", "値": "1"}


def test_note_rows_below_a_table_become_note_records_not_a_new_table() -> None:
    """表の下の注記の行を別の表・表頭にせず、注記の記録にする（#3566）。データは合成。

    表は A1:C4、空の行を挟んで A6・A7 に注記を 1 つずつ置く。注記の印をセルに含む行の記録は、
    その注記の文を本文に含める（1 行 = 1 記録）。
    """
    note_1 = "※1 合成の注記の一つ目で、表の値の条件を説明する文です。"
    note_2 = "※2 合成の注記の二つ目の文です。"
    outcome = convert(
        _xlsx_bytes(
            {
                "表": [
                    ["項目", "値", "区分"],
                    ["項目A", "10※1", "甲"],
                    ["項目B", "20", "乙"],
                    ["項目C", "30", "丙"],
                    [],
                    [note_1],
                    [note_2],
                ]
            }
        ),
        "",
        "excel_to_jsonl",
        None,
    )
    assert outcome.converted is True and outcome.derived_bytes is not None
    raw = [json.loads(line) for line in outcome.derived_bytes.decode("utf-8").splitlines()]
    assert [item["record_type"] for item in raw] == [
        "document",
        "sheet",
        "row",
        "row",
        "row",
        "note",
        "note",
    ]
    sheet = raw[1]
    assert sheet["header_row"] == 1
    assert [column["name"] for column in sheet["columns"]] == ["項目", "値", "区分"]
    assert all(note_1 not in str(column) for column in sheet["columns"])
    assert [(item["record_id"], item["text"], item["marker"]) for item in raw[5:]] == [
        ("sheet-000-r6", note_1, "※1"),
        ("sheet-000-r7", note_2, "※2"),
    ]
    assert raw[5]["location"] == {"row_start": 6, "row_end": 6, "cell_range": "A6:A6"}
    assert raw[2]["notes"] == [note_1]
    assert "notes" not in raw[3]

    # JSONL から戻した記録と要素: 注記は文の要素、注記の印のある行の本文は注記を含む。
    loaded = _loaded(outcome.derived_bytes)
    assert [note["text"] for note in _sheet(loaded)["notes"]] == [note_1, note_2]
    extraction = structured_records_extraction(
        _records(outcome.derived_bytes),
        source_parser="test",
        parser_backend="test",
        parser_version="0",
    )
    by_id = {element.element_id: element for element in extraction.elements}
    assert by_id["sheet-000-r2"].text == f"項目: 項目A / 値: 10※1 / 区分: 甲\n{note_1}"
    assert by_id["sheet-000-r3"].text == "項目: 項目B / 値: 20 / 区分: 乙"
    note = by_id["sheet-000-r6"]
    assert note.text == note_1
    assert note.content_kind == "text"
    assert note.metadata["sheet_block_kind"] == "note"
    assert note.metadata["structured_record_id"] == "sheet-000-r6"
