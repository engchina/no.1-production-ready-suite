"""帳票（方眼紙）の版式と、1 シートの複数の表の読み方の検証（#2516）。

テストのデータはすべて合成（ラベル・値は匿名の「項目A」「値A」）。帳票は openpyxl で
テストの中で作る。
"""

from __future__ import annotations

import io
from typing import Any

import openpyxl
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from rag_parser_core.sheet_records import sheet_records_extraction
from rag_parser_core.structured_records import (
    FormFieldLine,
    SheetLine,
    StructuredRecordsDocument,
    parse_structured_records,
    structured_records_extraction,
    structured_records_to_sheet_records,
)

from app.converters import convert, read_excel

_FILL = PatternFill("solid", fgColor="DDDDDD")


def _grid(worksheet: Worksheet, *, width: float = 2.5, columns: int = 30) -> None:
    """方眼紙（狭い列の格子）にする。"""
    for col in range(1, columns + 1):
        worksheet.column_dimensions[get_column_letter(col)].width = width


def _put(worksheet: Worksheet, ref: str, value: Any, *, fill: bool = False) -> None:
    first = ref.split(":", 1)[0]
    worksheet[first] = value
    if ":" in ref:
        worksheet.merge_cells(ref)
    if fill:
        worksheet[first].fill = _FILL


def _bytes(workbook: openpyxl.Workbook) -> bytes:
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _form_workbook() -> openpyxl.Workbook:
    """ラベルの右・下の値、縦に結合したラベル、単位のセル、空欄、項目のまとまりを持つ帳票。"""
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "様式"
    _grid(worksheet)
    _put(worksheet, "B1:AC1", "様式A")
    _put(worksheet, "B3:E3", "項目A")
    _put(worksheet, "F3:N3", "値A")
    _put(worksheet, "P3:S3", "項目B")
    _put(worksheet, "T3:AC3", "2026-01-02")
    _put(worksheet, "B4:E5", "項目C")
    _put(worksheet, "F4:AC4", "値C1")
    _put(worksheet, "F5:AC5", "値C2")
    _put(worksheet, "B6:E6", "項目D")
    _put(worksheet, "F6:J6", 1200)
    worksheet["F6"].number_format = "#,##0"
    _put(worksheet, "K6:L6", "円")
    _put(worksheet, "B7:E7", "項目E")
    worksheet.merge_cells("F7:N7")  # 記入されていない欄
    _put(worksheet, "B9:E9", "項目F")
    _put(worksheet, "B10:AC11", "値F の文")
    _put(worksheet, "B13:D14", "区分G")
    _put(worksheet, "E13:H13", "項目G1")
    _put(worksheet, "I13:P13", "値G1")
    _put(worksheet, "E14:H14", "項目G2")
    _put(worksheet, "I14:P14", "値G2")
    return workbook


def _records(derived: bytes) -> StructuredRecordsDocument:
    document = parse_structured_records(derived)
    assert isinstance(document, StructuredRecordsDocument), document
    return document


def _convert(source: bytes, **options: Any) -> StructuredRecordsDocument:
    outcome = convert(source, "", "excel_to_jsonl", None, options=options)
    assert outcome.converted is True, outcome.warnings
    assert outcome.derived_bytes is not None
    return _records(outcome.derived_bytes)


def _sheets(document: StructuredRecordsDocument) -> list[SheetLine]:
    return [line for line in document.lines if isinstance(line, SheetLine)]


def _fields(document: StructuredRecordsDocument) -> list[FormFieldLine]:
    return [line for line in document.lines if isinstance(line, FormFieldLine)]


def test_auto_detects_form_layout_with_confidence_and_features() -> None:
    document = _convert(_bytes(_form_workbook()))
    (sheet,) = _sheets(document)
    assert sheet.mode == "form"
    detection = sheet.layout_detection
    assert detection is not None
    assert detection.layout == "form"
    assert detection.method == "detected"
    assert detection.confidence >= 0.6
    assert set(detection.features) == {"merged", "narrow", "misaligned", "label_value"}
    assert detection.features["merged"] == 1.0
    assert detection.features["narrow"] == 1.0
    assert "結合セル" in detection.reason
    # 判定の根拠は診断にも残す（抽出の parser_artifacts の sheet_diagnostics に入る）。
    assert any(item.code == "layout_detected_form" for item in sheet.diagnostics)
    assert sheet.header_row is None


def test_form_fields_keep_label_value_and_cell_ranges() -> None:
    document = _convert(_bytes(_form_workbook()))
    fields = [
        (
            line.label,
            line.value,
            line.label_cell,
            line.value_cell,
            line.location.cell_range,
            line.section_path,
        )
        for line in _fields(document)
    ]
    assert fields == [
        ("項目A", "値A", "B3:E3", "F3:N3", "B3:N3", ["様式A"]),
        ("項目B", "2026-01-02", "P3:S3", "T3:AC3", "P3:AC3", ["様式A"]),
        # 縦に結合したラベルは、右の 2 行の値をまとめる。
        ("項目C", "値C1 値C2", "B4:E5", "F4:AC5", "B4:AC5", ["様式A"]),
        # 表示の書式と、右の単位のセル。
        ("項目D", "1,200 円", "B6:E6", "F6:L6", "B6:L6", ["様式A"]),
        # 右に値の無いラベルは、すぐ下の値と組にする。
        ("項目F", "値F の文", "B9:E9", "B10:AC11", "B9:AC11", ["様式A"]),
        # 縦に結合し、右に複数の項目を持つ単位は、項目のまとまりの見出し。
        ("項目G1", "値G1", "E13:H13", "I13:P13", "E13:P13", ["様式A", "区分G"]),
        ("項目G2", "値G2", "E14:H14", "I14:P14", "E14:P14", ["様式A", "区分G"]),
    ]
    first = _fields(document)[0]
    assert first.location.row_start == 3 and first.location.row_end == 3
    (sheet,) = _sheets(document)
    # 記入されていない欄（値の無い結合セル）は記録にせず、数だけを診断に残す。
    assert ("form_empty_fields", "1") in {(item.code, item.detail) for item in sheet.diagnostics}
    # 1 行に複数の項目があっても記録の ID は重ならない（列を付ける）。
    ids = [line.record_id for line in document.lines]
    assert len(ids) == len(set(ids))
    assert first.record_id == "sheet-000-r3-B"


def test_hinted_label_rows_pair_with_values_below_and_become_table_rows() -> None:
    """塗りつぶしのラベルが横に並び、下に同じ位置の値の行があれば上下で組にする。"""
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "様式"
    _grid(worksheet, columns=16)
    for ref, value in (("B2:F2", "項目H"), ("G2:K2", "項目I"), ("L2:P2", "項目J")):
        _put(worksheet, ref, value, fill=True)
    for ref, value in (("B3:F3", "値H"), ("G3:K3", "値I"), ("L3:P3", "値J")):
        _put(worksheet, ref, value)
    _put(worksheet, "B5:F5", "列K", fill=True)
    _put(worksheet, "G5:K5", "列L", fill=True)
    for row, (key, number) in enumerate((("k1", 10), ("k2", 20)), start=6):
        _put(worksheet, f"B{row}:F{row}", key)
        _put(worksheet, f"G{row}:K{row}", number)
    document = _convert(_bytes(workbook))
    (sheet,) = _sheets(document)
    assert sheet.mode == "form"
    assert [(line.label, line.value, line.location.cell_range) for line in _fields(document)] == [
        ("項目H", "値H", "B2:F3"),
        ("項目I", "値I", "G2:K3"),
        ("項目J", "値J", "L2:P3"),
    ]
    # 値の行が 2 行以上なら、ラベルの行を表頭にした表の行にする。
    loaded = structured_records_to_sheet_records(document)
    rows = [block for block in loaded.sheets[0].blocks if block.kind == "row"]
    assert [(block.values, block.cell_range) for block in rows] == [
        ({"列K": "k1", "列L": "10"}, "B6:K6"),
        ({"列K": "k2", "列L": "20"}, "B7:K7"),
    ]


def test_plain_table_and_threshold_keep_table_layout() -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "表"
    worksheet.append(["列A", "列B", "列C"])
    for index in range(5):
        worksheet.append([f"a{index}", index, index * 2])
    (sheet,) = _sheets(_convert(_bytes(workbook)))
    assert sheet.mode == "table"
    assert sheet.layout_detection is not None
    assert sheet.layout_detection.layout == "table"
    assert sheet.layout_detection.confidence >= 0.9
    assert "閾値" in sheet.layout_detection.reason
    assert not any(item.code.startswith("layout_detected") for item in sheet.diagnostics)

    # 閾値を上げると、帳票の版式でも帳票にしない（空の行で分かれた表として読む）。
    form_sheets = _sheets(_convert(_bytes(_form_workbook()), form_detection_threshold=1.0))
    assert form_sheets
    for form_sheet in form_sheets:
        assert form_sheet.mode != "form"
        assert form_sheet.layout_detection is not None
        assert form_sheet.layout_detection.layout != "form"


def test_configured_form_mode_reads_without_narrow_columns() -> None:
    workbook = _form_workbook()
    worksheet = workbook.active
    for col in range(1, 31):
        worksheet.column_dimensions[get_column_letter(col)].width = 8.43
    (sheet,) = _sheets(document := _convert(_bytes(workbook), mode="form"))
    assert sheet.mode == "form"
    assert sheet.layout_detection is not None
    assert sheet.layout_detection.method == "configured"
    assert sheet.layout_detection.confidence == 1.0
    assert _fields(document)[0].label == "項目A"


def test_sheet_with_multiple_tables_is_split_with_own_headers() -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "表"
    worksheet.append(["表の題名1"])
    worksheet.append(["列A", "列B", "列C"])
    for index in range(3):
        worksheet.append([f"a{index}", index, index * 2])
    worksheet.append([])
    worksheet.append([])
    worksheet.append(["表の題名2"])
    worksheet.append(["項目X", "項目Y"])
    for index in range(3):
        worksheet.append([f"x{index}", f"y{index}"])
    document = _convert(_bytes(workbook))
    sheets = _sheets(document)
    assert [
        (line.name, line.table_index, line.region, line.title, line.header_row) for line in sheets
    ] == [
        ("表", 1, "A1:C5", "表の題名1", 2),
        ("表", 2, "A8:B12", "表の題名2", 9),
    ]
    assert [column.name for column in sheets[1].columns] == ["項目X", "項目Y"]
    loaded = structured_records_to_sheet_records(document)
    second = loaded.sheets[1].blocks[0]
    assert second.values == {"項目X": "x0", "項目Y": "y0"}
    assert second.row_start == 10

    # 要素は表の題名を節に持ち、metadata に表の番号と範囲を残す（親子階層の親を表ごとに分ける）。
    extraction = sheet_records_extraction(
        loaded, source_parser="test", parser_backend="test", parser_version="v1"
    )
    records = [item for item in extraction.elements if item.content_kind == "record"]
    assert records[0].section_path == ["表", "表の題名1"]
    assert records[-1].section_path == ["表", "表の題名2"]
    assert records[-1].metadata["sheet_table_index"] == 2
    assert records[-1].metadata["sheet_region"] == "A8:B12"
    ids = [item.element_id for item in extraction.elements]
    assert len(ids) == len(set(ids))


def test_side_by_side_tables_are_split_by_empty_columns() -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "横"
    worksheet.append(["列P", "列Q", None, "列R", "列S"])
    worksheet.append(["p1", 1, None, "r1", 10])
    worksheet.append(["p2", 2, None, "r2", 20])
    sheets = _sheets(_convert(_bytes(workbook)))
    assert [(line.region, [column.name for column in line.columns]) for line in sheets] == [
        ("A1:B3", ["列P", "列Q"]),
        ("D1:E3", ["列R", "列S"]),
    ]


def test_blank_row_inside_a_table_keeps_one_table() -> None:
    """空の 1 行で区切った表の続き（表頭らしい行で始まらない）は前の表につなぐ。"""
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "表"
    worksheet.append(["列A", "列B"])
    worksheet.append(["a1", 1])
    worksheet.append(["a2", 2])
    worksheet.append([])
    worksheet.append(["a3", 3])
    worksheet.append(["a4", 4])
    (sheet,) = _sheets(document := _convert(_bytes(workbook)))
    assert sheet.table_index is None and sheet.region is None
    loaded = structured_records_to_sheet_records(document)
    assert [block.row_start for block in loaded.sheets[0].blocks] == [2, 3, 5, 6]


def test_split_can_be_turned_off_and_is_skipped_with_header_row() -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "表"
    worksheet.append(["列A", "列B", "列C"])
    worksheet.append(["a1", 1, 2])
    worksheet.append([])
    worksheet.append([])
    worksheet.append(["項目X", "項目Y"])
    worksheet.append(["x1", "y1"])
    source = _bytes(workbook)
    assert len(_sheets(_convert(source))) == 2
    assert len(_sheets(_convert(source, split_tables=False))) == 1
    assert len(_sheets(_convert(source, header_row=1))) == 1


def test_form_and_table_in_one_sheet_are_read_separately() -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "様式"
    _grid(worksheet, columns=20)
    _put(worksheet, "B1:E1", "項目A")
    _put(worksheet, "F1:M1", "値A")
    _put(worksheet, "B2:E2", "項目B")
    _put(worksheet, "F2:T2", "値B")
    _put(worksheet, "B3:E3", "項目C")
    _put(worksheet, "F3:J3", "値C")
    _put(worksheet, "K3:N3", "項目D")
    _put(worksheet, "O3:T3", "値D")
    worksheet["B6"], worksheet["C6"], worksheet["D6"] = "列A", "列B", "列C"
    for row in range(7, 17):
        worksheet[f"B{row}"] = f"a{row}"
        worksheet[f"C{row}"] = row
        worksheet[f"D{row}"] = row * 10
    document = _convert(_bytes(workbook))
    sheets = _sheets(document)
    assert [(line.table_index, line.mode, line.region) for line in sheets] == [
        (1, "form", "B1:T3"),
        (2, "table", "B6:D16"),
    ]
    assert [(line.label, line.value) for line in _fields(document)] == [
        ("項目A", "値A"),
        ("項目B", "値B"),
        ("項目C", "値C"),
        ("項目D", "値D"),
    ]
    assert sheets[1].header_row == 6


def test_form_records_become_elements_with_cell_metadata() -> None:
    outcome = convert(_bytes(_form_workbook()), "", "excel_to_jsonl", None)
    assert outcome.derived_bytes is not None
    extraction = structured_records_extraction(
        _records(outcome.derived_bytes),
        source_parser="test",
        parser_backend="test",
        parser_version="v1",
    )
    first = next(item for item in extraction.elements if item.content_kind == "record")
    assert first.text == "項目A: 値A"
    assert first.section_path == ["様式", "様式A"]
    assert first.metadata["sheet_block_kind"] == "form_field"
    assert first.metadata["sheet_layout"] == "form"
    assert first.metadata["cell_range"] == "B3:N3"
    assert first.metadata["form_label_cell"] == "B3:E3"
    assert first.metadata["form_value_cell"] == "F3:N3"
    assert first.metadata["structured_record_id"] == first.element_id
    assert extraction.parser_artifacts["form_field_count"] == 7


def test_read_excel_can_be_called_without_the_service() -> None:
    """変換の本体は原本の bytes と選択肢だけで呼べる（会話の添付などでも使う）。"""
    result = read_excel(_bytes(_form_workbook()), options={"mode": "auto"})
    assert result.failure is None
    assert result.document is not None
    assert result.document.sheets[0].mode == "form"
    assert read_excel(b"").failure == "excel_empty"
