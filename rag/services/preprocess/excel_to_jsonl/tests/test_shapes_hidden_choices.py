"""帳票の図形の文字・非表示の行と列・選択の記号の読み方の検証（#2526）。

テストのデータはすべて合成（ラベル・値は匿名の「項目A」「値A」）。ブックは openpyxl でテストの中で
作り、図形は xlsx の zip に最小の drawing の XML を書き足して作る。
"""

from __future__ import annotations

import io
import struct
import zipfile
from types import SimpleNamespace
from typing import Any

import openpyxl
import pytest
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from pr_backend_core.spreadsheet import reader, shapes
from pr_backend_core.spreadsheet.choices import read_choice
from rag_parser_core.sheet_records import (
    SheetRecords,
    SheetRecordsDocument,
    sheet_records_extraction,
)
from rag_parser_core.structured_records import (
    FormFieldLine,
    PreambleLine,
    StructuredRecordsDocument,
    parse_structured_records,
    structured_records_to_sheet_records,
)

from app.converters import convert, read_excel


def _bytes(workbook: openpyxl.Workbook) -> bytes:
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _put(worksheet: Worksheet, ref: str, value: Any) -> None:
    worksheet[ref.split(":", 1)[0]] = value
    if ":" in ref:
        worksheet.merge_cells(ref)


def _form_sheet(worksheet: Worksheet) -> None:
    """方眼紙（狭い列の格子）の帳票。"""
    for col in range(1, 31):
        worksheet.column_dimensions[get_column_letter(col)].width = 2.5
    _put(worksheet, "B1:AC1", "様式A")
    _put(worksheet, "B3:E3", "項目A")
    _put(worksheet, "F3:N3", "値A")
    _put(worksheet, "B4:E4", "項目B")
    worksheet.merge_cells("F4:N4")  # 図形の文字を重ねる空欄
    _put(worksheet, "B5:E5", "項目C")
    _put(worksheet, "F5:N5", "値C")


def _document(source: bytes, **options: Any) -> SheetRecordsDocument:
    result = read_excel(source, options=options)
    assert result.document is not None, result.failure
    return result.document


def _sheet(document: SheetRecordsDocument, index: int = 0) -> SheetRecords:
    return document.sheets[index]


def _codes(sheet: SheetRecords) -> dict[str, str | None]:
    return {item.code: item.detail for item in sheet.diagnostics}


# ---- 図形・テキストボックスの文字 ----------------------------------------------------------

_DRAWING_NS = (
    'xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"'
)


def _marker(tag: str, row: int, col: int, row_off: int = 0, col_off: int = 0) -> str:
    return (
        f"<xdr:{tag}><xdr:col>{col}</xdr:col><xdr:colOff>{col_off}</xdr:colOff>"
        f"<xdr:row>{row}</xdr:row><xdr:rowOff>{row_off}</xdr:rowOff></xdr:{tag}>"
    )


def _shape(*paragraphs: str) -> str:
    body = "".join(f"<a:p><a:r><a:t>{text}</a:t></a:r></a:p>" for text in paragraphs)
    return (
        '<xdr:sp><xdr:nvSpPr><xdr:cNvPr id="2" name="TextBox 1"/><xdr:cNvSpPr txBox="1"/>'
        "</xdr:nvSpPr><xdr:spPr/>"
        f"<xdr:txBody><a:bodyPr/>{body}</xdr:txBody></xdr:sp>"
    )


def _two_cell(start: tuple[int, int], end: tuple[int, int, int, int], body: str) -> str:
    """0 始まりの from（行・列）と to（行・列・行のオフセット・列のオフセット）のアンカー。"""
    return (
        "<xdr:twoCellAnchor>"
        + _marker("from", *start)
        + _marker("to", *end)
        + body
        + "<xdr:clientData/></xdr:twoCellAnchor>"
    )


def _with_drawing(source: bytes, anchors: list[str], *, sheet: int = 1) -> bytes:
    """xlsx の zip に、シートの drawing（anchors）を書き足す。"""
    drawing = f'<?xml version="1.0" encoding="UTF-8"?><xdr:wsDr {_DRAWING_NS}>'
    drawing += "".join(anchors) + "</xdr:wsDr>"
    rels = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rIdShape1" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing" '
        f'Target="../drawings/drawing{sheet}.xml"/></Relationships>'
    )
    package = zipfile.ZipFile(io.BytesIO(source))
    buffer = io.BytesIO()
    with package, zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
        for info in package.infolist():
            payload = package.read(info.filename)
            if info.filename == f"xl/worksheets/sheet{sheet}.xml":
                payload = payload.replace(
                    b"</worksheet>",
                    b'<drawing xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/'
                    b'relationships" r:id="rIdShape1"/></worksheet>',
                )
            if info.filename == "[Content_Types].xml":
                payload = payload.replace(
                    b"</Types>",
                    f'<Override PartName="/xl/drawings/drawing{sheet}.xml" ContentType='
                    '"application/vnd.openxmlformats-officedocument.drawing+xml"/></Types>'.encode(),
                )
            target.writestr(info, payload)
        target.writestr(f"xl/drawings/drawing{sheet}.xml", drawing)
        target.writestr(f"xl/worksheets/_rels/sheet{sheet}.xml.rels", rels)
    return buffer.getvalue()


def test_text_box_over_blank_field_becomes_form_value() -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "様式"
    _form_sheet(worksheet)
    alternate = (
        '<mc:AlternateContent><mc:Choice Requires="a14">'
        + _shape("注意の文", "2 行目")
        + "</mc:Choice><mc:Fallback>"
        + _shape("注意の文", "2 行目")
        + "</mc:Fallback></mc:AlternateContent>"
    )
    source = _with_drawing(
        _bytes(workbook),
        [
            # F4:N4 の空欄に重ねたテキストボックス（to の列のオフセットが 0 なので O 列は
            # 入れない）。
            _two_cell((3, 5), (3, 14, 100, 0), _shape("図形の値")),
            # 帳票の下の、セルの無い場所のテキストボックス（mc:Fallback は読まない）。
            _two_cell((7, 1), (9, 13, 50, 50), alternate),
        ],
    )

    sheet = _sheet(_document(source, mode="form"))
    fields = {(block.label, block.value): block for block in sheet.blocks}
    field = fields[("項目B", "図形の値")]
    assert field.value_cell == "F4:N4"
    assert field.shape_cells == ["F4:N4"]
    note = fields[(None, "注意の文\n2 行目")]
    assert note.cell_range == "B8:N10"
    assert note.shape_cells == ["B8:N10"]
    # 図形の無い項目には印を付けない。
    assert fields[("項目A", "値A")].shape_cells == []


def test_text_box_is_read_by_table_layout_as_preamble() -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "一覧"
    # 1・2 行目は空け、表は 3 行目から。
    for index, row in enumerate(
        (["列A", "列B", "列C"], ["a1", 1, 2], ["a2", 3, 4], ["a3", 5, 6]), start=3
    ):
        for col, value in enumerate(row, start=1):
            worksheet.cell(row=index, column=col, value=value)
    # A2:C2 の表の説明（to の行のオフセットが 0 なので 3 行目は入れない）。
    source = _with_drawing(_bytes(workbook), [_two_cell((1, 0), (2, 2, 0, 10), _shape("表の説明"))])

    sheet = _sheet(_document(source))
    assert sheet.mode == "table"
    assert sheet.header_row == 3
    assert [(row.text, row.shape_cells) for row in sheet.preamble] == [("表の説明", ["A2:C2"])]
    assert [block.values for block in sheet.blocks][0] == {"列A": "a1", "列B": "1", "列C": "2"}
    assert all(block.shape_cells == [] for block in sheet.blocks)

    # JSONL の書き出しと読み戻しで、図形のセル範囲が残る。
    outcome = convert(source, "", "excel_to_jsonl", None)
    assert outcome.derived_bytes is not None
    records = parse_structured_records(outcome.derived_bytes)
    assert isinstance(records, StructuredRecordsDocument)
    preamble = [line for line in records.lines if isinstance(line, PreambleLine)]
    assert preamble[0].shape_cells == ["A2:C2"]
    assert b'"shape_cells":[]' not in outcome.derived_bytes
    restored = structured_records_to_sheet_records(records)
    assert restored.sheets[0].preamble[0].shape_cells == ["A2:C2"]
    extraction = sheet_records_extraction(
        restored, source_parser="test", parser_backend="test", parser_version="1"
    )
    assert extraction.elements[0].metadata["sheet_shape_cells"] == "A2:C2"


def test_text_box_beside_table_is_kept_as_its_own_item() -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    for row in (["列A", "列B"], ["a1", 1], ["a2", 2], ["a3", 3]):
        worksheet.append(row)
    # 表の列と重ならない右の離れた場所のテキストボックス（1 行だけの領域）。
    source = _with_drawing(
        _bytes(workbook), [_two_cell((0, 5), (0, 8, 10, 10), _shape("欄外の注意"))]
    )

    tables = _document(source).sheets
    assert [table.mode for table in tables] == ["table", "form"]
    assert [(block.value, block.shape_cells) for block in tables[1].blocks] == [
        ("欄外の注意", ["F1:I1"])
    ]


def test_text_box_over_filled_cells_is_placed_in_free_cell_or_diagnosed() -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "様式"
    _form_sheet(worksheet)
    source = _with_drawing(
        _bytes(workbook),
        [
            # 値のある B3:E3 と空いた O3 にかかる: 空いた最初のセル（O3）に置く。
            _two_cell((2, 1), (2, 14, 10, 10), _shape("重なる文")),
            # すべてのセルに値がある B5:E5 だけにかかる: 置けないので診断に残す。
            _two_cell((4, 1), (4, 4, 10, 10), _shape("置けない文")),
        ],
    )
    sheet = _sheet(_document(source, mode="form"))
    # O3 は値A のすぐ右なので、値の続きになる。
    placed = [block for block in sheet.blocks if "O3" in block.shape_cells]
    assert [(block.label, block.value) for block in placed] == [("項目A", "値A 重なる文")]
    diagnostics = [(item.code, item.cell, item.detail) for item in sheet.diagnostics]
    assert ("shape_text_not_placed", "B5:E5", "置けない文") in diagnostics


def test_xls_shape_text_count_from_biff_records(monkeypatch: pytest.MonkeyPatch) -> None:
    def record(kind: int, data: bytes = b"") -> bytes:
        return struct.pack("<HH", kind, len(data)) + data

    def sheet_stream(texts: int, notes: int) -> bytes:
        body = record(0x0809, b"\x00" * 8)
        body += record(0x01B6) * texts + record(0x001C) * notes
        # 埋め込みのグラフ（BOF〜EOF）の中の記録もシートの記録として読む。
        body += record(0x0809, b"\x00" * 8) + record(0x000A)
        return body + record(0x000A)

    first, second = sheet_stream(texts=3, notes=1), sheet_stream(texts=1, notes=1)

    def boundsheet(position: int, kind: int = 0) -> bytes:
        return record(0x0085, struct.pack("<IBB", position, 0, kind) + b"\x01\x00a")

    globals_size = len(record(0x0809, b"\x00" * 8) + boundsheet(0) * 3 + record(0x000A))
    stream = (
        record(0x0809, b"\x00" * 8)
        + boundsheet(globals_size)
        + boundsheet(0, kind=2)  # グラフのシートは xlrd と同じく数えない
        + boundsheet(globals_size + len(first))
        + record(0x000A)
        + first
        + second
    )
    from xlrd import compdoc

    monkeypatch.setattr(
        compdoc,
        "CompDoc",
        lambda data: SimpleNamespace(
            get_named_stream=lambda name: stream if name == "Workbook" else None
        ),
    )
    assert shapes.xls_shape_text_counts(b"synthetic") == [2, 0]


def test_xls_reports_shapes_and_skips_hidden_columns(monkeypatch: pytest.MonkeyPatch) -> None:
    import xlrd

    cells = {(0, 0): "項目A", (0, 1): "補助", (0, 2): "値A", (1, 0): "項目B", (1, 2): "値B"}
    sheet = SimpleNamespace(
        name="様式",
        visibility=0,
        nrows=2,
        ncols=3,
        merged_cells=[],
        defcolwidth=None,
        standardwidth=None,
        colinfo_map={1: SimpleNamespace(width=2048, hidden=1)},
        rowinfo_map={},
        cell=lambda row, col: SimpleNamespace(
            ctype=xlrd.XL_CELL_TEXT if (row, col) in cells else xlrd.XL_CELL_EMPTY,
            value=cells.get((row, col), ""),
        ),
        cell_xf_index=lambda row, col: 0,
    )
    book = SimpleNamespace(sheets=lambda: [sheet], xf_list=[], font_list=[])
    monkeypatch.setattr(xlrd, "open_workbook", lambda **kwargs: book)
    monkeypatch.setattr(reader, "xls_shape_text_counts", lambda data: [2])

    grids, failure = reader._read_xls(b"synthetic")
    assert failure is None and grids is not None
    assert grids[0].cells == {(1, 1): "項目A", (1, 3): "値A", (2, 1): "項目B", (2, 3): "値B"}
    assert grids[0].hidden_cells == {(1, 2)}
    assert ("shape_text_not_supported_xls", None, "2") in grids[0].notes

    grids, _ = reader._read_xls(b"synthetic", include_hidden=True)
    assert grids is not None and grids[0].cells[(1, 2)] == "補助"


# ---- 非表示の行・列 --------------------------------------------------------------------


def _table_with_hidden() -> openpyxl.Workbook:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "一覧"
    for row in (
        ["列A", "列B", "補助", "列D"],
        ["a1", 1, "x1", 10],
        ["a2", 2, "x2", 20],
        ["a3", 3, "x3", 30],
    ):
        worksheet.append(row)
    worksheet.column_dimensions["C"].hidden = True
    worksheet.row_dimensions[3].hidden = True
    return workbook


def test_hidden_rows_and_columns_are_skipped_by_default() -> None:
    source = _bytes(_table_with_hidden())

    sheet = _sheet(_document(source))
    assert [column.name for column in sheet.columns] == ["列A", "列B", "列D"]
    assert [block.values for block in sheet.blocks] == [
        {"列A": "a1", "列B": "1", "列D": "10"},
        {"列A": "a3", "列B": "3", "列D": "30"},
    ]
    assert _codes(sheet)["hidden_rows_skipped"] == "3"
    assert _codes(sheet)["hidden_columns_skipped"] == "C"

    every = _sheet(_document(source, include_hidden_cells=True))
    assert [column.name for column in every.columns] == ["列A", "列B", "補助", "列D"]
    assert len(every.blocks) == 3
    assert "hidden_rows_skipped" not in _codes(every)


def test_rows_hidden_by_auto_filter_are_read() -> None:
    workbook = _table_with_hidden()
    workbook.active.auto_filter.ref = "A1:D4"
    sheet = _sheet(_document(_bytes(workbook)))
    # フィルターで隠れた行（3 行目）は読む。隠した列（C）は読まない。
    assert [block.values["列A"] for block in sheet.blocks] == ["a1", "a2", "a3"]
    assert "補助" not in [column.name for column in sheet.columns]
    assert _codes(sheet)["hidden_filter_rows_read"] == "3"
    assert "hidden_rows_skipped" not in _codes(sheet)


def test_hidden_helper_column_is_not_read_into_form_fields() -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "様式"
    _form_sheet(worksheet)
    worksheet["AD3"] = "補助の値"
    worksheet.column_dimensions["AD"].hidden = True

    sheet = _sheet(_document(_bytes(workbook), mode="form"))
    assert "補助の値" not in [block.value for block in sheet.blocks]
    assert _codes(sheet)["hidden_columns_skipped"] == "AD"


# ---- 選択の記号 --------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        ("■A □B", "A"),
        ("□A ■B □C", "B"),
        ("☑A ☐B ☑C", "A、C"),
        ("✓ A ☐ B", "A"),
        ("●A ○B", "A"),
        ("[x] A [ ] B", "A"),
        ("■A／□B", "A"),
        ("区分 ■A □B", "区分 A"),
        # 変えない: 記号 1 つ・選んだものが無い・すべて選んだ・記号だけ・記号の無い文。
        ("■概要", None),
        ("□A □B", None),
        ("■A ■B", None),
        ("✓", None),
        ("□ □", None),
        ("値A", None),
        ("○A ○B", None),
    ],
)
def test_read_choice(value: str, expected: str | None) -> None:
    assert read_choice(value) == expected


def test_choice_marks_become_selected_value_in_form_and_jsonl() -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "様式"
    _form_sheet(worksheet)
    _put(worksheet, "F4:N4", "■有 □無")
    # 記号と選択肢を別のセルに置いた帳票（値の続きとして 1 つの値になる）。
    _put(worksheet, "B6:E6", "項目D")
    _put(worksheet, "F6", "□")
    _put(worksheet, "G6:H6", "甲")
    _put(worksheet, "I6", "☑")
    _put(worksheet, "J6:K6", "乙")
    source = _bytes(workbook)

    sheet = _sheet(_document(source, mode="form"))
    fields = {block.label: block for block in sheet.blocks}
    assert (fields["項目B"].value, fields["項目B"].value_original) == ("有", "■有 □無")
    assert (fields["項目D"].value, fields["項目D"].value_original) == ("乙", "□ 甲 ☑ 乙")
    assert fields["項目A"].value_original is None
    assert _codes(sheet)["form_choices_read"] == "2"

    outcome = convert(source, "", "excel_to_jsonl", None, options={"mode": "form"})
    assert outcome.derived_bytes is not None
    records = parse_structured_records(outcome.derived_bytes)
    assert isinstance(records, StructuredRecordsDocument)
    lines = {line.label: line for line in records.lines if isinstance(line, FormFieldLine)}
    assert (lines["項目B"].value, lines["項目B"].value_original) == ("有", "■有 □無")
    extraction = sheet_records_extraction(
        structured_records_to_sheet_records(records),
        source_parser="test",
        parser_backend="test",
        parser_version="1",
    )
    element = next(item for item in extraction.elements if item.text == "項目B: 有")
    assert element.metadata["form_value_original"] == "■有 □無"
