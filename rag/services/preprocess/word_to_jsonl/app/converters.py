"""Word(.doc/.docx)→構造化レコード JSONL 前処理マイクロサービスの変換実装（#2457）。

Word の本文を見出しの階層で読み、見出しの下の段落・リスト・表と、表の前後の説明を意味のまとまりの
記録（``passage``）にして、構造化レコードの JSONL（`rag_parser_core.structured_records`）にする。
表示の行ごとには分けない。parser のサービス（構造化レコード JSONL）はこの記録を 1 つずつ要素にし、
固定分割（``fixed_delimiter``）で 1 記録 1 chunk になる。

- 形式は先頭のバイトで決める（拡張子は見ない）。OLE の複合文書は、
  ``WordDocument`` の stream と FIB の識別子で Word であることを確かめ、プロファイルを分けた
  LibreOffice（マクロの実行を無効・時間の上限・同時に 1 つ・一時ディレクトリの清掃・子プロセスの
  後始末）で .docx にしてから読む。ZIP は .docx（本文の content type が wordprocessingml の文書）を
  そのまま読み、マクロ付き（.docm / .dotm）・テンプレート（.dotx）は VBA の部品を除いた .docx に
  してから読む（マクロは実行しない。診断 ``macros_not_executed`` を残す）。暗号化・破損・Word で
  ない・依存の欠如・時間の超過・空の出力は、変換せずに理由のコード（``word_<理由>``）を返す
  （成功扱いにしない）。
- 見出しは段落のアウトラインのレベル、無ければ段落のスタイル（基のスタイルもたどる）の名前
  （``heading N`` / ``Title``）とアウトラインのレベルで決める。番号付きの段落はリストの項目
  （番号は Word の番号の書式で表示）、見出しに番号があれば見出しの前に付ける。
- まとまり: 見出しから次の見出しまでの本文を 1 つの記録にする。段落の直後の表・リストはその段落
  （導入の文）とまとめ、表の直後の段落（次の表・リストの導入でないもの）は表の注記として表と
  まとめる。本文が ``RECORD_TARGET_CHARS`` を超える節は、このまとまりの境目で複数の記録に分け、
  各記録に ``part``（まとまりの ID・順番・節の全体の位置）を持たせる。1 つの表・段落が大きいときは
  表の行（表頭を繰り返す）・文の境目で分ける（文より長いときは文字数で分け、診断を残す）。
- 位置は要素化した .docx の本文の順の番号（段落・表）で、.doc は派生物の位置（``location_basis``
  = ``derived``）として原本と派生物の sha256 を残す。原本に無いページ番号は作らない。
- 変更履歴は挿入を残し削除を除いた本文にする（診断を残す）。画像・脚注・文末脚注は読まずに、件数の
  診断を残す。ヘッダー・フッターは読まない。
"""

from __future__ import annotations

import contextlib
import hashlib
import io
import os
import re
import shutil
import signal
import subprocess
import tempfile
import threading
import zipfile
from collections.abc import Iterable, Iterator
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Literal

from lxml import etree
from rag_parser_core.preprocess import ConvertOutcome
from rag_parser_core.source import SourceProfile
from rag_parser_core.structured_records import (
    STRUCTURED_RECORDS_CONTENT_TYPE,
    ContentLine,
    ContextLine,
    DocumentLine,
    PassageLine,
    RecordDiagnostic,
    RecordPart,
    SectionLine,
    SourceConversion,
    WordListItem,
    WordLocation,
    WordParagraph,
    WordTable,
    dump_structured_records,
    word_passage_text,
)

WordSourceType = Literal["doc", "docx"]

CONVERTER_NAME = "word_to_jsonl"
# 表の縦の結合の続きを上の行の値で埋めた（#3565）ので v2。
CONVERTER_VERSION = "v2"
PROFILE = "word_to_jsonl"
# OLE2 の複合文書（.doc）と ZIP（.docx）の先頭のバイト。
_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_ZIP_MAGIC = b"PK\x03\x04"
# Word の FIB の識別子（WordDocument の stream の先頭）と、暗号化の bit（FibBase の fEncrypted）。
_FIB_IDENT = 0xA5EC
_FIB_ENCRYPTED = 0x0100
# 受け付ける大きさ（原本・.docx の展開後の合計・ZIP の要素の数）。
MAX_INPUT_BYTES = 64 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 512 * 1024 * 1024
MAX_ZIP_MEMBERS = 20_000
# LibreOffice の変換の時間の上限（秒）。前処理の HTTP の timeout より短くする。
CONVERSION_TIMEOUT_SECONDS = 120
# 1 記録の本文の目安（見出しの行を含む文字数）。超える節はまとまりの境目で分ける。
RECORD_TARGET_CHARS = 1500

_DOCX_MAIN_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"
_TEMPLATE_MAIN_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml"
)
_MACRO_MAIN_TYPES = frozenset(
    {
        "application/vnd.ms-word.document.macroEnabled.main+xml",
        "application/vnd.ms-word.template.macroEnabledTemplate.main+xml",
    }
)
_CONTENT_TYPES_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"
_XML_PARSER = etree.XMLParser(resolve_entities=False, no_network=True, huge_tree=False)

# LibreOffice の利用者のプロファイル（変換ごとに新しく作る）。マクロを実行しない。
_SCRIPTING = "/org.openoffice.Office.Common/Security/Scripting"
_MACRO_DISABLED_REGISTRY = (
    '<?xml version="1.0" encoding="UTF-8"?>\n'
    '<oor:items xmlns:oor="http://openoffice.org/2001/registry"'
    ' xmlns:xs="http://www.w3.org/2001/XMLSchema"'
    ' xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">\n'
    f'<item oor:path="{_SCRIPTING}"><prop oor:name="MacroSecurityLevel" oor:op="fuse">'
    "<value>3</value></prop></item>\n"
    f'<item oor:path="{_SCRIPTING}"><prop oor:name="DisableMacrosExecution" oor:op="fuse">'
    "<value>true</value></prop></item>\n"
    "</oor:items>\n"
)
# 変換は同時に 1 つ（LibreOffice の CPU・メモリを取り合わない）。
_CONVERSION_SLOT = threading.BoundedSemaphore(1)
# 文の境目（段落が大きすぎるときに分ける）。
_SENTENCE_END = re.compile(r"(?<=[。．！？!?])\s*|(?<=\.)\s+")


def _w(tag: str) -> str:
    return f"{{{_W_NS}}}{tag}"


_W_P = _w("p")
_W_TBL = _w("tbl")
_W_TR = _w("tr")
_W_TC = _w("tc")
_W_SDT = _w("sdt")
_W_SDT_CONTENT = _w("sdtContent")
_W_CUSTOM_XML = _w("customXml")
_W_VAL = _w("val")
# 本文の文字に含めない要素（削除・移動元・互換の代替・書式）。
_TEXT_SKIP_TAGS = frozenset(
    {_w("del"), _w("moveFrom"), _MC_FALLBACK, _w("pPr"), _w("rPr"), _w("tblPr"), _w("tcPr")}
)


class WordRejected(Exception):
    """変換しない理由（``reason`` は ``word_<理由>`` のコード）。"""

    def __init__(self, reason: str) -> None:
        super().__init__(reason)
        self.reason = reason


def soffice_path() -> str | None:
    """利用可能な LibreOffice 実行ファイルを返す(未導入なら None)。"""
    return shutil.which("soffice") or shutil.which("libreoffice")


def convert(
    source_bytes: bytes,
    content_type: str,
    preprocess_profile: str,
    source_profile: SourceProfile | None,
) -> ConvertOutcome:
    """選択プリセットで変換する。変換できないときは理由のコードを付けて変換しない。"""
    del content_type
    if preprocess_profile != PROFILE:
        return ConvertOutcome.passthrough(
            reason=f"preprocess_unsupported_profile:{preprocess_profile}"
        )
    try:
        return _word_to_records(source_bytes, source_profile)
    except WordRejected as exc:
        return ConvertOutcome.passthrough(reason=exc.reason)


def _extension(source_profile: SourceProfile | None) -> str:
    extension = (source_profile.extension if source_profile is not None else None) or ""
    return extension.strip().lower().lstrip(".")


def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _word_to_records(source_bytes: bytes, source_profile: SourceProfile | None) -> ConvertOutcome:
    if not source_bytes:
        raise WordRejected("word_empty")
    if len(source_bytes) > MAX_INPUT_BYTES:
        raise WordRejected("word_input_too_large")
    diagnostics: list[_Diagnostic] = []
    conversion: SourceConversion | None = None
    if source_bytes.startswith(_OLE_MAGIC):
        diagnostics.extend(inspect_ole_word(source_bytes))
        docx_bytes = doc_to_docx(source_bytes)
        if inspect_docx_package(docx_bytes) != "document":
            docx_bytes = without_macros(docx_bytes)
        conversion = SourceConversion(
            original_type="doc",
            original_sha256=_sha256(source_bytes),
            derived_type="docx",
            derived_sha256=_sha256(docx_bytes),
            converter="libreoffice",
        )
    elif source_bytes.startswith(_ZIP_MAGIC):
        kind = inspect_docx_package(source_bytes)
        docx_bytes = source_bytes
        if kind != "document":
            # マクロ付き（.docm / .dotm）・テンプレート（.dotx）は、VBA を読まずに本文だけを読む。
            docx_bytes = without_macros(source_bytes)
            if kind == "macro":
                diagnostics.append(_Diagnostic("macros_not_executed"))
    else:
        raise WordRejected("word_unsupported_format")
    reader = _DocxReader.load(docx_bytes)
    lines = reader.lines(
        source_type="doc" if conversion is not None else "docx",
        conversion=conversion,
        diagnostics=diagnostics,
    )
    if not any(isinstance(line, PassageLine) for line in lines):
        raise WordRejected("word_no_content")
    warnings = tuple(item.warning() for item in reader.diagnostics)
    return ConvertOutcome(
        converted=True,
        converter_name=CONVERTER_NAME,
        converter_version=CONVERTER_VERSION,
        derived_bytes=dump_structured_records(lines),
        derived_content_type=STRUCTURED_RECORDS_CONTENT_TYPE,
        warnings=warnings,
    )


# ---- 形式の確認 ---------------------------------------------------------------------------


@dataclass(frozen=True)
class _Diagnostic:
    code: str
    count: int | None = None

    def warning(self) -> str:
        return f"word_{self.code}" + (f":{self.count}" if self.count is not None else "")

    def record(self) -> RecordDiagnostic:
        return RecordDiagnostic(
            code=self.code, detail=str(self.count) if self.count is not None else None
        )


def inspect_ole_word(data: bytes) -> list[_Diagnostic]:
    """OLE の複合文書が暗号化されていない Word（.doc）かを確かめる。"""
    import olefile

    try:
        ole = olefile.OleFileIO(io.BytesIO(data))
    except Exception as exc:  # olefile は壊れた入力で OSError 以外も出す
        raise WordRejected("word_corrupt") from exc
    with ole:
        # 暗号化した .docx（Office の暗号化）は EncryptedPackage を持つ OLE の複合文書になる。
        if ole.exists("EncryptedPackage") or ole.exists("EncryptionInfo"):
            raise WordRejected("word_encrypted")
        if not ole.exists("WordDocument"):
            raise WordRejected("word_not_word_document")
        try:
            header = ole.openstream("WordDocument").read(12)
        except Exception as exc:
            raise WordRejected("word_corrupt") from exc
        macros = ole.exists("Macros")
    if len(header) < 12 or int.from_bytes(header[0:2], "little") != _FIB_IDENT:
        raise WordRejected("word_corrupt")
    if int.from_bytes(header[10:12], "little") & _FIB_ENCRYPTED:
        raise WordRejected("word_encrypted")
    # マクロは実行しない（LibreOffice のプロファイルで無効にし、.docx には残らない）。
    return [_Diagnostic("macros_not_executed")] if macros else []


def inspect_docx_package(data: bytes) -> str:
    """ZIP が Word の文書で、展開の大きさが上限の中かを確かめる。

    戻り値は本体の種類: ``document``（.docx）・``template``（.dotx）・``macro``（.docm / .dotm、
    または VBA の部品を持つ）。
    """
    try:
        package = zipfile.ZipFile(io.BytesIO(data))
    except (zipfile.BadZipFile, OSError) as exc:
        raise WordRejected("word_corrupt") from exc
    with package:
        infos = package.infolist()
        if len(infos) > MAX_ZIP_MEMBERS or sum(info.file_size for info in infos) > (
            MAX_UNCOMPRESSED_BYTES
        ):
            raise WordRejected("word_input_too_large")
        names = {info.filename for info in infos}
        if "[Content_Types].xml" not in names:
            raise WordRejected("word_not_word_document")
        try:
            types = etree.fromstring(package.read("[Content_Types].xml"), _XML_PARSER)
            broken = package.testzip()
        except (etree.XMLSyntaxError, zipfile.BadZipFile, OSError, ValueError) as exc:
            raise WordRejected("word_corrupt") from exc
    if broken is not None:
        raise WordRejected("word_corrupt")
    content_types = {
        item.get("ContentType", "") for item in types.iter(f"{{{_CONTENT_TYPES_NS}}}Override")
    }
    if not content_types & {_DOCX_MAIN_TYPE, _TEMPLATE_MAIN_TYPE, *_MACRO_MAIN_TYPES}:
        raise WordRejected("word_not_word_document")
    if content_types & _MACRO_MAIN_TYPES or any(_is_macro_part(name) for name in names):
        return "macro"
    if _DOCX_MAIN_TYPE not in content_types:
        return "template"
    return "document"


def _is_macro_part(name: str) -> bool:
    base = name.rsplit("/", 1)[-1].casefold()
    return base in {"vbaproject.bin", "vbadata.xml", "vbaproject.bin.rels"}


def without_macros(data: bytes) -> bytes:
    """マクロ付き・テンプレートの Word を、VBA の部品を除いた .docx の package にする。

    VBA の部品（vbaProject.bin・vbaData.xml とその関係）は読まずに捨て、本体の content type を
    文書（.docx）にする。マクロは実行しない（本文だけを python-docx で読む）。
    """
    source = zipfile.ZipFile(io.BytesIO(data))
    buffer = io.BytesIO()
    with source, zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            if _is_macro_part(info.filename):
                continue
            payload = source.read(info.filename)
            if info.filename == "[Content_Types].xml":
                payload = _document_content_types(payload)
            elif info.filename.endswith(".rels"):
                payload = _without_macro_relationships(payload)
            target.writestr(info, payload)
    return buffer.getvalue()


def _document_content_types(payload: bytes) -> bytes:
    root = etree.fromstring(payload, _XML_PARSER)
    for item in list(root):
        content_type = item.get("ContentType", "")
        if content_type in _MACRO_MAIN_TYPES or content_type == _TEMPLATE_MAIN_TYPE:
            item.set("ContentType", _DOCX_MAIN_TYPE)
        elif "vba" in content_type.casefold() or _is_macro_part(item.get("PartName", "")):
            root.remove(item)
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _without_macro_relationships(payload: bytes) -> bytes:
    root = etree.fromstring(payload, _XML_PARSER)
    for item in list(root):
        target = item.get("Target", "")
        if item.get("Type", "").casefold().endswith("/vbaproject") or _is_macro_part(target):
            root.remove(item)
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


# ---- .doc → .docx（LibreOffice）-------------------------------------------------------------


def doc_to_docx(data: bytes) -> bytes:
    """.doc を、変換ごとに分けたプロファイルの LibreOffice で .docx にする。"""
    soffice = soffice_path()
    if soffice is None:
        raise WordRejected("word_libreoffice_unavailable")
    if not _CONVERSION_SLOT.acquire(timeout=CONVERSION_TIMEOUT_SECONDS):
        raise WordRejected("word_conversion_busy")
    try:
        with tempfile.TemporaryDirectory(prefix="word-to-jsonl-") as tmp:
            root = Path(tmp)
            profile = root / "profile"
            (profile / "user").mkdir(parents=True)
            (profile / "user" / "registrymodifications.xcu").write_text(
                _MACRO_DISABLED_REGISTRY, encoding="utf-8"
            )
            work = root / "work"
            out = root / "out"
            for directory in (work, out, root / "tmp"):
                directory.mkdir()
            source = work / "input.doc"
            source.write_bytes(data)
            command = [
                soffice,
                f"-env:UserInstallation={profile.as_uri()}",
                "--headless",
                "--invisible",
                "--nologo",
                "--nodefault",
                "--nolockcheck",
                "--norestore",
                "--convert-to",
                "docx:MS Word 2007 XML",
                "--outdir",
                str(out),
                str(source),
            ]
            env = {
                "PATH": os.environ.get("PATH", "/usr/bin:/bin"),
                "HOME": str(root),
                "TMPDIR": str(root / "tmp"),
                "LANG": os.environ.get("LANG", "C.UTF-8"),
                "SAL_USE_VCLPLUGIN": "svp",
            }
            run_isolated(command, env=env, cwd=work, timeout=CONVERSION_TIMEOUT_SECONDS)
            produced = out / "input.docx"
            if not produced.is_file():
                raise WordRejected("word_conversion_no_output")
            docx = produced.read_bytes()
    finally:
        _CONVERSION_SLOT.release()
    if not docx:
        raise WordRejected("word_conversion_empty")
    return docx


def run_isolated(command: list[str], *, env: dict[str, str], cwd: Path, timeout: float) -> None:
    """新しいプロセスグループで実行し、終わったら（時間を超えたら）グループごと終了させる。"""
    try:
        process = subprocess.Popen(  # noqa: S603 - argv は固定で shell を通さない
            command,
            stdin=subprocess.DEVNULL,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            env=env,
            cwd=cwd,
            start_new_session=True,
        )
    except OSError as exc:
        raise WordRejected("word_conversion_failed") from exc
    try:
        returncode = process.wait(timeout=timeout)
    except subprocess.TimeoutExpired as exc:
        raise WordRejected("word_conversion_timeout") from exc
    finally:
        # LibreOffice が残した子プロセス（soffice.bin）も含めて終了させる。
        _kill_process_group(process)
    if returncode != 0:
        raise WordRejected("word_conversion_failed")


def _kill_process_group(process: subprocess.Popen[bytes]) -> None:
    with contextlib.suppress(ProcessLookupError, PermissionError):
        os.killpg(process.pid, signal.SIGKILL)
    with contextlib.suppress(subprocess.TimeoutExpired):
        process.wait(timeout=10)


# ---- .docx の本文を読む --------------------------------------------------------------------


@dataclass(frozen=True)
class _LevelFormat:
    num_format: str
    text: str
    start: int


@dataclass
class _Numbering:
    """番号の定義（numId → 抽象の番号の定義のレベルの書式）と、数えている番号。"""

    levels: dict[str, dict[int, _LevelFormat]] = field(default_factory=dict)
    counters: dict[str, list[int]] = field(default_factory=dict)

    def next_marker(self, num_id: str, level: int) -> tuple[bool, str | None]:
        """(番号付きか, 表示)。記号のリストは (False, None)。"""
        formats = self.levels.get(num_id, {})
        current = formats.get(level)
        if current is None or current.num_format in {"bullet", "none"}:
            return False, None
        counters = self.counters.setdefault(
            num_id, [formats[i].start - 1 if i in formats else 0 for i in range(9)]
        )
        counters[level] += 1
        for deeper in range(level + 1, 9):
            counters[deeper] = formats[deeper].start - 1 if deeper in formats else 0
        marker = current.text
        for index in range(9):
            placeholder = f"%{index + 1}"
            if placeholder in marker:
                fmt = formats[index].num_format if index in formats else "decimal"
                marker = marker.replace(placeholder, _format_number(max(counters[index], 1), fmt))
        return True, marker.strip() or None


def _roman(value: int) -> str:
    numerals = (
        (1000, "M"),
        (900, "CM"),
        (500, "D"),
        (400, "CD"),
        (100, "C"),
        (90, "XC"),
        (50, "L"),
        (40, "XL"),
        (10, "X"),
        (9, "IX"),
        (5, "V"),
        (4, "IV"),
        (1, "I"),
    )
    out = []
    for amount, numeral in numerals:
        while value >= amount:
            out.append(numeral)
            value -= amount
    return "".join(out)


def _letters(value: int) -> str:
    # Word の lowerLetter は a..z の後 aa, bb, …（同じ文字を繰り返す）。
    return chr(ord("a") + (value - 1) % 26) * ((value - 1) // 26 + 1)


def _format_number(value: int, num_format: str) -> str:
    if num_format == "decimalZero":
        return f"{value:02d}"
    if num_format == "decimalFullWidth":
        return str(value).translate(str.maketrans("0123456789", "０１２３４５６７８９"))
    if num_format == "lowerLetter":
        return _letters(value)
    if num_format == "upperLetter":
        return _letters(value).upper()
    if num_format == "lowerRoman":
        return _roman(value).lower()
    if num_format == "upperRoman":
        return _roman(value)
    if num_format in {"decimalEnclosedCircle", "decimalEnclosedCircleChinese"} and value <= 20:
        return chr(0x2460 + value - 1)
    return str(value)


def _val(element: Any, tag: str) -> str | None:
    child = element.find(_w(tag)) if element is not None else None
    return child.get(_W_VAL) if child is not None else None


def _int(value: str | None, default: int = 0) -> int:
    try:
        return int(value) if value is not None else default
    except ValueError:
        return default


def _level_format(lvl: Any) -> _LevelFormat:
    return _LevelFormat(
        num_format=_val(lvl, "numFmt") or "decimal",
        text=_val(lvl, "lvlText") or "",
        start=_int(_val(lvl, "start"), 1),
    )


def _numbering(document: Any) -> _Numbering:
    try:
        root = document.part.numbering_part.element
    except (KeyError, NotImplementedError, AttributeError):
        return _Numbering()
    abstract: dict[str, dict[int, _LevelFormat]] = {}
    for item in root.iterchildren(_w("abstractNum")):
        abstract[item.get(_w("abstractNumId"), "")] = {
            _int(lvl.get(_w("ilvl"))): _level_format(lvl) for lvl in item.iterchildren(_w("lvl"))
        }
    numbering = _Numbering()
    for num in root.iterchildren(_w("num")):
        formats = dict(abstract.get(_val(num, "abstractNumId") or "", {}))
        for override in num.iterchildren(_w("lvlOverride")):
            level = _int(override.get(_w("ilvl")))
            lvl = override.find(_w("lvl"))
            if lvl is not None:
                formats[level] = _level_format(lvl)
            start = _val(override, "startOverride")
            if start is not None and level in formats:
                base = formats[level]
                formats[level] = _LevelFormat(base.num_format, base.text, _int(start, 1))
        numbering.levels[num.get(_w("numId"), "")] = formats
    return numbering


@dataclass(frozen=True)
class _Style:
    name: str
    based_on: str | None
    outline_level: int | None
    num_id: str | None
    num_level: int


def _styles(document: Any) -> tuple[dict[str, _Style], str | None]:
    """段落のスタイル（ID → 名前・基のスタイル・アウトラインのレベル・番号）と既定のスタイル。"""
    styles: dict[str, _Style] = {}
    default: str | None = None
    for item in document.styles.element.iterchildren(_w("style")):
        if item.get(_w("type")) not in (None, "paragraph"):
            continue
        style_id = item.get(_w("styleId"), "")
        ppr = item.find(_w("pPr"))
        num_pr = ppr.find(_w("numPr")) if ppr is not None else None
        outline = _val(ppr, "outlineLvl")
        styles[style_id] = _Style(
            name=(_val(item, "name") or "").casefold(),
            based_on=_val(item, "basedOn"),
            outline_level=_int(outline) if outline is not None else None,
            num_id=_val(num_pr, "numId") if num_pr is not None else None,
            num_level=_int(_val(num_pr, "ilvl")) if num_pr is not None else 0,
        )
        if item.get(_w("default")) in {"1", "true"}:
            default = style_id
    return styles, default


_HEADING_STYLE = re.compile(r"heading\s*([1-9])")


def _collect_text(node: Any, out: list[str]) -> None:
    tag = node.tag
    if not isinstance(tag, str) or tag in _TEXT_SKIP_TAGS:
        return
    if tag == _w("t"):
        out.append(node.text or "")
    elif tag in {_w("tab"), _w("ptab")}:
        out.append(" ")
    elif tag in {_w("br"), _w("cr")}:
        out.append("\n")
    elif tag == _w("noBreakHyphen"):
        out.append("-")
    elif tag == _W_P and out:
        # テキストボックスの中の段落は改行で区切る。
        out.append("\n")
    for child in node:
        _collect_text(child, out)


def _normalize(text: str) -> str:
    lines = (" ".join(line.split()) for line in text.replace(" ", " ").split("\n"))
    return "\n".join(line for line in lines if line)


def _paragraph_text(paragraph: Any) -> str:
    out: list[str] = []
    for child in paragraph:
        _collect_text(child, out)
    return _normalize("".join(out))


def _body_items(container: Any) -> Iterator[Any]:
    """本文の直下の段落と表（内容コントロール・customXml の中もたどる）。"""
    for child in container:
        if child.tag in {_W_P, _W_TBL}:
            yield child
        elif child.tag == _W_SDT:
            content = child.find(_W_SDT_CONTENT)
            if content is not None:
                yield from _body_items(content)
        elif child.tag == _W_CUSTOM_XML:
            yield from _body_items(child)


def _row_cells(row: Any) -> Iterator[Any]:
    for child in row:
        if child.tag == _W_TC:
            yield child
        elif child.tag == _W_SDT:
            content = child.find(_W_SDT_CONTENT)
            if content is not None:
                yield from _row_cells(content)


def _table_rows(table: Any) -> Iterator[Any]:
    for child in table:
        if child.tag == _W_TR:
            yield child
        elif child.tag == _W_SDT:
            content = child.find(_W_SDT_CONTENT)
            if content is not None:
                yield from _table_rows(content)


def _cell_text(cell: Any) -> str:
    parts: list[str] = []
    for item in _body_items(cell):
        if item.tag == _W_P:
            text = _paragraph_text(item)
        else:
            # 表の中の表は、行を「 / 」、セルを「 | 」でつないだ文字にする。
            text = " / ".join(
                " | ".join(value for value in row if value) for row, _ in _table_grid(item)
            )
        if text:
            parts.append(text)
    return " ".join(" ".join(parts).split())


def _on(value: str | None) -> bool:
    return value not in {"0", "false", "off"}


def _table_grid(table: Any) -> list[tuple[list[str], bool]]:
    """(行のセルの文字, 表頭の行か)。横の結合は同じ値で埋め、縦の結合の続きは上の行の値で埋める。

    縦の結合（``w:vMerge``）の続きのセルには、同じ列の上の行（結合の始まり）の文字を写す
    （HTML の ``rowspan`` を覆う行に写すのと同じ読み方。#3565）。空にすると Markdown の表では
    空欄と区別できず、行の記録がその行の区分を失う。結合は Word の XML で決まるので推測ではない。
    """
    rows: list[tuple[list[str], bool]] = []
    for row in _table_rows(table):
        trpr = row.find(_w("trPr"))
        header_flag = trpr.find(_w("tblHeader")) if trpr is not None else None
        above = rows[-1][0] if rows else []
        values = [""] * _int(_val(trpr, "gridBefore"))
        for cell in _row_cells(row):
            tcpr = cell.find(_w("tcPr"))
            span = max(1, _int(_val(tcpr, "gridSpan"), 1))
            merge = tcpr.find(_w("vMerge")) if tcpr is not None else None
            continued = merge is not None and merge.get(_W_VAL) != "restart"
            if continued:
                column = len(values)
                values.extend(
                    above[index] if index < len(above) else ""
                    for index in range(column, column + span)
                )
            else:
                values.extend([_cell_text(cell)] * span)
        rows.append((values, header_flag is not None and _on(header_flag.get(_W_VAL))))
    return rows


@dataclass
class _Section:
    record_id: str | None
    level: int
    path: list[str]
    blocks: list[WordParagraph | WordListItem | WordTable] = field(default_factory=list)


Block = WordParagraph | WordListItem | WordTable


@dataclass
class _DocxReader:
    document: Any
    numbering: _Numbering
    styles: dict[str, _Style]
    default_style: str | None
    diagnostics: list[_Diagnostic] = field(default_factory=list)

    @classmethod
    def load(cls, data: bytes) -> _DocxReader:
        import docx

        try:
            document = docx.Document(io.BytesIO(data))
            styles, default = _styles(document)
            numbering = _numbering(document)
        except Exception as exc:  # python-docx は壊れた部品で様々な例外を出す
            raise WordRejected("word_corrupt") from exc
        return cls(document=document, numbering=numbering, styles=styles, default_style=default)

    # ---- 段落の種類 ------------------------------------------------------------------------

    def _style_chain(self, paragraph: Any) -> Iterator[_Style]:
        ppr = paragraph.find(_w("pPr"))
        style_id = _val(ppr, "pStyle") or self.default_style
        seen: set[str] = set()
        while style_id and style_id not in seen and style_id in self.styles:
            seen.add(style_id)
            style = self.styles[style_id]
            yield style
            style_id = style.based_on

    def _heading_level(self, paragraph: Any) -> int | None:
        ppr = paragraph.find(_w("pPr"))
        outline = _val(ppr, "outlineLvl")
        if outline is not None:
            level = _int(outline, 9)
            return level + 1 if level < 9 else None
        for style in self._style_chain(paragraph):
            if match := _HEADING_STYLE.fullmatch(style.name):
                return int(match.group(1))
            if style.name == "title":
                return 0
            if style.outline_level is not None:
                return style.outline_level + 1 if style.outline_level < 9 else None
        return None

    def _numbering_of(self, paragraph: Any) -> tuple[str, int] | None:
        ppr = paragraph.find(_w("pPr"))
        num_pr = ppr.find(_w("numPr")) if ppr is not None else None
        if num_pr is not None:
            num_id = _val(num_pr, "numId")
            if num_id is not None:
                return None if num_id == "0" else (num_id, _int(_val(num_pr, "ilvl")))
        for style in self._style_chain(paragraph):
            if style.num_id is not None:
                return None if style.num_id == "0" else (style.num_id, style.num_level)
        return None

    # ---- 本文 -----------------------------------------------------------------------------

    def lines(
        self,
        *,
        source_type: WordSourceType,
        conversion: SourceConversion | None,
        diagnostics: list[_Diagnostic],
    ) -> list[DocumentLine | ContextLine | ContentLine]:
        """JSONL の行（document → 見出しとまとまりを本文の順に）。

        ``diagnostics`` は形式の確認の診断（document の行と warning に足す）。
        """
        body = self.document.element.body
        self._collect_diagnostics(body)
        emitter = _Emitter(source_type=source_type)
        body_index = paragraph_index = table_index = 0
        for item in _body_items(body):
            body_index += 1
            if item.tag == _W_TBL:
                table_index += 1
                table = self._table(item, body_index=body_index, table_index=table_index)
                if table is not None:
                    emitter.add_block(table)
                continue
            paragraph_index += 1
            text = _paragraph_text(item)
            numbered = self._numbering_of(item)
            ordered, marker = (
                self.numbering.next_marker(*numbered) if numbered is not None else (False, None)
            )
            if not text:
                continue
            level = self._heading_level(item)
            if level is not None:
                heading = f"{marker} {text}" if marker else text
                emitter.add_heading(
                    " ".join(heading.split()),
                    level=level,
                    body_index=body_index,
                    paragraph_index=paragraph_index,
                )
            elif numbered is not None:
                emitter.add_block(
                    WordListItem(
                        text=text,
                        body_index=body_index,
                        paragraph_index=paragraph_index,
                        level=min(max(numbered[1], 0), 8),
                        ordered=ordered,
                        marker=marker,
                    )
                )
            else:
                emitter.add_block(
                    WordParagraph(text=text, body_index=body_index, paragraph_index=paragraph_index)
                )
        content = emitter.finish()
        self.diagnostics = [*diagnostics, *self.diagnostics, *emitter.diagnostics()]
        document_line = DocumentLine(
            record_id="document",
            source_type=source_type,
            converter=CONVERTER_NAME,
            converter_version=CONVERTER_VERSION,
            location_basis="derived" if conversion is not None else "original",
            conversion=conversion,
            diagnostics=[item.record() for item in self.diagnostics],
        )
        return [document_line, *content]

    def _table(self, table: Any, *, body_index: int, table_index: int) -> WordTable | None:
        grid = _table_grid(table)
        width = max((len(values) for values, _ in grid), default=0)
        header_rows = 0
        for _, is_header in grid:
            if not is_header:
                break
            header_rows += 1
        rows = [[*values, *([""] * (width - len(values)))] for values, _ in grid]
        if not any(any(row) for row in rows):
            return None
        # 空の行も残し、行の番号を元の表と同じにする（本文にするときに空の行は書かない）。
        return WordTable(
            body_index=body_index,
            table_index=table_index,
            header_rows=header_rows,
            rows=rows,
            row_start=1,
            row_end=len(rows),
        )

    def _collect_diagnostics(self, body: Any) -> None:
        images = sum(1 for _ in body.iter(_w("drawing"))) + sum(1 for _ in body.iter(_w("pict")))
        if images:
            self.diagnostics.append(_Diagnostic("images_not_extracted", images))
        if any(True for _ in body.iter(_w("ins"), _w("del"), _w("moveFrom"), _w("moveTo"))):
            self.diagnostics.append(_Diagnostic("tracked_changes_accepted"))
        for code, partname, tag in (
            ("footnotes_not_extracted", "/word/footnotes.xml", "footnote"),
            ("endnotes_not_extracted", "/word/endnotes.xml", "endnote"),
        ):
            count = self._note_count(partname, tag)
            if count:
                self.diagnostics.append(_Diagnostic(code, count))

    def _note_count(self, partname: str, tag: str) -> int:
        for part in self.document.part.package.iter_parts():
            if str(part.partname) != partname:
                continue
            try:
                root = etree.fromstring(part.blob, _XML_PARSER)
            except etree.XMLSyntaxError:
                return 0
            # id が 0 以下の注は区切り線などの既定の注（本文ではない）。
            return sum(
                1
                for note in root.iterchildren(_w(tag))
                if _int(note.get(_w("id")), 0) > 0 and _paragraph_text(note)
            )
        return 0


# ---- 見出しとまとまりの記録 ----------------------------------------------------------------


def _location(blocks: Iterable[Block]) -> WordLocation:
    items = list(blocks)
    paragraphs = [block.paragraph_index for block in items if not isinstance(block, WordTable)]
    return WordLocation(
        body_start=min(block.body_index for block in items),
        body_end=max(block.body_index for block in items),
        paragraph_start=min(paragraphs) if paragraphs else None,
        paragraph_end=max(paragraphs) if paragraphs else None,
        table_indices=list(
            dict.fromkeys(block.table_index for block in items if isinstance(block, WordTable))
        ),
    )


@dataclass
class _Emitter:
    """見出しの階層をたどり、節ごとの本文を意味のまとまりの記録にする。"""

    source_type: WordSourceType
    lines: list[ContextLine | ContentLine] = field(default_factory=list)
    stack: list[_Section] = field(default_factory=list)
    current: _Section = field(default_factory=lambda: _Section(None, -1, []))
    sections: int = 0
    passages: int = 0
    length_splits: int = 0
    heading_location: WordLocation | None = None

    def add_block(self, block: Block) -> None:
        self.current.blocks.append(block)

    def add_heading(
        self, heading: str, *, level: int, body_index: int, paragraph_index: int
    ) -> None:
        self._flush(next_level=level)
        while self.stack and self.stack[-1].level >= level:
            self.stack.pop()
        parent = self.stack[-1] if self.stack else None
        self.sections += 1
        section = _Section(
            record_id=f"s{self.sections:04d}",
            level=level,
            path=[*(parent.path if parent is not None else []), heading],
        )
        location = WordLocation(
            body_start=body_index,
            body_end=body_index,
            paragraph_start=paragraph_index,
            paragraph_end=paragraph_index,
        )
        self.lines.append(
            SectionLine(
                record_id=str(section.record_id),
                source_type=self.source_type,
                parent_id=parent.record_id if parent is not None else None,
                heading=heading,
                level=level,
                heading_path=section.path,
                location=location,
            )
        )
        self.stack.append(section)
        self.current = section
        self.heading_location = location

    def finish(self) -> list[ContextLine | ContentLine]:
        self._flush(next_level=None)
        return self.lines

    def diagnostics(self) -> list[_Diagnostic]:
        return (
            [_Diagnostic("paragraph_split_by_length", self.length_splits)]
            if self.length_splits
            else []
        )

    # ---- 節の本文の記録 --------------------------------------------------------------------

    def _flush(self, *, next_level: int | None) -> None:
        section = self.current
        if not section.blocks:
            # 本文も下位の見出しも無い見出しは、見出しだけの記録にする（見出しを失わない）。
            has_child = next_level is not None and next_level > section.level
            if (
                section.record_id is not None
                and self.heading_location is not None
                and not has_child
            ):
                self._passage(section, [], location=self.heading_location, part=None)
            return
        records = self._pack(section)
        location = _location(section.blocks)
        group_id = f"{section.record_id or 'body'}-content"
        for index, blocks in enumerate(records, start=1):
            part = (
                RecordPart(
                    group_id=group_id, index=index, count=len(records), group_location=location
                )
                if len(records) > 1
                else None
            )
            self._passage(section, blocks, location=_location(blocks), part=part)
        section.blocks = []

    def _passage(
        self,
        section: _Section,
        blocks: list[Block],
        *,
        location: WordLocation,
        part: RecordPart | None,
    ) -> None:
        self.passages += 1
        self.lines.append(
            PassageLine(
                record_id=f"p{self.passages:04d}",
                source_type=self.source_type,
                parent_id=section.record_id,
                section_path=section.path,
                location=location,
                blocks=blocks,
                part=part,
            )
        )

    def _text_length(self, section: _Section, blocks: list[Block]) -> int:
        line = PassageLine(
            record_id="measure",
            source_type=self.source_type,
            section_path=section.path,
            location=WordLocation(body_start=1, body_end=1),
            blocks=blocks,
        )
        return len(word_passage_text(line))

    def _pack(self, section: _Section) -> list[list[Block]]:
        """まとまり（導入の文 + 表・リスト + 表の注記）を、目安の文字数まで記録に詰める。"""
        records: list[list[Block]] = []
        current: list[Block] = []
        for unit in _units(section.blocks):
            if self._text_length(section, [*current, *unit]) <= RECORD_TARGET_CHARS:
                current.extend(unit)
                continue
            if current:
                records.append(current)
                current = []
            if self._text_length(section, unit) <= RECORD_TARGET_CHARS:
                current = list(unit)
                continue
            # 1 つのまとまりが大きい: 表は行、リストは項目、段落は文の境目で分けて詰める。
            for atom in self._atoms(section, unit):
                if current and (self._text_length(section, [*current, atom]) > RECORD_TARGET_CHARS):
                    records.append(current)
                    current = []
                current.append(atom)
        if current:
            records.append(current)
        return records

    def _atoms(self, section: _Section, unit: list[Block]) -> list[Block]:
        atoms: list[Block] = []
        for block in unit:
            if self._text_length(section, [block]) <= RECORD_TARGET_CHARS:
                atoms.append(block)
            elif isinstance(block, WordTable):
                atoms.extend(self._split_table(section, block))
            else:
                atoms.extend(self._split_text(section, block))
        return atoms

    def _split_table(self, section: _Section, table: WordTable) -> list[Block]:
        header = table.rows[: table.header_rows] if table.header_rows else table.rows[:1]
        body = table.rows[len(header) :]
        parts: list[Block] = []
        rows: list[list[str]] = []
        offset = table.row_start + len(header)
        start = offset

        def piece(selected: list[list[str]], first: int) -> WordTable:
            return WordTable(
                body_index=table.body_index,
                table_index=table.table_index,
                header_rows=len(header),
                rows=[*header, *selected],
                row_start=first,
                row_end=first + len(selected) - 1,
            )

        for index, row in enumerate(body):
            candidate = [*rows, row]
            if rows and self._text_length(section, [piece(candidate, start)]) > RECORD_TARGET_CHARS:
                parts.append(piece(rows, start))
                rows = [row]
                start = offset + index
            else:
                rows = candidate
        if rows:
            parts.append(piece(rows, start))
        return parts

    def _split_text(self, section: _Section, block: WordParagraph | WordListItem) -> list[Block]:
        budget = max(200, RECORD_TARGET_CHARS - len(section.path[-1] if section.path else "") - 8)
        pieces: list[str] = []
        current = ""
        for sentence in (item for item in _SENTENCE_END.split(block.text) if item):
            while len(sentence) > budget:
                # 文の境目の無い長い文は文字数で分ける（診断を残す）。
                if current:
                    pieces.append(current)
                    current = ""
                pieces.append(sentence[:budget])
                sentence = sentence[budget:]
                self.length_splits += 1
            if current and len(current) + len(sentence) > budget:
                pieces.append(current)
                current = sentence
            else:
                current += sentence
        if current:
            pieces.append(current)
        return [block.model_copy(update={"text": text}) for text in pieces]


def _units(blocks: list[Block]) -> list[list[Block]]:
    """分けない単位: 段落の直後の表・リストはその段落と、表の直後の注記の段落は表とまとめる。"""
    units: list[list[Block]] = []
    index = 0
    while index < len(blocks):
        block = blocks[index]
        unit: list[Block] = []
        if isinstance(block, WordParagraph):
            following = blocks[index + 1] if index + 1 < len(blocks) else None
            if following is None or isinstance(following, WordParagraph):
                units.append([block])
                index += 1
                continue
            unit.append(block)
            index += 1
        block = blocks[index]
        if isinstance(block, WordListItem):
            while index < len(blocks) and isinstance(blocks[index], WordListItem):
                unit.append(blocks[index])
                index += 1
        else:
            unit.append(block)
            index += 1
            note = blocks[index] if index < len(blocks) else None
            after = blocks[index + 1] if index + 1 < len(blocks) else None
            if isinstance(note, WordParagraph) and (
                after is None or isinstance(after, WordParagraph)
            ):
                unit.append(note)
                index += 1
        units.append(unit)
    return units
