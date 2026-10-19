"""テストの合成の Word 文書（業務の内容は書かない。#2457）。

`synthetic_docx()` は python-docx で作る .docx。`tests/data/synthetic.doc` は同じ内容を
LibreOffice で .doc（Word 97-2003）にしたもので、次で作り直せる（LibreOffice が要る）::

    uv run --locked python -m tests.fixtures tests/data
"""

from __future__ import annotations

import io
import subprocess
import sys
import tempfile
from pathlib import Path

import docx
from docx.document import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

DATA_DIR = Path(__file__).parent / "data"
SYNTHETIC_DOC = DATA_DIR / "synthetic.doc"


def save(document: Document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def list_number_id(document: Document) -> str:
    """既定のテンプレートの「List Number」の番号の定義の ID。"""
    return str(document.styles["List Number"].element.pPr.numPr.numId.val)


def set_numbering(paragraph: Paragraph, num_id: str, level: int) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), num_id)
    num_pr.append(ilvl)
    num_pr.append(num)
    ppr.append(num_pr)


def mark_header_row(row: object) -> None:
    tr_pr = row._tr.get_or_add_trPr()  # type: ignore[attr-defined]
    flag = OxmlElement("w:tblHeader")
    tr_pr.append(flag)


def synthetic_document() -> Document:
    """見出し・段落・番号のリスト・記号のリスト・表（表頭・結合セル）・注記を持つ合成の文書。"""
    document = docx.Document()
    document.add_heading("サンプル文書", 0)
    document.add_paragraph("この文書はテスト用の合成の文書です。")
    document.add_heading("1 概要", 1)
    document.add_paragraph("この節は手順の前提を説明します。")
    document.add_paragraph("次の順に操作します。")
    document.add_paragraph("画面 A を開く。", style="List Number")
    document.add_paragraph("項目 B を選ぶ。", style="List Number")
    nested = document.add_paragraph("詳細の値を確かめる。", style="List Number")
    set_numbering(nested, list_number_id(document), 1)
    document.add_paragraph("保存する。", style="List Number")
    document.add_paragraph("記号の項目", style="List Bullet")
    document.add_heading("1.1 設定の値", 2)
    document.add_paragraph("設定の値は次の表のとおりです。")
    table = document.add_table(rows=4, cols=3)
    for r, values in enumerate(
        [
            ["項目", "既定値", "範囲"],
            ["値 A", "10", "1〜100"],
            ["値 B", "有効", "有効 / 無効"],
            ["値 C", "", ""],
        ]
    ):
        for c, value in enumerate(values):
            table.cell(r, c).text = value
    mark_header_row(table.rows[0])
    # 値 C の行は「既定値」と「範囲」を横に結合する。
    merged = table.cell(3, 1).merge(table.cell(3, 2))
    merged.text = "設定しない"
    document.add_paragraph("注: 値は再起動の後に有効になります。")
    document.add_heading("2 見出しだけの節", 1)
    document.add_heading("3 改行のある段落", 1)
    paragraph = document.add_paragraph("1 行目")
    paragraph.add_run().add_break()
    paragraph.add_run("2 行目")
    return document


def synthetic_docx() -> bytes:
    return save(synthetic_document())


def build_doc_fixture(directory: Path) -> Path:
    """`synthetic_docx()` を LibreOffice で .doc にする（fixture を作り直すときだけ使う）。"""
    directory.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        source = Path(tmp) / "synthetic.docx"
        source.write_bytes(synthetic_docx())
        profile = (Path(tmp) / "profile").as_uri()
        subprocess.run(  # noqa: S603 - 固定の argv
            [
                "soffice",
                f"-env:UserInstallation={profile}",
                "--headless",
                "--convert-to",
                "doc:MS Word 97",
                "--outdir",
                str(directory),
                str(source),
            ],
            check=True,
            capture_output=True,
            timeout=120,
        )
    return directory / "synthetic.doc"


if __name__ == "__main__":
    print(build_doc_fixture(Path(sys.argv[1]) if len(sys.argv) > 1 else DATA_DIR))
