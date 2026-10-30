"""Word(.doc/.docx)→構造化レコード JSONL 変換の検証（#2457）。

テストのデータはすべて合成（業務の文書は使わない。`tests/fixtures.py`）。LibreOffice（soffice）は
CI に無いことがあるため、.doc の実変換のテストは soffice があるときだけ動かし、ほかは変換を
差し替えて決定論で確かめる（office_to_pdf のテストと同じ扱い）。.doc の形式の確認（OLE・FIB）は、
LibreOffice で作って commit した合成の .doc（`tests/data/synthetic.doc`）で確かめる。
"""

from __future__ import annotations

import io
import json
import os
import shutil
import time
import zipfile
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import docx
import olefile
import pytest
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from rag_parser_core.structured_records import (
    JSONL_RECORD_BOUNDARY,
    STRUCTURED_RECORDS_CONTENT_TYPE,
    PassageLine,
    SectionLine,
    StructuredRecordsDocument,
    WordListItem,
    WordParagraph,
    WordTable,
    parse_structured_records,
    structured_records_extraction,
    word_passage_text,
)

from app import converters
from app.converters import WordRejected, convert
from tests.fixtures import SYNTHETIC_DOC, mark_header_row, save, synthetic_docx


def _convert(data: bytes, extension: str | None = None) -> Any:
    profile = SimpleNamespace(extension=extension) if extension is not None else None
    return convert(data, "application/octet-stream", "word_to_jsonl", profile)  # type: ignore[arg-type]


def _records(data: bytes, extension: str | None = None) -> StructuredRecordsDocument:
    outcome = _convert(data, extension)
    assert outcome.converted, outcome.warnings
    assert outcome.derived_content_type == STRUCTURED_RECORDS_CONTENT_TYPE
    document = parse_structured_records(outcome.derived_bytes)
    assert isinstance(document, StructuredRecordsDocument), document
    return document


def _passages(document: StructuredRecordsDocument) -> list[PassageLine]:
    return [line for line in document.lines if isinstance(line, PassageLine)]


def _sections(document: StructuredRecordsDocument) -> list[SectionLine]:
    return [line for line in document.lines if isinstance(line, SectionLine)]


def _reason(data: bytes, extension: str | None = None) -> str:
    outcome = _convert(data, extension)
    assert outcome.converted is False
    assert outcome.derived_bytes is None
    return outcome.warnings[0]


def _fake_doc_to_docx(monkeypatch: pytest.MonkeyPatch, docx_bytes: bytes) -> None:
    monkeypatch.setattr(converters, "doc_to_docx", lambda data: docx_bytes)


# ---- JSONL の契約 ------------------------------------------------------------------------------


def test_docx_output_is_strict_jsonl_without_boundary_marker() -> None:
    outcome = _convert(synthetic_docx())
    data = outcome.derived_bytes
    assert data.endswith(b"\n") and not data.startswith(b"\xef\xbb\xbf")
    lines = data.decode("utf-8").split("\n")
    assert lines[-1] == ""
    for raw in lines[:-1]:
        payload = json.loads(raw)
        assert isinstance(payload, dict)
        assert payload["schema"] == "production-ready.structured-records"
        assert payload["schema_version"] == 1
        assert payload["source_type"] == "docx"
    # 段落の中の改行は JSON のエスケープのまま 1 物理行に収め、境界のマーカーは JSONL に置かない。
    assert "1 行目\\n2 行目" in data.decode("utf-8")
    assert JSONL_RECORD_BOUNDARY.encode() not in data
    document = parse_structured_records(data)
    assert isinstance(document, StructuredRecordsDocument)
    assert document.document.location_basis == "original"
    assert document.document.conversion is None


def test_heading_hierarchy_and_parents() -> None:
    document = _records(synthetic_docx())
    sections = {line.heading: line for line in _sections(document)}
    assert sections["サンプル文書"].level == 0
    assert sections["サンプル文書"].parent_id is None
    assert sections["1 概要"].parent_id == sections["サンプル文書"].record_id
    assert sections["1.1 設定の値"].parent_id == sections["1 概要"].record_id
    assert sections["1.1 設定の値"].heading_path == ["サンプル文書", "1 概要", "1.1 設定の値"]
    # 「2 見出しだけの節」の親は 1 階層上（1.1 ではなく表題）。
    assert sections["2 見出しだけの節"].parent_id == sections["サンプル文書"].record_id
    for passage in _passages(document):
        parent = next(line for line in _sections(document) if line.record_id == passage.parent_id)
        assert passage.section_path == parent.heading_path


def test_lists_keep_lead_in_and_markers_in_one_record() -> None:
    document = _records(synthetic_docx())
    passage = next(p for p in _passages(document) if p.section_path[-1] == "1 概要")
    items = [block for block in passage.blocks if isinstance(block, WordListItem)]
    assert [item.text for item in items] == [
        "画面 A を開く。",
        "項目 B を選ぶ。",
        "詳細の値を確かめる。",
        "保存する。",
        "記号の項目",
    ]
    assert [item.marker for item in items[:2]] == ["1.", "2."]
    assert items[2].level == 1
    assert items[3].marker == "3."
    assert items[4].ordered is False and items[4].marker is None
    text = word_passage_text(passage)
    assert text.splitlines()[0] == "1 概要"
    assert "次の順に操作します。\n1. 画面 A を開く。" in text
    assert passage.part is None


def test_table_with_lead_in_note_header_and_merged_cells() -> None:
    document = _records(synthetic_docx())
    passage = next(p for p in _passages(document) if p.section_path[-1] == "1.1 設定の値")
    kinds = [block.kind for block in passage.blocks]
    # 表の前の導入の文・表・表の後の注記が 1 つの記録（表示の行ごとには分けない）。
    assert kinds == ["paragraph", "table", "paragraph"]
    table = passage.blocks[1]
    assert isinstance(table, WordTable)
    assert table.header_rows == 1
    assert table.table_index == 1
    assert (table.row_start, table.row_end) == (1, 4)
    # 横の結合は同じ値で埋める。
    assert table.rows[3] == ["値 C", "設定しない", "設定しない"]
    assert passage.location.table_indices == [1]
    assert passage.location.body_start < table.body_index < passage.location.body_end
    text = word_passage_text(passage)
    assert "| 項目 | 既定値 | 範囲 |\n| --- | --- | --- |\n| 値 A | 10 | 1〜100 |" in text
    assert text.endswith("注: 値は再起動の後に有効になります。")


def test_vertical_merge_continuation_carries_the_merged_value() -> None:
    """縦の結合の続きのセルは結合の始まりの値で埋め、行ごとに区分を持つ（#3565）。"""
    document = docx.Document()
    document.add_heading("表の節", 1)
    table = document.add_table(rows=5, cols=4)
    for r, values in enumerate(
        [
            ["分類", "名称", "番号", "設定"],
            ["", "名称 k1", "101", "有効"],
            ["", "名称 k2", "102", "無効"],
            ["分類 L", "名称 l1", "201", "有効"],
            ["", "名称 l2", "202", "有効"],
        ]
    ):
        for c, value in enumerate(values):
            table.cell(r, c).text = value
    mark_header_row(table.rows[0])
    table.cell(1, 0).merge(table.cell(2, 0)).text = "分類 K"
    # 2 列 × 2 行の結合（横と縦の両方）も、続きの行の各列に値を写す。
    table.cell(3, 2).merge(table.cell(4, 3)).text = "共通"
    passage = _passages(_records(save(document)))[0]
    block = passage.blocks[0]
    assert isinstance(block, WordTable)
    assert block.rows[1] == ["分類 K", "名称 k1", "101", "有効"]
    assert block.rows[2] == ["分類 K", "名称 k2", "102", "無効"]
    assert block.rows[3] == ["分類 L", "名称 l1", "共通", "共通"]
    assert block.rows[4] == ["", "名称 l2", "共通", "共通"]
    text = word_passage_text(passage)
    assert "| 分類 K | 名称 k2 | 102 | 無効 |" in text
    # 前処理の版を上げた（既存の文書は取込み直すまで前の版の本文のまま）。
    assert _records(save(document)).document.converter_version == "v2"


def test_vertical_merge_in_nested_table_and_split_table_parts() -> None:
    """表の中の表と、行で分けた大きい表の続きの部分でも、縦の結合の値を写す。"""
    document = docx.Document()
    document.add_heading("表の節", 1)
    table = document.add_table(rows=81, cols=2)
    table.cell(0, 0).text = "区分"
    table.cell(0, 1).text = "説明"
    mark_header_row(table.rows[0])
    for row in range(1, 81):
        table.cell(row, 1).text = f"{row:03d} " + "説明の文。" * 6
    table.cell(1, 0).merge(table.cell(80, 0)).text = "区分 A"
    tables = [
        b for p in _passages(_records(save(document))) for b in p.blocks if isinstance(b, WordTable)
    ]
    assert len(tables) > 1
    # 続きの部分の行も、上の部分で始まった結合の値を持つ。
    assert all(row[0] == "区分 A" for t in tables for row in t.rows[1:])

    nested_document = docx.Document()
    nested_document.add_heading("表の節", 1)
    outer = nested_document.add_table(rows=1, cols=1)
    inner = outer.cell(0, 0).add_table(rows=2, cols=2)
    inner.cell(0, 0).merge(inner.cell(1, 0)).text = "区分 B"
    inner.cell(0, 1).text = "1"
    inner.cell(1, 1).text = "2"
    block = _passages(_records(save(nested_document)))[0].blocks[0]
    assert isinstance(block, WordTable)
    assert block.rows[0] == ["区分 B | 1 / 区分 B | 2"]


def test_heading_only_section_and_content_before_first_heading() -> None:
    document = docx.Document()
    document.add_paragraph("見出しの前の本文")
    document.add_heading("A", 1)
    document.add_heading("A-1", 2)
    document.add_paragraph("本文")
    document.add_heading("B", 1)
    records = _records(save(document))
    passages = _passages(records)
    assert passages[0].parent_id is None and passages[0].section_path == []
    assert word_passage_text(passages[0]) == "見出しの前の本文"
    # 「A」は下位の見出しを持つので見出しだけの記録を作らず、「B」は見出しだけの記録にする。
    assert [p.section_path for p in passages[1:]] == [["A", "A-1"], ["B"]]
    assert passages[2].blocks == []
    assert word_passage_text(passages[2]) == "B"


def test_positions_count_empty_paragraphs_and_do_not_invent_pages() -> None:
    document = docx.Document()
    document.add_heading("節", 1)
    document.add_paragraph("")
    document.add_paragraph("3 番目の段落")
    records = _records(save(document))
    passage = _passages(records)[0]
    block = passage.blocks[0]
    assert isinstance(block, WordParagraph)
    assert (block.body_index, block.paragraph_index) == (3, 3)
    raw = save(document)
    assert b"page" not in _convert(raw).derived_bytes


def test_outline_level_and_numbered_heading() -> None:
    document = docx.Document()
    heading = document.add_paragraph("アウトラインの見出し")
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), "1")
    heading._p.get_or_add_pPr().append(outline)
    document.add_paragraph("本文")
    sections = _sections(_records(save(document)))
    assert [(s.heading, s.level) for s in sections] == [("アウトラインの見出し", 2)]


def test_extraction_elements_carry_heading_path_and_positions() -> None:
    document = _records(synthetic_docx())
    extraction = structured_records_extraction(
        document,
        source_parser="structured_records_jsonl",
        parser_backend="structured_records_jsonl",
        parser_version="v1",
    )
    assert extraction.document_type == "Word"
    assert extraction.parser_artifacts["source_format"] == "docx"
    assert extraction.parser_artifacts["structured_records_location_basis"] == "original"
    element = next(e for e in extraction.elements if e.section_path[-1:] == ["1.1 設定の値"])
    assert element.content_kind == "text"
    assert element.parent_id is not None
    assert element.metadata["structured_record_id"] == element.element_id
    assert element.metadata["word_table_indices"] == "1"
    assert element.metadata["word_block_kinds"] == "paragraph,table"
    assert element.metadata["section_level"] == 3
    assert "page_number" not in element.model_dump(exclude_none=True)


# ---- 大きい節を分ける -----------------------------------------------------------------------


def test_large_section_is_split_with_part_ids_and_parent_location() -> None:
    document = docx.Document()
    document.add_heading("大きい節", 1)
    for index in range(12):
        document.add_paragraph(f"段落 {index:02d}。" + "本文の文。" * 40)
    records = _records(save(document))
    passages = _passages(records)
    assert len(passages) > 1
    group = {p.part.group_id for p in passages if p.part is not None}
    assert len(group) == 1
    assert [p.part.index for p in passages if p.part] == list(range(1, len(passages) + 1))
    assert all(p.part is not None and p.part.count == len(passages) for p in passages)
    whole = passages[0].part.group_location  # type: ignore[union-attr]
    assert (whole.body_start, whole.body_end) == (2, 13)
    for passage in passages:
        assert len(word_passage_text(passage)) <= converters.RECORD_TARGET_CHARS
        # 各部分も見出しを文脈として先頭に持ち、親（見出しの section）を指す。
        assert word_passage_text(passage).startswith("大きい節\n")
        assert passage.parent_id == _sections(records)[0].record_id
    # 段落は途中で分けず、元の順のまま並ぶ。
    texts = [b.text for p in passages for b in p.blocks if isinstance(b, WordParagraph)]
    assert [text[:5] for text in texts] == [f"段落 {i:02d}" for i in range(12)]


def test_large_table_is_split_by_rows_repeating_header() -> None:
    document = docx.Document()
    document.add_heading("大きい表", 1)
    document.add_paragraph("表の導入")
    table = document.add_table(rows=81, cols=2)
    table.cell(0, 0).text = "番号"
    table.cell(0, 1).text = "説明"
    for row in range(1, 81):
        table.cell(row, 0).text = f"{row:03d}"
        table.cell(row, 1).text = "説明の文。" * 6
    passages = _passages(_records(save(document)))
    tables = [b for p in passages for b in p.blocks if isinstance(b, WordTable)]
    assert len(tables) > 1
    assert all(t.rows[0] == ["番号", "説明"] for t in tables)
    assert tables[0].row_start == 2
    for before, after in zip(tables, tables[1:], strict=False):
        assert after.row_start == before.row_end + 1
    assert tables[-1].row_end == 81
    assert isinstance(passages[0].blocks[0], WordParagraph)
    assert passages[0].blocks[0].text == "表の導入"


def test_long_sentence_is_split_by_length_with_diagnostic() -> None:
    document = docx.Document()
    document.add_heading("長い文", 1)
    document.add_paragraph("あ" * 4000)
    outcome = _convert(save(document))
    assert outcome.converted
    assert any(w.startswith("word_paragraph_split_by_length:") for w in outcome.warnings)
    records = parse_structured_records(outcome.derived_bytes)
    assert isinstance(records, StructuredRecordsDocument)
    texts = [b.text for p in _passages(records) for b in p.blocks if isinstance(b, WordParagraph)]
    assert "".join(texts) == "あ" * 4000


# ---- 診断（読まない部分）----------------------------------------------------------------------


def test_tracked_changes_keep_insertions_and_drop_deletions() -> None:
    document = docx.Document()
    document.add_heading("変更履歴", 1)
    paragraph = document.add_paragraph("残す")
    inserted = OxmlElement("w:ins")
    inserted.set(qn("w:id"), "1")
    inserted.set(qn("w:author"), "test")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "挿入"
    run.append(text)
    inserted.append(run)
    deleted = OxmlElement("w:del")
    deleted.set(qn("w:id"), "2")
    deleted.set(qn("w:author"), "test")
    drun = OxmlElement("w:r")
    dtext = OxmlElement("w:delText")
    dtext.text = "削除"
    drun.append(dtext)
    deleted.append(drun)
    paragraph._p.append(inserted)
    paragraph._p.append(deleted)
    outcome = _convert(save(document))
    assert "word_tracked_changes_accepted" in outcome.warnings
    records = parse_structured_records(outcome.derived_bytes)
    assert isinstance(records, StructuredRecordsDocument)
    assert word_passage_text(_passages(records)[0]) == "変更履歴\n残す挿入"
    assert [d.code for d in records.document.diagnostics] == ["tracked_changes_accepted"]


def _with_part(data: bytes, name: str, content: bytes, content_type: str, rel_type: str) -> bytes:
    """.docx に部品（脚注など）を足す。"""
    source = zipfile.ZipFile(io.BytesIO(data))
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            payload = source.read(info.filename)
            if info.filename == "[Content_Types].xml":
                payload = payload.replace(
                    b"</Types>",
                    f'<Override PartName="/{name}" ContentType="{content_type}"/></Types>'.encode(),
                )
            if info.filename == "word/_rels/document.xml.rels":
                target_name = name.split("/", 1)[1]
                payload = payload.replace(
                    b"</Relationships>",
                    f'<Relationship Id="rIdTest" Type="{rel_type}" Target="{target_name}"/>'
                    "</Relationships>".encode(),
                )
            target.writestr(info, payload)
        target.writestr(name, content)
    return buffer.getvalue()


def test_footnotes_and_images_are_reported_not_extracted() -> None:
    document = docx.Document()
    document.add_heading("注と画像", 1)
    paragraph = document.add_paragraph("本文")
    drawing = OxmlElement("w:drawing")
    run = OxmlElement("w:r")
    run.append(drawing)
    paragraph._p.append(run)
    w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    footnotes = (
        f'<w:footnotes xmlns:w="{w}">'
        '<w:footnote w:id="-1"><w:p><w:r><w:t>区切り</w:t></w:r></w:p></w:footnote>'
        '<w:footnote w:id="1"><w:p><w:r><w:t>脚注の文</w:t></w:r></w:p></w:footnote>'
        "</w:footnotes>"
    ).encode()
    data = _with_part(
        save(document),
        "word/footnotes.xml",
        footnotes,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes",
    )
    outcome = _convert(data)
    assert outcome.converted
    assert "word_images_not_extracted:1" in outcome.warnings
    assert "word_footnotes_not_extracted:1" in outcome.warnings
    assert "脚注の文".encode() not in outcome.derived_bytes


# ---- .doc（原本と派生物）--------------------------------------------------------------------


def test_doc_is_converted_and_records_original_and_derived(monkeypatch: pytest.MonkeyPatch) -> None:
    derived = synthetic_docx()
    _fake_doc_to_docx(monkeypatch, derived)
    original = SYNTHETIC_DOC.read_bytes()
    records = _records(original, extension=".doc")
    assert records.source_type == "doc"
    assert records.document.location_basis == "derived"
    conversion = records.document.conversion
    assert conversion is not None
    assert conversion.original_type == "doc" and conversion.derived_type == "docx"
    import hashlib

    assert conversion.original_sha256 == hashlib.sha256(original).hexdigest()
    assert conversion.derived_sha256 == hashlib.sha256(derived).hexdigest()
    assert conversion.converter == "libreoffice"
    assert all(line.source_type == "doc" for line in records.lines)
    extraction = structured_records_extraction(
        records,
        source_parser="structured_records_jsonl",
        parser_backend="structured_records_jsonl",
        parser_version="v1",
    )
    assert extraction.parser_artifacts["structured_records_location_basis"] == "derived"
    assert extraction.parser_artifacts["structured_records_source_conversion"]["original_type"] == (
        "doc"
    )
    assert {e.metadata["word_location_basis"] for e in extraction.elements} == {"derived"}


@pytest.mark.skipif(
    shutil.which("soffice") is None, reason="LibreOffice が無い（CI は差し替えで確かめる）"
)
def test_doc_real_libreoffice_conversion_keeps_structure() -> None:
    records = _records(SYNTHETIC_DOC.read_bytes(), extension=".doc")
    assert records.document.location_basis == "derived"
    headings = [line.heading for line in _sections(records)]
    assert headings == [
        "サンプル文書",
        "1 概要",
        "1.1 設定の値",
        "2 見出しだけの節",
        "3 改行のある段落",
    ]
    table_passage = next(p for p in _passages(records) if p.section_path[-1] == "1.1 設定の値")
    assert [b.kind for b in table_passage.blocks] == ["paragraph", "table", "paragraph"]


def test_doc_without_libreoffice_is_rejected(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(converters, "soffice_path", lambda: None)
    assert _reason(SYNTHETIC_DOC.read_bytes(), ".doc") == "word_libreoffice_unavailable"


def test_docx_without_libreoffice_still_converts(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(converters, "soffice_path", lambda: None)
    assert _convert(synthetic_docx()).converted


def _with_fib_flags(data: bytes, mask: int) -> bytes:
    """合成の .doc の WordDocument の stream の FIB の flags に bit を立てる。"""
    buffer = io.BytesIO(data)
    with olefile.OleFileIO(buffer, write_mode=True) as ole:
        stream = bytearray(ole.openstream("WordDocument").read())
        flags = int.from_bytes(stream[10:12], "little") | mask
        stream[10:12] = flags.to_bytes(2, "little")
        ole.write_stream("WordDocument", bytes(stream))
    return buffer.getvalue()


def test_encrypted_doc_is_rejected() -> None:
    encrypted = _with_fib_flags(SYNTHETIC_DOC.read_bytes(), 0x0100)
    assert _reason(encrypted, ".doc") == "word_encrypted"


def test_encrypted_ooxml_package_is_rejected(monkeypatch: pytest.MonkeyPatch) -> None:
    class _Encrypted:
        def __enter__(self) -> _Encrypted:
            return self

        def __exit__(self, *args: object) -> None:
            return None

        def exists(self, name: str) -> bool:
            return name in {"EncryptedPackage", "EncryptionInfo"}

    monkeypatch.setattr(olefile, "OleFileIO", lambda *args, **kwargs: _Encrypted())
    # 暗号化した .docx は OLE の複合文書になる（拡張子に関係なく先頭のバイトで判定する）。
    assert _reason(SYNTHETIC_DOC.read_bytes(), ".docx") == "word_encrypted"


def test_ole_that_is_not_word_is_rejected() -> None:
    data = SYNTHETIC_DOC.read_bytes()
    name = "WordDocument".encode("utf-16-le")
    assert name in data
    assert _reason(data.replace(name, "WordDocumenX".encode("utf-16-le")), ".doc") == (
        "word_not_word_document"
    )


def test_corrupt_doc_is_rejected() -> None:
    assert _reason(SYNTHETIC_DOC.read_bytes()[:1024], ".doc") == "word_corrupt"


def test_doc_with_bad_fib_is_rejected() -> None:
    data = bytearray(SYNTHETIC_DOC.read_bytes())
    buffer = io.BytesIO(bytes(data))
    with olefile.OleFileIO(buffer, write_mode=True) as ole:
        stream = bytearray(ole.openstream("WordDocument").read())
        stream[0:2] = b"\x00\x00"
        ole.write_stream("WordDocument", bytes(stream))
    assert _reason(buffer.getvalue(), ".doc") == "word_corrupt"


# ---- 拒否（成功扱いにしない）------------------------------------------------------------------


def test_empty_and_unsupported_inputs_are_rejected() -> None:
    assert _reason(b"") == "word_empty"
    assert _reason(b"plain text") == "word_unsupported_format"


def test_too_large_input_is_rejected(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(converters, "MAX_INPUT_BYTES", 10)
    assert _reason(synthetic_docx()) == "word_input_too_large"


def _replace_member(
    data: bytes, name: str, transform: Any, extra: dict[str, bytes] | None = None
) -> bytes:
    source = zipfile.ZipFile(io.BytesIO(data))
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            payload = source.read(info.filename)
            target.writestr(info, transform(payload) if info.filename == name else payload)
        for extra_name, payload in (extra or {}).items():
            target.writestr(extra_name, payload)
    return buffer.getvalue()


def _macro_enabled_docx() -> bytes:
    """合成の .docx を、マクロ付き（.docm）の形にする（本体の content type・VBA の部品と関係）。"""
    macro_type = b"application/vnd.ms-word.document.macroEnabled.main+xml"
    vba_type = b"application/vnd.ms-office.vbaProject"
    data = _replace_member(
        synthetic_docx(),
        "[Content_Types].xml",
        lambda payload: payload.replace(
            b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
            macro_type,
        ).replace(
            b"</Types>", b'<Default Extension="bin" ContentType="' + vba_type + b'"/></Types>'
        ),
        {"word/vbaProject.bin": b"synthetic-vba"},
    )
    return _replace_member(
        data,
        "word/_rels/document.xml.rels",
        lambda payload: payload.replace(
            b"</Relationships>",
            b'<Relationship Id="rIdVba" Type="http://schemas.microsoft.com/office/2006/'
            b'relationships/vbaProject" Target="vbaProject.bin"/></Relationships>',
        ),
    )


def test_macro_enabled_word_is_read_without_macros() -> None:
    """.docm は VBA の部品を捨てて本文だけを読む（マクロは実行しない。#2457 の判断 5 の変更）。"""
    data = _macro_enabled_docx()
    assert converters.inspect_docx_package(data) == "macro"
    outcome = _convert(data, ".docm")
    assert outcome.converted, outcome.warnings
    assert "word_macros_not_executed" in outcome.warnings
    assert b"synthetic-vba" not in outcome.derived_bytes
    records = parse_structured_records(outcome.derived_bytes)
    assert isinstance(records, StructuredRecordsDocument)
    assert records.source_type == "docx"
    assert [d.code for d in records.document.diagnostics] == ["macros_not_executed"]
    assert [line.heading for line in _sections(records)][:2] == ["サンプル文書", "1 概要"]
    stripped = zipfile.ZipFile(io.BytesIO(converters.without_macros(data)))
    assert not any("vba" in name.casefold() for name in stripped.namelist())
    assert b"vbaProject" not in stripped.read("word/_rels/document.xml.rels")
    assert b"macroEnabled" not in stripped.read("[Content_Types].xml")


def test_vba_part_alone_marks_package_as_macro() -> None:
    with_vba = _replace_member(synthetic_docx(), "", lambda p: p, {"word/vbaProject.bin": b"x"})
    assert converters.inspect_docx_package(with_vba) == "macro"
    assert _convert(with_vba).converted


def test_template_is_read_as_document() -> None:
    template = _replace_member(
        synthetic_docx(),
        "[Content_Types].xml",
        lambda payload: payload.replace(
            b"wordprocessingml.document.main", b"wordprocessingml.template.main"
        ),
    )
    assert converters.inspect_docx_package(template) == "template"
    outcome = _convert(template, ".dotx")
    assert outcome.converted
    assert "word_macros_not_executed" not in outcome.warnings


def test_corrupt_and_non_word_packages_are_rejected() -> None:
    assert _reason(synthetic_docx()[:200]) == "word_corrupt"
    no_main = _replace_member(
        synthetic_docx(),
        "[Content_Types].xml",
        lambda payload: payload.replace(b"wordprocessingml.document.main", b"spreadsheetml.x"),
    )
    assert _reason(no_main) == "word_not_word_document"


def test_document_without_content_is_rejected() -> None:
    assert _reason(save(docx.Document())) == "word_no_content"


# ---- LibreOffice の隔離（時間・後始末・出力）--------------------------------------------------


def test_run_isolated_times_out_and_kills_process_group(tmp_path: Path) -> None:
    marker = tmp_path / "child.pid"
    command = ["/bin/sh", "-c", f"sleep 30 & echo $! > {marker}; wait"]
    started = time.monotonic()
    with pytest.raises(WordRejected) as raised:
        converters.run_isolated(command, env=dict(os.environ), cwd=tmp_path, timeout=0.5)
    assert raised.value.reason == "word_conversion_timeout"
    assert time.monotonic() - started < 10
    child = int(marker.read_text().strip())
    # 子プロセス（LibreOffice の soffice.bin にあたる）も残さない。
    deadline = time.monotonic() + 5
    while time.monotonic() < deadline:
        try:
            os.kill(child, 0)
        except ProcessLookupError:
            break
        time.sleep(0.05)
    else:
        pytest.fail("子プロセスが残った")


def test_run_isolated_reports_failure(tmp_path: Path) -> None:
    with pytest.raises(WordRejected) as raised:
        converters.run_isolated(["/bin/false"], env=dict(os.environ), cwd=tmp_path, timeout=5)
    assert raised.value.reason == "word_conversion_failed"


def _fake_soffice(monkeypatch: pytest.MonkeyPatch, write: bytes | None) -> list[list[str]]:
    calls: list[list[str]] = []

    def fake_run(command: list[str], *, env: dict[str, str], cwd: Path, timeout: float) -> None:
        calls.append(command)
        profile = Path(command[1].split("=", 1)[1].removeprefix("file://"))
        registry = (profile / "user" / "registrymodifications.xcu").read_text(encoding="utf-8")
        assert "DisableMacrosExecution" in registry and "<value>true</value>" in registry
        assert env["HOME"] != os.environ.get("HOME")
        if write is not None:
            out = Path(command[command.index("--outdir") + 1])
            (out / "input.docx").write_bytes(write)

    monkeypatch.setattr(converters, "soffice_path", lambda: "/usr/bin/soffice")
    monkeypatch.setattr(converters, "run_isolated", fake_run)
    return calls


def test_doc_conversion_without_output_is_rejected(monkeypatch: pytest.MonkeyPatch) -> None:
    calls = _fake_soffice(monkeypatch, None)
    assert _reason(SYNTHETIC_DOC.read_bytes(), ".doc") == "word_conversion_no_output"
    assert len(calls) == 1
    # 変換のたびに一時ディレクトリ（プロファイル）を作り、終わったら消す。
    profile = Path(calls[0][1].split("=", 1)[1].removeprefix("file://"))
    assert not profile.exists()


def test_doc_conversion_with_empty_output_is_rejected(monkeypatch: pytest.MonkeyPatch) -> None:
    _fake_soffice(monkeypatch, b"")
    assert _reason(SYNTHETIC_DOC.read_bytes(), ".doc") == "word_conversion_empty"


def test_doc_conversion_with_broken_output_is_rejected(monkeypatch: pytest.MonkeyPatch) -> None:
    _fake_soffice(monkeypatch, b"PK\x03\x04broken")
    assert _reason(SYNTHETIC_DOC.read_bytes(), ".doc") == "word_corrupt"


def test_doc_conversion_with_fake_output_is_converted(monkeypatch: pytest.MonkeyPatch) -> None:
    calls = _fake_soffice(monkeypatch, synthetic_docx())
    records = _records(SYNTHETIC_DOC.read_bytes(), ".doc")
    assert records.document.location_basis == "derived"
    command = calls[0]
    assert "--headless" in command and "--norestore" in command
    assert command[command.index("--convert-to") + 1].startswith("docx")


def test_other_profile_passes_through() -> None:
    outcome = convert(synthetic_docx(), "application/octet-stream", "excel_to_jsonl", None)
    assert outcome.converted is False
    assert any("preprocess_unsupported_profile" in w for w in outcome.warnings)


def test_conversion_is_deterministic() -> None:
    data = synthetic_docx()
    assert _convert(data).derived_bytes == _convert(data).derived_bytes


def test_health_reports_libreoffice(monkeypatch: pytest.MonkeyPatch) -> None:
    from app import main

    monkeypatch.setattr(main, "soffice_path", lambda: None)
    health = main._health()
    assert health.status == "degraded"
    assert health.supported_profiles == ["word_to_jsonl"]
    monkeypatch.setattr(main, "soffice_path", lambda: "/usr/bin/soffice")
    assert main._health().status == "ok"
