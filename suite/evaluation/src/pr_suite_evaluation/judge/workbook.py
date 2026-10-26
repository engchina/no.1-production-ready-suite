"""人手の判定の記入用の Excel のブック（.xlsx）の書き出しと読み込み（#3418）。

ブックは人が読む・書くためのもので、機械の正本は項目のファイル（`*.items.json`）。ブックの `_meta`
の sheet に項目のファイルの sha256 を書き、読み込みで照らす（違う項目のブックを混ぜない）。

- `ラベル`: 1 行 = 項目 × 観点。記入者は判定・誤りの型・メモ・記入者の列だけを書く。
- `誤りの型と対照`: #2577 の誤りの型と対照の一覧・holdout の目標の件数・記入済みの件数（式）。
- `判定の基準`: rubric（版・sha256・観点の基準・付け方）。
- `合成の対照`: 故意に壊した回答（合成。人は判定しない。誤りにならない壊し方なら除外の列に x）。
"""

from __future__ import annotations

import re
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

LABEL_SHEET = "ラベル"
TARGET_SHEET = "誤りの型と対照"
RUBRIC_SHEET = "判定の基準"
SYNTHETIC_SHEET = "合成の対照"
META_SHEET = "_meta"
WORKBOOK_SCHEMA_VERSION = 1
# Excel の 1 つのセルの文字数の上限（32767）より少し小さくする。
MAX_CELL_CHARS = 32_000
TRUNCATED_MARK = "\n…（以下省略。全文は項目のファイル）"
LABEL_CHOICES = ("pass", "fail", "cannot_judge")
EXCLUDE_MARK = "x"
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_HEADER_FILL = PatternFill("solid", fgColor="DDE5F0")
_INPUT_FILL = PatternFill("solid", fgColor="FFF7D6")
_WRAP = Alignment(wrap_text=True, vertical="top")


@dataclass(frozen=True)
class Column:
    key: str
    header: str
    width: float
    # 記入者が書く列（色を付ける）。
    editable: bool = False


LABEL_COLUMNS: tuple[Column, ...] = (
    Column("item_id", "項目の ID", 22),
    Column("aspect", "観点", 20),
    Column("split", "区分（dev / holdout）", 12),
    Column("question", "質問", 40),
    Column("answer", "回答の本文", 60),
    Column("evidence", "根拠の抜粋", 60),
    Column("standard_answer", "標準回答", 40),
    Column("rubric", "観点の基準", 40),
    Column("label", "判定（pass / fail / cannot_judge）", 16, editable=True),
    Column("error_type", "誤りの型・対照（任意）", 22, editable=True),
    Column("note", "メモ（判断の理由）", 30, editable=True),
    Column("annotator", "記入者（匿名の ID）", 14, editable=True),
)
SYNTHETIC_COLUMNS: tuple[Column, ...] = (
    Column("item_id", "合成の項目の ID", 28),
    Column("base_item_id", "元の項目の ID", 22),
    Column("kind", "壊し方", 20),
    Column("fail_aspects", "fail にした観点", 24),
    Column("description", "変えた点", 40),
    Column("answer", "回答（壊した後）", 60),
    Column("exclude", "除外（誤りにならないなら x）", 14, editable=True),
)
HEADER_TO_KEY = {column.header: column.key for column in (*LABEL_COLUMNS, *SYNTHETIC_COLUMNS)}


class WorkbookError(ValueError):
    """記入用のブックを読めない・形が違う（利用者へ返す失敗）。"""


@dataclass
class WorkbookContent:
    """読んだブック（ラベルの行・合成の除外・メタ）。"""

    meta: dict[str, str]
    rows: list[dict[str, str]]
    excluded_synthetic: set[str] = field(default_factory=set)


def cell_text(value: str | None, limit: int = MAX_CELL_CHARS) -> str:
    """セルに書く文字（制御文字を除き、上限を超えたら省略の印を付ける）。"""
    text = _ILLEGAL.sub("", value or "")
    if len(text) > limit:
        return text[: limit - len(TRUNCATED_MARK)] + TRUNCATED_MARK
    return text


def _text(sheet: Worksheet, row: int, column: int, value: str) -> Any:
    cell = sheet.cell(row=row, column=column)
    cell.value = cell_text(value)
    # 「=」で始まる回答・質問を式にしない（文字として書く。式は数える列だけに置く）。
    cell.data_type = "s"
    return cell


def _put(sheet: Worksheet, row: int, column: int, value: object) -> Any:
    if isinstance(value, str):
        return _text(sheet, row, column, value)
    cell = sheet.cell(row=row, column=column)
    cell.value = value if isinstance(value, int | float) else str(value)
    return cell


def _header(sheet: Worksheet, columns: Sequence[Column]) -> None:
    for index, column in enumerate(columns, start=1):
        cell = _text(sheet, 1, index, column.header)
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
        cell.alignment = _WRAP
        sheet.column_dimensions[get_column_letter(index)].width = column.width
    sheet.freeze_panes = "C2"


def column_letter(columns: Sequence[Column], key: str) -> str:
    for index, column in enumerate(columns, start=1):
        if column.key == key:
            return get_column_letter(index)
    raise KeyError(key)


def write_workbook(
    path: Path,
    *,
    meta: Mapping[str, str],
    label_rows: Sequence[Mapping[str, str]],
    extra_columns: Sequence[str] = (),
    target_rows: Sequence[Sequence[object]],
    aspect_ids: Sequence[str],
    rubric_rows: Sequence[Sequence[str]],
    synthetic_rows: Sequence[Mapping[str, str]] = (),
    error_type_ids: Sequence[str] = (),
) -> None:
    """記入用（と裁定用）のブックを書く。`extra_columns` は判定の列の前に足す参考の列の見出し。"""
    book = Workbook()
    labels = book.active
    if labels is None:  # pragma: no cover - Workbook() は必ず sheet を 1 つ持つ
        raise WorkbookError("ブックを作れません。")
    labels.title = LABEL_SHEET
    insert_at = next(i for i, column in enumerate(LABEL_COLUMNS) if column.key == "label")
    columns = [
        *LABEL_COLUMNS[:insert_at],
        *(Column(f"extra:{header}", header, 16) for header in extra_columns),
        *LABEL_COLUMNS[insert_at:],
    ]
    _header(labels, columns)
    for row_index, row in enumerate(label_rows, start=2):
        for column_index, column in enumerate(columns, start=1):
            cell = _put(labels, row_index, column_index, row.get(column.key, ""))
            cell.alignment = _WRAP
            if column.editable:
                cell.fill = _INPUT_FILL
        labels.row_dimensions[row_index].height = 150
    last = max(len(label_rows) + 1, 2)
    label_letter = column_letter(columns, "label")
    choices = DataValidation(type="list", formula1=f'"{",".join(LABEL_CHOICES)}"', allow_blank=True)
    choices.add(f"{label_letter}2:{label_letter}{last}")
    labels.add_data_validation(choices)
    if error_type_ids:
        error_letter = column_letter(columns, "error_type")
        types = DataValidation(
            type="list", formula1=f'"{",".join(error_type_ids)}"', allow_blank=True
        )
        types.add(f"{error_letter}2:{error_letter}{last}")
        labels.add_data_validation(types)

    targets = book.create_sheet(TARGET_SHEET)
    split_letter = column_letter(columns, "split")
    error_letter = column_letter(columns, "error_type")
    aspect_letter = column_letter(columns, "aspect")
    sheet_ref = f"'{LABEL_SHEET}'"
    headers = ("区分", "ID", "内容", "主に fail になる観点", "holdout の目標（以上）", "記入済み")
    for index, header in enumerate(headers, start=1):
        cell = _text(targets, 1, index, header)
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
    for row_index, values in enumerate(target_rows, start=2):
        for column_index, value in enumerate(values, start=1):
            _put(targets, row_index, column_index, value).alignment = _WRAP
        # 記入済みの行（holdout。項目 × 観点の行の数）。
        targets.cell(row=row_index, column=len(headers)).value = (
            f"=COUNTIFS({sheet_ref}!${error_letter}:${error_letter},B{row_index},"
            f'{sheet_ref}!${split_letter}:${split_letter},"holdout")'
        )
    start = len(target_rows) + 3
    aspect_headers = ("観点", "holdout の行", "pass", "fail", "cannot_judge", "未記入")
    for index, header in enumerate(aspect_headers, start=1):
        cell = _text(targets, start, index, header)
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
    for offset, aspect_id in enumerate(aspect_ids, start=1):
        line = start + offset
        _text(targets, line, 1, aspect_id)
        base = (
            f"{sheet_ref}!${aspect_letter}:${aspect_letter},A{line},"
            f'{sheet_ref}!${split_letter}:${split_letter},"holdout"'
        )
        targets.cell(row=line, column=2).value = f"=COUNTIFS({base})"
        label_range = f"{sheet_ref}!${label_letter}:${label_letter}"
        for position, choice in enumerate(LABEL_CHOICES, start=3):
            targets.cell(
                row=line, column=position
            ).value = f'=COUNTIFS({base},{label_range},"{choice}")'
        targets.cell(row=line, column=6).value = f"=B{line}-C{line}-D{line}-E{line}"
    for letter, width in zip("ABCDEF", (12, 24, 50, 30, 16, 12), strict=True):
        targets.column_dimensions[letter].width = width

    rubric = book.create_sheet(RUBRIC_SHEET)
    for row_index, values in enumerate(rubric_rows, start=1):
        for column_index, value in enumerate(values, start=1):
            _text(rubric, row_index, column_index, value).alignment = _WRAP
    rubric.column_dimensions["A"].width = 24
    rubric.column_dimensions["B"].width = 100

    if synthetic_rows:
        synthetic = book.create_sheet(SYNTHETIC_SHEET)
        _header(synthetic, SYNTHETIC_COLUMNS)
        for row_index, row in enumerate(synthetic_rows, start=2):
            for column_index, column in enumerate(SYNTHETIC_COLUMNS, start=1):
                cell = _put(synthetic, row_index, column_index, row.get(column.key, ""))
                cell.alignment = _WRAP
                if column.editable:
                    cell.fill = _INPUT_FILL

    hidden = book.create_sheet(META_SHEET)
    for row_index, (key, value) in enumerate(meta.items(), start=1):
        _text(hidden, row_index, 1, key)
        _text(hidden, row_index, 2, value)
    hidden.sheet_state = "hidden"
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)


def _value(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _sheet_rows(sheet: Any, name: str, path: Path) -> list[dict[str, str]]:
    rows = sheet.iter_rows(values_only=True)
    try:
        header = [_value(item) for item in next(rows)]
    except StopIteration as error:
        raise WorkbookError(f"{path.name} の {name} が空です。") from error
    keys = [HEADER_TO_KEY.get(item) for item in header]
    result: list[dict[str, str]] = []
    for values in rows:
        row = {key: _value(value) for key, value in zip(keys, values, strict=False) if key}
        if any(row.values()):
            result.append(row)
    return result


def read_workbook(path: Path) -> WorkbookContent:
    """記入した（か裁定した）ブックを読む。"""
    try:
        book = load_workbook(path, read_only=True, data_only=True)
    except (OSError, ValueError, KeyError) as error:
        raise WorkbookError(f"ブックを読めません: {path.name}") from error
    try:
        names = set(book.sheetnames)
        if LABEL_SHEET not in names or META_SHEET not in names:
            raise WorkbookError(
                f"記入用のブックの形ではありません（{LABEL_SHEET}・{META_SHEET} の sheet が無い）: "
                f"{path.name}"
            )
        meta = {
            _value(key): _value(value)
            for key, value, *_ in book[META_SHEET].iter_rows(values_only=True)
            if key is not None
        }
        rows = _sheet_rows(book[LABEL_SHEET], LABEL_SHEET, path)
        excluded: set[str] = set()
        if SYNTHETIC_SHEET in names:
            for row in _sheet_rows(book[SYNTHETIC_SHEET], SYNTHETIC_SHEET, path):
                if row.get("exclude", "").lower() == EXCLUDE_MARK and row.get("item_id"):
                    excluded.add(row["item_id"])
        return WorkbookContent(meta=meta, rows=rows, excluded_synthetic=excluded)
    finally:
        book.close()
