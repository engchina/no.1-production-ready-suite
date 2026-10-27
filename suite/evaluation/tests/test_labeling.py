"""人手の判定の記入用のブックの書き出し・取り込みと、故意に壊した回答（#3418・#2577）。

書き出し → 記入（テストの中だけの合成の記入。人の判定ではない）→ 取り込み（記入者の間の κ と
食い違い）→ 裁定 → LabelSet → calibrate の通しを確かめる。判定のモデルは台本で、実サービスは
呼ばない。
"""

from __future__ import annotations

import json
from collections.abc import Callable, Mapping
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook
from pr_evaluation_scoring.perturbation import (
    PERTURBATIONS,
    add_unsupported_claim,
    drop_condition,
    parse_kinds,
    reverse_steps,
    swap_values,
)
from test_rag_mcp_target import FakeRagMcp, _evidence, _mcp, _output

from pr_suite_evaluation.cli import main
from pr_suite_evaluation.dataset import load_dataset
from pr_suite_evaluation.experiment import ExperimentRecord, ScorerConfig, run_experiment
from pr_suite_evaluation.judge.calibration import (
    CalibrationReport,
    CalibrationThresholds,
    LabelSet,
    calibrate,
    judge_labels,
    load_labels,
)
from pr_suite_evaluation.judge.client import JudgeReply
from pr_suite_evaluation.judge.judging import JudgeSettings
from pr_suite_evaluation.judge.labeling import (
    ERROR_TYPE_IDS,
    InterAnnotatorReport,
    assign_splits,
    build_labeling_items,
    load_items,
)
from pr_suite_evaluation.judge.rubric import load_rubric
from pr_suite_evaluation.judge.workbook import (
    LABEL_COLUMNS,
    LABEL_SHEET,
    META_SHEET,
    RUBRIC_SHEET,
    SYNTHETIC_SHEET,
    TARGET_SHEET,
    cell_text,
)
from pr_suite_evaluation.targets.rag_mcp import RagMcpTarget, RagTargetConfig

RUBRIC = load_rubric()
MODEL = "synthetic-judge"

# 合成のケース（業務の内容を含まない）。数値・番号つきの手順・条件の文を持つ回答を混ぜる。
NUMBERS = ("設定 S の上限は 30 です。下限は 5 です。", "設定 S の上限は 30、下限は 5。")
STEPS = (
    "手続き P は次の順に進めます。\n1. 申請書を作る\n2. 上長の確認を受ける\n3. 担当者に送る",
    "申請書を作り、上長の確認を受けてから担当者に送る。",
)
CONDITION = (
    "記録 R は 5 年間保管します。ただし、更新した場合は更新の日から数えます。",
    "5 年間。更新した場合は更新の日から数える。",
)
TEMPLATES = (NUMBERS, STEPS, CONDITION)


def _cases() -> list[dict[str, Any]]:
    cases: list[dict[str, Any]] = []
    for index in range(1, 13):
        answer, standard = TEMPLATES[index % 3]
        case: dict[str, Any] = {
            "id": f"c{index:02d}",
            "query": f"合成の質問 {index}",
            "category": "document_answerable" if index % 2 else "knowledge_missing",
            "standard_answer": standard,
        }
        if index == 12:
            # 標準回答の無いケースは claim_support だけを判定する。
            del case["standard_answer"]
        cases.append(case)
    return cases


def _record(cases: list[dict[str, Any]], dataset_path: Path) -> ExperimentRecord:
    outputs = []
    for index, case in enumerate(cases, start=1):
        output = _output(
            f"trace-{case['id']}",
            evidence=[_evidence("doc-s", f"k{index}", f"合成の根拠 {index}。上限は 30。")],
        )
        output["answer"] = TEMPLATES[index % 3][0]
        outputs.append(output)
    config = RagTargetConfig(mcp_url="http://rag/api/mcp", knowledge_base_ids=("kb",))
    return run_experiment(
        load_dataset(dataset_path),
        RagMcpTarget(_mcp(FakeRagMcp(outputs)), config),
        target_config=config.provenance(),
        scorer=ScorerConfig(target_kind="rag_mcp"),
        log=lambda _: None,
    )


@pytest.fixture
def exported(tmp_path: Path) -> dict[str, Path]:
    """評価の結果を保存し、label-template でブックと項目のファイルを書く。"""
    cases = _cases()
    dataset = tmp_path / "set.json"
    dataset.write_text(json.dumps({"cases": cases}, ensure_ascii=False), "utf-8")
    record = tmp_path / "a.json"
    record.write_text(_record(cases, dataset).model_dump_json(), "utf-8")
    workbook = tmp_path / "labels.xlsx"
    args = ["label-template", f"A={record}", "--dataset", str(dataset), "--perturb", "all"]
    assert main([*args, "--output", str(workbook)]) == 0
    return {
        "dataset": dataset,
        "record": record,
        "workbook": workbook,
        "items": tmp_path / "labels.items.json",
    }


def _columns(sheet: Any) -> dict[str, int]:
    header = [cell.value for cell in sheet[1]]
    return {
        column.key: header.index(column.header) + 1
        for column in LABEL_COLUMNS
        if column.header in header
    }


Filler = Callable[[str, str], tuple[str, str]]


def _fill(source: Path, target: Path, annotator: str, fill: Filler) -> None:
    """記入者の代わりに、テストの中だけの合成の判定をブックに書く（人の判定ではない）。"""
    book = load_workbook(source)
    sheet = book[LABEL_SHEET]
    columns = _columns(sheet)
    for row in range(2, sheet.max_row + 1):
        item_id = sheet.cell(row=row, column=columns["item_id"]).value
        aspect = sheet.cell(row=row, column=columns["aspect"]).value
        label, error_type = fill(str(item_id), str(aspect))
        sheet.cell(row=row, column=columns["label"]).value = label
        sheet.cell(row=row, column=columns["error_type"]).value = error_type
        sheet.cell(row=row, column=columns["note"]).value = f"{annotator} の合成のメモ"
        sheet.cell(row=row, column=columns["annotator"]).value = annotator
    book.save(target)


def _first(item_id: str, aspect: str) -> tuple[str, str]:
    case = item_id.split(":")[1]
    if case == "c07":
        return "cannot_judge", ""
    if aspect == "claim_support" and case in {"c03", "c06"}:
        return "fail", "unsupported_claim"
    return "pass", "correct_paraphrase" if case == "c01" else ""


def _second(item_id: str, aspect: str) -> tuple[str, str]:
    case = item_id.split(":")[1]
    if aspect == "claim_support" and case == "c05":
        return "fail", "unsupported_claim"
    if aspect == "claim_support" and case == "c09":
        return "cannot_judge", ""
    return _first(item_id, aspect)


class StubJudge:
    """台本の判定のモデル: 合成の項目は壊した観点を fail（1 件だけ見逃す）、ほかは pass。"""

    def __init__(self, labels: LabelSet, *, flaky: bool = False) -> None:
        self._items = {(item.question, item.answer): item for item in labels.items}
        self._flaky = flaky
        self.calls = 0

    @property
    def model(self) -> str:
        return MODEL

    def complete(
        self, *, instructions: str, prompt: str, schema: Mapping[str, Any], schema_name: str
    ) -> JudgeReply:
        self.calls += 1
        payload = json.loads(prompt)
        item = self._items[(payload["question"], payload["answer"])]
        missed = item.id.endswith("~add_unsupported_claim") and "c02" in item.id
        aspects = []
        for aspect_id in payload["aspects"]:
            label = "pass"
            if item.origin == "synthetic" and item.labels.get(aspect_id) == "fail" and not missed:
                label = "fail"
            if item.origin == "human" and item.labels.get(aspect_id) == "fail":
                label = "fail"
            if self._flaky and aspect_id == "claim_support" and item.id.endswith("c01"):
                label = "fail"
            aspects.append({"aspect_id": aspect_id, "label": label, "reason": "合成"})
        return JudgeReply(text=json.dumps({"aspects": aspects}), model=MODEL, usage={})


# ---- 故意に壊した回答 ---------------------------------------------------------------------


def test_swap_values_swaps_two_distinct_numbers_but_not_list_marks_or_ids() -> None:
    result = swap_values("1. 上限は 30 で、下限は 5 です。参照 [E3]。", None, "k")
    assert result is not None
    assert result.answer == "1. 上限は 5 で、下限は 30 です。参照 [E3]。"
    assert swap_values("上限は 30 です。", None, "k") is None
    assert swap_values("値は 30、別の値も 30。", None, "k") is None


def test_drop_condition_drops_the_condition_sentence_closest_to_the_standard_answer() -> None:
    answer, standard = CONDITION
    result = drop_condition(answer, standard, "k")
    assert result is not None
    assert result.answer == "記録 R は 5 年間保管します。"
    # 標準回答が無い・1 文だけ・条件の文が標準回答と重ならないなら壊さない。
    assert drop_condition(answer, None, "k") is None
    assert drop_condition("ただし、更新した場合は数え直す。", standard, "k") is None
    assert drop_condition("値は 5。ただし、別の場合は除く。", "全く違う内容", "k") is None


def test_add_unsupported_claim_inserts_a_fixed_claim_deterministically() -> None:
    first = add_unsupported_claim("段落 1\n\n段落 2", None, "item-a")
    again = add_unsupported_claim("段落 1\n\n段落 2", None, "item-a")
    assert first is not None and again is not None
    assert first.answer == again.answer
    lines = first.answer.split("\n\n")
    assert (lines[0], lines[2]) == ("段落 1", "段落 2")
    assert add_unsupported_claim("  ", None, "k") is None


def test_reverse_steps_reverses_numbered_steps_and_keeps_the_numbers() -> None:
    result = reverse_steps(STEPS[0], None, "k")
    assert result is not None
    assert result.answer.split("\n")[1:] == [
        "1. 担当者に送る",
        "2. 上長の確認を受ける",
        "3. 申請書を作る",
    ]
    assert reverse_steps("1. 同じ\n2. 同じ", None, "k") is None
    assert reverse_steps("手順はありません。", None, "k") is None


def test_perturbation_kinds_and_fail_aspects() -> None:
    assert parse_kinds("all") == list(PERTURBATIONS)
    assert parse_kinds("swap_values, reverse_steps") == ["swap_values", "reverse_steps"]
    assert parse_kinds(None) == []
    with pytest.raises(ValueError, match="知らない壊し方"):
        parse_kinds("unknown")
    assert {aspect for item in PERTURBATIONS.values() for aspect in item.fail_aspects} <= set(
        RUBRIC.aspect_ids
    )


# ---- 区分と項目 ---------------------------------------------------------------------------


def test_splits_are_deterministic_stratified_and_shared_by_the_same_case(tmp_path: Path) -> None:
    strata = {f"c{index:02d}": "x" if index % 2 else "y" for index in range(20)}
    splits = assign_splits(strata, holdout_ratio=0.7, salt="s")
    assert splits == assign_splits(strata, holdout_ratio=0.7, salt="s")
    for stratum in ("x", "y"):
        units = [unit for unit, value in strata.items() if value == stratum]
        assert sum(splits[unit] == "dev" for unit in units) == 3
    assert assign_splits(strata, holdout_ratio=0.7, salt="other") != splits
    with pytest.raises(ValueError, match="holdout の割合"):
        assign_splits(strata, holdout_ratio=0, salt="s")

    cases = _cases()
    dataset = tmp_path / "set.json"
    dataset.write_text(json.dumps({"cases": cases}, ensure_ascii=False), "utf-8")
    record = _record(cases, dataset)
    items = build_labeling_items(
        [("A", record), ("C", record)],
        load_dataset(dataset),
        RUBRIC,
        perturb=["swap_values"],
    )
    by_id = {item.id: item for item in items.items}
    for case in cases:
        assert by_id[f"A:{case['id']}"].split == by_id[f"C:{case['id']}"].split
    assert by_id["A:c12"].aspects == ["claim_support"]
    synthetic = items.synthetic_items
    assert synthetic and all(item.split == by_id[item.id.split("~")[0]].split for item in synthetic)
    assert all(item.synthetic_labels for item in synthetic)


# ---- 書き出し・記入・取り込み・裁定・校正の通し -----------------------------------------------


def test_label_template_writes_the_workbook_and_the_items_file(exported: dict[str, Path]) -> None:
    items = load_items(exported["items"])
    assert items.rubric["sha256"] == RUBRIC.sha256
    human = items.human_items
    assert len(human) == 12
    book = load_workbook(exported["workbook"])
    assert book.sheetnames == [LABEL_SHEET, TARGET_SHEET, RUBRIC_SHEET, SYNTHETIC_SHEET, META_SHEET]
    assert book[META_SHEET].sheet_state == "hidden"
    sheet = book[LABEL_SHEET]
    columns = _columns(sheet)
    # 1 行 = 項目 × 観点（標準回答の無い c12 は claim_support だけ）。
    assert sheet.max_row - 1 == sum(len(item.aspects) for item in human) == 11 * 3 + 1
    row = {key: sheet.cell(row=2, column=index).value for key, index in columns.items()}
    assert row["split"] in {"dev", "holdout"}
    assert row["label"] is None and row["annotator"] is None
    assert "合成の根拠" in str(row["evidence"])
    assert str(row["rubric"]).startswith(("標準回答", "回答"))
    targets = book[TARGET_SHEET]
    ids = [targets.cell(row=index, column=2).value for index in range(2, 2 + len(ERROR_TYPE_IDS))]
    assert ids == list(ERROR_TYPE_IDS)
    assert str(targets.cell(row=2, column=6).value).startswith("=COUNTIFS(")
    synthetic = book[SYNTHETIC_SHEET]
    assert synthetic.max_row - 1 == len(items.synthetic_items) > 0


def test_cell_text_strips_control_characters_and_truncates() -> None:
    assert cell_text("a\x01b") == "ab"
    assert cell_text("x" * 100, limit=60).endswith("全文は項目のファイル）")


def test_round_trip_from_two_annotators_to_calibration(
    exported: dict[str, Path], tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    first, second = tmp_path / "a1.xlsx", tmp_path / "a2.xlsx"
    _fill(exported["workbook"], first, "a1", _first)
    _fill(exported["workbook"], second, "a2", _second)
    items_arg = ["--items", str(exported["items"])]
    report_path, adjudication = tmp_path / "iaa.json", tmp_path / "adj.xlsx"
    args = ["label-import", *items_arg, "--annotations", str(first), str(second)]
    outputs = ["--report", str(report_path), "--adjudication-output", str(adjudication)]
    assert main([*args, *outputs]) == 0
    report = InterAnnotatorReport.model_validate_json(report_path.read_text("utf-8"))
    support = report.aspects["claim_support"]
    # c07 は 2 人とも cannot_judge、c09 は a2 が cannot_judge → 比べるのは 10 行。
    # a1 の fail は c03・c06、a2 の fail は c03・c05・c06 → 一致 9/10、κ = (0.9-0.62)/0.38。
    assert support.pairs[0].compared == 10
    assert support.pairs[0].agreement == 0.9
    assert support.pairs[0].kappa == round((0.9 - 0.62) / 0.38, 4)
    assert support.disagreements == 2
    assert report.aspects["answer_correctness"].disagreements == 0
    assert sorted(item.item_id for item in report.disagreements) == ["A:c05", "A:c09"]
    assert "要裁定 2 行" in capsys.readouterr().err

    # 裁定しないと LabelSet を書かない（食い違いが残る）。
    labels_path = tmp_path / "labels.json"
    assert main([*args, "--output", str(labels_path)]) == 2
    assert "判定の決まらない行が 2 行" in capsys.readouterr().err

    # 裁定（テストの中の合成の裁定）: 要裁定の行を埋め、合成の対照を 1 つ除外する。
    book = load_workbook(adjudication)
    sheet = book[LABEL_SHEET]
    header = [cell.value for cell in sheet[1]]
    columns = _columns(sheet)
    state = header.index("状態") + 1
    filled = 0
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row=row, column=state).value == "要裁定":
            case = str(sheet.cell(row=row, column=columns["item_id"]).value)
            label = "fail" if case.endswith("c05") else "pass"
            sheet.cell(row=row, column=columns["label"]).value = label
            sheet.cell(row=row, column=columns["annotator"]).value = "adj1"
            filled += 1
    assert filled == 2
    synthetic = book[SYNTHETIC_SHEET]
    excluded = str(synthetic.cell(row=2, column=1).value)
    synthetic.cell(row=2, column=7).value = "x"
    book.save(adjudication)

    command = [*args, "--adjudicated", str(adjudication), "--output", str(labels_path)]
    assert main(command) == 0
    labels = load_labels(labels_path)
    by_id = {item.id: item for item in labels.items}
    # 全観点が cannot_judge の c07 は落とし、合成の対照は除外したもの以外を残す。
    assert "A:c07" not in by_id
    assert excluded not in by_id
    assert by_id["A:c05"].labels["claim_support"] == "fail"
    assert by_id["A:c05"].annotator_labels == {
        "a1": {
            "requirement_coverage": "pass",
            "answer_correctness": "pass",
            "claim_support": "pass",
        },
        "a2": {
            "requirement_coverage": "pass",
            "answer_correctness": "pass",
            "claim_support": "fail",
        },
    }
    assert by_id["A:c05"].labeler == "adj1"
    assert by_id["A:c03"].error_types == ["unsupported_claim"]
    assert by_id["A:c12"].labels == {"claim_support": "pass"}
    synthetic_items = [item for item in labels.items if item.origin == "synthetic"]
    assert synthetic_items and all(
        set(item.labels.values()) == {"fail"} and item.perturbation for item in synthetic_items
    )
    assert labels.provenance["adjudicator"] == "adj1"
    assert labels.provenance["inter_annotator"]["claim_support"]["disagreements"] == 2

    # calibrate: 合成の対照は holdout の合否に入れず、synthetic の区分に別に集計する。
    settings = JudgeSettings()
    judge = StubJudge(labels)
    runs = [judge_labels(labels, RUBRIC, judge, settings) for _ in range(2)]
    flaky = judge_labels(labels, RUBRIC, StubJudge(labels, flaky=True), settings)
    calibration = calibrate(
        labels,
        RUBRIC,
        judgements=runs[0],
        judge_provenance={"model": MODEL},
        settings=settings,
        thresholds=CalibrationThresholds(min_items=3),
        repeat_judgements=[runs[1], flaky],
    )
    support_report = calibration.aspects["claim_support"]
    human_holdout = [
        item
        for item in labels.items
        if item.origin == "human" and item.split == "holdout" and "claim_support" in item.labels
    ]
    assert support_report.splits["holdout"].total == len(human_holdout)
    assert support_report.splits["holdout"].kappa in {1.0, None}
    synthetic_split = support_report.splits["synthetic"]
    assert synthetic_split.confusion["fail"]["pass"] >= 1
    assert synthetic_split.false_pass_rate is not None and synthetic_split.false_pass_rate > 0
    assert calibration.provenance["labels"]["origins"]["synthetic"] == len(synthetic_items)
    assert calibration.provenance["repeats"] == 3
    # 3 回のうち 1 回だけ c01 の claim_support が変わる → 揺れた項目に数える。
    variance = calibration.variance["claim_support"]
    assert variance.runs == 3 and variance.unstable == 1
    assert variance.unstable_items == ["A:c01"]
    assert calibration.variance["answer_correctness"].unstable == 0

    # 保存した校正の結果から、モデルを呼ばずに集計だけやり直す（分散も残る）。
    saved = tmp_path / "calibration.json"
    saved.write_text(calibration.model_dump_json(), "utf-8")
    again = tmp_path / "again.json"
    rerun = ["calibrate", str(labels_path), "--judgements", str(saved), "--min-items", "3"]
    assert main([*rerun, "--output", str(again)]) == 0
    restored = CalibrationReport.model_validate_json(again.read_text("utf-8"))
    assert restored.variance["claim_support"].unstable_items == ["A:c01"]
    assert "unstable_rate=" in capsys.readouterr().err


def test_label_import_rejects_mismatched_or_malformed_workbooks(
    exported: dict[str, Path], tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    first = tmp_path / "a1.xlsx"
    _fill(exported["workbook"], first, "a1", _first)
    items_arg = ["--items", str(exported["items"])]
    # 記入者は 2 人以上。
    single = ["label-import", *items_arg, "--annotations", str(first), "--report"]
    assert main([*single, str(tmp_path / "r.json")]) == 2
    assert "2 つ以上" in capsys.readouterr().err
    # 語彙に無い判定。
    bad = tmp_path / "bad.xlsx"
    _fill(exported["workbook"], bad, "a2", lambda item, aspect: ("maybe", ""))
    args = ["label-import", *items_arg, "--annotations", str(first), str(bad)]
    assert main([*args, "--report", str(tmp_path / "r.json")]) == 2
    assert "判定は pass / fail / cannot_judge" in capsys.readouterr().err
    # 別の項目のファイルのブックは混ぜない。
    raw = json.loads(exported["items"].read_text("utf-8"))
    raw["items"][0]["question"] = "変えた質問"
    other = tmp_path / "other.items.json"
    other.write_text(json.dumps(raw, ensure_ascii=False), "utf-8")
    second = tmp_path / "a2.xlsx"
    _fill(exported["workbook"], second, "a2", _first)
    mismatch = ["label-import", "--items", str(other), "--annotations", str(first), str(second)]
    assert main([*mismatch, "--report", str(tmp_path / "r.json")]) == 2
    assert "項目のファイルのブックではありません" in capsys.readouterr().err
    # 全員が一致していれば、裁定なしで LabelSet を書ける。記入者の ID は <ID>=<ブック> でも渡せる。
    labels_path = tmp_path / "labels.json"
    agreed = ["label-import", *items_arg, "--annotations", f"a1={first}", f"a2={second}"]
    assert main([*agreed, "--output", str(labels_path)]) == 0
    labels = load_labels(labels_path)
    assert {item.labeler for item in labels.items if item.origin == "human"} == {"a1+a2"}


def test_label_set_fields_match_shared_contract() -> None:
    """LabelSet の形は platform の共有の名前と同じ（RAG の判定の校正の画面と往復する。#3578）。"""
    from pr_evaluation_scoring.labeling import LABEL_SET_FIELDS, LABEL_SET_ITEM_FIELDS

    from pr_suite_evaluation.judge.calibration import LabeledItem

    assert tuple(LabelSet.model_fields) == LABEL_SET_FIELDS
    assert tuple(LabeledItem.model_fields) == LABEL_SET_ITEM_FIELDS
